# test_sad_corr_gui.py — вікно «Кореляційний аналіз»: Кендалл, Холм, ДІ, пропуски, експорт (Xvfb)
import os, sys, tempfile, shutil
sys.path.insert(0, os.environ.get("SAD_SRC", os.path.dirname(os.path.abspath(__file__))))
os.environ.setdefault("DISPLAY", ":99")
import numpy as np
from scipy import stats
import tkinter as tk
FAIL = []
def check(c, msg):
    print(("  [OK]   " if c else "  [FAIL] ") + msg, flush=True)
    if not c: FAIL.append(msg)
import main as m
import sad_breeding_ui as U
LOG = {"info": [], "warn": []}
mb = tk.messagebox
mb.showinfo = lambda t="", msg="", **k: LOG["info"].append((t, msg))
mb.showwarning = lambda t="", msg="", **k: LOG["warn"].append((t, msg))
mb.showerror = lambda t="", msg="", **k: LOG["warn"].append((t, msg))
mb.askyesno = lambda *a, **k: False
U.open_file_cross = lambda p: None
TMP = tempfile.mkdtemp(prefix="sadcg_"); SAVE = {"p": None}
tk.filedialog.asksaveasfilename = lambda **k: SAVE["p"]
root = tk.Tk(); app = m.SADTk(root); root._sad_app = app; root.update()
rng = np.random.default_rng(3)

def toplevels(w=root):
    out = [w] if isinstance(w, tk.Toplevel) else []
    for c in w.winfo_children(): out += toplevels(c)
    return out
def widgets(w, cls):
    out = [w] if w.winfo_class() == cls else []
    for c in w.winfo_children(): out += widgets(c, cls)
    return out
def find_win(frag):
    r = [t for t in toplevels() if frag in t.title()]; return r[-1] if r else None
def fill(win, cols, names=None):
    for j, col in enumerate(cols):
        if names: win.header_vars[j].set(names[j])
        for i, v in enumerate(col):
            while i >= len(win.entries): win._add_row_widgets(len(win.entries)) if hasattr(win, "_add_row_widgets") else None
            win.entries[i][j].delete(0, tk.END); win.entries[i][j].insert(0, "" if v is None else str(v))
def clear(win):
    for r_ in win.entries:
        for e in r_: e.delete(0, tk.END)

w = m.CorrelationWindow(root, {}); root.update()
print("== 1. Діалог параметрів: Кендалл і Холм")
captured = {}
def fake_wait(dlg):
    radios = {b.cget("text").split(" ")[0]: b for b in widgets(dlg, "Radiobutton")}
    captured["methods"] = sorted(k for k in radios if k in ("Авто", "Пірсон", "Спірмен", "Кендалл"))
    captured["corr"] = [b.cget("text").split(" ")[0] for b in widgets(dlg, "Radiobutton")][4:]
    radios["Кендалл"].invoke()
    [b for b in widgets(dlg, "Button") if "Виконати" in b.cget("text")][0].invoke()
w.win.wait_window = fake_wait
n = 25; base = rng.normal(size=n)
cols = [base, 0.7 * base + rng.normal(0, .6, n), -0.5 * base + rng.normal(0, .8, n), rng.normal(size=n)]
fill(w, [[f"{v:.3f}" for v in c] for c in cols], ["Урожай", "Маса", "Вологість", "Шум"])
w._run_analysis(); root.update()
check(set(captured["methods"]) == {"Авто", "Пірсон", "Спірмен", "Кендалл"}, f"у діалозі є всі методи: {captured['methods']}")
check(captured["corr"][0] == "Холм" and "Бонферроні" in captured["corr"], f"поправки: {captured['corr']} (Холм першим — рекомендована)")
rw = find_win("Кореляційний аналіз — результати")
check(rw is not None, "вікно результатів відкрилось")
labels_txt = " ".join(l.cget("text") for l in widgets(rw, "Label"))
check("Кендалл τ-b" in labels_txt and "Холм" in labels_txt, "у шапці вказано метод «Кендалл τ-b» і поправку «Холм»")
M = w._res_M
check(M["method"] == "kendall" and M["correction"] == "holm" and len(M["pairs"]) == 6, "виконано Кендалл + Холм для 6 пар")
ref = stats.kendalltau(cols[0], cols[1], variant="b")
check(abs(M["r"][0, 1] - ref[0]) < 1e-12 and abs(M["p_raw"][0, 1] - ref[1]) < 1e-12, f"τ(Урожай, Маса) = {M['r'][0, 1]:.4f} збігається зі scipy")
tabs = {b.cget("text").strip(): b for b in widgets(rw, "Button")}
tb = [b for k, b in tabs.items() if "Таблиця пар" in k][0]; tb.invoke(); root.update()
tv = widgets(rw, "Treeview")
check(len(tv) == 1 and len(tv[0].get_children()) == 6, "вкладка «Таблиця пар»: 6 рядків")
check(w._res_pairs_sec["headers"][3] == "τ-b" and "ДІ від" in w._res_pairs_sec["headers"], "заголовки: τ-b, довірчий інтервал, p сире й з поправкою")
rows = w._res_pairs_sec["rows"]; check(all(r[4] < r[3] < r[5] for r in rows), "для кожної пари довірчий інтервал містить оцінку")
check(w._res_pairs_sec["rows"][0][-1] in ("істотний", "не істотний") and rows[0][3] == max(rows, key=lambda r: abs(r[3]))[3], "рядки впорядковано за |τ|; є висновок")
check(w._hm_fig is not None and w._sc_fig is not None, "теплова карта й матриця розсіювання побудовані для Кендалла")
import sad_breeding_ui as _U
check(_U.fcell(0.00001, "p", True) == "<0,0001" and _U.fcell(0.0312, "p", True) == "0,0312" and _U.fcell(0.0312, "p", False) == "0.0312",
      "p-значення форматуються як «<0,0001» / «0,0312»")
SAVE["p"] = os.path.join(TMP, "pairs.xlsx")
[b for b in widgets(rw, "Button") if b.cget("text") == "📊 Excel"][0].invoke()
import openpyxl; wb = openpyxl.load_workbook(SAVE["p"]); ws = wb.worksheets[0]
check(all(ws.cell(row=r_, column=1).value for r_ in range(2, 8)) and ws["A8"].value is None and isinstance(ws["D2"].value, float)
      and ws.max_row > 8, "Excel із таблиці пар: 6 рядків пар, далі порожній рядок і примітки; r — число, а не текст")
SAVE["p"] = os.path.join(TMP, "pairs.docx")
[b for b in widgets(rw, "Button") if b.cget("text") == "📄 Word"][0].invoke()
import docx; d = docx.Document(SAVE["p"]); check(len(d.tables) == 1 and len(d.tables[0].rows) == 7, "Word із таблиці пар: таблиця 6 пар")
rw.destroy()

print("\n== 2. Пропуски: пари зі спільних рядків (виправлена помилка зсуву)")
clear(w)
X = [str(v) for v in range(1, 11)]; Y = list(X); Y[2] = ""
Z = [str(v * v) for v in range(1, 11)]
fill(w, [X, Y, Z], ["X", "Y", "Z"] + [None] * 0)
LOG["info"].clear(); LOG["warn"].clear()
w._compute_and_show("pearson", 0.05, "none"); root.update()
M = w._res_M; check(abs(M["r"][0, 1] - 1.0) < 1e-12 and M["n"][0, 1] == 9, f"X–Y: r = {M['r'][0, 1]:.6f}, n = {M['n'][0, 1]} (раніше було б 0,99…)")
mk = np.array([0, 1, 3, 4, 5, 6, 7, 8, 9]); xz = np.arange(1, 11.0); zz = xz ** 2
check(abs(M["r"][1, 2] - stats.pearsonr(xz[mk], zz[mk])[0]) < 1e-12, "Y–Z зі спільними рядками збігається з ручним попарним виключенням")
check(M["n"][0, 2] == 10 and M["n"][0, 1] == 9, "n різні для різних пар (10 і 9) — рядки не зсунулись")
find_win("Кореляційний аналіз — результати").destroy()

print("\n== 3. Нечислові клітинки й порожня таблиця")
clear(w); fill(w, [[1, 2, 3, 4, 5, 6], [2, 4, "abc", 8, 10, 12], [3, 1, 2, 5, 4, 6]], ["A", "B", "C"])
LOG["warn"].clear(); w._compute_and_show("spearman", 0.05, "holm"); root.update()
check(any("Нечислові" in t and "рядок 3" in msg and "«B»" in msg for t, msg in LOG["warn"]), "нечислова клітинка названа (рядок 3, стовпець B) і виключена без зсуву рядків")
M = w._res_M; check(M["n"][0, 1] == 5 and abs(M["r"][0, 1] - 1.0) < 1e-12, "решту пари пораховано на спільних рядках")
find_win("Кореляційний аналіз — результати").destroy()
clear(w); LOG["warn"].clear(); w._compute_and_show("pearson", 0.05, "holm")
check(any("Замало" in t for t, _ in LOG["warn"]), "порожня таблиця → зрозуміле попередження")

print("\n== 4. Авто-вибір методу та інші поправки")
clear(w); nn = 30; a_ = rng.normal(size=nn)
fill(w, [[f"{v:.3f}" for v in a_], [f"{v:.3f}" for v in a_ + rng.normal(0, .5, nn)]], ["P", "Q"])
LOG["info"].clear(); w._compute_and_show("auto", 0.05, "bh"); root.update()
check(any("Пірсон" in msg for t, msg in LOG["info"]) and w._res_M["method"] == "pearson", "нормальні дані → авто обрав Пірсона")
find_win("Кореляційний аналіз — результати").destroy()
clear(w); skew = np.exp(rng.normal(size=nn)) ** 2
fill(w, [[f"{v:.3f}" for v in skew], [f"{v:.3f}" for v in skew * 2 + rng.normal(size=nn)]], ["P", "Q"])
LOG["info"].clear(); w._compute_and_show("auto", 0.05, "holm"); root.update()
check(any("Спірмен" in msg for t, msg in LOG["info"]) and w._res_M["method"] == "spearman", "сильно скошені дані → авто обрав Спірмена")
find_win("Кореляційний аналіз — результати").destroy()
print("\n== 5. Частинна кореляція у вікні")
clear(w); nn = 60; zc = rng.normal(size=nn); xc = zc + rng.normal(0, .6, nn); yc = zc + rng.normal(0, .6, nn); wc = rng.normal(size=nn)
fill(w, [[f"{v:.3f}" for v in c] for c in (xc, yc, wc, zc)], ["X", "Y", "W", "Z"])
btn = [b for b in widgets(w.win, "Button") if "Частинна" in b.cget("text")]
check(len(btn) == 1, "на панелі є кнопка «Частинна кореляція»")
state = {}
def fake_wait2(dlg):
    lfs = {lf.cget("text").split(" ")[0]: lf for lf in widgets(dlg, "Labelframe")}
    chk = lambda key, name: [c for c in widgets(lfs[key], "Checkbutton") if c.cget("text") == name][0]
    for nm in ("X", "Y", "W"): chk("Досліджувані", nm).invoke()
    chk("Контрольні", "Z").invoke()
    state["methods"] = sorted(r.cget("text").split(" ")[0] for r in widgets(dlg, "Radiobutton"))
    [b for b in widgets(dlg, "Button") if "Виконати" in b.cget("text")][0].invoke()
w.win.wait_window = fake_wait2
w._partial_dialog(); root.update()
check(state["methods"] == ["Пірсон", "Спірмен"], "у діалозі лише Пірсон і Спірмен (Кендалл для частинної не реалізовано)")
pw = find_win("Частинна кореляція — результати")
check(pw is not None, "вікно результатів частинної кореляції відкрилось")
sb = [b for b in widgets(pw, "Button") if b.cget("text").startswith("  ") and b.cget("text").strip()]
errs = []
for b in sb:
    try: b.invoke(); pw.update()
    except Exception as ex: errs.append(str(ex))
check(len(sb) == 5 and not errs, f"усі {len(sb)} розділів відображаються без помилок {errs[:1]}")
check(w._res_partial["n"] == nn and w._res_partial["controls"] == ["Z"] and len(w._res_partial["pairs"]) == 3, "n = 60, контроль Z, 3 пари")
b0 = sb[0]; b0.invoke(); pw.update(); tvs = widgets(pw, "Treeview"); check(tvs and len(tvs[0].get_children()) == 3, "таблиця «Частинні кореляції»: 3 рядки")
pw.destroy()
# перевірки введення
def fake_wait3(dlg):
    lfs = {lf.cget("text").split(" ")[0]: lf for lf in widgets(dlg, "Labelframe")}
    for key in ("Досліджувані", "Контрольні"):
        for c in widgets(lfs[key], "Checkbutton"):
            if c.cget("text") in ("X", "Y"): c.invoke()
    [b for b in widgets(dlg, "Button") if "Виконати" in b.cget("text")][0].invoke()
    state["dlg_alive"] = dlg.winfo_exists()
    [b for b in widgets(dlg, "Button") if "Скасувати" in b.cget("text")][0].invoke()
w.win.wait_window = fake_wait3; LOG["warn"].clear(); w._partial_dialog(); root.update()
check(any("не може бути і досліджуваною, і контрольною" in m_ for _, m_ in LOG["warn"]) and state.get("dlg_alive"), "перетин досліджуваних і контрольних ознак → попередження, діалог лишається відкритим")
clear(w); fill(w, [[1, 2, 3, 4], [2, 1, 4, 3]], ["A", "B"]); LOG["warn"].clear(); w._partial_dialog()
check(any("щонайменше 3 стовпці" in m_ for _, m_ in LOG["warn"]), "менше 3 стовпців → зрозуміле попередження")
clear(w); fill(w, [[1, 2, 3, 4, 5], [2, 1, 4, 3, 6], [5, 4, 3, 2, 1], [1, 3, 2, 5, 4]], ["A", "B", "C", "D"]); LOG["warn"].clear()
labels_, arrays_, _ = w._collect_columns(); w._run_partial(labels_, arrays_, [2, 3], [0, 1], "pearson", 0.05, "holm")
check(any("Замало" in m_ for _, m_ in LOG["warn"]) or any("n > k" in m_ or "потрібно" in m_ for _, m_ in LOG["warn"]), "замало спостережень для двох контрольних ознак → зрозуміла відмова без аварії")

print("\n== 6. Генотипні кореляції: кнопка, меню, вікно")
labels_all = []
def _rl(x):
    try:
        if x.winfo_class() == "Label" and x.cget("text"): labels_all.append(x.cget("text"))
    except Exception: pass
    for c in x.winfo_children(): _rl(c)
_rl(root); check(any("Генотипні кореляції" in s for s in labels_all), "пункт «Генотипні кореляції» є в меню розділу «Зв'язок змінних»")
btn_g = [b for b in widgets(w.win, "Button") if "Генотипні кореляції" in b.cget("text")]
check(len(btn_g) == 1, "у вікні кореляції є кнопка «Генотипні кореляції»")
btn_g[0].invoke(); root.update()
gw = find_win("Генотипні й фенотипні кореляції (rG, rP)")
check(gw is not None, "кнопка відкриває вікно введення (генотип × повторність × ознаки)")
import sad_breeding_win as BW
gc = BW.GenCorrWindow(root); root.update()
check(gc.header_names()[:3] == ["Генотип", "Повторність", "Ознака 1"], "заголовки за замовчуванням")
gc._add_col(); check(gc.header_names()[-1] == "Ознака 5" and len(gc.header_names()) == 7, "додавання стовпця: назва «Ознака 5»")
gc._del_col()
n0 = len(LOG["warn"]); gc._run(); check(len(LOG["warn"]) == n0 + 1 and "порожня" in LOG["warn"][-1][1], "порожня таблиця → зрозуміле попередження")
rg_ = np.random.default_rng(14); g_, r_ = 36, 3
sG = np.array([[1, .7, .3], [.7, 1, .2], [.3, .2, 1]]); ge = rg_.multivariate_normal(np.zeros(3), sG, size=g_)
data = []
for i in range(g_):
    for j in range(r_): data.append([f"G{i + 1}", f"R{j + 1}"] + [f"{10 + v:.3f}".replace(".", ",") for v in ge[i] + rg_.normal(0, 1.0, 3)])
while len(gc.entries) < len(data): gc._add_rows(10)
for hn, e in zip(["Генотип", "Повторність", "Висота", "Маса", "Урожай"], gc.header_entries): e.delete(0, tk.END); e.insert(0, hn)
for row, vals in zip(gc.entries, data):
    for e, v in zip(row, vals): e.delete(0, tk.END); e.insert(0, v)
gc.boot_var.set("500"); gc._run(); root.update()
rv = find_win("Генотипні й фенотипні кореляції — результати")
check(rv is not None, "результати відкрились")
sb = [b for b in widgets(rv, "Button") if b.cget("text").startswith("  ") and b.cget("text").strip()]
errs = []
for b in sb:
    try: b.invoke(); rv.update()
    except Exception as ex: errs.append(str(ex))
check(len(sb) == 6 and not errs, f"усі {len(sb)} розділів відображаються без помилок {errs[:1]}")
sb[0].invoke(); rv.update(); tv_ = widgets(rv, "Treeview"); check(tv_ and len(tv_[0].get_children()) == 3, "таблиця пар: 3 рядки (3 ознаки)")
rv.destroy()
gc.win.destroy(); gw.destroy()
data_bad = data[:-1]
gc2 = BW.GenCorrWindow(root); root.update()
while len(gc2.entries) < len(data_bad): gc2._add_rows(10)
for row, vals in zip(gc2.entries, data_bad):
    for e, v in zip(row, vals): e.delete(0, tk.END); e.insert(0, v)
n0 = len(LOG["warn"]); gc2._run(); check(len(LOG["warn"]) == n0 + 1 and "збалансованим" in LOG["warn"][-1][1], "незбалансований дослід → зрозуміле попередження")
gc2.win.destroy()

print("\n== 7. Зв'язок змінних різних типів: меню, кнопка, три режими")
labels_all = []
_rl(root); check(any("Зв'язок різних типів змінних" in s for s in labels_all), "пункт «Зв'язок різних типів змінних» є в меню розділу «Зв'язок змінних»")
bt = [b for b in widgets(w.win, "Button") if "Різні типи змінних" in b.cget("text")]
check(len(bt) == 1, "у вікні кореляції є кнопка «Різні типи змінних»")
bt[0].invoke(); root.update(); aw0 = find_win("Зв'язок змінних різних типів"); check(aw0 is not None, "кнопка відкриває вікно")
import sad_breeding_win as BW
aw = BW.AssocWindow(root); root.update()
check(aw.header_names()[:2] == ["Група", "Значення"], "режим η: заголовки «Група», «Значення»")
aw.mode_var.set(aw.MODES[2]); aw._on_mode(); check(aw.header_names() == ["Ознака A", "Ознака B", "Частота (необов'язково)"], "режим спряженості: заголовки змінюються")
aw.mode_var.set(aw.MODES[0]); aw._on_mode()
def put(win, data):
    for r_ in win.entries:
        for e in r_: e.delete(0, tk.END)
    while len(win.entries) < len(data): win._add_rows(10)
    for r_, vals in zip(win.entries, data):
        for e, v in zip(r_, vals): e.delete(0, tk.END); e.insert(0, v)
n0 = len(LOG["warn"]); aw._run(); check(len(LOG["warn"]) == n0 + 1 and "порожня" in LOG["warn"][-1][1], "порожня таблиця → зрозуміле попередження")
rg3 = np.random.default_rng(31); rows_e = []
for nm, mu in (("Сорт А", 20), ("Сорт Б", 24), ("Сорт В", 27)):
    for v in rg3.normal(mu, 2, 9): rows_e.append([nm, f"{v:.2f}".replace(".", ","), ""])
put(aw, rows_e); aw._run(); root.update(); rv = find_win("Кореляційне відношення η – результати")
check(rv is not None, "режим η: результати відкрились")
sb = [b for b in widgets(rv, "Button") if b.cget("text").startswith("  ") and b.cget("text").strip()]; errs = []
for b in sb:
    try: b.invoke(); rv.update()
    except Exception as ex: errs.append(str(ex))
check(len(sb) == 5 and not errs, f"режим η: {len(sb)} розділів без помилок {errs[:1]}"); rv.destroy()
aw.mode_var.set(aw.MODES[1]); aw._on_mode(); n0 = len(LOG["warn"]); aw._run()
check(len(LOG["warn"]) == n0 + 1 and "рівно 2" in LOG["warn"][-1][1], "режим «Дві групи» із трьома групами → пояснення («рівно 2», оберіть η)")
put(aw, [r_ for r_ in rows_e if r_[0] != "Сорт В"]); aw._run(); root.update(); rv = find_win("Дві групи – результати")
check(rv is not None, "режим «Дві групи»: результати відкрились")
sb = [b for b in widgets(rv, "Button") if b.cget("text").startswith("  ") and b.cget("text").strip()]; errs = []
for b in sb:
    try: b.invoke(); rv.update()
    except Exception as ex: errs.append(str(ex))
check(len(sb) == 5 and not errs, f"режим «Дві групи»: {len(sb)} розділів без помилок {errs[:1]}"); rv.destroy()
aw.mode_var.set(aw.MODES[2]); aw._on_mode()
put(aw, [["Сорт 1", "уражені", "12"], ["Сорт 1", "здорові", "38"], ["Сорт 2", "уражені", "25"], ["Сорт 2", "здорові", "25"], ["Сорт 3", "уражені", "6"], ["Сорт 3", "здорові", "44"]])
aw.boot_var.set("500"); aw._run(); root.update(); rv = find_win("Таблиця спряженості – результати")
check(rv is not None, "режим спряженості: результати відкрились")
sb = [b for b in widgets(rv, "Button") if b.cget("text").startswith("  ") and b.cget("text").strip()]; errs = []
for b in sb:
    try: b.invoke(); rv.update()
    except Exception as ex: errs.append(str(ex))
check(len(sb) == 8 and not errs, f"режим спряженості: {len(sb)} розділів без помилок {errs[:1]}"); rv.destroy()
put(aw, [["А", "так", "-3"], ["Б", "ні", "4"]]); n0 = len(LOG["warn"]); aw._run()
check(len(LOG["warn"]) == n0 + 1, "від'ємна частота → зрозуміле попередження")
aw.win.destroy(); aw0.destroy()

w.win.destroy(); shutil.rmtree(TMP, ignore_errors=True); root.destroy()
print("\n" + "=" * 60)
if FAIL:
    print(f"ПРОВАЛІВ: {len(FAIL)}"); [print("  -", f) for f in FAIL]; sys.exit(1)
print("УСІ ПЕРЕВІРКИ ВІКНА КОРЕЛЯЦІЇ ПРОЙДЕНО")
