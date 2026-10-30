# test_sad_report_gui.py – універсальний експорт звіту у Word/Excel (потрібен дисплей; Xvfb)
import os, sys, io, tempfile, shutil, zipfile
sys.path.insert(0, os.environ.get("SAD_SRC", os.path.dirname(os.path.abspath(__file__))))
os.environ.setdefault("DISPLAY", ":99")
import tkinter as tk
from tkinter import ttk
import matplotlib; matplotlib.use("Agg")
from matplotlib.figure import Figure
import sad_common as C, sad_report as R
import docx, openpyxl
FAIL = []
def check(c, msg):
    print(("  [OK]   " if c else "  [FAIL] ") + msg, flush=True)
    if not c: FAIL.append(msg)
TMP = tempfile.mkdtemp(prefix="sad_rep_")
LOG = {"info": [], "ask": [], "err": []}; ASK = {"ans": False}
R.messagebox.showinfo = lambda t="", m="", **k: LOG["info"].append((t, m))
R.messagebox.showerror = lambda t="", m="", **k: LOG["err"].append((t, m))
R.messagebox.askyesno = lambda t="", m="", **k: (LOG["ask"].append((t, m)), ASK["ans"])[1]
SAVE = {"path": None}; R.filedialog.asksaveasfilename = lambda **k: SAVE["path"]
root = tk.Tk(); root.withdraw()

def make_fig(title="Графік", n=5):
    f = Figure(figsize=(6, 3.5), dpi=100); ax = f.add_subplot(111); ax.bar(range(n), [i + 1 for i in range(n)]); ax.set_title(title); return f
def build_window():
    win = tk.Toplevel(root); win.title("Звіт ANOVA – Урожайність"); win.geometry("1000x800")
    tb = tk.Frame(win); tb.pack(fill=tk.X)
    tk.Button(tb, text="📋 Копіювати звіт").pack(side=tk.LEFT); tk.Label(tb, text="ПІДКАЗКА ПАНЕЛІ: не експортувати").pack(side=tk.LEFT); ttk.Combobox(tb, values=["a"]).pack(side=tk.LEFT)
    body = tk.Frame(win); body.pack(fill=tk.BOTH, expand=True)
    tk.Label(body, text="З В І Т   А Н А Л І З У", font=("Times New Roman", 14, "bold")).pack(anchor="w")
    tk.Label(body, text="Показник: урожайність, т/га", font=("Times New Roman", 12)).pack(anchor="w")
    tk.Label(body, text="⚠ Застосовано трансформацію ln(x).", font=("Times New Roman", 12), fg="#c62828").pack(anchor="w")
    tk.Label(body, text="Таблиця ANOVA", font=("Times New Roman", 12, "bold")).pack(anchor="w")
    frm, tv = C.make_tv(body, ["Джерело", "SS", "df", "F", "p"], [["Фактор A", "12,345", "2", "5,67", "0,0042"], ["Залишок", "30,1", "27", "", ""], ["Разом", "42,4", "29", "", "<0,0001"]]); frm.pack(fill=tk.X)
    tv.tag_configure("sig", background="#c8e6c9"); tv.item(tv.get_children()[0], tags=("sig",))
    tk.Label(body, text="Текстовий блок результатів", font=("Times New Roman", 12, "bold")).pack(anchor="w")
    txt = tk.Text(body, height=6); txt.pack(fill=tk.X)
    txt.insert("1.0", "Показник    Середнє    SD\nA1          12,5       1,2\nA2          14,1       1,4\nA3          15,3       1,1\n\nПримітка: різні літери – істотна різниця.")
    var = tk.StringVar(value="Динамічний підпис із змінної"); tk.Label(body, textvariable=var, font=("Times New Roman", 12)).pack(anchor="w")
    tk.Label(body, text="").pack()
    lf = tk.LabelFrame(body, text="Група розділу"); lf.pack(fill=tk.X)
    f2, tv2 = C.make_tv(lf, ["Рівень", "Середнє"], [["1", "0,5"], ["2", "0,7"]]); f2.pack()
    nb = ttk.Notebook(body); nb.pack(fill=tk.BOTH, expand=True)
    t1 = tk.Frame(nb); nb.add(t1, text="Гістограма"); C.embed_figure(make_fig("Гістограма урожайності"), t1)
    t2 = tk.Frame(nb); nb.add(t2, text="Таблиця груп"); f3, _ = C.make_tv(t2, ["Група", "n"], [["A", "10"], ["B", "10"]]); f3.pack()
    t3 = tk.Frame(nb); nb.add(t3, text="Порожня"); tk.Label(t3, text="")
    hidden = tk.Frame(body); hidden.pack(); C.embed_figure(make_fig("Прихований графік"), hidden); hidden.pack_forget()
    dlg = tk.Toplevel(win); tk.Label(dlg, text="ЦЕ ДІАЛОГ – не експортувати").pack()
    skip = tk.Frame(body); skip._sad_noexport = True; tk.Label(skip, text="ПОЗНАЧЕНО ЯК noexport").pack(); skip.pack()
    root.update()
    return win

print("== 1. Розбір чисел")
for s, v in (("0,0123", 0.0123), ("12", 12), ("−1 234,5", -1234.5), ("3.14", 3.14), ("-0,5", -0.5), ("1e-5", 1e-5), ("+7", 7)):
    check(R.parse_number(s) == v and type(R.parse_number(s)) is type(v), f"«{s}» → {v!r}")
for s in ("<0,0001", "–", "", "007", "12%", "A1", "1,2,3", "p<0.05", "nan", None, "0,5 ± 0,1", "1.2.3"):
    check(R.parse_number(s) is None, f"«{s}» лишається текстом")
check(R.parse_number("0") == 0 and R.parse_number("0,0") == 0.0, "нуль розпізнається")

print("\n== 2. Збирання вмісту вікна")
win = build_window(); items = R.collect(win)
kinds = [i["kind"] for i in items]; texts = " | ".join(i.get("text", "") for i in items)
check("ПІДКАЗКА ПАНЕЛІ" not in texts and "ЦЕ ДІАЛОГ" not in texts and "noexport" not in texts, "панель кнопок, діалог-нащадок і позначені _sad_noexport пропущено")
check(not any("Копіювати" in i.get("text", "") for i in items), "тексти кнопок не потрапили у звіт")
hd = [i for i in items if i["kind"] == "heading"]; check(hd[0]["text"] == "ЗВІТ АНАЛІЗУ" and hd[0]["level"] == 2, "жирний великий напис – заголовок 2-го рівня (розряджений текст стиснуто)")
check([h["text"] for h in hd if h["level"] == 3][:2] == ["Таблиця ANOVA", "Текстовий блок результатів"], "жирний звичайний напис – заголовок 3-го рівня")
check(any(i["kind"] == "text" and i["text"] == "Показник: урожайність, т/га" for i in items), "звичайний напис – абзац")
w_ = [i for i in items if i["kind"] == "text" and i["text"].startswith("⚠")]; check(w_ and w_[0]["warn"] is True, "червоний напис позначено як попередження")
check(any(i["kind"] == "text" and i["text"] == "Динамічний підпис із змінної" for i in items), "Label із textvariable: беруться актуальні значення змінної")
tb = [i for i in items if i["kind"] == "table"]; check(len(tb) == 4, f"таблиць зібрано {len(tb)}: ANOVA, із текстового блока, група розділу, вкладка «Таблиця груп»")
a = tb[0]; check(a["headers"] == ["Джерело", "SS", "df", "F", "p"] and a["rows"][0] == ["Фактор A", "12,345", "2", "5,67", "0,0042"] and a["rows"][2][4] == "<0,0001", "таблиця ANOVA: заголовки й значення як на екрані")
check(a["fills"] == ["C8E6C9", None, None], f"заливка рядка з тегом: {a['fills']}")
tt = tb[1]; check(tt["headers"] == ["Показник", "Середнє", "SD"] and tt["rows"] == [["A1", "12,5", "1,2"], ["A2", "14,1", "1,4"], ["A3", "15,3", "1,1"]], "Text зі стовпчиками, розділеними пробілами, розпізнано як таблицю")
check(any(i["kind"] == "text" and i["text"].startswith("Примітка") for i in items), "абзац після таблиці в Text лишається текстом")
figs = [i for i in items if i["kind"] == "figure"]; check(len(figs) == 2, f"графіків 2 (вкладка + прихована панель): {len(figs)}")
check([h["text"] for h in hd if h["level"] == 1] == ["Гістограма", "Таблиця груп"], "вкладки Notebook – заголовки 1-го рівня; порожня вкладка пропущена")
check(any(h["text"] == "Група розділу" for h in hd), "LabelFrame із вмістом – заголовок розділу")
order = [i.get("text") or i["kind"] for i in items if i["kind"] in ("heading", "table", "figure")]
check(order.index("Таблиця ANOVA") < order.index("Текстовий блок результатів") < order.index("Гістограма"), "порядок елементів збігається з порядком у вікні")
check(items == R._tidy(items), "результат «прибраний» (немає дублів і порожніх розділів)")
check(R.collect(tk.Toplevel(root)) == [], "порожнє вікно → порожній звіт")

print("\n== 2б. Бокова панель, підписи «ТАБЛИЦЯ N.», розряджені назви")
w3 = tk.Toplevel(root); side = tk.Frame(w3, bg="#2c3e50"); side.pack(side=tk.LEFT, fill=tk.Y); tk.Label(side, text="ЗВІТ").pack(); tk.Label(side, text="ГРАФІКИ").pack()
for nm in ("📄 Текстовий звіт", "📊 Графіки"):
    fr_ = tk.Frame(side); fr_.pack(); tk.Button(fr_, text=nm).pack(); tk.Label(fr_, text="підказка кнопки").pack()
cont = tk.Frame(w3); cont.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
tk.Label(cont, text="З В І Т   С Т А Т И С Т И Ч Н О Г О   А Н А Л І З У", font=("Times New Roman", 12)).pack()
tk.Label(cont, text="ТАБЛИЦЯ 1. Дисперсійний аналіз (ANOVA)", font=("Times New Roman", 13, "bold")).pack(); fa, _ = C.make_tv(cont, ["Джерело", "SS"], [["A", "1,5"]]); fa.pack()
tk.Label(cont, text="ТАБЛИЦЯ 2. Сила впливу факторів (% від SS)", font=("Times New Roman", 13, "bold")).pack(); fb, _ = C.make_tv(cont, ["Джерело", "%"], [["A", "50"]]); fb.pack()
tk.Label(cont, text="ТАБЛИЦЯ 3. Заголовок без таблиці далі", font=("Times New Roman", 13, "bold")).pack(); tk.Label(cont, text="Просто текст після заголовка").pack(); root.update()
it3 = R.collect(w3); tx3 = " | ".join(i.get("text", "") for i in it3)
check("ГРАФІКИ" not in tx3 and "ЗВІТ" not in tx3.replace("ЗВІТ СТАТИСТИЧНОГО", "") and "підказка кнопки" not in tx3, "бокова панель (підписи, кнопки, підказки кнопок) у звіт не потрапляє")
check(it3[0]["text"] == "ЗВІТ СТАТИСТИЧНОГО АНАЛІЗУ", f"розряджений заголовок стиснуто: «{it3[0]['text']}»")
caps = [i for i in it3 if i["kind"] == "caption"]; check([c["text"] for c in caps] == ["Дисперсійний аналіз (ANOVA)", "Сила впливу факторів (% від SS)"], "заголовки «ТАБЛИЦЯ N. …» перед таблицями стали підписами (без номера)")
check(any(i["kind"] == "heading" and i["text"].startswith("ТАБЛИЦЯ 3") for i in it3), "заголовок «ТАБЛИЦЯ 3. …» без таблиці далі лишається заголовком")
p3 = os.path.join(TMP, "c.docx"); R.export_docx(p3, it3, "Т"); d3 = docx.Document(p3); pp = [p.text for p in d3.paragraphs]
check("Таблиця 1 – Дисперсійний аналіз (ANOVA)" in pp and "Таблиця 2 – Сила впливу факторів (% від SS)" in pp and not any(x.startswith("ТАБЛИЦЯ 1") for x in pp), "у Word: «Таблиця 1 – Дисперсійний аналіз (ANOVA)», без подвійної нумерації")
check(sum(1 for x in pp if x.startswith("Таблиця ")) == 2, "підписів рівно стільки, скільки таблиць")
w3.destroy()
cvw = tk.Toplevel(root); cvx = tk.Canvas(cvw); cvx.pack(); tk.Label(cvx, text="Підпис осі в полотні", font=("Times New Roman", 9, "bold")); lb_ = tk.Label(cvx, text="Підпис осі в полотні", font=("Times New Roman", 9, "bold")); cvx.create_window(5, 5, window=lb_)
bd_ = tk.Frame(cvx); cvx.create_window(5, 40, window=bd_); tk.Label(bd_, text="Текст у прокручуваному тілі", font=("Times New Roman", 11)).pack(); root.update()
tx_cv = " | ".join(i.get("text", "") for i in R.collect(cvw)); check("Підпис осі" not in tx_cv and "Текст у прокручуваному тілі" in tx_cv, "підписи осей, вбудовані прямо в Canvas, пропущено; вміст звичайного прокручуваного тіла (Frame у Canvas) зберігається")
cvw.destroy()
dash = [{"kind": "heading", "text": "Звіт \u2014 результати", "level": 3}, {"kind": "text", "text": "A \u2014 B", "warn": False}, {"kind": "table", "headers": ["x\u2014y"], "rows": [["1 \u2014 2"]], "fills": [None], "widths": [50]}]
check(all("\u2014" not in (i.get("text", "") + "".join(i.get("headers", [])) + "".join(c for r in i.get("rows", []) for c in r)) for i in R._tidy(dash)), "довге тире замінюється коротким у тексті й таблицях")
pdash = os.path.join(TMP, "dash.docx"); R.export_docx(pdash, R._tidy(dash), "Звіт \u2014 назва", "ід \u2014 тест"); dd_ = docx.Document(pdash)
check("\u2014" not in " ".join(p.text for p in dd_.paragraphs) and dd_.paragraphs[0].text == "Звіт \u2013 назва", "довгого тире немає ні в заголовку документа, ні в підзаголовку")
tall = [{"kind": "figure", "fig": Figure(figsize=(7, 14), dpi=100)}]; tall[0]["fig"].add_subplot(111).plot([1, 2, 3]); ptall = os.path.join(TMP, "tall.docx"); R.export_docx(ptall, tall, "Висока"); dtall = docx.Document(ptall)
mh = (dtall.sections[0].page_height - dtall.sections[0].top_margin - dtall.sections[0].bottom_margin) / 360000.0 - 4.0
check(dtall.inline_shapes[0].height / 360000.0 <= mh + 0.01, f"надто високий рисунок зменшується до висоти сторінки ({dtall.inline_shapes[0].height / 360000.0:.1f} ≤ {mh:.1f} см), щоб не з'являлась порожня сторінка")
wfig = [{"kind": "table", "headers": [f"К{i}" for i in range(12)], "rows": [["1"] * 12], "fills": [None], "widths": [60] * 12}, {"kind": "figure", "fig": Figure(figsize=(9, 8), dpi=100)}]; wfig[1]["fig"].add_subplot(111).plot([1, 2, 3]); pwf = os.path.join(TMP, "wf.docx"); R.export_docx(pwf, wfig, "Альбомна"); dwf = docx.Document(pwf)
mh2 = (dwf.sections[0].page_height - dwf.sections[0].top_margin - dwf.sections[0].bottom_margin) / 360000.0 - 4.0
check(dwf.sections[0].orientation == 1 and dwf.inline_shapes[0].height / 360000.0 <= mh2 + 0.01 and mh2 < 14, f"альбомна орієнтація: рисунок вміщується ({dwf.inline_shapes[0].height / 360000.0:.1f} ≤ {mh2:.1f} см)")
print("\n== 3. Word")
docx_path = os.path.join(TMP, "r.docx"); st = R.export_docx(docx_path, items, "Звіт ANOVA – Урожайність", "тест")
d = docx.Document(docx_path)
check(st == {"tables": 4, "figures": 2, "max_cols": 5} and len(d.tables) == 4 and len(d.inline_shapes) == 2, f"у Word 4 таблиці і 2 рисунки: {len(d.tables)}, {len(d.inline_shapes)}")
check(all(t.style.name == "Table Grid" for t in d.tables), "усі таблиці зі стилем «Table Grid» – з сіткою")
t0 = d.tables[0]; check([c.text for c in t0.rows[0].cells] == ["Джерело", "SS", "df", "F", "p"] and [c.text for c in t0.rows[1].cells] == ["Фактор A", "12,345", "2", "5,67", "0,0042"], "вміст таблиці ANOVA у Word збігається з екраном")
xml = t0._tbl.xml; check('w:fill="D9E2F3"' in xml and xml.count("w:tblHeader") == 1 and "w:cantSplit" in xml, "заголовок таблиці зафарбовано, повторюється на сторінках, рядки не розриваються")
check('w:fill="C8E6C9"' in t0.rows[1]._tr.xml and 'w:fill="C8E6C9"' not in t0.rows[2]._tr.xml, "заливка рядка з тегом перенесена в Word лише на цей рядок")
paras = [p.text for p in d.paragraphs]; check("Таблиця 1 – Таблиця ANOVA" in paras and "Рисунок 1 – Гістограма урожайності" in paras and "Рисунок 2 – Прихований графік" in paras, "підписи «Таблиця N – …» і «Рисунок N – …» (з назви графіка)")
check(paras[0] == "Звіт ANOVA – Урожайність" and "тест" in paras[1], "заголовок і підзаголовок документа")
check(any(p.style.name.startswith("Heading") and "Таблиця ANOVA" == p.text for p in d.paragraphs), "заголовки вікна стали заголовками Word")
check(all((r.font.name or p.style.font.name) == "Times New Roman" for p in d.paragraphs if p.runs for r in p.runs if r.text.strip()) and d.styles["Normal"].font.name == "Times New Roman", "ефективний шрифт усіх текстових фрагментів – Times New Roman (явний або зі стилю)")
check(all(d.styles[f"Heading {i}"].font.name == "Times New Roman" and str(d.styles[f"Heading {i}"].font.color.rgb) == "000000" for i in (1, 2, 3)), "заголовки: Times New Roman, чорні (а не синій Calibri за замовчуванням)")
warn_runs = [r for p in d.paragraphs if p.text.startswith("⚠") for r in p.runs]; check(warn_runs and str(warn_runs[0].font.color.rgb) == "C62828", "попередження зберегло червоний колір")
check("Times New Roman" in d.sections[0].footer.paragraphs[0].text or "S.A.D." in d.sections[0].footer.paragraphs[0].text, "у колонтитулі позначено програму й час")
with zipfile.ZipFile(docx_path) as z: imgs = [n for n in z.namelist() if n.startswith("word/media/")]
check(len(imgs) == 2 and all(len(z_) > 3000 for z_ in [zipfile.ZipFile(docx_path).read(n) for n in imgs]), "рисунки – справжні PNG у файлі")
check(d.sections[0].orientation == 0, "до 9 стовпчиків – книжкова орієнтація")
wide_items = [{"kind": "table", "headers": [f"К{i}" for i in range(12)], "rows": [[str(i * j) for i in range(12)] for j in range(5)], "fills": [None] * 5, "widths": [60] * 12}]
wp = os.path.join(TMP, "w.docx"); R.export_docx(wp, wide_items, "Широка"); dw = docx.Document(wp)
check(dw.sections[0].orientation == 1 and dw.sections[0].page_width > dw.sections[0].page_height, "понад 9 стовпчиків – альбомна орієнтація")
tot_w = sum(c.width for c in dw.tables[0].rows[0].cells); avail = dw.sections[0].page_width - dw.sections[0].left_margin - dw.sections[0].right_margin
check(tot_w <= avail * 1.01, f"ширина таблиці вміщується на сторінку ({tot_w / 360000:.1f} ≤ {avail / 360000:.1f} см)")
many = [{"kind": "table", "headers": ["A", "B"], "rows": [[str(i), "x"] for i in range(200)], "fills": [None] * 200, "widths": [50, 50]}]
mp = os.path.join(TMP, "m.docx"); R.export_docx(mp, many, "Довга"); check(len(docx.Document(mp).tables[0].rows) == 201 and docx.Document(mp).tables[0]._tbl.xml.count("w:tblHeader") == 1, "довга таблиця (200 рядків): один повторюваний заголовок")
bad = [{"kind": "figure", "fig": "не фігура"}]; bp = os.path.join(TMP, "b.docx"); R.export_docx(bp, bad, "x"); check("не вдалося побудувати" in " ".join(p.text for p in docx.Document(bp).paragraphs), "зіпсований графік не ламає експорт – у документі примітка")

print("\n== 4. Excel")
xp = os.path.join(TMP, "r.xlsx"); R.export_xlsx(xp, items, "Звіт ANOVA – Урожайність", "тест"); wb = openpyxl.load_workbook(xp)
check(wb.sheetnames[0] == "Звіт" and "Графіки" in wb.sheetnames and len(wb.sheetnames) == 1 + 4 + 1, f"аркуші: {wb.sheetnames}")
sh = wb[wb.sheetnames[1]]; check([c.value for c in sh[1]] == ["Джерело", "SS", "df", "F", "p"] and sh["B2"].value == 12.345 and isinstance(sh["C2"].value, int) and sh["E4"].value == "<0,0001" and sh["E2"].value == 0.0042, "числа – числами (12,345 → 12.345, df – ціле), «<0,0001» – текстом")
check(sh["A2"].fill.fgColor.rgb[-6:] == "C8E6C9" and sh["A3"].fill.fgColor.rgb[-6:] != "C8E6C9" and sh["A1"].font.bold, "заливка рядка збережена, заголовок жирний")
check(sh.freeze_panes == "A2", "заголовок закріплено")
check(len({n for n in wb.sheetnames}) == len(wb.sheetnames) and all(len(n) <= 31 for n in wb.sheetnames), "назви аркушів унікальні й ≤ 31 символу")
wf = wb["Графіки"]; check(len(wf._images) == 2, f"аркуш «Графіки»: {len(wf._images)} зображення")
ws = wb["Звіт"]; allv = " ".join(str(c.value) for r in ws.iter_rows() for c in r if c.value); check("Показник: урожайність, т/га" in allv and "Таблиця 1 – Таблиця ANOVA" in allv and "аркуш" in allv, "аркуш «Звіт»: текст і посилання на аркуші таблиць")
long_h = [{"kind": "heading", "text": "Дуже довга назва розділу, що перевищує ліміт Excel для назви аркуша", "level": 3}] + [{"kind": "table", "headers": ["a"], "rows": [["1"]], "fills": [None], "widths": [50]}] * 3
lp = os.path.join(TMP, "l.xlsx"); R.export_xlsx(lp, long_h, "x"); lw = openpyxl.load_workbook(lp); check(len(lw.sheetnames) == 4 and len(set(lw.sheetnames)) == 4 and all(len(n) <= 31 for n in lw.sheetnames), "довгі й однакові назви розділів: аркуші унікальні й коротші за 31 символ")

print("\n== 5. Кнопки експорту у вікні")
win2 = build_window(); n0 = len(win2.pack_slaves()); bar = R.attach_export_bar(win2, "Звіт ANOVA – Урожайність")
check(bar is not None and win2.pack_slaves()[0] is bar and len(win2.pack_slaves()) == n0 + 1, "смужка експорту стоїть угорі вікна (перед усім іншим)")
check(R.attach_export_bar(win2) is None and len(win2.pack_slaves()) == n0 + 1, "повторний виклик не дублює смужку")
bt = [c.cget("text") for c in bar.winfo_children() if c.winfo_class() == "Button"]; check(bt == ["📄 Word", "📊 Excel"], f"кнопки: {bt}")
check(not any("Word" in i.get("text", "") or "Звіт повністю" in i.get("text", "") for i in R.collect(win2)), "сама смужка експорту у звіт не потрапляє")
SAVE["path"] = os.path.join(TMP, "btn.docx"); ASK["ans"] = False; [c for c in bar.winfo_children() if c.winfo_class() == "Button"][0].invoke()
check(os.path.exists(SAVE["path"]) and len(docx.Document(SAVE["path"]).tables) == 4 and LOG["ask"] and "Таблиць: 4, графіків: 2" in LOG["ask"][-1][1], "кнопка «Word»: файл створено, повідомлення «Таблиць: 4, графіків: 2»")
SAVE["path"] = os.path.join(TMP, "btn.xlsx"); [c for c in bar.winfo_children() if c.winfo_class() == "Button"][1].invoke()
check(os.path.exists(SAVE["path"]) and len(openpyxl.load_workbook(SAVE["path"]).sheetnames) == 6, "кнопка «Excel»: файл створено")
SAVE["path"] = None; n_ask = len(LOG["ask"]); [c for c in bar.winfo_children() if c.winfo_class() == "Button"][0].invoke(); check(len(LOG["ask"]) == n_ask, "скасування вибору файла нічого не створює")
empty = tk.Toplevel(root); n_i = len(LOG["info"]); R.export_window(empty, "docx"); check(len(LOG["info"]) == n_i + 1 and "немає результатів" in LOG["info"][-1][1], "порожнє вікно: пояснення замість порожнього файла")
gw = tk.Toplevel(root); tk.Label(gw, text="grid").grid(row=0, column=0); R.attach_export_bar(gw); check(gw._sad_export_bar.winfo_manager() == "place", "вікно на grid: смужка розміщується без конфлікту менеджерів геометрії")
check(R._safe_name("Звіт ANOVA: Урожайність, т/га") == "Звіт ANOVA_ Урожайність_ т_га" or "/" not in R._safe_name("Звіт ANOVA: Урожайність, т/га"), "назва файла очищена від недопустимих символів")
shutil.rmtree(TMP, ignore_errors=True)
print("\n" + "=" * 60)
if FAIL: print(f"ПРОВАЛІВ: {len(FAIL)}"); [print("  -", f) for f in FAIL]; sys.exit(1)
print("УСІ ПЕРЕВІРКИ ЕКСПОРТУ ЗВІТУ ПРОЙДЕНО")
