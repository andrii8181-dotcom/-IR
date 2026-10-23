# test_sad_mixed_sections.py — розділи змішаної моделі, крос-перевірка з класичним аналізом, експорт
import os, sys, math, tempfile, shutil, time
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import numpy as np
import sad_gxe_core as GX, sad_breeding_core as BC, sad_breeding_ui as U, sad_mixed_sections as MS
from sad_gxe_core import BreedingError
FAIL = []
def check(c, msg):
    print(("  [OK]   " if c else "  [FAIL] ") + msg, flush=True)
    if not c: FAIL.append(msg)
TMP = tempfile.mkdtemp(prefix="sadmx_")
def structure_ok(S, label):
    bad = []
    for s in S:
        if s["kind"] != "table": continue
        h = len(s["headers"])
        if len(s["nd"]) != h: bad.append(s["title"])
        if any(len(r) != h for r in s["rows"]): bad.append(s["title"] + " (рядок)")
        if s.get("tags") is not None and len(s["tags"]) != len(s["rows"]): bad.append(s["title"] + " (теги)")
    check(not bad, f"{label}: структура таблиць коректна {bad[:2]}")
def figs_ok(S, label):
    from matplotlib.figure import Figure
    import matplotlib
    msgs = []
    for s in S:
        if s["kind"] == "figure":
            f = Figure(figsize=s["size"])
            with matplotlib.rc_context(U._RC): m = s["build"](f)
            if m: msgs.append(m)
    check(not msgs, f"{label}: графіки будуються {msgs[:1]}")
def fmt(v): return f"{v:.3f}".replace(".", ",")

print("== 1. Збалансований RCBD: спадковість Каллісу = класична h² на рівні середніх")
rg = np.random.default_rng(3); g, r = 14, 4; ga = rg.normal(0, 2, g); rb = rg.normal(0, 1, r)
rows, gen, rep, y = [], [], [], []
for i in range(g):
    for k in range(r):
        v = 50 + ga[i] + rb[k] + rg.normal(0, 1.6)
        rows.append([f"G{i+1}", "", f"R{k+1}", "", fmt(v), "", ""]); gen.append(f"G{i+1}"); rep.append(f"R{k+1}"); y.append(v)
S, warns, res = MS.sections_mixed(rows, 4, "Урожай", mode="single", geno_random=True, rep_random=False)
cl = BC.genetic_parameters_single(gen, rep, [float(fmt(v).replace(",", ".")) for v in y])
h2 = next(s for s in S if s["key"] == "h2")["rows"][0][1]
check(abs(h2 - cl["h2_mean"]) < 2e-3, f"H² Каллісу {h2:.4f} = класична H² на рівні середніх {cl['h2_mean']:.4f} (розбіжність {abs(h2 - cl['h2_mean']):.1e})")
vr = next(s for s in S if s["key"] == "var")
check(abs(vr["rows"][0][1] - cl["var_g"]) < 2e-3 and abs(vr["rows"][-1][1] - cl["var_e"]) < 2e-3, "σ²g й σ²e збігаються з класичним ANOVA-розрахунком")
structure_ok(S, "RCBD"); figs_ok(S, "RCBD")
check({"var", "h2", "blup", "cat", "fixed", "help"} <= {s["key"] for s in S}, "очікувані розділи є")
bl = next(s for s in S if s["key"] == "blup")
check(len(bl["rows"]) == g and bl["rows"][0][0] == 1 and bl["tags"].count("good") == max(1, round(0.1 * g)), "BLUP: рядок на генотип, відсортовано за місцем, кращі 10 % виділено")
check(vr["rows"][0][5][1] in ("*", "**"), "LRT: генотипний ефект значущий")

print("\n== 2. Серія середовищ, незбалансована")
rg = np.random.default_rng(8); g, e, r = 30, 4, 2
ga = rg.normal(0, 1.5, g); ea = rg.normal(0, 2, e); ge = rg.normal(0, 1, (g, e))
rows = []
for i in range(g):
    for j in range(e):
        if rg.random() < 0.2: continue
        for k in range(r): rows.append([f"G{i+1}", f"E{j+1}", f"R{k+1}", "", fmt(30 + ga[i] + ea[j] + ge[i, j] + rg.normal(0, 1.1)), "", ""])
t0 = time.time(); S2, w2, res2 = MS.sections_mixed(rows, 4, "Урожай", mode="met", geno_random=True)
check({"var", "h2", "blup", "cat", "pred", "fixed", "help"} <= {s["key"] for s in S2}, f"розділи для серії середовищ (за {time.time() - t0:.1f} с)")
structure_ok(S2, "серія"); figs_ok(S2, "серія")
pr = next(s for s in S2 if s["key"] == "pred")
check(len(pr["rows"]) == g and len(pr["headers"]) == 1 + e, f"прогноз у середовищах: {len(pr['rows'])} × {len(pr['headers']) - 1}")
vr2 = next(s for s in S2 if s["key"] == "var")
check([vr2["rows"][i][0] for i in range(2)] == ["Генотип (σ²g)", "Генотип × середовище (σ²ge)"], "компоненти: генотип, G×E, …, залишок")
lrt_p = {r_[0]: r_[5] for r_ in vr2["rows"] if isinstance(r_[5], tuple)}
check(all(v[0] < 0.01 for v in lrt_p.values()), f"усі закладені ефекти значущі за LRT: { {k: round(v[0], 6) for k, v in lrt_p.items()} }")
em = MS.env_means(res2); check(len(em) == e, f"середні середовищ відновлено: {len(em)}")
# прогноз для наявної комбінації = середнє середовища + BLUP G + BLUP GE
ge_u = res2["u"]["ge"]; bt = MS.MX.blup_table(res2, "gen"); i0 = 0; g0 = bt["levels"][i0]; e0 = list(em)[1]
k = ge_u["levels"].index(f"{g0}|{e0}") if f"{g0}|{e0}" in ge_u["levels"] else None
exp = em[e0] + bt["blup"][i0] + (ge_u["blup"][k] if k is not None else 0)
row = next(r_ for r_ in pr["rows"] if r_[0] == g0)
check(abs(row[1 + 1] - exp) < 1e-9, "прогноз = середнє середовища + BLUP генотипу + BLUP взаємодії")

print("\n== 3. Фіксований генотип і неповні блоки")
rg = np.random.default_rng(12); g, nb = 10, 6; gt = rg.normal(0, 1.5, g); bt_ = rg.normal(0, 1.0, nb); rows = []
for i in range(g):
    for j in rg.choice(nb, 3, replace=False): rows.append([f"G{i+1}", "", "R1", f"B{j+1}", fmt(20 + gt[i] + bt_[j] + rg.normal(0, 0.5)), "", ""])
S3, w3, res3 = MS.sections_mixed(rows, 4, "Маса", mode="single", geno_random=False, use_blocks=True)
check({"var", "lsm", "fixed", "help"} <= {s["key"] for s in S3} and "blup" not in {s["key"] for s in S3}, "фіксований генотип → скориговані середні замість BLUP")
lsm = next(s for s in S3 if s["key"] == "lsm")
check(len(lsm["rows"]) == g and lsm["rows"][0][2] >= lsm["rows"][-1][2], "скориговані середні відсортовано за спаданням")
structure_ok(S3, "фіксований генотип")
check(any(r_[0].startswith("Блок") for r_ in next(s for s in S3 if s["key"] == "var")["rows"]), "блок у повторності враховано як випадковий ефект")

print("\n== 4. Попередження й помилки")
few = [[f"G{i % 4}", "", f"R{i // 4}", "", fmt(10 + (i % 4) + 0.3 * np.sin(i)), "", ""] for i in range(16)]
S6, w6, _ = MS.sections_mixed(few, 4, "Х", mode="single", geno_random=True)
check(any("Лише 4 генотипів" in w for w in w6), "мало генотипів → попередження")
nul = [[f"G{i % 12}", "", f"R{i // 12}", "", fmt(rg.normal(50, 1)), "", ""] for i in range(48)]
S7, w7, r7 = MS.sections_mixed(nul, 4, "Шум", mode="single", geno_random=True, rep_random=False)
check("gen" in r7["boundary"] and any("практично дорівнює нулю" in w for w in w7), "чистий шум: σ²g на межі, є попередження")
for bad, frag in [(([r_[:1] + [""] + r_[2:] for r_ in rows], "met"), "середовища"), (([[""] + r_[1:] for r_ in rows], "single"), "назви генотипу")]:
    try: MS.sections_mixed(bad[0], 4, "Х", mode=bad[1]); check(False, f"{frag}: мало впасти")
    except BreedingError as ex: check(frag in str(ex), f"помилка: {frag}")
try: MS.sections_mixed([r_[:4] + ["abc"] + r_[5:] for r_ in rows], 4, "Х", mode="single"); check(False, "нечислове")
except BreedingError as ex: check("не число" in str(ex), "нечислове значення → помилка з номером рядка")

print("\n== 5. Експорт")
import openpyxl, docx
px = os.path.join(TMP, "m.xlsx"); U.export_xlsx(px, S2, "Змішана модель"); wb = openpyxl.load_workbook(px)
check(len(wb.sheetnames) >= 5, f"Excel: {wb.sheetnames[:4]}…")
pd_ = os.path.join(TMP, "m.docx"); U.export_docx(pd_, "Змішана модель", S2, True); d = docx.Document(pd_)
check(len(d.tables) >= 5 and len(d.inline_shapes) >= 1, f"Word: таблиць {len(d.tables)}, рисунків {len(d.inline_shapes)}")
shutil.rmtree(TMP, ignore_errors=True)
print("\n" + "=" * 60)
if FAIL:
    print(f"ПРОВАЛІВ: {len(FAIL)}"); [print("  -", f) for f in FAIL]; sys.exit(1)
print("УСІ ПЕРЕВІРКИ РОЗДІЛІВ ЗМІШАНОЇ МОДЕЛІ ПРОЙДЕНО")
