# test_sad_breeding_sections.py — побудова розділів результатів, експорт у Excel/Word (без вікон)
import os, sys, math, tempfile, shutil
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import numpy as np
import sad_gxe_core as GX, sad_breeding_core as BC, sad_breeding_ui as U, sad_breeding as BR

FAIL = []
def check(c, msg):
    print(("  [OK]   " if c else "  [FAIL] ") + msg)
    if not c: FAIL.append(msg)
rng = np.random.default_rng(21)
TMP = tempfile.mkdtemp(prefix="sadbr_")

def structure_ok(sections, label):
    """Кожен рядок таблиці має стільки комірок, скільки заголовків і форматів; теги — по числу рядків."""
    bad = []
    for s in sections:
        if s["kind"] != "table": continue
        h = len(s["headers"])
        if len(s["nd"]) != h: bad.append(f"{s['title']}: nd {len(s['nd'])} ≠ {h}")
        for i, r in enumerate(s["rows"]):
            if len(r) != h: bad.append(f"{s['title']}: рядок {i} має {len(r)} ≠ {h}"); break
        if s.get("tags") is not None and len(s["tags"]) != len(s["rows"]): bad.append(f"{s['title']}: теги")
    check(not bad, f"{label}: структура всіх таблиць коректна" + (f" — {bad[:2]}" if bad else ""))

def figs_ok(sections, label):
    from matplotlib.figure import Figure
    import matplotlib
    msgs = []
    for s in sections:
        if s["kind"] != "figure": continue
        fig = Figure(figsize=s["size"])
        with matplotlib.rc_context(U._RC): m = s["build"](fig)
        if m: msgs.append(f"{s['title']}: {m}")
    check(not msgs, f"{label}: усі графіки будуються" + (f" — {msgs}" if msgs else ""))

print("== 1. Стабільність")
g, e, r = 8, 6, 3
gn = [f"Сорт {i+1}" for i in range(g)]; en = [f"{2019+j}" for j in range(e)]
X = 6 + rng.normal(0, 1, (g, 1)) + rng.normal(0, 1.5, (1, e)) + np.outer(rng.normal(0, .7, g), rng.normal(0, 1, e)) + rng.normal(0, .3, (g, e))
S, warns = BR.build_stability_sections(X, gn, en)
keys = [s["key"] for s in S]
check(keys[:2] == ["warn", "anova"] and {"param", "nonpar", "rankcorr", "summary", "ammi_axes", "ammi_gen", "ammi_env",
      "ammi1", "ammi2", "www", "www_tab", "aec", "aec_tab"} <= set(keys), f"усі очікувані розділи є ({len(S)})")
structure_ok(S, "стабільність"); figs_ok(S, "стабільність")
check(any("похибка за повторностями" in w.lower() or "похибка" in w.lower() for w in warns), "без повторностей — попередження про тести")
par = next(s for s in S if s["key"] == "param")
check(len(par["rows"]) == g and par["headers"][0] == "Генотип", "параметричні індекси: рядок на генотип")
rc = next(s for s in S if s["key"] == "rankcorr")
check(any("σ² Шукли" in n and "W²" in n for n in rc["notes"]), "кореляція рангів повідомляє, що Wricke і Shukla дублюють один одного")
sm = next(s for s in S if s["key"] == "summary")
check(set(sm["tags"]) <= {"good", "warn", "info", "bad"} and len(sm["tags"]) == g, "зведення класифікує кожен генотип")
# повні дані з повторностями → ANOVA і тести
gl, el, rl, yl = [], [], [], []
ga, ea, ge = rng.normal(0, 1, 7), rng.normal(0, 2, 5), rng.normal(0, .6, (7, 5))
for i in range(7):
    for j in range(5):
        for k in range(3):
            gl.append(f"S{i}"); el.append(f"E{j}"); rl.append(f"R{k}"); yl.append(10 + ga[i] + ea[j] + ge[i, j] + rng.normal(0, .5))
ca = GX.combined_anova(gl, el, rl, yl)
S2, w2 = BR.build_stability_sections(ca["means"], ca["gen_names"], ca["env_names"], ca["r"], ca["mse"], ca["df_err"], ca=ca)
an = next(s for s in S2 if s["key"] == "anova"); par2 = next(s for s in S2 if s["key"] == "param")
check(an["title"] == "ANOVA (за повторностями)" and len(an["rows"]) == 6, "ANOVA за повторностями: 5 джерел + загальна")
ip = par2["headers"].index("p(S²d=0)")
check(all(r[ip] == r[ip] for r in par2["rows"]), "з похибкою тест S²d виконано для всіх генотипів")
ax = next(s for s in S2 if s["key"] == "ammi_axes")
check(not np.isnan(ax["rows"][0][6]), "з похибкою F-тест осей AMMI виконано")
structure_ok(S2, "стабільність з повторностями")
Xn = X.copy(); Xn[1, 2] = np.nan
try: BR.build_stability_sections(Xn, gn, en); check(False, "NaN мав дати помилку")
except BC.BreedingError as ex: check("Сорт 2 × 2021" in str(ex), "пропуск у таблиці → зрозуміла помилка зі списком клітинок")

print("\n== 2. Генетичні параметри")
rows = []
ga = rng.normal(0, 2, 10)
for i in range(10):
    for k in range(3):
        rows.append([f"G{i}", "", f"R{k}", f"{50 + ga[i] + rng.normal(0, 1):.2f}".replace(".", ","), f"{20 + .5 * ga[i] + rng.normal(0, 1.5):.2f}", ""])
hdr = ["Генотип", "Середовище", "Повторність", "Висота", "Маса", "Порожня"]
Sg, bg = BR.sections_genpar(rows, hdr, False)
tab = Sg[0]
check(len(tab["rows"]) == 2 and [r[0] for r in tab["rows"]] == ["Висота", "Маса"], "порожній стовпець ознаки пропущено; 2 ознаки у зведенні")
check({s["key"] for s in Sg} >= {"genpar", "anova_Висота", "anova_Маса", "genpar_fig", "help"}, "ANOVA по кожній ознаці, графік, пояснення")
structure_ok(Sg, "генетичні параметри (1 середовище)"); figs_ok(Sg, "генетичні параметри")
check(any("середовищ" in w and "кількох" in w for w in bg), "є застереження про одне середовище")
rows_m = []
ga = rng.normal(0, 1.5, 12); ea = rng.normal(0, 2, 4)
for i in range(12):
    for j in range(4):
        for k in range(3):
            rows_m.append([f"G{i}", f"E{j}", f"R{k}", f"{30 + ga[i] + ea[j] + rng.normal(0, 1):.2f}", "", ""])
Sm, bm = BR.sections_genpar(rows_m, hdr, True, "10 %")
check(Sm[0]["rows"][0][-2] in ("низький", "середній", "високий") and Sm[0]["headers"][-2] == "Категорія H²", "кілька середовищ: таблиця з категоріями")
structure_ok(Sm, "генетичні параметри (кілька середовищ)")
check("10 %" in " ".join(Sm[0]["notes"]) and "1.76" in " ".join(Sm[0]["notes"]), "інтенсивність добору 10 % (k = 1,76) врахована")
try: BR.sections_genpar([r[:3] + ["a", "", ""] for r in rows], hdr, False); check(False, "нечислове мало дати помилку")
except BC.BreedingError as ex: check("не число" in str(ex), "нечислове значення → помилка з номером рядка")

print("\n== 3. Лінія × тестер та діалель")
def lt_rows(l=6, t=3, r=3):
    out = []; a = rng.normal(0, 1, l); b = rng.normal(0, .8, t)
    for i in range(l):
        for j in range(t):
            for k in range(r):
                out.append([f"L{i+1}", f"T{j+1}", f"R{k+1}", f"{15 + a[i] + b[j] + rng.normal(0, .4):.3f}", "", ""])
    return out
Sl, wl = BR.sections_line_tester(lt_rows(), 3, "Урожай")
check([s["key"] for s in Sl][:6] == ["anova", "gca_l", "gca_t", "sca", "var", "fig_gca"], "розділи лінія × тестер у правильному порядку")
structure_ok(Sl, "лінія × тестер"); figs_ok(Sl, "лінія × тестер")
anv = Sl[0]["rows"]; check(anv[-1][0] == "Загальна" and anv[-1][1] == 3 - 1 + 18 - 1 + (3 - 1) * (18 - 1), "загальні df = r·l·t − 1")
def dl_rows(p=5, r=3, method=2):
    out = []; g = rng.normal(0, 1, p)
    for i in range(p):
        for j in range(i, p):
            if method == 4 and i == j: continue
            for k in range(r):
                out.append([f"P{i+1}", f"P{j+1}", f"R{k+1}", f"{10 + g[i] + g[j] + rng.normal(0, .4):.3f}", "", ""])
    return out
for method, model in ((2, "I"), (2, "II"), (4, "II")):
    Sd, wd = BR.sections_diallel(dl_rows(5 if method == 2 else 6, 3, method), 3, "Урожай", method, model)
    structure_ok(Sd, f"діалель метод {method}, модель {model}"); figs_ok(Sd, f"діалель метод {method}")
check(Sd[0]["title"].startswith("ANOVA: Урожай (метод 4, модель II)"), "заголовок ANOVA містить метод і модель")
try: BR.sections_diallel(dl_rows(5, 3, 4), 3, "Урожай", 2, "I"); check(False, "метод 2 без батьків")
except BC.BreedingError as ex: check("батьківських форм" in str(ex), "метод 2 без батьків → зрозуміла помилка")

print("\n== 4. Стрес і шляховий аналіз")
rows_s = [[f"G{i+1}", f"{yp:.2f}", f"{ys:.2f}"] for i, (yp, ys) in enumerate(zip(rng.uniform(6, 12, 14), rng.uniform(3, 8, 14)))]
Ss, ws_ = BR.sections_stress(rows_s)
structure_ok(Ss, "стрес"); figs_ok(Ss, "стрес")
cors = next(s for s in Ss if s["key"] == "cors")
check(cors["tags"][0] == "good" and all(cors["rows"][i][3] >= cors["rows"][i + 1][3] for i in range(len(cors["rows"]) - 1)), "кореляції відсортовано за критерієм Фернандеса")
rows_p = [[f"{v:.3f}" for v in row] for row in np.column_stack([rng.normal(size=(40, 3)), rng.normal(size=40)])]
Sp, wp = BR.sections_path(rows_p, ["X1", "X2", "X3", "Y"])
structure_ok(Sp, "шляховий аналіз"); figs_ok(Sp, "шляховий аналіз")
eff = Sp[0]; k = 3
check(eff["headers"][6:9] == ["через X1", "через X2", "через X3"] and np.isnan(eff["rows"][0][6]), "матриця непрямих ефектів: діагональ порожня")

print("\n== 4б. Порівняння зі стандартом")
rows_c = []
eff = {"Стандарт": 0.0, "Сильний+": 4.0, "Сильний−": -4.0, "Нейтральний": 0.0, "Слабкий+": 0.8, "N6": 0.0, "N7": 0.0}
blk = rng.normal(0, 1, 4)
for nm, e_ in eff.items():
    for k in range(4): rows_c.append([nm, f"R{k+1}", f"{30 + e_ + blk[k] + rng.normal(0, 1):.2f}".replace(".", ","), "", ""])
Sc, wc = BR.sections_control(rows_c, 2, "Урожайність", ["Стандарт"])
structure_ok(Sc, "порівняння зі стандартом"); figs_ok(Sc, "порівняння зі стандартом")
tab = Sc[0]; vi = tab["headers"].index("Висновок")
ver = {r[0]: r[vi] for r in tab["rows"]}
check(ver["Сильний+"] == "достовірно вище стандарту" and ver["Сильний−"] == "достовірно нижче стандарту", "закладені великі ефекти виявлено в обидва боки")
check(ver["Нейтральний"] == "не відрізняється" and ver["N6"] == "не відрізняється", "номери без ефекту не позначено істотними")
check(tab["tags"][[r[0] for r in tab["rows"]].index("Сильний+")] == "good" and tab["tags"][[r[0] for r in tab["rows"]].index("Сильний−")] == "bad", "кольори: зелене — вище, червоне — нижче")
ip = tab["headers"].index("% до стандарту"); pc = {r[0]: r[ip] for r in tab["rows"]}
check(pc["Сильний+"] > 110 and pc["Сильний−"] < 90, f"% до стандарту: {pc['Сильний+']:.0f} % і {pc['Сильний−']:.0f} %")
check(Sc[1]["rows"][0][0] == 1 and Sc[1]["rows"][0][1] == "Сильний+", "рейтинг очолює номер із найбільшим % до стандарту")
Sg2, _ = BR.sections_control(rows_c, 2, "Урожайність", ["Стандарт"], alternative="greater")
check(Sg2[0]["rows"][[r[0] for r in Sg2[0]["rows"]].index("Сильний−")][vi] == "не відрізняється", "односторонній «більше»: нижчий номер не позначається")
for bad_args, frag in [((rows_c, 2, "Урожайність", ["Нема"]), "немає серед"), (([r[:2] + ["abc", "", ""] for r in rows_c], 2, "Урожайність", ["Стандарт"]), "не число")]:
    try: BR.sections_control(*bad_args); check(False, "мало впасти")
    except BC.BreedingError as ex: check(frag in str(ex), f"помилка «{frag}»")
print("\n== 5. Експорт Excel/Word (читаємо назад)")
import openpyxl, docx
allS = S2 + Sg + Sl
px = os.path.join(TMP, "r.xlsx"); U.export_xlsx(px, allS, "Тест")
wb = openpyxl.load_workbook(px)
check(len(wb.sheetnames) >= 8 and "Примітки" in wb.sheetnames, f"Excel: {len(wb.sheetnames)} аркушів + примітки")
check(all(len(n) <= 31 for n in wb.sheetnames) and len(set(wb.sheetnames)) == len(wb.sheetnames), "імена аркушів ≤ 31 символ і унікальні")
ws_ = next(w for w in wb.worksheets if w.title.startswith("Параметричні"))
check(isinstance(ws_["B2"].value, float) and ws_.freeze_panes == "A2", "числа лишаються числами, заголовок закріплено")
check(U.fcell((1.2346, "**"), 3, True) == "1,235**" and U.fcell((1.2346, "*"), 3, False) == "1.235*",
      "значущість поруч із числом: (1,2346; «**») → «1,235**»; без коми → «1.235*»")
check(U.fcell(float("nan"), 2, True) == "—" and U.fcell(None, 2, True) == "—" and U.fcell("текст", 2, True) == "текст" and U.fcell(5, None, True) == "5",
      "порожнє → «—»; рядок і ціле — без змін")
# SCA з очевидними сильними ефектами: у Excel мають з'явитись позначки значущості
big = []
sc_true = {("L1", "T1"): 3.0, ("L1", "T2"): -3.0, ("L2", "T1"): -3.0, ("L2", "T2"): 3.0, ("L3", "T1"): 0.0, ("L3", "T2"): 0.0}
for (l_, t_), s_ in sc_true.items():
    for k in range(3): big.append([l_, t_, f"R{k+1}", f"{10 + s_ + rng.normal(0, .2):.3f}", "", ""])
Sb, _ = BR.sections_line_tester(big, 3, "Тест")
pb = os.path.join(TMP, "big.xlsx"); U.export_xlsx(pb, Sb, "Тест")
wsb = next(w for w in openpyxl.load_workbook(pb).worksheets if w.title.startswith("SCA"))
stars = [c.value for row in wsb.iter_rows(min_row=2, min_col=2, max_col=3) for c in row if isinstance(c.value, str) and c.value.endswith("**")]
check(len(stars) >= 4, f"Excel: сильні SCA позначено «**» ({len(stars)} комірок)")
pdx = os.path.join(TMP, "r.docx"); U.export_docx(pdx, "Тест", S2 + Sl, True)
d = docx.Document(pdx)
check(len(d.tables) >= 10 and len(d.inline_shapes) >= 6, f"Word: таблиць {len(d.tables)}, рисунків {len(d.inline_shapes)}")
txt = [p.text for p in d.paragraphs]
check(any(t.startswith("Таблиця 1 —") for t in txt) and any(t.startswith("Рис. 1 —") for t in txt), "нумерація таблиць і рисунків")
check(any("," in c.text for c in d.tables[0].rows[1].cells[1:]), "десяткова кома у Word")
shutil.rmtree(TMP, ignore_errors=True)
print("\n" + "=" * 60)
if FAIL:
    print(f"ПРОВАЛІВ: {len(FAIL)}")
    for f in FAIL: print("  -", f)
    sys.exit(1)
print("УСІ ПЕРЕВІРКИ РОЗДІЛІВ ПРОЙДЕНО")
