# test_sad_assoc_sections.py – розділи результатів зв'язку змінних різних типів, експорт
import os, sys, math, tempfile, shutil
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import numpy as np
import sad_assoc_core as A, sad_assoc_sections as AS, sad_breeding_ui as U
from sad_gxe_core import BreedingError
from matplotlib.figure import Figure
import matplotlib
FAIL = []
def check(c, msg):
    print(("  [OK]   " if c else "  [FAIL] ") + msg, flush=True)
    if not c: FAIL.append(msg)
rng = np.random.default_rng(20)
def fmt(v): return f"{v:.3f}".replace(".", ",")
def struct_ok(S): return not [s["title"] for s in S if s["kind"] == "table" and (any(len(r) != len(s["headers"]) for r in s["rows"]) or len(s["nd"]) != len(s["headers"]))]
def figs_ok(S):
    bad = []
    for s in S:
        if s["kind"] == "figure":
            f = Figure(figsize=s["size"])
            with matplotlib.rc_context(U._RC): m = s["build"](f)
            if m: bad.append(m)
    return not bad

print("== 1. η (група × кількісна ознака)")
rows = []
for nm, mu, n in (("Сорт А", 20, 8), ("Сорт Б", 24, 9), ("Сорт В", 21, 7), ("Сорт Г", 27, 10)):
    for v in rng.normal(mu, 2, n): rows.append([nm, fmt(v), ""])
S, w, e = AS.sections_eta(rows, "Урожайність")
check([s["key"] for s in S] == ["eta", "np", "groups", "fig", "help"], "розділи: η, непараметричний аналог, групи, графік, пояснення")
check(struct_ok(S) and figs_ok(S), "структура таблиць коректна, графік будується")
vals = {h: v for h, v in zip(S[0]["headers"], S[0]["rows"][0])}
check(abs(vals["η²"] - e["eta2"]) < 1e-12 and vals["Величина"] == "великий" and vals["p (ANOVA)"] < 0.001 and vals["η²"] > 0.4, f"η² = {vals['η²']:.3f} (великий ефект, p < 0,001)")
check(len(S[2]["rows"]) == 4 and [r[0] for r in S[2]["rows"]] == ["Сорт А", "Сорт Б", "Сорт В", "Сорт Г"], "описові показники для кожної групи в порядку появи")
rh = [["Н", fmt(v), ""] for v in rng.normal(10, 0.3, 12)] + [["В", fmt(v), ""] for v in rng.normal(12, 6, 12)] + [["М", fmt(v), ""] for v in rng.normal(11, 0.3, 3)]
S2, w2, e2 = AS.sections_eta(rh, "Х")
check(any("Дисперсії груп різні" in x for x in w2) and any("лише 3" in x for x in w2) and any("у 3 рази" in x for x in w2), "попередження: різні дисперсії, мала група, нерівні обсяги")
for fn, frag in [(lambda: AS.sections_eta([["А", "1", ""], ["", "2", ""]], "Х"), "не вказано групу"), (lambda: AS.sections_eta([["А", "", ""]], "Х"), "немає значення"),
                 (lambda: AS.sections_eta([["А", "abc", ""], ["А", "1", ""]], "Х"), "не число"), (lambda: AS.sections_eta([["А", "1", ""], ["А", "2", ""]], "Х"), "щонайменше 2 групи")]:
    try: fn(); check(False, f"«{frag}» мало дати помилку")
    except BreedingError as ex: check(frag in str(ex), f"помилка «{frag}»")

print("\n== 2. Дві групи")
r2 = [["Контроль", fmt(v), ""] for v in rng.normal(10, 2, 20)] + [["Обробка", fmt(v), ""] for v in rng.normal(13, 2, 22)]
S3, w3, e3 = AS.sections_two_groups(r2, "Маса", n_boot=300)
check([s["key"] for s in S3] == ["rank", "pb", "groups", "fig", "help"] and struct_ok(S3) and figs_ok(S3), "розділи й графік")
check(e3["rrb"] < -0.5 and S3[0]["rows"][0][8] == "великий" and S3[0]["rows"][0][0] == "Контроль", f"r_rb = {e3['rrb']:.2f}: група 1 (Контроль) переважно менша → від'ємний знак, великий ефект")
try: AS.sections_two_groups(rows, "X"); check(False, "4 групи мали дати помилку")
except BreedingError as ex: check("рівно 2" in str(ex) and "η" in str(ex), "більше двох груп → пояснення, що обрати η")

print("\n== 3. Таблиця спряженості")
obs = []
for a, b, k in (("Сорт 1", "уражені", 12), ("Сорт 1", "здорові", 38), ("Сорт 2", "уражені", 25), ("Сорт 2", "здорові", 25), ("Сорт 3", "уражені", 6), ("Сорт 3", "здорові", 44)):
    obs += [[a, b, ""]] * k
freq = [["Сорт 1", "уражені", "12"], ["Сорт 1", "здорові", "38"], ["Сорт 2", "уражені", "25"], ["Сорт 2", "здорові", "25"], ["Сорт 3", "уражені", "6"], ["Сорт 3", "здорові", "44"]]
Sa, wa, ca = AS.sections_contingency(obs, "Сорт", "Стан", n_boot=300, n_mc=3000)
Sb, wb, cb = AS.sections_contingency(freq, "Сорт", "Стан", n_boot=300, n_mc=3000)
check(np.array_equal(ca["table"], cb["table"]) and abs(ca["chi2"] - cb["chi2"]) < 1e-12, "рядки-спостереження і готові частоти дають ту саму таблицю й χ²")
check([s["key"] for s in Sa] == ["table", "shares", "expected", "resid", "stats", "stack", "heat", "help"] and struct_ok(Sa) and figs_ok(Sa), "розділи й обидва графіки")
tb = Sa[0]; check(tb["rows"][-1][0] == "Усього" and tb["rows"][-1][-1] == 150 and tb["rows"][0][-1] == 50, "підсумкові рядок і стовпець у таблиці частот")
sh = Sa[1]; check(all(abs(sum(r_[1:]) - 100) < 1e-9 for r_ in sh["rows"]), "частки в кожному рядку дають 100 %")
rs = Sa[3]; check(any(isinstance(c_, tuple) and c_[1] == "*" for r_ in rs["rows"] for c_ in r_[1:]), "значущі залишки позначено зірочкою (межа Бонферроні)")
stt = Sa[4]; names = [r_[0] for r_ in stt["rows"]]
check("χ² Пірсона" in names and "V Крамера" in names and "V Крамера, скоригований на зміщення" in names and "T Чупрова" in names and "C Пірсона / C_max" in names and not any("φ" in n for n in names) and not any("Відношення" in n for n in names),
      "3×2: є χ², V, V скоригований, T, C; φ і OR (лише для 2×2) відсутні")
check(any("Монте-Карло" in n for n in names), "для таблиці 3×2 точний тест – Монте-Карло")
two = AS.sections_contingency([["А", "так", "30"], ["А", "ні", "10"], ["Б", "так", "15"], ["Б", "ні", "25"]], "Група", "Результат", n_boot=200)[0]
nn = [r_[0] for r_ in two[4]["rows"]]; check("φ (2×2, зі знаком)" in nn and "Відношення шансів (OR)" in nn and "χ² із поправкою Єйтса (2×2)" in nn and any("Фішера" in n for n in nn), "для 2×2 є φ, OR, поправка Єйтса й точний тест Фішера")
zr = AS.sections_contingency([["А", "так", "30"], ["А", "ні", "10"], ["Б", "так", "0"], ["Б", "ні", "0"], ["В", "так", "15"], ["В", "ні", "25"]], "Група", "Результат", n_boot=200)
check([r_[0] for r_ in zr[0][0]["rows"]] == ["А", "В", "Усього"] and any("нульовою сумою" in x for x in zr[1]), "рядок із нульовою сумою виключено, підписи узгоджені, є попередження")
for fn, frag in [(lambda: AS.sections_contingency([], "A", "B"), "порожня"), (lambda: AS.sections_contingency([["а", "", ""]], "A", "B"), "обидві ознаки"),
                 (lambda: AS.sections_contingency([["а", "б", "x"]], "A", "B"), "не число"), (lambda: AS.sections_contingency([["а", "б", "1,5"], ["в", "г", "3"]], "A", "B"), "цілими")]:
    try: fn(); check(False, f"«{frag}» мало дати помилку")
    except BreedingError as ex: check(frag in str(ex), f"помилка «{frag}»")

print("\n== 4. Експорт")
import openpyxl, docx
tmp = tempfile.mkdtemp(); U.export_xlsx(os.path.join(tmp, "a.xlsx"), Sa, "Спряженість"); wb = openpyxl.load_workbook(os.path.join(tmp, "a.xlsx"))
ws = next(w_ for w_ in wb.worksheets if w_.title.startswith("Скориговані")); vals = [c_.value for row in ws.iter_rows(min_row=2, min_col=2) for c_ in row]
check(any(isinstance(v, str) and v.endswith("*") for v in vals), "Excel: залишки зі зірочкою збережено як текст")
U.export_docx(os.path.join(tmp, "a.docx"), "Спряженість", Sa, True); d = docx.Document(os.path.join(tmp, "a.docx"))
check(len(d.tables) == 5 and len(d.inline_shapes) == 2, f"Word: таблиць {len(d.tables)}, рисунків {len(d.inline_shapes)}")
pcell = [c_.text for c_ in d.tables[4].rows[1].cells][4]; check(pcell.startswith("<0,0001") or pcell.startswith("0,"), f"Word: p у форматі статті «{pcell}»")
shutil.rmtree(tmp, ignore_errors=True)
print("\n" + "=" * 60)
if FAIL:
    print(f"ПРОВАЛІВ: {len(FAIL)}"); [print("  -", f) for f in FAIL]; sys.exit(1)
print("УСІ ПЕРЕВІРКИ РОЗДІЛІВ АСОЦІАЦІЇ ПРОЙДЕНО")
