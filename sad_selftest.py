# sad_selftest.py — S.A.D.: самоперевірка програми (для CI та перевірки зібраного .exe)
# -*- coding: utf-8 -*-
"""
Запуск:  SAD.exe --selftest [--selftest-out файл.log]   (або python main.py --selftest)
Перевіряє, що в зібраному пакеті є всі модулі й дані: імпорт усіх модулів, репрезентативні розрахунки (scipy, numpy),
побудова графіків (matplotlib), експорт Word/Excel, відкриття всіх вікон. Код виходу 0 — усе гаразд.
"""
import os, sys, io, tempfile, traceback, datetime as dt

MODULES = ['sad_app', 'sad_assoc_core', 'sad_assoc_sections', 'sad_breeding', 'sad_breeding_core', 'sad_breeding_ui', 'sad_breeding_win', 'sad_chart_style_dlg', 'sad_clip', 'sad_cluster_pca', 'sad_common',
           'sad_control_core', 'sad_corr_core', 'sad_corr_sections', 'sad_correlation', 'sad_descriptive_ttest', 'sad_design_core',
           'sad_design_sections', 'sad_gencorr_core', 'sad_gencorr_sections', 'sad_gxe_core', 'sad_het_ord_sections',
           'sad_heterosis_core', 'sad_homogeneous', 'sad_journal_trial', 'sad_manova', 'sad_mixed_core', 'sad_mixed_sections',
           'sad_ordinal_core', 'sad_phenology', 'sad_phenology_dialogs', 'sad_phenology_journal', 'sad_phenology_results', 'sad_phenology_win', 'sad_qr_dialog', 'sad_qrlabels', 'sad_regression', 'sad_report', 'sad_repeated', 'sad_scheme_constructor', 'sad_selection_core',
           'sad_selection_sections', 'sad_stability_ancova', 'sad_weather', 'sad_weather_charts', 'sad_weather_core', 'sad_weather_export', 'sad_weather_indices', 'sad_worldmap']
WINDOWS = ["CorrelationWindow", "DescriptiveWindow", "TTestWindow", "RegressionWindow", "ClusterWindow", "PCAWindow",
           "RepeatedMeasuresWindow", "MixedRepeatedWindow", "StabilityWindow", "AncovaWindow", "ManovaWindow", "FieldJournalWindow",
           "SchemeConstructorWindow", "WeatherWindow", "BreedingStabilityWindow", "GenParWindow", "CombAbilityWindow", "StressWindow",
           "PathWindow", "ControlWindow", "MixedWindow", "DesignWindow", "SelectionWindow", "HeterosisWindow", "ScoresWindow",
           "GenCorrWindow", "AssocWindow", "PhenologyWindow"]


class _Log:
    def __init__(self, path):
        self.path = path; self.lines = []; self.fail = 0
    def ok(self, msg): self._add("OK   ", msg)
    def bad(self, msg, exc=None):
        self.fail += 1
        self._add("ПРОВАЛ", msg + (f": {type(exc).__name__}: {exc}" if exc else ""))
        if exc: self.lines.append("".join(traceback.format_exception(type(exc), exc, exc.__traceback__))[-1500:])
    def _add(self, tag, msg):
        line = f"[{tag}] {msg}"; self.lines.append(line)
        try: print(line, flush=True)
        except Exception: pass
    def flush(self):
        if not self.path: return
        try:
            with open(self.path, "w", encoding="utf-8") as f: f.write("\n".join(self.lines) + "\n")
        except Exception: pass


def _check(log, name, fn):
    try:
        fn(); log.ok(name)
    except Exception as e:
        log.bad(name, e)


def run(main_module, out_path=None):
    log = _Log(out_path)
    log.lines.append(f"S.A.D. самоперевірка · {dt.datetime.now():%Y-%m-%d %H:%M:%S} · Python {sys.version.split()[0]} · "
                     f"{'зібрана програма' if getattr(sys, 'frozen', False) else 'вихідний код'}")
    import importlib
    # 1. імпорт усіх модулів
    for name in MODULES:
        _check(log, f"імпорт {name}", lambda n=name: importlib.import_module(n))
    # 2. розрахунки
    import numpy as np
    def t_control():
        import sad_control_core as C
        r = C.compare_to_control(["St"] * 4 + ["A"] * 4 + ["B"] * 4 + ["D"] * 4, np.r_[np.arange(4), np.arange(4) + 2, np.arange(4) + .5, np.arange(4) - 3.] + np.tile([.1, -.2, .3, 0], 4), "St")
        assert len(r["rows"]) == 3 and 0 <= r["rows"][0]["p_dunnett"] <= 1
    def t_mixed():
        import sad_mixed_core as M
        rg = np.random.default_rng(1); g = [f"G{i % 8}" for i in range(40)]; y = rg.normal(size=40) + np.array([int(x[1:]) for x in g]) * .3
        r = M.fit_mixed(y, {"gen": g}, [], ["gen"]); assert r["sigma2"]["gen"] >= 0
    def t_gxe():
        import sad_gxe_core as G
        rg = np.random.default_rng(2); X = 5 + rg.normal(size=(6, 1)) + rg.normal(size=(1, 5)) + rg.normal(0, .3, (6, 5))
        G.stability_indices(X); G.ammi(X); gg = G.gge(X); G.which_won_where(gg); G.aec_mean_stability(gg)
    def t_breeding():
        import sad_breeding_core as B
        rg = np.random.default_rng(3); gen = [f"G{i}" for i in range(6) for _ in range(3)]; rep = [f"R{k}" for _ in range(6) for k in range(3)]
        B.genetic_parameters_single(gen, rep, rg.normal(size=18) + np.repeat(np.arange(6), 3) * .5)
    def t_corr():
        import sad_corr_core as C
        rg = np.random.default_rng(4); a = [rg.normal(size=25) for _ in range(4)]
        C.correlation_matrix(a, list("ABCD"), "kendall"); C.partial_matrix(a, list("ABCD"), [3], [0, 1, 2])
    def t_gencorr():
        import sad_gencorr_core as G
        rg = np.random.default_rng(5); g = 12; ge = rg.normal(size=(g, 2)); gen, rep, Y = [], [], []
        for i in range(g):
            for j in range(3): gen.append(f"G{i}"); rep.append(f"R{j}"); Y.append(ge[i] + rg.normal(0, .8, 2))
        G.genotypic_correlations(gen, rep, np.array(Y), ["x", "y"], n_boot=30)
    def t_assoc():
        import sad_assoc_core as AC
        rg = np.random.default_rng(9)
        AC.eta_squared([rg.normal(m, 1, 8) for m in (0, 1, 2)]); AC.two_group_effects(rg.normal(0, 1, 12), rg.normal(.5, 1, 12), n_boot=50)
        AC.contingency_analysis([[12, 5, 3], [4, 9, 8]], n_boot=50, n_mc=500)
    def t_indices():
        import calendar as _cal
        import sad_weather_core as W, sad_weather_indices as WI
        s0 = dt.date(1988, 1, 1); n = 365 * 36 + 9; rg = np.random.default_rng(12); dts = [s0 + dt.timedelta(days=i) for i in range(n)]
        doy = np.array([d.timetuple().tm_yday for d in dts], float); t = 10 + 11 * np.sin(2 * np.pi * (doy - 110) / 365) + rg.normal(0, 2, n)
        ser = W.DailySeries(dts, {"tmean": t, "tmax": t + 5, "tmin": t - 5, "prec": np.where(rg.random(n) < .3, rg.gamma(1.4, 5.0, n), 0.),
                                  "rad": np.full(n, 15.0), "wind": np.full(n, 2.5), "rh": np.full(n, 70.0)}, {"grid_elevation_m": 200.0})
        et = float(WI.et0_pm(21.5, 12.3, 22.07, 2.078, 187, 50 + 48 / 60.0, 100.0, ea=1.409)); assert abs(et - 3.88) < 0.01, et         # FAO-56, приклад 18
        sp = WI.spi_table(ser, ref=(1991, 2020), scales=(1, 3), study=(dt.date(2022, 1, 1), dt.date(2022, 12, 31)))
        assert len(sp["rows"]) == 12 and all(r["spi"][3] == r["spi"][3] for r in sp["rows"])
        st_ = W.default_settings(); st_["lat"] = 48.75; rep = W.build_report(ser.slice(dt.date(2020, 1, 1), dt.date(2023, 12, 31)), st_, spi=sp)
        assert rep["intl"]["pet_method"] == "pm" and rep["intl"]["vine"] and rep["intl"]["aridity"]
    def t_phenology():
        import datetime as _d
        import sad_phenology as PH
        assert PH.parse_date_flex("12 травня", 2026)[0] == _d.date(2026, 5, 12) and PH.parse_date_flex("31.04", 2026)[0] is None
        phs = [p for p in PH.preset_phases("pome") if p["key"] in ("bloom0", "bloom50", "bloom1")]; rows = []
        for c, e in (("A", 0), ("B", 3), ("C", 9)):
            for r in (1, 2, 3):
                b = _d.date(2026, 5, 1) + _d.timedelta(days=e + r % 2)
                rows.append({"name": c, "rep": str(r), "dates": {"bloom0": PH.fmt_date(b), "bloom50": PH.fmt_date(b + _d.timedelta(days=4)), "bloom1": PH.fmt_date(b + _d.timedelta(days=10))}})
        ctx = PH.analyze({2026: rows}, phs, "A", 3.0)
        assert ctx["n_obs"] == 27 and ctx["vs"][("C", "bloom0")][1] == "пізній" and ctx["anova"]["bloom0"]["p"] < 0.001 and ctx["anova"]["bloom0"]["letters"]["A"] == "a"
        an = ctx["anova"]["bloom0"]; assert ctx["compare"]["bloom0"]["method"] == "lsd" and an["h2"] > 0.8
        for _m in ("tukey", "sk", "control", "kw", "bonferroni", "duncan"):
            _r = PH.compare_means(ctx["records"], "bloom0", an, _m, "A"); assert _r["groups"] and "error" not in _r, _m
        assert PH.scott_knott({"a": 0.0, "b": 0.1, "c": 20.0, "d": 20.1}, 0.25, 12, 3)["a"] == 0 and abs(PH.holm_adjust([0.01, 0.04, 0.03])[1] - 0.06) < 1e-12
        assert PH.onset_from_visits([(_d.date(2026, 5, 1), 0), (_d.date(2026, 5, 5), 20), (_d.date(2026, 5, 9), 80)], 50)[0] == _d.date(2026, 5, 7)
        import sad_phenology_journal as PJ, sad_journal_trial as JT                       # журнал → фенологія й аналізи
        pl_ = [{"row": 1, "position": i + 1, "role": "recorded", "replication": 1 + i // 4, "factors": {"Сорт": 1 + (i // 2) % 2}} for i in range(8)]
        jj = {"factor_defs": [{"name": "Сорт", "levels": ["А", "Б"]}], "variant_names": [], "plants": pl_,
              "records": {"Ф": {"kind": "date", "preset": "pome", "phase_key": "bloom0", "phase_name": "Початок цвітіння", "values": {f"1:{i + 1}": f"{10 + i}.05.2026" for i in range(8)}}}}
        res = PJ.from_journal_years({"2026": jj}); assert res["preset"] == "pome" and len(res["years_data"][2026]) == 4, res
        _, rows_ = JT.aggregate_journal_data(JT._plants_from_saved(pl_), jj["factor_defs"], [], jj["records"], ["Ф"]); assert rows_[0]["Ф"] > 120, rows_
    def t_report():
        import tempfile, shutil, os
        import sad_report as RP
        from matplotlib.figure import Figure
        fig = Figure(figsize=(4, 3)); fig.add_subplot(111).plot([1, 2, 3]); fig.axes[0].set_title("Тест")
        items = [{"kind": "heading", "text": "Розділ", "level": 3}, {"kind": "text", "text": "Абзац", "warn": False},
                 {"kind": "table", "headers": ["Джерело", "SS", "p"], "rows": [["A", "1,50", "0,012"], ["Залишок", "3,2", ""]], "fills": ["C8E6C9", None], "widths": [100, 60, 60]},
                 {"kind": "figure", "fig": fig}]
        d = tempfile.mkdtemp()
        try:
            RP.export_docx(os.path.join(d, "t.docx"), items, "Тест", "підзаголовок"); RP.export_xlsx(os.path.join(d, "t.xlsx"), items, "Тест")
            import docx, openpyxl
            dd = docx.Document(os.path.join(d, "t.docx")); assert len(dd.tables) == 1 and dd.tables[0].style.name == "Table Grid" and len(dd.inline_shapes) == 1
            wb = openpyxl.load_workbook(os.path.join(d, "t.xlsx")); assert wb.worksheets[1]["B2"].value == 1.5 and "Графіки" in wb.sheetnames
        finally: shutil.rmtree(d, ignore_errors=True)
    def t_qrlabels():
        import tempfile, itertools, shutil
        import sad_qrlabels as QL
        fd = [{"name": "Сорт", "levels": ["Голден", "Гала"]}, {"name": "Добриво", "levels": ["0", "N60"]}]
        plants = [{"row": b, "position": j, "role": "recorded", "replication": b, "factors": {"Сорт": c[0], "Добриво": c[1]}}
                  for b in (1, 2) for j, c in enumerate(itertools.product((1, 2), (1, 2)), 1)]
        data = QL.build_experiment({"factor_defs": fd, "cfg": {"model": "area"}, "plants": plants}, "K7M2", "Тест"); QL.fill_payloads(data, "descriptive", "m", "S")
        assert len(data["plots"]) == 8 and all(QL.parse_payload(p["payload"])["valid"] for p in data["plots"])
        d = tempfile.mkdtemp()
        try:
            files = QL.save_all(d, data, "S", {}, dpi=100); assert len(files) == 4 + 8, len(files)
            assert open(files[0], "rb").read(4) == b"%PDF" and b"DejaVuSans" in open(files[0], "rb").read()      # кирилицю вбудовано
            assert open(files[1], "rb").read(4) == b"%PDF" and open(files[4], "rb").read(4) == b"\x89PNG"
            QL.fill_payloads(data, "descriptive", "m", "R58x40"); rp = os.path.join(d, "roll.pdf")      # рулон для принтера етикеток
            assert QL.render_labels_pdf(rp, data, "R58x40", {"mono": True, "printer_dpi": 203}) == 8 and open(rp, "rb").read(4) == b"%PDF"
            img = QL.render_label_png(data, data["plots"][0], "R58x40", {"mono": True}, dpi=203); assert img.mode == "1" and img.size == (464, 320), img.size
        finally: shutil.rmtree(d, ignore_errors=True)
    def t_worldmap():
        import sad_worldmap as WM
        D = WM.MapData.get()
        assert len(D.pop) > 7000 and D.country_at(30.2109, 48.7543) == "Україна"
        n = D.nearest(30.2109, 48.7543); assert n and "Умань" in D.full_label(n[0])
        assert D.search("умань") and D.search("kyiv")
        poly = WM.clip_polygon([-1e6, 1e6, 1e6, -1e6], [-1e6, -1e6, 1e6, 1e6], 0, 0, 100, 80); assert len(poly) == 4
    def t_clip():
        import sad_clip as SC
        h = SC.build_html_table(["А", "Б"], [["1", "2"]]); hd, frag = SC.parse_cf_html(SC.cf_html(h)); assert frag == h and 'border="1"' in h
    def t_chart_style():
        import matplotlib; from matplotlib.figure import Figure
        import sad_weather_core as W, sad_weather_charts as CH
        s0 = dt.date(2022, 1, 1); n = 730; rg = np.random.default_rng(10); dts = [s0 + dt.timedelta(days=i) for i in range(n)]
        t = 9 + 11 * np.sin(2 * np.pi * np.array([d.timetuple().tm_yday for d in dts]) / 365)
        stt = W.default_settings(); stt["lat"] = 48.75                      # широта потрібна для вінограрських індексів
        rep = W.build_report(W.DailySeries(dts, {"tmean": t, "tmax": t + 4, "tmin": t - 4, "prec": rg.random(n) * 5, "rh": 55 + 25 * rg.random(n)}), stt)
        for kind, _ in CH.CHARTS:
            f = Figure(); res = CH.build_chart(f, kind, rep, "Тест", {"general": {"width_cm": 18, "font_family": "DejaVu Sans"}, "colors": {kind: {}}})
            if kind == "spi": assert isinstance(res, str) and "SPI" in res, res          # SPI потребує довгого ряду: діаграма пояснює, що робити
            else: assert res is None, (kind, res)
    def t_scipy_misc():
        from scipy import stats, optimize, sparse, special, spatial
        from scipy.sparse.linalg import splu
        stats.t.sf(2, 10); special.roots_genlaguerre(10, 1.0); spatial.ConvexHull(np.random.default_rng(6).normal(size=(10, 2)))
        splu(sparse.eye(5, format="csc")); optimize.brentq(lambda x: x - 1, 0, 3)
    for n, f in (("розрахунок: тест Дуннетта", t_control), ("розрахунок: змішана модель REML", t_mixed), ("розрахунок: AMMI/GGE/стабільність", t_gxe),
                 ("розрахунок: генетичні параметри", t_breeding), ("розрахунок: кореляції (Кендалл, частинна)", t_corr),
                 ("розрахунок: генотипні кореляції", t_gencorr), ("розрахунок: η, бісеріальні, спряженість", t_assoc), ("scipy: stats/special/sparse/spatial/optimize", t_scipy_misc),
                 ("міжнародні показники: ET₀ FAO-56, SPI (scipy gamma.fit), Вінклер, Югліна, індекс сухості", t_indices),
                 ("фенологія: розбір дат, зведення, ANOVA з НІР і літерами, візити, дані з польового журналу", t_phenology),
                 ("експорт звіту: Word (таблиці з сіткою, графіки) і Excel", t_report),
                 ("QR-таблички: ділянки, QR (segno), PDF із кирилицею, PNG, Excel, маніфест", t_qrlabels),
                 ("карта світу: дані, країна, найближче місто, пошук, відсікання", t_worldmap), ("копіювання таблиці: HTML-формат буфера", t_clip),
                 ("діаграми: усі 8 зі стилем користувача", t_chart_style)):
        _check(log, n, f)
    # 3. matplotlib і експорт
    def t_figs():
        import matplotlib; from matplotlib.figure import Figure
        f = Figure(figsize=(4, 3)); ax = f.add_subplot(111); ax.plot([1, 2, 3], label="Тест"); ax.set_title("Кирилиця: ГТК, σ², ρ"); ax.legend()
        b = io.BytesIO(); f.savefig(b, format="png", dpi=100); assert b.getvalue()[:4] == b"\x89PNG"
        from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg  # noqa
    def t_export():
        import sad_breeding_ui as U, sad_corr_sections as CS
        rg = np.random.default_rng(7); arr = [rg.normal(size=40) + 0 for _ in range(4)]
        S, w, M = CS.sections_partial(arr, list("ABCD"), [3], [0, 1, 2])
        d = tempfile.mkdtemp()
        U.export_xlsx(os.path.join(d, "t.xlsx"), S, "Тест"); U.export_docx(os.path.join(d, "t.docx"), "Тест", S, True)
        assert os.path.getsize(os.path.join(d, "t.docx")) > 5000 and os.path.getsize(os.path.join(d, "t.xlsx")) > 2000
    def t_weather():
        import sad_weather_core as W, sad_weather_export as X, sad_weather_charts as CH
        s0 = dt.date(2023, 1, 1); n = 730; rg = np.random.default_rng(8); dts = [s0 + dt.timedelta(days=i) for i in range(n)]
        doy = np.array([d.timetuple().tm_yday for d in dts], float); t = 9 + 11 * np.sin(2 * np.pi * (doy - 110) / 365)
        ser = W.DailySeries(dts, {"tmean": t, "tmax": t + 4, "tmin": t - 4, "prec": np.where(rg.random(n) < .3, 5., 0.)})
        ser.meta.update({"retrieved_at": "2026-01-01T12:00:00"})
        rep = W.build_report(ser, W.default_settings())
        png = X.figure_png(rep, "walter", "Тест", CH.build_chart, dpi=80); assert png
        d = tempfile.mkdtemp(); meta = X.make_meta(rep, "Тест", 48.7, 30.2, "nasa")
        pub = X.publication_text(rep, "Тест", 48.7, 30.2, "nasa")
        X.export_xlsx(os.path.join(d, "w.xlsx"), rep, meta, pub); X.export_docx(os.path.join(d, "w.docx"), rep, meta, pub, [("Діаграма", png)])
    for n, f in (("matplotlib: PNG, кирилиця, TkAgg", t_figs), ("експорт Word і Excel (частинна кореляція)", t_export), ("погода: звіт, діаграма, Word, Excel", t_weather)):
        _check(log, n, f)
    # 4. вікна
    try:
        import tkinter as tk
        from tkinter import messagebox
        for n in ("showinfo", "showwarning", "showerror"): setattr(messagebox, n, lambda *a, **k: None)
        messagebox.askyesno = lambda *a, **k: False
        root = tk.Tk(); root.withdraw()
        try:
            app = main_module.SADTk(root); root._sad_app = app; log.ok("головне вікно SADTk створено")
        except Exception as e:
            log.bad("головне вікно SADTk", e)
        for name in WINDOWS:
            cls = getattr(main_module, name, None)
            if cls is None: log.bad(f"вікно {name}: клас не знайдено в main"); continue
            try:
                try: w = cls(root, {})
                except TypeError: w = cls(root)
                root.update()
                try: w.win.destroy()
                except Exception: pass
                log.ok(f"вікно {name}")
            except Exception as e:
                log.bad(f"вікно {name}", e)
        try:
            import sad_worldmap as WM
            mp = WM.WorldMap(root, width=700, height=450); mp.size = lambda: (700, 450)       # головне вікно приховане – розмір задаємо явно
            mp.goto(30.2, 48.75, 7.0); assert mp.stats["items"] > 20, mp.stats
            mp.goto(31, 49, 5.5); assert mp.stats["items"] > 100, mp.stats; mp.destroy(); log.ok("віджет карти: малювання")
        except Exception as e:
            log.bad("віджет карти: малювання", e)
        try:
            from sad_scheme_constructor import SchemeConstructorWindow
            sc = SchemeConstructorWindow(root, {}); sc._model_var.set("area"); sc._refresh_model_panels()
            sc.factor_defs = [{"name": "Сорт", "levels": ["A", "B", "C"]}, {"name": "Добриво", "levels": ["0", "N60"]}]; sc._unlock_step2()
            def _btns(fr):
                out = []
                def walk(x):
                    for c in x.winfo_children():
                        if c.winfo_class() == "Button" and "Сформувати план" in str(c.cget("text")): out.append(str(c.cget("state")))
                        walk(c)
                walk(fr); return out
            assert _btns(sc._step3_area_content) == ["normal"], _btns(sc._step3_area_content)          # Крок 3 відкрито після введення розмірів
            sc._av["reps"].set("1"); assert _btns(sc._step3_area_content) == ["disabled"]               # некоректні повторності блокують Крок 3
            sc._av["reps"].set("4"); assert _btns(sc._step3_area_content) == ["normal"] and "24 ділянки" in sc._area_size_status.cget("text")
            sc._area_plan = None
            import itertools as _it
            scheme_ = {"factor_defs": sc.factor_defs, "cfg": {"model": "area", "design_used": "rcbd"},
                       "plants": [{"row": b, "position": j, "role": "recorded", "replication": b, "factors": {"Сорт": c[0], "Добриво": c[1]}}
                                  for b in (1, 2) for j, c in enumerate(_it.product((1, 2, 3), (1, 2)), 1)]}
            from sad_qr_dialog import QRLabelsDialog
            dlg = QRLabelsDialog(sc.win, scheme_, {}, None); dlg._rebuild(); assert dlg.data is not None and len(dlg.data["plots"]) == 12 and dlg._photo is not None   # перегляд (PIL.ImageTk) працює
            dlg.win.destroy()
            sc.win.destroy(); log.ok("конструктор схеми: повторності відкривають Крок 3 (площинна модель); діалог QR-табличок із переглядом")
        except Exception as e:
            log.bad("конструктор схеми: повторності відкривають Крок 3 (площинна модель); діалог QR-табличок із переглядом", e)
        try: root.destroy()
        except Exception: pass
    except Exception as e:
        log.bad("tkinter: створення вікон", e)
    log.lines.append(f"ПІДСУМОК: {'УСЕ ГАРАЗД' if log.fail == 0 else str(log.fail) + ' ПРОВАЛ(ИВ)'}")
    try: print(log.lines[-1], flush=True)
    except Exception: pass
    log.flush()
    return 0 if log.fail == 0 else 1
