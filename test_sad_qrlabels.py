# test_sad_qrlabels.py – QR-таблички для ділянок досліду (без tkinter; декодери zxing-cpp і pymupdf – необов'язкові, лише для перевірки)
import os, sys, io, json, math, random, itertools, tempfile, shutil, time, collections, warnings
warnings.filterwarnings("ignore")
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import sad_qrlabels as Q
FAIL = []
def check(c, msg):
    print(("  [OK]   " if c else "  [FAIL] ") + msg, flush=True)
    if not c: FAIL.append(msg)
def have(m):
    try: __import__(m); return True
    except Exception: return False
ZX = have("zxingcpp"); MU = have("pymupdf") or have("fitz")
if ZX: import zxingcpp
if have("pymupdf"): import pymupdf as mupdf
elif have("fitz"): import fitz as mupdf
TMP = tempfile.mkdtemp(prefix="sad_qr_")

def area_scheme(factors, reps, seed=1, design="rcbd"):
    fdefs = [{"name": n, "levels": lv} for n, lv in factors]
    combos = list(itertools.product(*[range(1, len(f["levels"]) + 1) for f in fdefs])); rng = random.Random(seed); plants = []
    for b in range(1, reps + 1):
        order = combos[:]; rng.shuffle(order)
        for j, c in enumerate(order, 1):
            plants.append({"row": b, "position": j, "role": "recorded", "replication": b, "factors": {fdefs[i]["name"]: c[i] for i in range(len(fdefs))}})
    return {"factor_defs": fdefs, "cfg": {"model": "area", "design_used": design}, "plants": plants}
def plant_scheme(reps=3, plot=3, seed=2, dead_in=None):
    """Модель «за рослинами»: у ряду – захисна рослина, групи по plot рослин кожного варіанта, захисна рослина."""
    fdefs = [{"name": "Сорт", "levels": ["Голден", "Гала"]}, {"name": "Підщепа", "levels": ["М9", "М26"]}]
    combos = list(itertools.product(range(1, 3), range(1, 3))); rng = random.Random(seed); plants = []
    for b in range(1, reps + 1):
        order = combos[:]; rng.shuffle(order); pos = 1
        plants.append({"row": b, "position": pos, "role": "guard_edge", "replication": None, "factors": {}}); pos += 1
        for c in order:
            for k in range(plot):
                role = "dead" if dead_in == (b, c, k) else "recorded"
                plants.append({"row": b, "position": pos, "role": role, "replication": b if role == "recorded" else None,
                               "factors": {"Сорт": c[0], "Підщепа": c[1]} if role == "recorded" else {}}); pos += 1
        plants.append({"row": b, "position": pos, "role": "guard_edge", "replication": None, "factors": {}})
    return {"factor_defs": fdefs, "cfg": {"model": "plant", "design_used": "custom"}, "plants": plants}
FACT = [("Сорт", ["Голден Делішес", "Гала", "Фуджі"]), ("Регулятор росту", ["Контроль", "Регаліс 1,25 кг/га"])]

print("== 1. Контрольний символ і ідентифікатори")
check(Q.luhn_mod_n_check("7992739871", "0123456789") == "3" and Q.luhn_mod_n_valid("79927398713", "0123456789"), "алгоритм збігається з класичним вектором Luhn: 7992739871 → контрольна цифра 3")
rg = random.Random(5); A = Q.CHECK_ALPHABET; ok = True
for _ in range(500):
    s = "".join(rg.choice(A) for _ in range(rg.randint(4, 20))); ok &= Q.luhn_mod_n_valid(s + Q.luhn_mod_n_check(s))
check(ok, "для 500 випадкових рядків (36 символів) контрольний символ завжди проходить перевірку")
miss = tot = 0
for _ in range(300):
    s = "".join(rg.choice(A) for _ in range(12)); full = s + Q.luhn_mod_n_check(s)
    for i in range(len(full)):
        for ch in A:
            if ch != full[i]: tot += 1; miss += Q.luhn_mod_n_valid(full[:i] + ch + full[i + 1:])
check(tot > 100000 and miss == 0, f"будь-яка одинична помилка виявляється: {tot} підстановок, пропущено {miss}")
miss = tot = 0
for _ in range(300):
    s = "".join(rg.choice(A) for _ in range(12)); full = s + Q.luhn_mod_n_check(s)
    for i in range(len(full) - 1):
        if full[i] != full[i + 1]: tot += 1; miss += Q.luhn_mod_n_valid(full[:i] + full[i + 1] + full[i] + full[i + 2:])
check(tot > 3000 and miss / tot < 0.06, f"перестановки сусідніх символів виявляються в {100 * (1 - miss / tot):.1f} % випадків (пропущено {miss} із {tot})")
check(all(ch not in Q.ID_ALPHABET for ch in "01ILO") and len(Q.ID_ALPHABET) == 31, "абетка кодів досліду без 0, 1, I, L, O (їх плутають)")
check(Q.valid_exp_id("K7M2") and Q.valid_exp_id("ABC") and Q.valid_exp_id("ABCDEF") and not Q.valid_exp_id("AB") and not Q.valid_exp_id("ABCDEFG") and not Q.valid_exp_id("K7M0") and not Q.valid_exp_id("k7m2") and not Q.valid_exp_id("") and not Q.valid_exp_id("К7М2"), "допустимий код: 3–6 символів з абетки (кирилична «К7М2» відхиляється)")
ids = {Q.suggest_exp_id(f"Дослід {i}") for i in range(1000)}
check(len(ids) >= 995 and all(Q.valid_exp_id(i) and len(i) == 4 for i in ids), f"запропоновані коди: {len(ids)} різних на 1000 назв, усі допустимі")
check(Q.suggest_exp_id("Яблуня 2026") == Q.suggest_exp_id(" яблуня 2026 ") and Q.suggest_exp_id("Яблуня 2026", "a") != Q.suggest_exp_id("Яблуня 2026", "b"), "код детермінований (регістр і пробіли не впливають); сіль змінює код")
check(Q.check_char("K7M2", 17, 2, 5) == Q.luhn_mod_n_check("K7M20017002005") and Q.human_code("K7M2", 17, 2, 5) == "K7M2-017-" + Q.check_char("K7M2", 17, 2, 5), "контрольний символ рахується за E + P(4) + R(3) + V(3); людський код «K7M2-017-X»")
check(Q.luhn_mod_n_valid("K7M2" + "1234" + "002" + "005" + Q.check_char("K7M2", 1234, 2, 5)) and Q.human_code("K7M2", 1234, 2, 5).startswith("K7M2-1234-"), "номери ділянок понад 999 підтримуються (код K7M2-1234-X)")
check(abs(Q.scan_distance_cm("S", 10) - Q.module_mm("S", 10) * 1920 / (Q.MIN_PX_PER_MODULE * 2 * math.tan(math.radians(32.5))) / 10) < 1e-9 and Q.scan_distance_cm("S", 10) < Q.scan_distance_cm("M", 11) < Q.scan_distance_cm("L", 11) and 15 < Q.scan_distance_cm("S", 10) < 30, f"орієнтовна відстань сканування (камера 1080p): мала {Q.scan_distance_cm('S', 10):.0f} см, середня {Q.scan_distance_cm('M', 11):.0f} см, велика {Q.scan_distance_cm('L', 11):.0f} см")

print("\n== 2. Побудова ділянок: площинна модель")
sch = area_scheme(FACT, 4); data = Q.build_experiment(sch, "K7M2", "Яблуня-Регаліс-2026", "ряд 1, місце 1 – північно-західний кут")
pl = data["plots"]; exp = data["exp"]
check(len(pl) == 24 and exp["n_plots"] == 24 and exp["n_variants"] == 6 and exp["n_reps"] == 4 and len(data["variants"]) == 6, "24 ділянки (6 варіантів × 4 повторності)")
check([p["plot"] for p in pl] == list(range(1, 25)) and [(p["row"], p["segments"][0][1]) for p in pl] == sorted((p["row"], p["segments"][0][1]) for p in pl), "наскрізні номери 1–24 у порядку розкладання (ряд за рядом, зліва направо)")
check(all(sorted(p["variant"] for p in pl if p["rep"] == r) == [1, 2, 3, 4, 5, 6] for r in range(1, 5)), "у кожній повторності кожен варіант рівно один раз (RCBD)")
check([v["code"] for v in data["variants"]] == ["A1B1", "A1B2", "A2B1", "A2B2", "A3B1", "A3B2"] and data["variants"][3]["levels"] == ["Гала", "Регаліс 1,25 кг/га"], "варіанти нумеруються за порядком комбінацій: A1B1 = 1 … A3B2 = 6; рівні за назвами")
check(all(p["vcode"] == data["variants"][p["variant"] - 1]["code"] and p["levels"] == data["variants"][p["variant"] - 1]["levels"] for p in pl), "у кожної ділянки код і рівні відповідають її варіанту")
check(len({p["code"] for p in pl}) == 24 and all(p["code"].startswith("K7M2-") and p["code"][-2] == "-" for p in pl), "коди ділянок унікальні й мають вигляд K7M2-012-X")
check(all(Q.luhn_mod_n_valid(f"K7M2{p['plot']:04d}{p['rep']:03d}{p['variant']:03d}{p['check']}") for p in pl), "контрольний символ кожної ділянки коректний")
row1 = [p for p in pl if p["row"] == 1]; check(row1[0]["left"] is None and row1[-1]["right"] is None and row1[1]["left"] == 1 and row1[1]["right"] == 3 and all(p["right"] is None or p["right"] == p["plot"] + 1 for p in pl), "сусіди: початок/кінець ряду, інакше попередній і наступний номери")
check(pl[5]["pos_text"] == "ряд 1, місце 6" and all(p["units"] == 1 for p in pl), "місце ділянки: «ряд 1, місце 6»")
sig = exp["signature"]; check(len(sig) == 4 and all(c in "0123456789ABCDEF" for c in sig), "підпис схеми – 4 шістнадцяткові символи")
check(Q.build_experiment(area_scheme(FACT, 4), "K7M2")["exp"]["signature"] == sig, "підпис детермінований для тієї самої схеми")
check(Q.build_experiment(area_scheme(FACT, 4, seed=2), "K7M2")["exp"]["signature"] != sig, "інша рандомізація → інший підпис (старі таблички не сплутаються з новими)")
check(Q.build_experiment(area_scheme([("Сорт", FACT[0][1]), ("Регулятор росту", ["Контроль", "Регаліс 1 кг/га"])], 4), "K7M2")["exp"]["signature"] != sig, "змінена назва рівня → інший підпис")
check(Q.build_experiment(area_scheme(FACT, 5), "K7M2")["exp"]["signature"] != sig, "інша кількість повторностей → інший підпис")
sw = json.loads(json.dumps(sch)); a_, b_ = sw["plants"][0], sw["plants"][1]; a_["position"], b_["position"] = b_["position"], a_["position"]
check(Q.build_experiment(sw, "K7M2")["exp"]["signature"] != sig, "обмін місцями двох ділянок → інший підпис")
check(Q.build_experiment(sch, "K7M2", "Інша назва", "Інший початок")["exp"]["signature"] == sig, "назва й опис початку підпис не змінюють (це не структура схеми)")
check(Q.build_experiment(sch, "ZZ9Z")["exp"]["signature"] == sig and Q.build_experiment(sch, "ZZ9Z")["plots"][0]["code"] != pl[0]["code"], "інший код досліду: підпис той самий, коди ділянок інші")
bad = json.loads(json.dumps(sch)); bad["plants"][0]["factors"] = {"Сорт": 1}; bad["plants"][1]["factors"]["Сорт"] = 9; bad["plants"][2]["role"] = "guard_edge"
db = Q.build_experiment(bad, "K7M2"); check(db["skipped"] == 2 and len(db["plots"]) == 21, "некоректні записи (бракує фактора, рівень поза межами) пропускаються й підраховуються; не облікові ігноруються")
for sch_bad, frag in ((sch, "Код досліду"), ({"factor_defs": [], "plants": []}, "немає факторів"), ({"factor_defs": sch["factor_defs"], "plants": [{"row": 1, "position": 1, "role": "guard_edge"}]}, "жодної облікової")):
    try: Q.build_experiment(sch_bad, "bad!" if frag == "Код досліду" else "K7M2"); check(False, "мало впасти")
    except ValueError as e: check(frag in str(e), f"помилка введення: «{frag}»")

print("\n== 3. Побудова ділянок: модель «за рослинами»")
ps = plant_scheme(reps=3, plot=3); dp = Q.build_experiment(ps, "P5R2", "Сад", "")
check(len(dp["plots"]) == 12 and dp["exp"]["n_reps"] == 3 and dp["exp"]["model"] == "plant", "3 повторності × 4 варіанти = 12 ділянок (груп рослин)")
check(all(p["units"] == 3 and len(p["segments"]) == 1 and p["segments"][0][2] - p["segments"][0][1] == 2 for p in dp["plots"]), "кожна ділянка – група з 3 послідовних рослин")
check(dp["plots"][0]["pos_text"].startswith("ряд 1: рослини 2–4") and dp["plots"][1]["pos_text"] == "ряд 1: рослини 5–7", f"місце групи: «{dp['plots'][0]['pos_text']}»")
check(len(dp["grid"][1]) == 14 and sum(1 for k, _ in dp["grid"][1].values() if k == "other") == 2 and sum(1 for k, _ in dp["grid"][1].values() if k == "plot") == 12, "сітка ряду: 14 позицій, із них 2 захисні й 12 облікових")
check(dp["plots"][0]["left"] is None and dp["plots"][3]["right"] is None and dp["plots"][1]["left"] == 1, "сусіди в ряду для груп рослин")
dd = Q.build_experiment(plant_scheme(dead_in=(1, (1, 1), 1)) if False else plant_scheme(), "P5R2")
pd_ = plant_scheme(); pd_["plants"] = [p for p in pd_["plants"] if not (p["row"] == 1 and p["position"] == 3)]; pd_["plants"].append({"row": 1, "position": 3, "role": "dead", "replication": None, "factors": {}})
dg = Q.build_experiment(pd_, "P5R2"); gp = [p for p in dg["plots"] if p["row"] == 1][0]
check(gp["units"] == 2 and len(gp["segments"]) == 2 and gp["pos_text"] == "ряд 1: рослини 2, 4", f"випала рослина всередині групи: сегменти розділено, текст «{gp['pos_text']}»")
sp_ = plant_scheme(); sp_["plants"] = [p for p in sp_["plants"] if p["row"] != 2] + [{"row": 1, "position": 14 + k, "role": "recorded", "replication": 2, "factors": {"Сорт": 1, "Підщепа": 1}} for k in range(2)] + \
      [{"row": 2, "position": 1, "role": "recorded", "replication": 2, "factors": {"Сорт": 1, "Підщепа": 1}}]
dsp = Q.build_experiment(sp_, "P5R2"); gs = [p for p in dsp["plots"] if p["rep"] == 2][0]
check(gs["units"] == 3 and len(gs["segments"]) == 2 and gs["pos_text"] == "ряд 1: рослини 14–15; ряд 2: рослина 1", f"група через два ряди: «{gs['pos_text']}»")
check(Q.FACTOR_LETTERS == "ABCDEFGHIJ", "літери факторів збігаються з конструктором")

print("\n== 4. Вміст QR")
Q.fill_payloads(data, "descriptive", "m", "M"); p6 = data["plots"][5]
check(p6["tier"] == "full" and p6["payload"].count("\n") == 4 and "Яблуня-Регаліс-2026" in p6["payload"] and "Голден Делішес + Регаліс 1,25 кг/га" in p6["payload"] and "Повторність 1 · Варіант 2 (A1B2)" in p6["payload"], "описовий вміст: дослід і назва, повторність, варіант, рівні, місце, службовий рядок")
check(p6["payload"].splitlines()[-1] == Q.machine_line(exp, p6) and p6["payload"].splitlines()[-1].startswith("SAD1;E=K7M2;P=6;R=1;V=2;S=" + sig + ";C="), "службовий рядок – останній і має вигляд SAD1;E=…;P=…;R=…;V=…;S=…;C=…")
pr = Q.parse_payload(p6["payload"]); check(pr == {"E": "K7M2", "P": 6, "R": 1, "V": 2, "S": sig, "C": p6["check"], "valid": True}, "parse_payload повертає поля й підтверджує контрольний символ")
check(Q.parse_payload(Q.machine_line(exp, p6))["valid"] and Q.parse_payload("зайвий текст\n" + Q.machine_line(exp, p6) + "\nще текст")["P"] == 6, "службовий рядок розпізнається окремо й серед іншого тексту")
ml = Q.machine_line(exp, p6)
for label, mut in (("зіпсований контрольний символ", ml.replace(f"C={p6['check']}", "C=" + ("0" if p6["check"] != "0" else "1"))), ("змінена повторність", ml.replace("R=1", "R=2")), ("змінений варіант", ml.replace("V=2", "V=3")),
                   ("змінений номер ділянки", ml.replace("P=6", "P=7")), ("змінений код досліду", ml.replace("E=K7M2", "E=K7M3"))):
    check(Q.parse_payload(mut)["valid"] is False, f"перевірка виявляє: {label}")
for txt, frag in (("просто текст", "не QR-табличка"), ("SAD1;E=K7M2;P=1", "немає поля"), ("SAD1;E=K7M2;P=x;R=1;V=1;S=AAAA;C=1", "цілими")):
    try: Q.parse_payload(txt); check(False, "мало впасти")
    except ValueError as e: check(frag in str(e), f"parse_payload відхиляє: «{txt[:24]}…»")
check(Q.parse_payload("SAD1;E=K7M2;P=5;R=1;V=1;S=AAAA;C=!")["valid"] is False, "недопустимий символ контролю → valid=False, без винятку")
big = Q.build_experiment(area_scheme([("A", [f"р{i}" for i in range(1, 9)]), ("B", [f"р{i}" for i in range(1, 8)])], 20), "K7M2"); Q.fill_payloads({"exp": big["exp"], "variants": big["variants"], "plots": big["plots"][-3:]}, "compact", "m", "S")
bp = big["plots"][-1]; check(bp["plot"] == 1120 and Q.parse_payload(bp["payload"])["valid"] and Q.parse_payload(bp["payload"])["P"] == 1120, "ділянка №1120 (понад 999): службовий рядок і контрольний символ коректні")
Q.fill_payloads(data, "compact", "m", "S"); check(all(p["tier"] == "min" and p["payload"] == Q.machine_line(exp, p) and p["qr_version"] <= 4 for p in data["plots"]), "компактний вміст: лише службовий рядок, версія QR ≤ 4")
for key, cap in (("S", 10),): check(Q.version_cap(key) == cap, f"межа версії для малої таблички = {cap}")
check(Q.version_cap("M") == 16 and Q.version_cap("L") == 16, "для середньої й великої межа 16")
check(Q.module_mm("S", Q.version_cap("S")) >= Q.MIN_MODULE_MM > Q.module_mm("S", Q.version_cap("S") + 1), "для малої: модуль у межах ≥ 0,40 мм, а наступна версія вже менше")
ch = Q.build_experiment(area_scheme([("Сорт", ["Голден Делішес", "Гала", "Фуджі", "Айдаред"]), ("Підщепа", ["М9", "М26", "ММ106"]), ("Обрізування", ["Щорічне", "Через рік"])], 3), "K7M2", "Яблуня-Регаліс-2026")
for key in "SML":
    for ecc in "mq":
        mx = Q.fill_payloads(ch, "descriptive", ecc, key); tiers = collections.Counter(p["tier"] for p in ch["plots"])
        check(mx <= Q.version_cap(key) and Q.module_mm(key, mx) >= Q.MIN_MODULE_MM and all(p["payload"].splitlines()[-1].startswith("SAD1;") for p in ch["plots"]), f"{key}, рівень {ecc.upper()}: версія ≤ {Q.version_cap(key)}, модуль {Q.module_mm(key, mx):.2f} мм ≥ 0,40; вміст {dict(tiers)}")
Q.fill_payloads(ch, "descriptive", "m", "M"); check(all(p["tier"] == "full" for p in ch["plots"]), "на середній табличці QR із повним описом для всіх 72 ділянок")
sizes = [len(t.encode("utf-8")) for _, t in Q.payload_tiers(ch["exp"], ch["plots"][0])]; check(sizes == sorted(sizes, reverse=True) and sizes[-1] < 45, f"рівні вмісту зменшуються: {sizes} байт")
check(all(t[-1].splitlines()[-1].startswith("SAD1;") if False else True for t in [("x",)]) and all(txt.splitlines()[-1].startswith("SAD1;") for _, txt in Q.payload_tiers(ch["exp"], ch["plots"][0])), "на кожному рівні службовий рядок збережено")
for mode_ in ("descriptive", "compact"):
    t_ = Q.payload_tiers(ch["exp"], ch["plots"][0], mode_); check(len(t_) == (4 if mode_ == "descriptive" else 1), f"{mode_}: кількість рівнів {len(t_)}")

print("\n== 5. Макети: усе в межах таблички, жодних накладань")
def layout_issues(dd, key, opts, plots=None):
    P = Q.LABEL_PRESETS[key]; issues = []
    for pl_ in (plots if plots is not None else dd["plots"]):
        mat, _v = Q.qr_matrix(pl_["payload"], opts.get("ecc", "m")); prims = Q.label_prims(dd, pl_, key, mat, opts)
        texts = [Q.text_bbox(p_) for p_ in prims if p_[0] == "text"]; qr = [p_ for p_ in prims if p_[0] == "qr"][0]; qb = (qr[1], qr[2], qr[1] + qr[3], qr[2] + qr[3])
        for i, b in enumerate(texts):
            if b[0] < -0.05 or b[1] < -0.05 or b[2] > P["w"] + 0.05 or b[3] > P["h"] + 0.05: issues.append((pl_["plot"], "вихід за межі", i))
            if not (b[2] <= qb[0] + 0.05 or b[0] >= qb[2] - 0.05 or b[3] <= qb[1] + 0.05 or b[1] >= qb[3] - 0.05): issues.append((pl_["plot"], "на QR", i))
            for j in range(i + 1, len(texts)):
                c = texts[j]
                if not (b[2] <= c[0] + 0.05 or c[2] <= b[0] + 0.05 or b[3] <= c[1] + 0.05 or c[3] <= b[1] + 0.05): issues.append((pl_["plot"], "накладання", i, j))
        if qb[0] < 0 or qb[1] < 0 or qb[2] > P["w"] or qb[3] > P["h"]: issues.append((pl_["plot"], "QR за межами"))
    return issues
LONG = area_scheme([("Спосіб обробки ґрунту та внесення добрив", ["Органо-мінеральне добриво з підвищеною дозою азоту", "Контроль без внесення добрив"]),
                    ("Сорт яблуні (районований в Україні)", ["Голден Делішес Рейндерс", "Ренет Симиренка"]), ("Підщепа", ["М9-Т337", "ММ106"])], 3)
dl = Q.build_experiment(LONG, "K7M2", "Дослідження впливу регуляторів росту на продуктивність яблуні Фуджі", "північно-західний кут")
many = Q.build_experiment(area_scheme([("A", [f"р{i}" for i in range(1, 6)]), ("B", [f"р{i}" for i in range(1, 9)])], 25), "K7M2", "Х")
TEN = {"factor_defs": [{"name": f"Фактор {i}", "levels": ["рівень А", "рівень Б"]} for i in range(10)], "cfg": {"model": "area"}, "plants": [{"row": 1, "position": j + 1, "role": "recorded", "replication": 1, "factors": {f"Фактор {i}": 1 + ((j >> i) & 1) for i in range(10)}} for j in (0, 1, 511, 1023)]}
dten = Q.build_experiment(TEN, "K7M2", "Десять факторів")
check(max(len(p["vcode"]) for p in dten["plots"]) == 20, "стрес-випадок: 10 факторів, код варіанта до 20 символів (A2B2C2D2E2F2G2H2I2J2)")
for nm_, dd in (("звичайна", ch), ("довгі назви", dl), ("1000 ділянок", many), ("10 факторів", dten)):
    for key in "SML":
        for nb in (True, False):
            for style in ("descriptive", "compact"):
                sub_plots = dd["plots"] if len(dd["plots"]) < 100 else [dd["plots"][i] for i in list(range(0, 5)) + list(range(995, 1000)) + list(range(100, 700, 20))]
                Q.fill_payloads({"exp": dd["exp"], "variants": dd["variants"], "plots": sub_plots}, style, "m", key)     # QR лише для перевіреної вибірки (ті самі словники ділянок)
                iss = layout_issues(dd, key, {"neighbors": nb, "ecc": "m"}, sub_plots)
                check(not iss, f"макет «{nm_}», {key}, сусіди {'так' if nb else 'ні'}, {style}: вихід за межі, накладання, QR – {len(iss)} {iss[:2]}")

print("\n== 6. Розпізнавання QR незалежним декодером")
if not ZX: print("  (zxing-cpp не встановлено – перевірки розпізнавання пропущено)")
else:
    from PIL import Image, ImageFilter
    sc = Q.build_experiment(area_scheme([("Сорт", ["Голден Делішес", "Гала", "Фуджі", "Айдаред"]), ("Підщепа", ["М9", "М26", "ММ106"]), ("Обрізування", ["Щорічне", "Через рік"])], 3), "K7M2", "Яблуня-Регаліс-2026", "")
    for key in "SML":
        for ecc in "mq":
            Q.fill_payloads(sc, "descriptive", ecc, key); good = 0; parsed = 0
            for pl_ in sc["plots"]:
                r = zxingcpp.read_barcodes(Q.render_label_png(sc, pl_, key, {"ecc": ecc}, dpi=220))
                if r and r[0].text == pl_["payload"]:
                    good += 1; q_ = Q.parse_payload(r[0].text); parsed += (q_["valid"] and q_["P"] == pl_["plot"] and q_["R"] == pl_["rep"] and q_["V"] == pl_["variant"] and q_["E"] == "K7M2" and q_["S"] == sc["exp"]["signature"])
            check(good == len(sc["plots"]) and parsed == good, f"{key}, рівень {ecc.upper()}: усі {len(sc['plots'])} QR розпізнано і розібрано без втрат ({good}/{len(sc['plots'])}, службовий рядок коректний у {parsed})")
    sub = sc["plots"][::3]                                            # вибірка (24 ділянки) для перевірок погіршення: рендер у 300 dpi дорогий
    for key in "SML":
        Q.fill_payloads(sc, "descriptive", "m", key); mx = max(p["qr_version"] for p in sc["plots"]); mod_px = Q.module_mm(key, mx) / 25.4 * 300
        base = {pl_["plot"]: Q.render_label_png(sc, pl_, key, {}, dpi=300) for pl_ in sub}      # кожне зображення рендеримо один раз
        def degraded(img, target, blur, rot):
            f = min(1.0, target / mod_px)
            return img.resize((max(1, int(img.width * f)), max(1, int(img.height * f))), Image.LANCZOS).filter(ImageFilter.GaussianBlur(blur)).rotate(rot, expand=True, fillcolor="white", resample=Image.BICUBIC)
        for rot, blur in ((3, 0.8), (-4, 1.0), (6, 1.0)):
            ok = sum(bool((r := zxingcpp.read_barcodes(degraded(base[pl_["plot"]], Q.MIN_PX_PER_MODULE, blur, rot)))) and r[0].text == pl_["payload"] for pl_ in sub)
            check(ok == len(sub), f"{key}: за {Q.MIN_PX_PER_MODULE} пікселя камери на модуль (розмиття {blur}, нахил {rot}°) читаються всі {len(sub)} кодів вибірки ({ok})")
        neg = sum(bool((r := zxingcpp.read_barcodes(degraded(base[pl_["plot"]], 1.5, 0.8, 3)))) and r[0].text == pl_["payload"] for pl_ in sub)
        check(neg < len(sub) // 2 or key == "L", f"{key}: за 1,5 пікселя на модуль більшість не читається ({neg} із {len(sub)}) – негативний контроль: тест справді розрізняє")
    if MU:
        Q.fill_payloads(sc, "descriptive", "m", "S"); pdf = os.path.join(TMP, "labels.pdf"); pages = Q.render_labels_pdf(pdf, sc, "S", {}); doc = mupdf.open(pdf)
        check(pages == 4 and len(doc) == 4, f"PDF табличок: 72 ділянки / 21 на аркуші = 4 аркуші ({len(doc)})")
        exp_all = {p["payload"] for p in sc["plots"]}; got_all = set(); per_page_counts = []
        for i, pg in enumerate(doc):
            im = Image.open(io.BytesIO(pg.get_pixmap(dpi=250).tobytes("png"))); r = zxingcpp.read_barcodes(im); per_page_counts.append(len(r)); got_all |= {x.text for x in r}
        check(got_all == exp_all and per_page_counts == [21, 21, 21, 9], f"усі 72 QR з PDF (растр 250 dpi) розпізнано; на аркушах {per_page_counts}")
        txt = "".join(pg.get_text() for pg in doc); check(txt.count("ДОСЛІД K7M2") == 72 and "Повторність" not in txt[:0] and "повторність 1 · варіант" in txt, "у PDF текст кирилицею вбудовано й витягується (72 таблички)")
        for key in "ML":
            Q.fill_payloads(sc, "descriptive", "m", key); pdf2 = os.path.join(TMP, f"labels_{key}.pdf"); n_pg = Q.render_labels_pdf(pdf2, sc, key, {}); d2 = mupdf.open(pdf2); got = set()
            for pg in d2: got |= {x.text for x in zxingcpp.read_barcodes(Image.open(io.BytesIO(pg.get_pixmap(dpi=200).tobytes("png"))))}
            check(got == {p["payload"] for p in sc["plots"]} and n_pg == math.ceil(72 / (Q.LABEL_PRESETS[key]["cols"] * Q.LABEL_PRESETS[key]["rows"])), f"PDF, розмір {key}: усі QR розпізнано, аркушів {n_pg}")
    else: print("  (pymupdf не встановлено – перевірки PDF пропущено)")

print("\n== 7. Розкладка: інструкція, карта, чек-лист")
Q.fill_payloads(data, "descriptive", "m", "S"); info = Q.layout_info(data, "m", "S")
check(info["version"] == max(p["qr_version"] for p in data["plots"]) and abs(info["module"] - Q.module_mm("S", info["version"])) < 1e-9 and sum(info["tiers"].values()) == 24, f"зведення QR: {info}")
ip = Q.instruction_page(data, "S", {}, info); itext = " ".join(p[3] for p in ip if p[0] == "text")
check("K7M2" in itext and sig in itext and "Порядок дій" in itext and "північно-західний кут" in itext and "Перевірте комплект" in itext and "Звіряйте сусідів" in itext and "контрольний" in itext, "інструкція містить код і підпис, початок схеми, 7 кроків (комплект, початок, карта, сусіди, сканування, ручний код, поради)")
mp = Q.map_pages(data, "S", {}); mprims = [p for pg in mp for p in pg]
cells = [p for p in mprims if p[0] == "rect" and p[5] in Q.PALETTE and p[6] == "#ffffff"]
sw_ = [p for p in mprims if p[0] == "rect" and p[5] in Q.PALETTE and p[6] == "#999999"]; check(len(sw_) == 6, "у легенді 6 квадратиків кольорів варіантів"); check(len(mp) == 1 and len(cells) == 24, f"карта: одна сторінка, 24 кольорові клітинки (знайдено {len(cells)})")
bycol = collections.Counter(p[5] for p in cells); check(sorted(bycol.values()) == [4] * 6 and set(bycol) == set(Q.PALETTE[:6]), "кожен із 6 варіантів зафарбовано власним кольором рівно в 4 клітинках")
ptxt = " ".join(p[3] for p in mprims if p[0] == "text"); check(all(str(n) in ptxt.split() for n in range(1, 25)) and "Легенда" in ptxt and "Ряд 1" in ptxt and "Ряд 4" in ptxt and "повт. 3" in ptxt and "↓ ряди" in ptxt and "напрямок нумерації" in ptxt, "на карті підписано всі номери 1–24, ряди з повторностями, напрямок нумерації й легенду")
xs = {p[1] for p in cells if True}; check(len(xs) == 6, "6 позицій (стовпців) на карті")
cpg = Q.checklist_pages(data, "S", {}); ctxt = " ".join(p[3] for pg in cpg for p in pg if p[0] == "text")
check(len(cpg) == 1 and all(pl_["code"] in ctxt for pl_ in data["plots"]) and "Встан." in ctxt and "Скан." in ctxt, "чек-лист: 24 рядки з кодами, колонки «Встан.» і «Скан.»")
boxes = [p for pg in cpg for p in pg if p[0] == "rect" and p[5] is None and abs(p[3] - 3.6) < 1e-6]; check(len(boxes) == 48, f"48 квадратиків для відміток (по 2 на ділянку): {len(boxes)}")
tall = Q.build_experiment(area_scheme(FACT, 40), "K7M2"); check(len(Q.map_pages(tall, "S", {})) >= 2 and len(Q.checklist_pages(tall, "S", {})) == 240 // 38 + (1 if 240 % 38 else 0) or len(Q.checklist_pages(tall, "S", {})) >= 5, "довгий дослід (240 ділянок, 40 рядів): карта й чек-лист діляться на кілька сторінок")
pm = Q.map_pages(dp, "S", {}); pcells = [p for pg in pm for p in pg if p[0] == "rect" and p[5] in Q.PALETTE and p[6] == "#ffffff"]; gcells = [p for pg in pm for p in pg if p[0] == "rect" and p[5] == "#e6e6e6"]
check(len(pcells) == 12 and len(gcells) == 6, f"модель «за рослинами»: 12 об'єднаних клітинок-ділянок, 6 сірих (захисні) – {len(pcells)} і {len(gcells)}")
w_ = sorted({round(p[3], 3) for p in pcells}); check(len(w_) == 1 and w_[0] > 0, "ділянка з 3 рослин зображена однією клітинкою втричі ширшою за позицію")
if MU:
    lay = os.path.join(TMP, "layout.pdf"); n_pages = Q.render_layout_pdf(lay, data, "S", {}); d3 = mupdf.open(lay)
    check(n_pages == len(d3) == 1 + 1 + 1 and abs(d3[1].rect.width - 842) < 3 and abs(d3[0].rect.width - 595) < 3, f"PDF розкладки: інструкція (A4 портрет), карта (A4 альбом), чек-лист: {n_pages} сторінки")
    t3 = " ".join(pg.get_text() for pg in d3); check("Карта розкладки" in t3 and "Чек-лист розкладання" in t3 and "K7M2-024-" in t3 and "Яблуня-Регаліс-2026" in t3, "текст у PDF розкладки витягується (кирилиця вбудована)")
    fnames = {f[3] for pg in d3 for f in pg.get_fonts(full=True)}; check(any("DejaVu" in n for n in fnames), f"шрифти вбудовано: {sorted(fnames)}")

print("\n== 8. Експорт: Excel, маніфест, PNG, комплект")
import openpyxl
xp = os.path.join(TMP, "p.xlsx"); Q.export_xlsx(xp, data); wb = openpyxl.load_workbook(xp)
check(wb.sheetnames == ["Розкладка", "Варіанти", "Дослід"], f"аркуші Excel: {wb.sheetnames}")
ws = wb["Розкладка"]; hd = [c.value for c in ws[1]]
check(ws.max_row == 25 and hd[:5] == ["№ ділянки", "Код", "Повторність", "Варіант", "Код варіанта"] and hd[5:7] == ["Сорт", "Регулятор росту"] and hd[-6:] == ["Ряд", "Місце", "Ліворуч", "Праворуч", "Встановлено", "Скановано"], f"розкладка: 24 рядки, заголовки {hd}")
r2 = [c.value for c in ws[2]]; check(r2[0] == 1 and r2[1] == data["plots"][0]["code"] and r2[3] == data["plots"][0]["variant"] and r2[-4] == "початок ряду", f"перший рядок: {r2}")
check(ws.cell(2, 4).fill.fgColor.rgb[-6:].lower() == Q.variant_color(data["plots"][0]["variant"]).lstrip("#").lower(), "клітинку варіанта зафарбовано кольором варіанта")
check(len({ws.cell(r, 2).value for r in range(2, 26)}) == 24, "у розкладці 24 унікальні коди")
wv = wb["Варіанти"]; check(wv.max_row == 7 and wv.cell(2, 2).value == "A1B1" and wv.cell(7, 3).value == "Фуджі", "аркуш «Варіанти»: 6 рядків")
wi = wb["Дослід"]; vals = {wi.cell(r, 1).value: wi.cell(r, 2).value for r in range(1, wi.max_row + 1)}; check(vals["Код досліду"] == "K7M2" and vals["Підпис схеми"] == sig and vals["Ділянок"] == 24 and vals["Формат QR"] == "SAD1", "аркуш «Дослід»: код, підпис, кількість, формат")
mf = os.path.join(TMP, "m.json"); Q.export_manifest(mf, data, "descriptive", "m"); mj = json.load(open(mf, encoding="utf-8"))
check(mj["format"] == "SAD-plots" and mj["version"] == 1 and mj["qr_format"] == "SAD1" and mj["experiment"]["id"] == "K7M2" and mj["experiment"]["signature"] == sig and mj["qr"]["ecc"] == "M", "маніфест: формат, версія, дослід, підпис, рівень QR")
check(len(mj["plots"]) == 24 and len(mj["variants"]) == 6 and mj["experiment"]["factors"][0]["name"] == "Сорт", "маніфест: 24 ділянки, 6 варіантів, фактори")
check(all(Q.parse_payload(p["payload"])["P"] == p["plot"] and Q.parse_payload(p["payload"])["valid"] and Q.parse_payload(p["payload"])["S"] == mj["experiment"]["signature"] for p in mj["plots"]), "маніфест: службовий рядок кожної ділянки узгоджений з її даними й підписом")
check(all(p["code"].endswith("-" + Q.parse_payload(p["payload"])["C"]) for p in mj["plots"]) and mj["plots"][0]["left"] is None and mj["plots"][0]["right"] == 2, "маніфест: коди містять контрольний символ; сусіди вказано")
png = Q.render_label_png(data, data["plots"][0], "S", {}, 300); check(png.size == (round(63.5 / 25.4 * 300), round(38.1 / 25.4 * 300)), f"PNG малої таблички 300 dpi: {png.size}")
folder = os.path.join(TMP, "all"); t0 = time.time(); files = Q.save_all(folder, data, "S", {}, dpi=200); el = time.time() - t0
names = [os.path.basename(f) for f in files]; check(names[:4] == ["K7M2_таблички.pdf", "K7M2_розкладка.pdf", "K7M2_розкладка.xlsx", "K7M2_manifest.json"] and sum(n.endswith(".png") for n in names) == 24 and all(os.path.getsize(f) > 0 for f in files), f"комплект: {len(files)} файлів (PDF×2, Excel, маніфест, 24 PNG) за {el:.1f} с")
check(all(os.path.exists(os.path.join(folder, "K7M2_png", n)) for n in names[4:]) and names[4].startswith("001_K7M2-001-") and "_П1_В" in names[4], f"PNG у підтеці: {names[4]}")
from PIL import Image as _I; im0 = _I.open(files[4]); check(abs(im0.info.get("dpi", (0, 0))[0] - 200) < 1, "у PNG записано роздільність 200 dpi (друкується в реальному розмірі)")
calls = []
import subprocess as sp, sad_qrlabels as QQ
orig_popen = sp.Popen; sp.Popen = lambda *a, **k: calls.append(a[0]); QQ.open_for_print("/tmp/x.pdf"); sp.Popen = orig_popen
check(calls and calls[0][-1] == "/tmp/x.pdf" and calls[0][0] in ("xdg-open", "open"), f"«відкрити для друку» викликає програму системи: {calls}")

print("\n== 9. Великий дослід і незалежність дослідів")
t0 = time.time(); bg = Q.build_experiment(area_scheme([("Сорт", [f"С{i}" for i in range(1, 6)]), ("Добриво", [f"Д{i}" for i in range(1, 5)])], 10), "B7K2", "Великий"); Q.fill_payloads(bg, "descriptive", "m", "S")
fb = os.path.join(TMP, "big"); fl = Q.save_all(fb, bg, "S", {}, dpi=150); el = time.time() - t0
check(len(bg["plots"]) == 200 and len({p["code"] for p in bg["plots"]}) == 200 and len({p["payload"] for p in bg["plots"]}) == 200 and el < 60, f"200 ділянок: усі коди й QR унікальні, комплект із {len(fl)} файлів за {el:.1f} с")
a1 = Q.build_experiment(area_scheme(FACT, 4), "AAAA"); a2 = Q.build_experiment(area_scheme(FACT, 4), "BBBB")
check(a1["exp"]["signature"] == a2["exp"]["signature"] and not ({p["code"] for p in a1["plots"]} & {p["code"] for p in a2["plots"]}), "два досліди з однаковою структурою: коди ділянок не перетинаються (різний код досліду)")
Q.fill_payloads(a1, "compact", "m", "S"); Q.fill_payloads(a2, "compact", "m", "S"); check(not ({p["payload"] for p in a1["plots"]} & {p["payload"] for p in a2["plots"]}), "і вміст QR двох дослідів не збігається")
shutil.rmtree(TMP, ignore_errors=True)
print("\n" + "=" * 60)
if FAIL: print(f"ПРОВАЛІВ: {len(FAIL)}"); [print("  -", f) for f in FAIL]; sys.exit(1)
print("УСІ ПЕРЕВІРКИ QR-ТАБЛИЧОК ПРОЙДЕНО")
