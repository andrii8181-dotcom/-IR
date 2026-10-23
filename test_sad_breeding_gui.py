# test_sad_breeding_gui.py — вікна розділу «Селекція та генетика» (Xvfb)
import os, sys, math, tempfile, shutil, time, datetime as dt
sys.path.insert(0, os.environ.get("SAD_SRC", os.path.dirname(os.path.abspath(__file__))))
os.environ.setdefault("DISPLAY", ":99")
import numpy as np
import tkinter as tk
FAIL = []
def check(c, msg):
    print(("  [OK]   " if c else "  [FAIL] ") + msg, flush=True)
    if not c: FAIL.append(msg)
import main as m
import sad_breeding_win as BW, sad_breeding as BR, sad_gxe_core as GX, sad_breeding_ui as U
LOG = {"info": [], "warn": [], "error": [], "ask": []}
mb = tk.messagebox
mb.showinfo = lambda t="", msg="", **k: LOG["info"].append((t, msg))
mb.showwarning = lambda t="", msg="", **k: LOG["warn"].append((t, msg))
mb.showerror = lambda t="", msg="", **k: LOG["error"].append((t, msg))
mb.askyesno = lambda t="", msg="", **k: (LOG["ask"].append((t, msg)) or False)
U.open_file_cross = lambda p: None
TMP = tempfile.mkdtemp(prefix="sadbw_"); SAVE = {"p": None}
tk.filedialog.asksaveasfilename = lambda **k: SAVE["p"]
root = tk.Tk(); app = m.SADTk(root); root._sad_app = app; root.update()
rng = np.random.default_rng(5)

def toplevels(w=root):
    out = []
    if isinstance(w, tk.Toplevel): out.append(w)
    for c in w.winfo_children(): out.extend(toplevels(c))
    return out
def buttons(w):
    out = []
    try:
        if w.winfo_class() == "Button": out.append(w)
    except Exception: pass
    for c in w.winfo_children(): out.extend(buttons(c))
    return out
def find_win(frag):
    r = [t for t in toplevels() if frag in t.title()]
    return r[-1] if r else None
def fill(win, rows, headers=None):
    if headers:
        for e, h in zip(win.header_entries, headers): e.delete(0, tk.END); e.insert(0, h)
    while len(win.entries) < len(rows): win._add_rows(10)
    for r_, vals in zip(win.entries, rows):
        for e, v in zip(r_, vals): e.delete(0, tk.END); e.insert(0, v)
def click_all(viewer_win, prefix=""):
    """Натискає кожну кнопку розділу; повертає (кількість, список помилок)."""
    btns = [b for b in buttons(viewer_win) if b.cget("text").startswith("  ") and b.cget("text").strip()]
    errs = []
    for b in btns:
        try: b.invoke(); viewer_win.update()
        except Exception as ex: errs.append(f"{b.cget('text').strip()}: {ex}")
    return len(btns), errs

print("== 1. Меню")
labels = []
def rl(w):
    try:
        if w.winfo_class() == "Label" and w.cget("text"): labels.append(w.cget("text"))
    except Exception: pass
    for c in w.winfo_children(): rl(c)
rl(root)
check(any("Селекція та генетика" in t for t in labels), "розділ «Селекція та генетика» на головній сторінці")
for nm in ("Аналіз стабільності", "Генетичні параметри", "Комбінаційна здатність", "Стресостійкість", "Шляховий аналіз"):
    check(any(nm in t for t in labels), f"пункт «{nm}»")
check(not any(t == "Спеціальні" for t in labels), "старий розділ «Спеціальні» прибрано (його єдиний пункт перенесено)")

print("\n== 2. Генетичні параметри")
w = m.GenParWindow(root); root.update()
rows = []; ga = rng.normal(0, 2, 10)
for i in range(10):
    for k in range(3): rows.append([f"G{i+1}", "", f"R{k+1}", f"{50 + ga[i] + rng.normal(0, 1):.2f}".replace(".", ","), f"{20 + .4 * ga[i] + rng.normal(0, 1.2):.2f}", ""])
fill(w, rows, ["Генотип", "Середовище", "Повторність", "Висота", "Маса", "Порожня"])
w._run(); root.update()
v = find_win("Генетичні параметри — результати")
check(v is not None, "вікно результатів відкрилось")
n, errs = click_all(v)
check(n == 5 and not errs, f"усі {n} розділів (параметри, 2 × ANOVA, графік, пояснення) відображаються без помилок {errs[:1]}")
check(any("Висота" in b.cget("text") for b in buttons(v)) and any("Маса" in b.cget("text") for b in buttons(v)), "ANOVA для кожної ознаки в меню розділів")
SAVE["p"] = os.path.join(TMP, "gp.xlsx")
[b for b in buttons(v) if "Excel" in b.cget("text")][0].invoke()
import openpyxl; wb = openpyxl.load_workbook(SAVE["p"]); check("Генетичні параметри" in wb.sheetnames, f"Excel із вікна: {wb.sheetnames[:4]}…")
SAVE["p"] = os.path.join(TMP, "gp.docx")
[b for b in buttons(v) if "Word" in b.cget("text")][0].invoke()
import docx; d = docx.Document(SAVE["p"]); check(len(d.tables) >= 3 and len(d.inline_shapes) >= 1, f"Word із вікна: таблиць {len(d.tables)}, рисунків {len(d.inline_shapes)}")
v.destroy()
# вставка з буфера із заголовками
w.win.clipboard_clear(); w.win.clipboard_append("Ген\tСер\tПовт\tУрожай\n1\tA\t1\t5,5\n1\tA\t2\t5,7\n")
for e in w.entries[0]: e.delete(0, tk.END)
w.entries[0][0].focus_set(); w._paste()
check(w.header_entries[3].get() == "Урожай" and w.entries[0][3].get() == "5,5" and w.entries[1][3].get() == "5,7",
      "вставка з буфера розпізнала рядок заголовків і заповнила дані")
# помилки
w2 = m.GenParWindow(root); root.update(); n0 = len(LOG["warn"]); w2._run()
check(len(LOG["warn"]) == n0 + 1 and "порожня" in LOG["warn"][-1][1], "порожня таблиця → зрозуміле попередження")
fill(w2, [["G1", "", "R1", "abc", "", ""], ["G1", "", "R2", "1", "", ""], ["G2", "", "R1", "2", "", ""]]); n0 = len(LOG["warn"]); w2._run()
check(len(LOG["warn"]) == n0 + 1 and "не число" in LOG["warn"][-1][1], "нечислове значення → попередження з номером рядка")
w2.win.destroy(); w.win.destroy()

print("\n== 3. Комбінаційна здатність")
w = m.CombAbilityWindow(root); root.update()
check(w.header_entries[0].get() == "Лінія" and str(w.method_cb.cget("state")) == "disabled", "схема «Лінія × тестер»: заголовки й вимкнені параметри діалелю")
rows = []; a = rng.normal(0, 1, 6); b = rng.normal(0, .8, 3)
for i in range(6):
    for j in range(3):
        for k in range(3): rows.append([f"L{i+1}", f"T{j+1}", f"R{k+1}", f"{15 + a[i] + b[j] + rng.normal(0, .4):.3f}", "", ""])
fill(w, rows); w._refresh_traits(); w._run(); root.update()
v = find_win("Лінія × тестер")
check(v is not None, "лінія × тестер: результати відкрились")
n, errs = click_all(v); check(n >= 8 and not errs, f"лінія × тестер: {n} розділів, без помилок {errs[:1]}"); v.destroy()
w.design_var.set("Діалель (Гріффінг)"); w._on_design()
check(w.header_entries[0].get() == "Батько 1" and str(w.method_cb.cget("state")) == "readonly", "перемикання на діалель змінює заголовки й вмикає метод/модель")
for r_ in w.entries:
    for e in r_: e.delete(0, tk.END)
rows = []; g = rng.normal(0, 1, 5)
for i in range(5):
    for j in range(i, 5):
        for k in range(3): rows.append([f"P{i+1}", f"P{j+1}", f"R{k+1}", f"{10 + g[i] + g[j] + rng.normal(0, .4):.3f}", "", ""])
fill(w, rows); w.model_var.set("II — випадкова"); w._run(); root.update()
v = find_win("Діалель")
check(v is not None, "діалель (метод 2, модель II): результати відкрились")
n, errs = click_all(v); check(n >= 7 and not errs, f"діалель: {n} розділів без помилок {errs[:1]}")
check(any("модель II" in b.cget("text") for b in buttons(v)), "у назві ANOVA вказано метод і модель"); v.destroy()
w.method_var.set("4 (лише F₁)"); w.model_var.set("I — фіксована"); w._run(); root.update()
v = find_win("Діалель")
labs = []
def collect(x):
    try:
        if x.winfo_class() == "Label": labs.append(x.cget("text"))
    except Exception: pass
    for c in x.winfo_children(): collect(c)
collect(v)
check(any("Метод 4 не використовує батьківські форми" in t for t in labs), "метод 4 на даних із батьками: працює, а попередження про виключення батьків показано")
v.destroy(); w.win.destroy()

print("\n== 3б. Порівняння зі стандартом")
w = m.ControlWindow(root); root.update()
rows = []
for nm, e_ in (("Стандарт", 0.0), ("Н1", 3.0), ("Н2", 0.0), ("Н3", -3.0)):
    for k in range(4): rows.append([nm, f"R{k+1}", f"{30 + e_ + rng.normal(0, 1):.2f}", "", ""])
fill(w, rows); w._refresh_entries(); w._refresh_traits()
check(set(w.ctl_cb.cget("values")) == {"Стандарт", "Н1", "Н2", "Н3"}, "список номерів для вибору стандарту підтягується з таблиці")
n0 = len(LOG["warn"]); w._run(); check(len(LOG["warn"]) == n0 + 1 and "стандарт" in LOG["warn"][-1][1].lower(), "без стандарту → зрозуміле попередження")
w.ctl_var.set("Стандарт"); w._run(); root.update()
v = find_win("Порівняння зі стандартом — Ознака 1")
check(v is not None, "результати відкрились")
n, errs = click_all(v); check(n == 4 and not errs, f"{n} розділи без помилок {errs[:1]}"); v.destroy()
w.alt_var.set("більше стандарту"); w._run(); root.update(); v2 = find_win("Порівняння зі стандартом — Ознака 1"); check(v2 is not None, "односторонній тест працює"); v2.destroy()
w.win.destroy()

print("\n== 3в. Змішана модель")
w = m.MixedWindow(root); root.update()
rows = []; ga = rng.normal(0, 1.5, 15); ea = rng.normal(0, 2, 3)
for i in range(15):
    for j in range(3):
        if (i + j) % 7 == 0: continue
        for k in range(2): rows.append([f"G{i+1}", f"E{j+1}", f"R{k+1}", "", f"{40 + ga[i] + ea[j] + rng.normal(0, 1):.2f}", "", ""])
rows.append(["G1", "E1", "R3", "", "", "", ""])                       # порожнє значення = пропущена ділянка
fill(w, rows); w._refresh_traits(); w._run(); root.update()
v = find_win("Змішана модель — Ознака 1")
check(v is not None, "результати відкрились (пропущена ділянка з порожнім значенням проігнорована)")
n, errs = click_all(v); check(n == 7 and not errs, f"{n} розділів без помилок {errs[:1]}"); v.destroy()
w.geno_var.set("фіксований (скориговані середні)"); w.mode_var.set("Один дослід"); w._run(); root.update()
v = find_win("Змішана модель — Ознака 1"); check(v is not None, "фіксований генотип: працює"); v.destroy()
w.blk_var.set(True); n0 = len(LOG["warn"]); w._run(); check(len(LOG["warn"]) == n0 + 1 and "Блок" in LOG["warn"][-1][1], "блоки ввімкнено без даних → зрозуміле попередження")
w.win.destroy()

print("\n== 3г. Генератор схем")
w = m.DesignWindow(root); root.update()
fill(w, [[f"N{i+1}", ""] for i in range(36)] + [["St1", "ст"], ["St2", "ст"]]); w._update_hint()
check("36 нових" in w.hint.cget("text") and "6" in w.hint.cget("text"), f"підказка допустимих k: «{w.hint.cget('text')}»")
w.k_var.set("6"); w.r_var.set("2"); w.seed_var.set("5"); w._run(); root.update()
v = find_win("План досліду — Альфа-решітка"); check(v is not None, "альфа-решітка згенерована")
n, errs = click_all(v); check(n == 5 and not errs, f"{n} розділів без помилок {errs[:1]}"); v.destroy()
w.k_var.set("7"); n0 = len(LOG["warn"]); w._run(); check(len(LOG["warn"]) == n0 + 1 and "Допустимі розміри" in LOG["warn"][-1][1], "недопустимий розмір блоку → пояснення з переліком допустимих")
w.kind_var.set("Аугментована"); w._on_kind(); w.k_var.set("4"); w._run(); root.update()
v = find_win("План досліду — Аугментована"); check(v is not None, "аугментована схема згенерована (2 стандарти, 4 блоки)"); v.destroy()
w.k_var.set("abc"); n0 = len(LOG["warn"]); w._run(); check(len(LOG["warn"]) == n0 + 1, "нечислові параметри → попередження")
w.k_var.set("4")
# кнопка «Заповнити»: імена й стандарти
w._fill_names(); dlg = [t for t in toplevels(w.win) if "Згенерувати" in t.title()][0]
ents = [e for e in dlg.winfo_children()[0].winfo_children() if isinstance(e, tk.Entry)]
ents[0].delete(0, tk.END); ents[0].insert(0, "20"); ents[1].delete(0, tk.END); ents[1].insert(0, "3")
[b for b in buttons(dlg) if "Заповнити" in b.cget("text")][0].invoke(); root.update()
check(w.entries[0][0].get() == "N1" and w.entries[19][0].get() == "N20" and w.entries[20][0].get() == "St1" and w.entries[20][1].get() == "ст", "генерація назв: N1…N20 і стандарти St1…St3 позначені «ст»")
w.win.destroy()
w = m.MixedWindow(root); root.update()
rows = []; gt = rng.normal(0, 1, 12)
for i in range(12):
    for k in range(3): rows.append([f"G{i+1}" if i else "Ст", "", f"R{k+1}", f"B{(i + k) % 4 + 1}", f"{20 + gt[i] + (3 if i == 3 else 0) + rng.normal(0, .5):.2f}", "", ""])
fill(w, rows); w.geno_var.set("фіксований (скориговані середні)"); w.mode_var.set("Один дослід"); w.blk_var.set(True); w.ctl_var.set("Ст"); w._run(); root.update()
v = find_win("Змішана модель — Ознака 1"); labs2 = [b.cget("text") for b in buttons(v)]
check(any("Порівняння зі стандартом" in t for t in labs2) and any("Ефективність" in t for t in labs2), "змішана модель із полем «Стандарти»: з'явились порівняння зі стандартом і ефективність блоків")
n, errs = click_all(v); check(not errs, f"усі {n} розділів відображаються без помилок {errs[:1]}"); v.destroy(); w.win.destroy()

print("\n== 3д. Індекси добору")
w = m.SelectionWindow(root); root.update()
rg_ = np.random.default_rng(15); Gt_ = np.array([[5.0, -2.5, 1.0], [-2.5, 4.0, 0.0], [1.0, 0.0, 3.0]]); Et_ = np.diag([7.0, 5.0, 6.0])
gv_ = rg_.multivariate_normal(np.zeros(3), Gt_, 30); rows = []
for i in range(30):
    for j in range(3):
        yv = 50 + gv_[i] + rg_.multivariate_normal(np.zeros(3), Et_); rows.append([f"G{i+1}", f"R{j+1}"] + [f"{v:.3f}" for v in yv] + [""])
fill(w, rows, ["Генотип", "Повторність", "Урожайність", "Білок", "Стійкість", ""])
w.w_var.set("1; 1; 0.5"); w._run(); root.update(); v = find_win("Індекси добору — результати")
check(v is not None, "результати відкрились (порожній 4-й стовпець проігноровано, а ваги — на 3 ознаки)")
n, errs = click_all(v); check(n == 8 and not errs, f"{n} розділів без помилок {errs[:1]}"); v.destroy()
w.d_var.set("2; -1; 0"); w._run(); root.update(); v = find_win("Індекси добору — результати")
check(v is not None, "з бажаними приростами вікно відкрилось")
[b for b in buttons(v) if "Очікуваний відгук" in b.cget("text")][0].invoke(); v.update()
def tree_first_col(x):
    out = []
    if isinstance(x, tk.ttk.Treeview): out += [str(x.item(i, "values")[0]) for i in x.get_children()]
    for c in x.winfo_children(): out += tree_first_col(c)
    return out
check("Песек–Бейкер" in tree_first_col(v), "таблиця відгуку містить рядок «Песек–Бейкер»"); v.destroy()
w.w_var.set("1; 1"); n0 = len(LOG["warn"]); w._run(); check(len(LOG["warn"]) == n0 + 1 and "вкажіть 3" in LOG["warn"][-1][1], "кількість ваг ≠ кількості ознак → зрозуміле попередження")
w.w_var.set("1; 1; 1"); w.p_var.set("150"); n0 = len(LOG["warn"]); w._run(); check(len(LOG["warn"]) == n0 + 1, "частка відбору 150 % → попередження")
w.win.destroy()

print("\n== 3е. Гетерозис")
w = m.HeterosisWindow(root); root.update()
rows = []
for nm, mu, p1, p2 in (("П1", 40.0, "", ""), ("П2", 50.0, "", ""), ("Ст", 47.0, "", ""), ("Н1", 62.0, "П1", "П2"), ("Н2", 45.0, "П1", "П2")):
    for k in range(4): rows.append([nm, p1, p2, f"R{k+1}", f"{mu + rng.normal(0, 1):.3f}", "", ""])
fill(w, rows); w._refresh_traits(); w._refresh_names()
check(set(w.std_cb.cget("values")) == {"П1", "П2", "Ст", "Н1", "Н2"}, "список генотипів для вибору стандарту підтягується з таблиці")
w.std_var.set("Ст"); w._run(); root.update(); v = find_win("Гетерозис — Ознака 1")
check(v is not None, "результати відкрились"); n, errs = click_all(v); check(n == 3 and not errs, f"{n} розділи без помилок {errs[:1]}"); v.destroy()
w.dir_var.set("менше — краще"); w._run(); root.update(); v = find_win("Гетерозис — Ознака 1"); check(v is not None, "напрям «менше — краще» працює"); v.destroy()
fill(w, [["Х", "П1", "", "R1", "1", "", ""]]); n0 = len(LOG["warn"]); w._run(); check(len(LOG["warn"]) == n0 + 1, "некоректні дані → зрозуміле попередження")
w.win.destroy()

print("\n== 3ж. Бальні ознаки й AUDPC")
w = m.ScoresWindow(root); root.update()
rows = []
for nm, lo_, hi_ in (("Стійкий", 1, 4), ("Середній", 3, 7), ("Вразливий", 6, 10)):
    for k in range(12): rows.append([nm, "R1", f"{rng.integers(lo_, hi_)}", "", ""])
fill(w, rows); w._refresh_traits(); w._run(); root.update(); v = find_win("Бальні ознаки — Бал")
check(v is not None, "порівняння за балами: результати відкрились"); n, errs = click_all(v); check(n == 5 and not errs, f"{n} розділів без помилок {errs[:1]}"); v.destroy()
w.mode_var.set("AUDPC (розвиток у часі)"); w._on_mode()
check(w.header_entries[2].get() == "01.06.2024", "перемикання режиму змінює заголовки на дати обліків")
for r_ in w.entries:
    for e in r_: e.delete(0, tk.END)
rows = []
for nm, rate in (("Стійкий", 2.0), ("Вразливий", 11.0)):
    for k in range(8): rows.append([nm, f"R{k+1}"] + [f"{min(100, rate * t + rng.normal(0, 1) + 1):.1f}" for t in (0, 7, 14)])
fill(w, rows); w.ymax_var.set("100"); w._run(); root.update(); v = find_win("AUDPC — результати")
check(v is not None, "AUDPC: результати відкрились"); n, errs = click_all(v); check(n == 5 and not errs, f"{n} розділів без помилок {errs[:1]}"); v.destroy()
w.header_entries[3].delete(0, tk.END); w.header_entries[3].insert(0, "вчора"); n0 = len(LOG["warn"]); w._run(); check(len(LOG["warn"]) == n0 + 1 and "вчора" in LOG["warn"][-1][1], "нерозпізнаний момент обліку → попередження з його текстом")
w.win.destroy()

print("\n== 4. Стресостійкість і шляховий аналіз")
w = m.StressWindow(root); root.update()
rows = [[f"G{i+1}", f"{yp:.2f}", f"{ys:.2f}"] for i, (yp, ys) in enumerate(zip(rng.uniform(6, 12, 12), rng.uniform(3, 8, 12)))]
fill(w, rows); w._run(); root.update(); v = find_win("Стресостійкість — результати")
check(v is not None, "стрес: результати відкрились")
n, errs = click_all(v); check(n >= 5 and not errs, f"стрес: {n} розділів без помилок {errs[:1]}"); v.destroy()
fill(w, [["G1", "-1", "2"]]); 
for r_ in w.entries[1:]:
    for e in r_: e.delete(0, tk.END)
n0 = len(LOG["warn"]); w._run(); check(len(LOG["warn"]) == n0 + 1, "замало генотипів чи недодатні значення → попередження")
w.win.destroy()
w = m.PathWindow(root); root.update()
check(len(w.header_entries) == 6, "шляховий аналіз: 5 ознак + Y"); w._add_col(); check(len(w.header_entries) == 7 and all(len(r) == 7 for r in w.entries), "додавання стовпця працює")
w._del_col()
Z = rng.normal(size=(50, 3)); y = Z @ [.7, -.3, .4] + rng.normal(0, 1, 50)
rows = [[f"{a:.3f}" for a in row] + ["", ""] + [f"{yy:.3f}"] for row, yy in zip(Z, y)]
fill(w, rows, ["Висота", "Маса 1000", "Довжина колоса", "", "", "Урожай"]); w._run(); root.update()
v = find_win("Шляховий аналіз — результати")
check(v is not None, "шляховий аналіз: результати відкрились (порожні стовпці X4, X5 проігноровано)")
n, errs = click_all(v); check(n >= 5 and not errs, f"шляховий аналіз: {n} розділів без помилок {errs[:1]}")
v.destroy(); w.win.destroy()

print("\n== 5. Аналіз стабільності: розширення")
w = m.BreedingStabilityWindow(root, {}); root.update()
check(all(hasattr(w, a) for a in ("r_var", "mse_var", "df_var", "svp_var")), "додаткові поля r / MSE / df / GGE на панелі")
check("ДОДАТКОВІ АНАЛІЗИ" in w.HELP_TEXT and "Eberhart" in w.HELP_TEXT.replace("Ебергарт", "Eberhart") or "EBERHART" in w.HELP_TEXT, "довідка містить опис додаткових аналізів")
g_, e_ = 6, 5
Xm = 6 + rng.normal(0, 1, (g_, 1)) + rng.normal(0, 1.4, (1, e_)) + np.outer(rng.normal(0, .7, g_), rng.normal(0, 1, e_)) + rng.normal(0, .25, (g_, e_))
def put_matrix(win, M):
    while win.cols_n < 1 + M.shape[1]: win._add_col()
    for r_ in win.entries:
        for e in r_: e.delete(0, tk.END)
    for i in range(M.shape[0]):
        win.entries[i][0].insert(0, f"Сорт {i+1}")
        for j in range(M.shape[1]): win.entries[i][1 + j].insert(0, f"{M[i, j]:.4f}")
put_matrix(w, Xm); w._run(); root.update()
v = find_win("Аналіз стабільності — Результати")
check(v is not None, "вікно результатів стабільності відкрилось")
plus = [b for b in buttons(v) if b.cget("text").strip().startswith("➕")]
check(len(plus) >= 14, f"додано {len(plus)} нових розділів у бічне меню")
# нижні розділи мають бути ДОСЯЖНІ прокруткою бічного меню
def canvases(x):
    out = [x] if x.winfo_class() == "Canvas" else []
    for c in x.winfo_children(): out += canvases(c)
    return out
v.geometry("1100x640"); v.update()
cv = [c for c in canvases(plus[0].master.master) if c.winfo_height() > 50][0]
cv.yview_moveto(0.0); v.update()
top_y = plus[0].winfo_rooty() - cv.winfo_rooty()
before = plus[-1].winfo_rooty() - cv.winfo_rooty()
cv.yview_moveto(1.0); v.update()
last = plus[-1]; ly = last.winfo_rooty() - cv.winfo_rooty()
check(0 <= top_y < cv.winfo_height() and before > cv.winfo_height(), f"без прокрутки останній розділ поза межами видимого (y = {before} > {cv.winfo_height()}) — тож прокрутка потрібна")
check(0 <= ly and ly + last.winfo_height() <= cv.winfo_height() + 2, f"після прокрутки донизу «{last.cget('text').strip()}» повністю видно (y = {ly}, область {cv.winfo_height()})")
cv.yview_moveto(0.0); v.update()
errs = []
for b in plus:
    try: b.invoke(); v.update()
    except Exception as ex: errs.append(f"{b.cget('text').strip()}: {ex}")
check(not errs, f"усі нові розділи (таблиці й графіки) відображаються без помилок {errs[:1]}")
old = [b for b in buttons(v) if "Звіт" in b.cget("text") or "GGE Biplot" in b.cget("text")]
for b in old: b.invoke(); v.update()
check(len(old) == 2, "старі розділи (звіт ER, GGE biplot) збережено й працюють")
er = w._st_data["er_rows"]
check(all(("b не відрізняється" in r[4] or "b >" in r[4] or "b <" in r[4] or "н/д" in r[4]) for r in er)
      and not any("Стабільний (bi≈1" in r[4] for r in er), "висновок ER — за t-тестом, а не за довільними порогами")
SAVE["p"] = os.path.join(TMP, "st.xlsx"); [b for b in buttons(v) if "Excel — нові" in b.cget("text")][0].invoke()
wb = openpyxl.load_workbook(SAVE["p"]); check(len(wb.sheetnames) >= 10, f"Excel нових розділів: {len(wb.sheetnames)} аркушів")
SAVE["p"] = os.path.join(TMP, "st.docx"); [b for b in buttons(v) if "Word — нові" in b.cget("text")][0].invoke()
d = docx.Document(SAVE["p"]); check(len(d.tables) >= 10 and len(d.inline_shapes) >= 4, f"Word нових розділів: таблиць {len(d.tables)}, рисунків {len(d.inline_shapes)}")
v.destroy()

print("\n== 6. Довга таблиця → похибка, ANOVA, строгі тести")
gl, el, rl_, yl = [], [], [], []
ga, ea, ge = rng.normal(0, 1, 6), rng.normal(0, 2, 5), rng.normal(0, .6, (6, 5))
for i in range(6):
    for j in range(5):
        for k in range(3):
            gl.append(f"Сорт {i+1}"); el.append(f"{2019+j}"); rl_.append(f"R{k+1}"); yl.append(10 + ga[i] + ea[j] + ge[i, j] + rng.normal(0, .5))
txt = "Генотип\tСередовище\tПовторність\tЗначення\n" + "\n".join(f"{a}\t{b}\t{c}\t{v_:.3f}".replace(".", ",") for a, b, c, v_ in zip(gl, el, rl_, yl))
w.win.clipboard_clear(); w.win.clipboard_append(txt); n0 = len(LOG["info"]) + len(LOG["warn"]); w._import_long(); root.update()
check(len(LOG["info"]) + len(LOG["warn"]) == n0 + 1, "імпорт повідомив підсумок (рядки, MSE, Бартлетт)")
check(w.r_var.get() == "3" and float(w.mse_var.get()) > 0 and w.df_var.get() == str(5 * 5 * 2), f"поля заповнено: r = {w.r_var.get()}, MSE = {w.mse_var.get()}, df = {w.df_var.get()}")
check(w.entries[0][0].get() == "Сорт 1" and w.env_vars[0].get() == "2019" and w.cols_n == 6, "таблиця середніх заповнена, середовища названо")
w._run(); root.update(); v = find_win("Аналіз стабільності — Результати")
titles = [b.cget("text").strip() for b in buttons(v)]
check(any("ANOVA (за повторностями)" in t for t in titles), "є розділ «ANOVA (за повторностями)» із комбінованого аналізу")
for b in [b for b in buttons(v) if "Параметричні" in b.cget("text")]: b.invoke(); v.update()
v.destroy()
# змінили таблицю після імпорту → ANOVA за повторностями не застосовується
w.entries[0][1].delete(0, tk.END); w.entries[0][1].insert(0, "99")
w._run(); root.update(); v = find_win("Аналіз стабільності — Результати")
titles = [b.cget("text").strip() for b in buttons(v)]
check(any("ANOVA (за середніми)" in t for t in titles) and not any("за повторностями" in t for t in titles),
      "таблицю змінено після імпорту → ANOVA за повторностями не застосовується (щоб не змішувати різні дані)")
v.destroy()
# пропуск у таблиці
w.entries[1][2].delete(0, tk.END); w._run(); root.update(); v = find_win("Аналіз стабільності — Результати")
titles = [b.cget("text").strip() for b in buttons(v)]
check(any("Додаткові аналізи недоступні" in t for t in titles), "порожня клітинка → додаткові аналізи недоступні з поясненням, старі розділи працюють")
for b in [b for b in buttons(v) if "Додаткові аналізи недоступні" in b.cget("text")]: b.invoke(); v.update()
v.destroy()
# погані r/MSE
w.r_var.set("x"); n0 = len(LOG["warn"]); w._run(); root.update(); v = find_win("Аналіз стабільності — Результати")
check(v is not None and any("Додаткові аналізи недоступні" in b.cget("text") for b in buttons(v)), "нечисловий r → зрозуміла відмова додаткових аналізів")
v.destroy(); w.win.destroy()
n0 = len(LOG["warn"])
w = m.BreedingStabilityWindow(root, {}); w.win.clipboard_clear(); w.win.clipboard_append("A\tB\tC\n1\t2\t3"); w._import_long()
check(len(LOG["warn"]) == n0 + 1, "некоректний вміст буфера для довгої таблиці → попередження")
w.win.destroy()

shutil.rmtree(TMP, ignore_errors=True); root.destroy()
print("\n" + "=" * 60)
if FAIL:
    print(f"ПРОВАЛІВ: {len(FAIL)}"); [print("  -", f) for f in FAIL]; sys.exit(1)
print("УСІ ПЕРЕВІРКИ ВІКОН ПРОЙДЕНО")
