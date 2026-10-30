# sad_qrlabels.py – S.A.D. v2.1: QR-таблички для ділянок досліду (без tkinter)
# -*- coding: utf-8 -*-
"""
Для кожної ділянки (повторність × варіант) формується QR-код і табличка, яку дослідник сканує телефоном біля ділянки:
звичайна камера показує зрозумілий текст (дослід, повторність, варіант, рівні факторів, місце), а майбутній польовий журнал
розпізнає службовий рядок формату SAD1 (див. QR_FORMAT.md) і одразу знає, до якої ділянки вносити значення.

Захист від помилок розміщення: унікальний код досліду; наскрізний номер ділянки в порядку розкладання; контрольний символ
(Luhn mod 36); підпис схеми (змінилась схема – старі таблички виявляються); колір варіанта; сусіди ліворуч/праворуч;
карта розкладки й чек-лист.
"""
import os, io, re, json, math, hashlib, datetime as _dt
import itertools
import segno
from reportlab.pdfgen import canvas as _rl
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

FACTOR_LETTERS = "ABCDEFGHIJ"
ID_ALPHABET = "23456789ABCDEFGHJKMNPQRSTUVWXYZ"           # без 0, 1, I, L, O – щоб не плутати при ручному введенні
CHECK_ALPHABET = "0123456789ABCDEFGHIJKLMNOPQRSTUVWXYZ"
FORMAT = "SAD1"
PT = 25.4 / 72.0                                           # міліметрів в одному пункті
PALETTE = ["#1f77b4", "#d62728", "#2ca02c", "#ff7f0e", "#9467bd", "#8c564b", "#e377c2", "#17becf", "#bcbd22", "#7f7f7f", "#0b3d91", "#a00000"]


# ═══════════════════════════════════════════════════════════════
# 1. Ідентифікатори й контрольний символ
# ═══════════════════════════════════════════════════════════════
def luhn_mod_n_check(s, alphabet=CHECK_ALPHABET):
    """Контрольний символ за алгоритмом Luhn mod N (виявляє будь-яку одиничну помилку та більшість перестановок сусідніх символів)."""
    n = len(alphabet); factor = 2; total = 0
    for ch in reversed(s):
        addend = factor * alphabet.index(ch); factor = 1 if factor == 2 else 2
        total += addend // n + addend % n
    return alphabet[(n - total % n) % n]


def luhn_mod_n_valid(s, alphabet=CHECK_ALPHABET):
    n = len(alphabet); factor = 1; total = 0
    for ch in reversed(s):
        addend = factor * alphabet.index(ch); factor = 2 if factor == 1 else 1
        total += addend // n + addend % n
    return total % n == 0


def check_char(exp_id, plot, rep, variant):
    """Контрольний символ ділянки: за рядком E + P(4 цифри) + R(3) + V(3) – фіксована ширина, щоб додаток міг перерахувати."""
    return luhn_mod_n_check(f"{exp_id}{plot:04d}{rep:03d}{variant:03d}")


def valid_exp_id(s):
    return bool(s) and 3 <= len(s) <= 6 and all(ch in ID_ALPHABET for ch in s)


def suggest_exp_id(name, salt=""):
    """Детермінований код досліду з 4 символів за назвою (однакова назва → той самий код)."""
    h = hashlib.sha1((name.strip().lower() + "|" + salt).encode("utf-8")).digest()
    return "".join(ID_ALPHABET[b % len(ID_ALPHABET)] for b in h[:4])


def human_code(exp_id, plot, rep, variant):
    return f"{exp_id}-{plot:03d}-{check_char(exp_id, plot, rep, variant)}"


def scheme_signature(factor_defs, plots):
    """Підпис схеми (4 шістнадцяткові символи): змінилась схема – підпис інший, тож старі таблички не сплутаються з новими."""
    canon = json.dumps({"f": [[fd["name"], list(fd["levels"])] for fd in factor_defs],
                        "p": [[p["plot"], p["rep"], p["variant"], p["row"], p["segments"][0][1]] for p in plots]},
                       ensure_ascii=False, sort_keys=True)
    return hashlib.sha1(canon.encode("utf-8")).hexdigest()[:4].upper()


# ═══════════════════════════════════════════════════════════════
# 2. Ділянки зі збереженої схеми (обидві моделі)
# ═══════════════════════════════════════════════════════════════
def _segments(cells):
    """[(ряд, позиція)] → [(ряд, від, до)] послідовні відрізки в межах ряду."""
    out = []
    for r, p in sorted(cells):
        if out and out[-1][0] == r and p == out[-1][2] + 1: out[-1] = (r, out[-1][1], p)
        else: out.append((r, p, p))
    return out


def _pos_text(segs, per_plot_cells):
    """Місце ділянки словами: «ряд 1, місце 6» (площинна модель) або «ряд 2: рослини 5–6, 8» (група рослин; кілька рядів – через «;»)."""
    if per_plot_cells == 1:
        r, a, _ = segs[0]; return f"ряд {r}, місце {a}"
    by_row = {}
    for r, a, b in segs: by_row.setdefault(r, []).append(f"{a}–{b}" if a != b else str(a))
    parts = []
    for r, rng in list(by_row.items())[:3]:
        parts.append(f"ряд {r}: " + ("рослина " if len(rng) == 1 and "–" not in rng[0] else "рослини ") + ", ".join(rng))
    return "; ".join(parts) + (f" (+{len(by_row) - 3})" if len(by_row) > 3 else "")


def build_experiment(scheme, exp_id, name="", origin="", created=None):
    """Формує опис досліду й ділянок зі словника схеми (формат .sadp). Повертає dict(exp=…, variants=…, plots=…, grid=…, skipped=n)."""
    if not valid_exp_id(exp_id): raise ValueError("Код досліду має містити 3–6 символів із набору 2–9, A–Z (без I, L, O).")
    fdefs = scheme.get("factor_defs") or []
    if not fdefs: raise ValueError("У схемі немає факторів.")
    cfg = scheme.get("cfg", {}); nlev = [len(fd["levels"]) for fd in fdefs]
    variants = []; var_no = {}
    for n, idxs in enumerate(itertools.product(*[range(1, k + 1) for k in nlev]), 1):
        vcode = "".join(f"{FACTOR_LETTERS[i]}{idxs[i]}" for i in range(len(idxs)))
        variants.append({"n": n, "code": vcode, "idxs": list(idxs), "levels": [fdefs[i]["levels"][idxs[i] - 1] for i in range(len(idxs))]})
        var_no[idxs] = n
    groups = {}; skipped = 0; grid_rows = {}
    for p in scheme.get("plants", []):
        row, pos = int(p["row"]), int(p["position"]); rec = p.get("role") == "recorded"
        if not rec:
            grid_rows.setdefault(row, {})[pos] = ("other", None); continue
        f = p.get("factors") or {}
        try: idxs = tuple(int(f[fd["name"]]) for fd in fdefs)
        except (KeyError, TypeError, ValueError): skipped += 1; continue
        if any(not (1 <= i <= k) for i, k in zip(idxs, nlev)): skipped += 1; continue
        groups.setdefault((int(p.get("replication") or 1), idxs), []).append((row, pos))
    if not groups: raise ValueError("У схемі немає жодної облікової ділянки (рядків з кодами факторів).")
    ordered = sorted(groups.items(), key=lambda kv: (min(kv[1]), kv[0][0], var_no[kv[0][1]]))
    plots = []
    for no, ((rep, idxs), cells) in enumerate(ordered, 1):
        segs = _segments(cells); v = var_no[idxs]
        plots.append({"plot": no, "rep": rep, "variant": v, "vcode": variants[v - 1]["code"], "levels": variants[v - 1]["levels"],
                      "row": segs[0][0], "segments": segs, "units": len(cells), "pos_text": _pos_text(segs, len(cells)),
                      "left": None, "right": None})
        for r, a, b in segs:
            for q in range(a, b + 1): grid_rows.setdefault(r, {})[q] = ("plot", no)
    by_row = {}
    for pl in plots: by_row.setdefault(pl["row"], []).append(pl)
    for lst in by_row.values():
        lst.sort(key=lambda x: x["segments"][0][1])
        for i, pl in enumerate(lst):
            pl["left"] = lst[i - 1]["plot"] if i > 0 else None; pl["right"] = lst[i + 1]["plot"] if i + 1 < len(lst) else None
    for pl in plots:
        pl["check"] = check_char(exp_id, pl["plot"], pl["rep"], pl["variant"]); pl["code"] = human_code(exp_id, pl["plot"], pl["rep"], pl["variant"])
    sig = scheme_signature(fdefs, plots)
    exp = {"id": exp_id, "name": name.strip(), "created": created or _dt.date.today().isoformat(), "model": cfg.get("model", "plant"),
           "design": cfg.get("design_used", ""), "origin": origin.strip(), "signature": sig, "n_plots": len(plots), "n_variants": len(variants),
           "n_reps": max(pl["rep"] for pl in plots), "factors": [{"letter": FACTOR_LETTERS[i], "name": fd["name"], "levels": list(fd["levels"])} for i, fd in enumerate(fdefs)],
           "trait": cfg.get("trait_name", ""), "per_plot_units": 1 if cfg.get("model") == "area" else max(pl["units"] for pl in plots)}
    for pl in plots: pl["payload"], pl["tier"] = "", ""
    grid = {r: dict(sorted(c.items())) for r, c in sorted(grid_rows.items())}
    return {"exp": exp, "variants": variants, "plots": plots, "grid": grid, "skipped": skipped}


# ═══════════════════════════════════════════════════════════════
# 3. Вміст QR
# ═══════════════════════════════════════════════════════════════
def machine_line(exp, pl):
    return f"{FORMAT};E={exp['id']};P={pl['plot']};R={pl['rep']};V={pl['variant']};S={exp['signature']};C={pl['check']}"


def parse_payload(text):
    """Розбирає вміст QR (описовий чи службовий) → dict(E,P,R,V,S,C, valid). Кидає ValueError, якщо це не табличка S.A.D."""
    m = re.search(rf"{FORMAT};([^\n\r]*)", text)
    if not m: raise ValueError("Це не QR-табличка S.A.D. (немає службового рядка SAD1).")
    kv = dict(part.split("=", 1) for part in m.group(1).split(";") if "=" in part)
    for k in ("E", "P", "R", "V", "S", "C"):
        if k not in kv: raise ValueError(f"У службовому рядку немає поля {k}.")
    try: p, r, v = int(kv["P"]), int(kv["R"]), int(kv["V"])
    except ValueError: raise ValueError("Номери ділянки, повторності й варіанта мають бути цілими.")
    try: ok = luhn_mod_n_valid(f"{kv['E']}{p:04d}{r:03d}{v:03d}{kv['C']}")
    except ValueError: ok = False
    return {"E": kv["E"], "P": p, "R": r, "V": v, "S": kv["S"], "C": kv["C"], "valid": ok}


def _short(s, n):
    s = " ".join(str(s).split()); return s if len(s) <= n else s[:n - 1] + "…"


def payload_tiers(exp, pl, style="descriptive"):
    """Варіанти вмісту від найповнішого до мінімального (щоб за потреби зменшити щільність QR; кирилиця займає по 2 байти)."""
    mach = machine_line(exp, pl)
    if style == "compact": return [("min", mach)]
    L1 = f"Дослід {exp['id']}"
    L2 = f"Повторність {pl['rep']} · Варіант {pl['variant']} ({pl['vcode']})"
    L3 = _short(" + ".join(pl["levels"]), 56)
    L4 = f"№{pl['plot']} з {exp['n_plots']} · {_short(pl['pos_text'], 34)}"
    named = L1 + (f" · {_short(exp['name'], 24)}" if exp["name"] else "")
    return [("full", "\n".join([named, L2, L3, L4, mach])), ("no_name", "\n".join([L1, L2, L3, L4, mach])),
            ("no_levels", "\n".join([L1, L2, L4, mach])), ("min", mach)]


MIN_MODULE_MM = 0.40                                   # менший модуль QR гірше читається телефоном на відкритому повітрі


MIN_PX_PER_MODULE = 3.2        # виміряно: за ≥ 3 пікселів камери на модуль QR (з розмиттям і нахилом) читаються 100 % кодів, за 2,7 – вже ні


def scan_distance_cm(key, version, frame_px=1920, fov_deg=65.0):
    """Орієнтовна найбільша відстань, з якої QR читається камерою з frame_px пікселів по ширині кадру (за MIN_PX_PER_MODULE)."""
    return module_mm(key, version) * frame_px / (MIN_PX_PER_MODULE * 2 * math.tan(math.radians(fov_deg / 2))) / 10.0


def version_cap(key, ecc="m"):
    """Найбільша версія QR, за якої модуль на табличці ≥ MIN_MODULE_MM (не більше 16). Для рулонних етикеток межа – 3 точки
    принтера 203 dpi (0,375 мм): термопринтер друкує такий модуль чітко, і це поширена рекомендація для QR."""
    lim = 3 * 25.4 / 203.0 - 1e-9 if LABEL_PRESETS[key].get("media") == "roll" else MIN_MODULE_MM
    v = 1
    while v < 16 and module_mm(key, v + 1) >= lim: v += 1
    return v


def make_qr(text, ecc="m"):
    return segno.make(text, error=ecc.lower(), boost_error=False, micro=False)


def fit_payload(exp, pl, style="descriptive", ecc="m", max_version=10):
    """Обирає найповніший вміст, що вкладається в max_version. → (текст, рівень, qr)."""
    tiers = payload_tiers(exp, pl, style)
    for tier, text in tiers:
        q = make_qr(text, ecc)
        if q.version <= max_version: return text, tier, q
    return tiers[-1][1], tiers[-1][0], make_qr(tiers[-1][1], ecc)


def fill_payloads(data, style="descriptive", ecc="m", key="S"):
    """Заповнює plots[i]['payload'], ['tier'], ['qr_version'] для обраного розміру таблички; повертає найбільшу версію QR."""
    mx = 0; max_version = version_cap(key, ecc)
    for pl in data["plots"]:
        text, tier, q = fit_payload(data["exp"], pl, style, ecc, max_version)
        pl["payload"], pl["tier"], pl["qr_version"] = text, tier, q.version; mx = max(mx, q.version)
    return mx


def qr_matrix(text, ecc="m"):
    q = make_qr(text, ecc)
    return [[bool(c) for c in row] for row in q.matrix], q.version


# ═══════════════════════════════════════════════════════════════
# 4. Шрифти й вимірювання тексту (DejaVu із matplotlib – містить кирилицю)
# ═══════════════════════════════════════════════════════════════
_FONTS = {}
FONT, FONT_B = "SADSans", "SADSans-Bold"


def font_paths():
    if not _FONTS:
        import matplotlib
        d = os.path.join(matplotlib.get_data_path(), "fonts", "ttf")
        _FONTS.update(regular=os.path.join(d, "DejaVuSans.ttf"), bold=os.path.join(d, "DejaVuSans-Bold.ttf"))
        pdfmetrics.registerFont(TTFont(FONT, _FONTS["regular"])); pdfmetrics.registerFont(TTFont(FONT_B, _FONTS["bold"]))
    return _FONTS


def tw(text, size, bold=False):
    """Ширина тексту в мм."""
    font_paths(); return pdfmetrics.stringWidth(text, FONT_B if bold else FONT, size) * PT


def fit_size(text, size, max_w, bold=False, min_size=5.0):
    while size > min_size and tw(text, size, bold) > max_w: size -= 0.25
    return size


def wrap_text(text, size, max_w, max_lines, bold=False):
    words = text.split(); lines = []; cur = ""
    for w in words:
        t = (cur + " " + w).strip()
        if tw(t, size, bold) <= max_w: cur = t; continue
        if cur: lines.append(cur)
        while tw(w, size, bold) > max_w and len(w) > 1:
            k = len(w)
            while k > 1 and tw(w[:k], size, bold) > max_w: k -= 1
            lines.append(w[:k]); w = w[k:]
        cur = w
    if cur: lines.append(cur)
    if len(lines) > max_lines:
        lines = lines[:max_lines]; last = lines[-1]
        while last and tw(last + "…", size, bold) > max_w: last = last[:-1]
        lines[-1] = last + "…"
    return lines


def _lum(hexs):
    h = hexs.lstrip("#"); r, g, b = (int(h[i:i + 2], 16) / 255 for i in (0, 2, 4)); return 0.2126 * r + 0.7152 * g + 0.0722 * b


def variant_color(v): return PALETTE[(v - 1) % len(PALETTE)]
def on_color(col): return "#ffffff" if _lum(col) < 0.55 else "#111111"


# ═══════════════════════════════════════════════════════════════
# 5. Макети табличок (примітиви в мм, початок координат – лівий верхній кут)
# ═══════════════════════════════════════════════════════════════
LABEL_PRESETS = {
    "S": {"name": "Мала 63,5×38,1 мм (21 на аркуші)", "w": 63.5, "h": 38.1, "cols": 3, "rows": 7, "qr": 27.0, "layout": "h", "k": 1.0},
    "M": {"name": "Середня 99,1×67,7 мм (8 на аркуші)", "w": 99.1, "h": 67.7, "cols": 2, "rows": 4, "qr": 44.0, "layout": "h", "k": 1.55},
    "L": {"name": "Велика 100×141 мм (4 на аркуші)", "w": 100.0, "h": 141.0, "cols": 2, "rows": 2, "qr": 76.0, "layout": "v", "k": 1.6},
}

# ── Рулонні етикетки для принтерів етикеток (Xprinter, TSC, Godex, Zebra…): одна етикетка – одна сторінка PDF розміром з етикетку
ROLL_SIZES = [(40, 30), (50, 30), (58, 30), (58, 40), (60, 40), (70, 50), (75, 50), (100, 50), (100, 70), (100, 100), (100, 150)]
ROLL_W_RANGE = (20.0, 118.0)          # ширина носія, мм (типово для 4-дюймових принтерів: 20–118 мм; ширина друку – до 108 мм)
ROLL_H_RANGE = (15.0, 300.0)
PRINTER_PRINT_W = 108.0               # найбільша ширина друку 4-дюймового принтера, мм
ROLL_DPI = (203, 300)


def _num(x): return f"{x:g}".replace(".", ",")


def make_roll_preset(w, h):
    """Параметри рулонної етикетки w × h мм (w – поперек стрічки, як у налаштуваннях принтера). Компонування й розмір QR підбираються
    за пропорціями тими самими правилами, що й для аркушевих табличок (k – масштаб відносно малої таблички 63,5 × 38,1 мм)."""
    w, h = float(w), float(h)
    if w / h >= 1.2:                                                  # горизонтальна: QR ліворуч, текст праворуч
        layout = "h"; k = max(0.6, min(2.2, min(h / 38.1, w / 63.5)))
        qr = min(h - 5.0 * k - 0.6 * k - 4.2 * k - 1.0 * k, 0.45 * w)
    else:                                                             # вертикальна: QR угорі, текст під ним
        layout = "v"; k = max(0.75, min(2.0, min(w / 100.0, h / 141.0) * 1.6))
        qr = min(w - 6.0 * k, h - 5.0 * k - 2.0 * k - 33.6 * k)
    return {"name": f"Рулон {_num(w)}×{_num(h)} мм (принтер етикеток)", "w": w, "h": h, "cols": 1, "rows": 1, "qr": round(max(qr, 10.0), 2),
            "layout": layout, "k": round(k, 3), "media": "roll"}


def roll_key(w, h): return f"R{float(w):g}x{float(h):g}"


def register_roll(w, h):
    """Додає рулонний розмір до переліку (якщо його ще немає). → ключ. ValueError – розмір поза межами носія."""
    w, h = float(w), float(h)
    if not (ROLL_W_RANGE[0] <= w <= ROLL_W_RANGE[1]): raise ValueError(f"Ширина етикетки має бути від {_num(ROLL_W_RANGE[0])} до {_num(ROLL_W_RANGE[1])} мм.")
    if not (ROLL_H_RANGE[0] <= h <= ROLL_H_RANGE[1]): raise ValueError(f"Висота етикетки має бути від {_num(ROLL_H_RANGE[0])} до {_num(ROLL_H_RANGE[1])} мм.")
    key = roll_key(w, h)
    if key not in LABEL_PRESETS: LABEL_PRESETS[key] = make_roll_preset(w, h)
    return key


def parse_roll_key(key):
    """'R58x40' → (58.0, 40.0) або None."""
    import re as _re
    m = _re.fullmatch(r"R(\d+(?:\.\d+)?)x(\d+(?:\.\d+)?)", str(key or ""))
    return (float(m.group(1)), float(m.group(2))) if m else None


def is_roll(key): return LABEL_PRESETS.get(key, {}).get("media") == "roll"


for _w, _h in ROLL_SIZES: register_roll(_w, _h)


def snap_qr(x, y, size, n_modules, dpi):
    """Підганяє QR під сітку точок принтера: модуль – ціла кількість точок (не менше 2), лівий верхній кут – на точці сітки;
    QR центрується в межах початкової ділянки. n_modules – сторона матриці без тихої зони. → (x, y, size)."""
    dot = 25.4 / float(dpi); n = n_modules + 8
    m = max(2, int(math.floor(size / n / dot + 1e-9))) * dot; new = m * n
    x2 = round((x + (size - new) / 2.0) / dot) * dot; y2 = round((y + (size - new) / 2.0) / dot) * dot
    return x2, y2, new


def module_mm(key, version):
    """Розмір модуля QR, мм. Для рулонних етикеток – після прив'язки до сітки точок принтера 203 dpi (найгрубша поширена роздільність),
    щоб підбір вмісту QR (version_cap) враховував фактичний модуль."""
    P = LABEL_PRESETS[key]; m = P["qr"] / (17 + 4 * version + 8)
    if P.get("media") == "roll":
        dot = 25.4 / 203.0; m = max(2, int(math.floor(m / dot + 1e-9))) * dot
    return m


def _neighbors_text(exp, pl, by_no):
    def f(n):
        if n is None: return None
        o = by_no[n]; return f"№{n} (В{o['variant']})"
    l, r = f(pl["left"]), f(pl["right"])
    return "← " + (l or "початок ряду") + "      " + (r or "кінець ряду") + " →"


def _fit_lines(levels_long, levels_short, size, width, y, limit, step):
    """Рядки опису рівнів: спершу «Фактор: рівень», а якщо не вміщується за висотою – лише назви рівнів; не вище за limit."""
    max_lines = max(1, int((limit - y) / step))
    long_lines = wrap_text("; ".join(levels_long), size, width, 99)
    if len(long_lines) <= max_lines: return long_lines
    return wrap_text(" + ".join(levels_short.split(" + ")), size, width, max_lines)


def label_prims(data, pl, key, qr, opts=None):
    """Список примітивів однієї таблички: ('rect', x, y, w, h, fill, stroke, lw) ('text', x, y_base, s, pt, bold, align, color) ('qr', x, y, size, matrix)."""
    opts = opts or {}; exp = data["exp"]; by_no = {p["plot"]: p for p in data["plots"]}
    P = LABEL_PRESETS[key]; w, h, k = P["w"], P["h"], P["k"]
    col = variant_color(pl["variant"]); fg = on_color(col); ink = "#111111"; grey = "#555555"
    if opts.get("mono"):                                              # термодрук: лише чорне й біле (сірі заливки й сірий текст друкуються растром)
        col, fg, ink, grey = "#000000", "#ffffff", "#000000", "#000000"
    prims = [("rect", 0, 0, w, h, "#ffffff", "#9a9a9a", 0.15)]
    band = 5.0 * k
    prims += [("rect", 0, 0, w, band, col, None, 0),
              ("text", 1.8 * k, band * 0.72, f"ДОСЛІД {exp['id']}", fit_size(f"ДОСЛІД {exp['id']}", 7.5 * k, w * 0.5, True), True, "l", fg),
              ("text", w - 1.8 * k, band * 0.72, f"№ {pl['plot']} з {exp['n_plots']}", fit_size(f"№ {pl['plot']} з {exp['n_plots']}", 7.5 * k, w * 0.42, True), True, "r", fg)]
    qs = P["qr"]; dpi = opts.get("snap_dpi")
    def qr_at(x, y, size):                                            # з прив'язкою до сітки точок принтера (рулонні етикетки)
        return snap_qr(x, y, size, len(qr), dpi) if dpi else (x, y, size)
    levels_short = " + ".join(pl["levels"]); levels_long = [f"{f['name']}: {lv}" for f, lv in zip(exp["factors"], pl["levels"])]
    foot = f"схема {exp['signature']}"
    if P["layout"] == "h":
        qx, qy, qs = qr_at(1.0 * k, band + 0.6 * k, qs)
        prims.append(("qr", qx, qy, qs, qr))
        prims.append(("text", qx + qs / 2, qy + qs + 2.9 * k, pl["code"], fit_size(pl["code"], 6.4 * k, qs, True), True, "c", ink))
        rx = qx + qs + 1.0 * k; rw = w - rx - 1.4 * k; y = band + 1.0 * k
        if k < 1.3:                                                   # мала табличка
            txt = f"П{pl['rep']} · В{pl['variant']}"; big = fit_size(txt, 19 * k, rw, True)
            y += big * PT * 0.85; prims.append(("text", rx, y, txt, big, True, "l", ink))
            y += 3.5 * k; cap = f"повторність {pl['rep']} · варіант {pl['variant']}"
            prims.append(("text", rx, y, cap, fit_size(cap, 6.0 * k, rw), False, "l", grey))
            y += 3.9 * k; prims.append(("text", rx, y, pl["vcode"], fit_size(pl["vcode"], 8.5 * k, rw, True), True, "l", ink))
            for ln in wrap_text(levels_short, 7.0 * k, rw, 2):
                y += 3.2 * k; prims.append(("text", rx, y, ln, 7.0 * k, False, "l", ink))
            pos = _short(pl["pos_text"], 40); prims.append(("text", rx, h - 1.5 * k, pos, fit_size(pos, 6.2 * k, rw), False, "l", grey))
            prims.append(("text", w - 1.2 * k, h - 1.5 * k, foot, 4.6 * k, False, "r", grey)) if tw(pos, 6.2 * k) + tw(foot, 4.6 * k) < rw - 1 else None
        else:                                                         # середня табличка
            for txt, sz in ((f"Повторність {pl['rep']}", 13.5 * k), (f"Варіант {pl['variant']}", 20 * k)):
                sz = fit_size(txt, sz, rw, True); y += sz * PT * 0.95; prims.append(("text", rx, y, txt, sz, True, "l", ink)); y += 0.8 * k
            vt = f"код варіанта {pl['vcode']}"; y += 3.0 * k; prims.append(("text", rx, y, vt, fit_size(vt, 8.0 * k, rw, True), True, "l", ink))
            yb = h - 1.8 * k; limit = (yb - 8.6 * k if opts.get("neighbors", True) else yb - 4.4 * k) - 4.2 * k
            for ln in _fit_lines(levels_long, levels_short, 7.4 * k, rw, y, limit, 3.4 * k):
                y += 3.4 * k; prims.append(("text", rx, y, ln, 7.4 * k, False, "l", ink))
            if opts.get("neighbors", True):
                nb = _neighbors_text(exp, pl, by_no); prims.append(("text", rx, yb - 8.6 * k, nb, fit_size(nb, 6.6 * k, rw), False, "l", ink))
            pos = _short(pl["pos_text"], 44); prims.append(("text", rx, yb - 4.4 * k, pos, fit_size(pos, 6.6 * k, rw), False, "l", grey))
            nm = _short(exp["name"], 34) if exp["name"] else ""; prims.append(("text", rx, yb, (nm + "  ·  " if nm else "") + foot, fit_size((nm + "  ·  " if nm else "") + foot, 6.0 * k, rw), False, "l", grey))
    else:                                                             # велика табличка (вертикальна)
        qx, qy, qs = qr_at((w - qs) / 2, band + 2.0 * k, qs)
        prims.append(("qr", qx, qy, qs, qr)); y = qy + qs + 1.0 * k
        prims.append(("text", w / 2, y + 2.8 * k, pl["code"], fit_size(pl["code"], 9.0 * k, w - 4, True), True, "c", ink)); y += 6.5 * k
        for txt, sz in ((f"Повторність {pl['rep']} · Варіант {pl['variant']}", 13.5 * k),):
            sz = fit_size(txt, sz, w - 6 * k, True); y += sz * PT * 0.95; prims.append(("text", w / 2, y, txt, sz, True, "c", ink)); y += 1.0 * k
        vt = f"код варіанта {pl['vcode']}"; y += 2.6 * k; prims.append(("text", w / 2, y, vt, fit_size(vt, 8.5 * k, w - 6 * k, True), True, "c", ink))
        yb = h - 1.8 * k; limit = (yb - 8.6 * k if opts.get("neighbors", True) else yb - 4.4 * k) - 4.4 * k
        for ln in _fit_lines(levels_long, levels_short, 7.8 * k, w - 6 * k, y, limit, 3.4 * k):
            y += 3.4 * k; prims.append(("text", w / 2, y, ln, 7.8 * k, False, "c", ink))
        if opts.get("neighbors", True):
            nb = _neighbors_text(exp, pl, by_no); prims.append(("text", w / 2, yb - 8.6 * k, nb, fit_size(nb, 6.8 * k, w - 4), False, "c", ink))
        pos = _short(pl["pos_text"], 60); prims.append(("text", w / 2, yb - 4.4 * k, pos, fit_size(pos, 6.8 * k, w - 4), False, "c", grey))
        nm = (_short(exp["name"], 40) + "  ·  " if exp["name"] else "") + foot; prims.append(("text", w / 2, yb, nm, fit_size(nm, 6.2 * k, w - 4), False, "c", grey))
    return prims


def text_bbox(pr):
    """(x1, y1, x2, y2) тексту в мм (наближено за висотою шрифту)."""
    _, x, yb, s, sz, bold, al, _c = pr; wd = tw(s, sz, bold); hh = sz * PT
    x1 = x if al == "l" else (x - wd / 2 if al == "c" else x - wd)
    return (x1, yb - hh * 0.78, x1 + wd, yb + hh * 0.22)


# ═══════════════════════════════════════════════════════════════
# 6. Малювання: PDF (вектор) і PNG
# ═══════════════════════════════════════════════════════════════
def _rgb(hexs):
    h = hexs.lstrip("#"); return tuple(int(h[i:i + 2], 16) / 255.0 for i in (0, 2, 4))


def _qr_runs(matrix):
    for r, row in enumerate(matrix):
        c = 0
        while c < len(row):
            if row[c]:
                s = c
                while c < len(row) and row[c]: c += 1
                yield r, s, c - s
            else: c += 1


def draw_prims_pdf(c, prims, ox, oy, page_h):
    """ox, oy – зсув таблички від лівого верхнього кута сторінки, мм."""
    mm = 72.0 / 25.4
    for pr in prims:
        kind = pr[0]
        if kind == "rect":
            _, x, y, w, h, fill, stroke, lw = pr
            if fill: c.setFillColorRGB(*_rgb(fill))
            if stroke: c.setStrokeColorRGB(*_rgb(stroke)); c.setLineWidth(max(lw, 0.05) * mm)
            c.rect((ox + x) * mm, page_h - (oy + y + h) * mm, w * mm, h * mm, stroke=1 if stroke else 0, fill=1 if fill else 0)
        elif kind == "text":
            _, x, yb, s, sz, bold, al, col = pr
            c.setFillColorRGB(*_rgb(col)); c.setFont(FONT_B if bold else FONT, sz)
            X, Y = (ox + x) * mm, page_h - (oy + yb) * mm
            (c.drawString if al == "l" else c.drawCentredString if al == "c" else c.drawRightString)(X, Y, s)
        elif kind == "qr":
            _, x, y, size, mat = pr
            n = len(mat) + 8; m = size / n
            c.setFillColorRGB(0, 0, 0)
            for r, s, ln in _qr_runs(mat):
                c.rect((ox + x + (4 + s) * m - 0.01) * mm, page_h - (oy + y + (4 + r + 1) * m + 0.01) * mm, (ln * m + 0.02) * mm, (m + 0.02) * mm, stroke=0, fill=1)
        elif kind == "line":
            _, x1, y1, x2, y2, col, lw = pr
            c.setStrokeColorRGB(*_rgb(col)); c.setLineWidth(lw * mm); c.line((ox + x1) * mm, page_h - (oy + y1) * mm, (ox + x2) * mm, page_h - (oy + y2) * mm)


def draw_prims_png(d, prims, dpi):
    from PIL import ImageFont
    s = dpi / 25.4; fp = font_paths(); fonts = {}
    def font(bold, size):
        key = (bold, round(size * dpi / 72))
        if key not in fonts: fonts[key] = ImageFont.truetype(fp["bold" if bold else "regular"], key[1])
        return fonts[key]
    for pr in prims:
        kind = pr[0]
        if kind == "rect":
            _, x, y, w, h, fill, stroke, lw = pr
            d.rectangle([round(x * s), round(y * s), round((x + w) * s) - 1, round((y + h) * s) - 1], fill=fill, outline=stroke, width=max(1, round(lw * s)) if stroke else 0)
        elif kind == "text":
            _, x, yb, t, sz, bold, al, col = pr
            d.text((x * s, yb * s), t, font=font(bold, sz), fill=col, anchor={"l": "ls", "c": "ms", "r": "rs"}[al])
        elif kind == "qr":
            _, x, y, size, mat = pr
            n = len(mat) + 8; tot = round(size * s); m = max(1, tot // n); off = (tot - m * n) // 2
            x0, y0 = round(x * s) + off, round(y * s) + off
            for r, st, ln in _qr_runs(mat):
                d.rectangle([x0 + (4 + st) * m, y0 + (4 + r) * m, x0 + (4 + st + ln) * m - 1, y0 + (4 + r + 1) * m - 1], fill="#000000")
        elif kind == "line":
            _, x1, y1, x2, y2, col, lw = pr
            d.line([x1 * s, y1 * s, x2 * s, y2 * s], fill=col, width=max(1, round(lw * s)))


def render_label_png(data, pl, key, opts=None, dpi=300):
    from PIL import Image, ImageDraw
    opts = opts or {}; P = LABEL_PRESETS[key]
    mat, _v = qr_matrix(pl["payload"], opts.get("ecc", "m"))
    if P.get("media") == "roll": opts = {**opts, "snap_dpi": dpi}
    img = Image.new("RGB", (round(P["w"] / 25.4 * dpi), round(P["h"] / 25.4 * dpi)), "white")
    prims = label_prims(data, pl, key, mat, opts)
    if P.get("media") == "roll": prims[0] = ("rect", 0, 0, P["w"], P["h"], "#ffffff", None, 0)
    draw_prims_png(ImageDraw.Draw(img), prims, dpi)
    if opts.get("mono"): img = img.convert("L").point(lambda v: 255 if v > 150 else 0, mode="1")    # 1-бітне зображення для термопринтера
    return img


def _page_grid(key):
    P = LABEL_PRESETS[key]; mx = (210 - P["cols"] * P["w"]) / 2; my = (297 - P["rows"] * P["h"]) / 2
    return P, mx, my


def render_labels_roll_pdf(path, data, key, opts=None):
    """PDF для принтера етикеток: кожна етикетка – окрема сторінка точно її розміру; QR прив'язано до сітки точок принтера.
    → кількість етикеток."""
    opts = dict(opts or {}); font_paths(); P = LABEL_PRESETS[key]; mm = 72.0 / 25.4
    dpi = int(opts.get("printer_dpi", 203)); opts["snap_dpi"] = dpi
    c = _rl.Canvas(path, pagesize=(P["w"] * mm, P["h"] * mm)); c.setTitle(f"Етикетки – дослід {data['exp']['id']}"); page_h = P["h"] * mm
    for pl in data["plots"]:
        mat, _v = qr_matrix(pl["payload"], opts.get("ecc", "m")); prims = label_prims(data, pl, key, mat, opts)
        prims[0] = ("rect", 0, 0, P["w"], P["h"], "#ffffff", None, 0)                # без рамки: межі задає сам носій
        draw_prims_pdf(c, prims, 0, 0, page_h); c.showPage()
    c.save(); return len(data["plots"])


def render_labels_pdf(path, data, key, opts=None):
    """PDF із табличками на аркушах A4 (для наклейкового чи звичайного паперу); для рулонного розміру – одна етикетка на сторінку.
    → кількість аркушів (для рулону – етикеток)."""
    if is_roll(key): return render_labels_roll_pdf(path, data, key, opts)
    opts = opts or {}; font_paths(); P, mx, my = _page_grid(key); per = P["cols"] * P["rows"]; plots = data["plots"]
    pages = max(1, math.ceil(len(plots) / per)); c = _rl.Canvas(path, pagesize=(210 * 72 / 25.4, 297 * 72 / 25.4)); c.setTitle(f"QR-таблички – дослід {data['exp']['id']}")
    page_h = 297 * 72 / 25.4
    for pg in range(pages):
        for i, pl in enumerate(plots[pg * per:(pg + 1) * per]):
            r, cc = divmod(i, P["cols"]); mat, _v = qr_matrix(pl["payload"], opts.get("ecc", "m"))
            prims = label_prims(data, pl, key, mat, opts)
            if not opts.get("cut_lines", True): prims[0] = ("rect", 0, 0, P["w"], P["h"], "#ffffff", None, 0)
            draw_prims_pdf(c, prims, mx + cc * P["w"], my + r * P["h"], page_h)
        c.setFillColorRGB(0.4, 0.4, 0.4); c.setFont(FONT, 6.5)
        c.drawString(8 * 72 / 25.4, 4.0 * 72 / 25.4, f"S.A.D. · дослід {data['exp']['id']} · аркуш {pg + 1} з {pages} · ділянки {pg * per + 1}–{min(len(plots), (pg + 1) * per)} · підпис схеми {data['exp']['signature']}")
        c.showPage()
    c.save(); return pages


# ═══════════════════════════════════════════════════════════════
# 7. Розкладка: інструкція, карта, чек-лист (PDF на звичайному папері)
# ═══════════════════════════════════════════════════════════════
def _t(x, y, s, sz, bold=False, al="l", col="#111111"): return ("text", x, y, s, sz, bold, al, col)
def _r(x, y, w, h, fill=None, stroke=None, lw=0.2): return ("rect", x, y, w, h, fill, stroke, lw)


def _trunc(s, size, max_w, bold=False):
    if tw(s, size, bold) <= max_w: return s
    while s and tw(s + "…", size, bold) > max_w: s = s[:-1]
    return s + "…"


def instruction_page(data, key, opts, info):
    exp = data["exp"]; W, H = 210, 297; m = 14; pr = []
    pr.append(_t(m, 18, "Розкладання QR-табличок на ділянках досліду", 15, True))
    pr.append(_t(m, 25, f"Дослід {exp['id']}" + (f" · {exp['name']}" if exp["name"] else "") + f" · створено {exp['created']}", 9, False, "l", "#444444"))
    pr.append(_r(m, 29, W - 2 * m, 0.4, "#1a4b8c"))
    rows = [("Ділянок (табличок)", str(exp["n_plots"])), ("Варіантів", str(exp["n_variants"])), ("Повторностей", str(exp["n_reps"])),
            ("Підпис схеми", exp["signature"]), ("Формат QR", f"{FORMAT}, версія {info['version']}, рівень {info['ecc'].upper()}, модуль {info['module']:.2f} мм"),
            ("Початок схеми", exp["origin"] or "ряд 1, місце 1 (див. карту)")]
    y = 37
    for a, b in rows:
        pr.append(_t(m, y, a + ":", 9.5, True)); pr.append(_t(m + 48, y, _trunc(b, 9.5, W - 2 * m - 48), 9.5)); y += 5.8
    y += 3; pr.append(_t(m, y, "Порядок дій", 11.5, True, "l", "#1a4b8c")); y += 6.5
    steps = [
        f"Перевірте комплект. На кожній табличці вгорі стоїть код досліду {exp['id']}, внизу – підпис схеми {exp['signature']}. Таблички з іншим кодом чи підписом належать іншому досліду або іншій версії схеми: відкладіть їх.",
        f"Знайдіть початок схеми: {exp['origin'] or 'ряд 1, місце 1 (лівий верхній кут карти)'}. Ділянки пронумеровано наскрізно вздовж рядів в одному напрямку (на карті – зліва направо), ряд за рядом.",
        "Розкладайте за картою й чек-листом: табличка з номером n встановлюється в клітинку з номером n. Колір смуги вгорі таблички відповідає варіанту (легенда – на карті): ряд кольорів на місцевості має збігатися з рядом кольорів на карті.",
        "Звіряйте сусідів: на табличці вказано номери ділянок ліворуч і праворуч (середній та великий розміри). Якщо поруч стоїть інша ділянка, табличку встановлено не там.",
        "Після встановлення відскануйте кожну табличку телефоном: камера покаже дослід, повторність, варіант і місце ділянки. Відмітьте у чек-листі «Встановлено» й «Скановано».",
        f"Якщо QR-код пошкоджено, скористайтеся кодом під QR (формат {exp['id']}-012-X): останній символ контрольний, тож польовий журнал виявить помилку ручного введення.",
        "Практичні поради: друкуйте на водостійкому матеріалі або ламінуйте; прикріплюйте табличку рівно, без перегинів і відблисків; не зменшуйте QR-код під час друку (масштаб 100 %); при мінімальному розмірі модуля менше 0,4 мм оберіть більшу табличку.",
    ]
    for i, st in enumerate(steps, 1):
        lines = wrap_text(st, 9.5, W - 2 * m - 7, 8)
        pr.append(_t(m, y, f"{i}.", 9.5, True))
        for ln in lines: pr.append(_t(m + 7, y, ln, 9.5)); y += 4.7
        y += 2.2
    y += 2; pr.append(_t(m, y, "Скільки табличок містить повний опис у QR", 11.5, True, "l", "#1a4b8c")); y += 6.2
    names = {"full": "дослід, назва, повторність, варіант, рівні, місце + службовий рядок", "no_name": "без назви досліду (дослід, повторність, варіант, рівні, місце)",
             "no_levels": "без назв рівнів (дослід, повторність, варіант, місце)", "min": "лише службовий рядок SAD1 (потрібен польовий журнал)"}
    for t, n in info["tiers"].items():
        pr.append(_t(m, y, f"{n} з {exp['n_plots']}", 9.5, True)); pr.append(_t(m + 26, y, _trunc(names.get(t, t), 9.5, W - 2 * m - 26), 9.5)); y += 5.2
    return pr


def map_pages(data, key, opts):
    """Карта розкладки: ряди зверху вниз, позиції зліва направо; клітинка – номер ділянки й варіант, колір – варіант."""
    exp, grid, plots = data["exp"], data["grid"], data["plots"]; W, H = 297, 210; m = 10; lab_w = 20; top = 32; bottom = 12
    by_no = {p["plot"]: p for p in plots}; maxpos = max(max(c) for c in grid.values())
    u = min(26.0, (W - 2 * m - lab_w) / maxpos); ch = max(5.5, min(15.0, u * 0.85))
    rows = sorted(grid); per_page = max(1, int((H - top - bottom) // (ch + 1.0)))
    legend_h = 8 + 6 * math.ceil(len(data["variants"]) / 3)
    pages = []; chunks = [rows[i:i + per_page] for i in range(0, len(rows), per_page)]
    for pi, chunk in enumerate(chunks):
        pr = [_t(m, 11, f"Карта розкладки – дослід {exp['id']}" + (f" · {exp['name']}" if exp["name"] else ""), 13, True),
              _t(m, 17, f"підпис схеми {exp['signature']} · ділянок {exp['n_plots']} · початок: {exp['origin'] or 'ряд 1, місце 1'}", 8.5, False, "l", "#444444"),
              _t(W - m, 11, f"аркуш {pi + 1} з {len(chunks)}" , 8.5, False, "r", "#444444"),
              _t(m, 24.5, "↓ ряди", 8, True, "l", "#1a4b8c"), _t(m + lab_w, 24.5, "→ напрямок нумерації позицій (в усіх рядах однаковий)", 8, True, "l", "#1a4b8c")]
        step = 1 if u >= 7 else (2 if u >= 4.5 else 5)
        for q in range(1, maxpos + 1):
            if q % step == 0 or q == 1: pr.append(_t(m + lab_w + (q - 0.5) * u, top - 2.0, str(q), 6.2, False, "c", "#666666"))
        for ri, r in enumerate(chunk):
            y0 = top + ri * (ch + 1.0); reps = {by_no[n]["rep"] for (kind, n) in grid[r].values() if kind == "plot"}
            pr.append(_t(m, y0 + ch * 0.45, f"Ряд {r}", 8.5, True)); 
            if len(reps) == 1: pr.append(_t(m, y0 + ch * 0.45 + 3.6, f"повт. {next(iter(reps))}", 6.5, False, "l", "#555555"))
            cells = sorted(grid[r].items()); i = 0
            while i < len(cells):
                pos, (kind, no) = cells[i]; j = i
                while j + 1 < len(cells) and cells[j + 1][0] == cells[j][0] + 1 and cells[j + 1][1] == (kind, no) and kind == "plot": j += 1
                x0 = m + lab_w + (pos - 1) * u; wd = (cells[j][0] - pos + 1) * u
                if kind == "plot":
                    pl = by_no[no]; col = variant_color(pl["variant"]); fg = on_color(col)
                    pr.append(_r(x0, y0, wd, ch, col, "#ffffff", 0.35))
                    if wd >= 11 and ch >= 8:
                        pr.append(_t(x0 + wd / 2, y0 + ch * 0.50, str(no), min(10.5, ch * 0.62), True, "c", fg)); pr.append(_t(x0 + wd / 2, y0 + ch * 0.86, f"В{pl['variant']}", min(7.5, ch * 0.42), False, "c", fg))
                    elif wd >= 6: pr.append(_t(x0 + wd / 2, y0 + ch * 0.66, str(no), min(7.0, wd * 0.55), True, "c", fg))
                else: pr.append(_r(x0, y0, wd, ch, "#e6e6e6", "#ffffff", 0.35))
                i = j + 1
        last = pi == len(chunks) - 1; used = top + len(chunk) * (ch + 1.0)
        if last and used + legend_h + 4 <= H - 6:
            _legend(pr, data, m, used + 6, W - 2 * m)
        pages.append(pr)
    used_last = top + len(chunks[-1]) * (ch + 1.0)
    if used_last + legend_h + 4 > H - 6:
        pr = [_t(m, 11, f"Карта розкладки – легенда варіантів · дослід {exp['id']}", 13, True)]; _legend(pr, data, m, 22, W - 2 * m); pages.append(pr)
    return pages


def _legend(pr, data, x, y, width):
    pr.append(_t(x, y, "Легенда: колір = варіант;  сірі клітинки – не облікові (захисні, випадіння, запилювачі)", 8.5, True, "l", "#1a4b8c")); y += 5.5
    colw = width / 3; sw = 4.6
    for i, v in enumerate(data["variants"]):
        cx = x + (i % 3) * colw; cy = y + (i // 3) * 6.0
        pr.append(_r(cx, cy - 3.6, sw, sw, variant_color(v["n"]), "#999999", 0.15))
        pr.append(_t(cx + sw + 1.8, cy, _trunc(f"В{v['n']} · {v['code']} · {' + '.join(v['levels'])}", 7.6, colw - sw - 4), 7.6))


def checklist_pages(data, key, opts):
    exp, plots = data["exp"], data["plots"]; W, H = 210, 297; m = 10; cw = [10, 36, 10, 52, 24, 28, 15, 15]; rh = 6.4; top = 30; per = int((H - top - 12) // rh) - 1
    heads = ["№", "Ряд / місце", "Повт.", "Варіант", "Код", "Сусіди", "Встан.", "Скан."]; chunks = [plots[i:i + per] for i in range(0, len(plots), per)]; pages = []
    for pi, chunk in enumerate(chunks):
        pr = [_t(m, 12, f"Чек-лист розкладання – дослід {exp['id']}" + (f" · {exp['name']}" if exp["name"] else ""), 12.5, True),
              _t(m, 18, f"підпис схеми {exp['signature']} · аркуш {pi + 1} з {len(chunks)} · «Встан.» – табличку встановлено, «Скан.» – перевірено скануванням", 7.8, False, "l", "#444444")]
        x = m
        for h_, w_ in zip(heads, cw): pr.append(_r(x, top - rh, w_, rh, "#d9e2f3", "#888888", 0.2)); pr.append(_t(x + w_ / 2, top - 2.0, h_, 7.6, True, "c")); x += w_
        by_no = {p["plot"]: p for p in plots}
        for i, p in enumerate(chunk):
            y = top + i * rh; x = m; col = variant_color(p["variant"])
            l = f"№{p['left']}" if p["left"] else "поч."; r_ = f"№{p['right']}" if p["right"] else "кін."
            vals = [str(p["plot"]), _trunc(p["pos_text"], 7.2, cw[1] - 2), str(p["rep"]), "", p["code"], f"← {l}  {r_} →", "", ""]
            for k, (v, w_) in enumerate(zip(vals, cw)):
                pr.append(_r(x, y, w_, rh, None, "#888888", 0.15))
                if k == 3:
                    pr.append(_r(x + 0.6, y + 0.9, 3.0, rh - 1.8, col, None, 0)); txt = _trunc(f"{p['vcode']} · {' + '.join(p['levels'])}", 7.0, w_ - 5.2)
                    pr.append(_t(x + 4.2, y + rh * 0.68, txt, 7.0))
                elif k in (6, 7): pr.append(_r(x + w_ / 2 - 1.8, y + rh / 2 - 1.8, 3.6, 3.6, None, "#333333", 0.25))
                elif v: pr.append(_t(x + (w_ / 2 if k in (0, 2) else 1.2), y + rh * 0.68, v, 7.2, k == 0, "c" if k in (0, 2) else "l"))
                x += w_
        pages.append(pr)
    return pages


def render_layout_pdf(path, data, key="S", opts=None, info=None):
    """PDF для звичайного паперу: інструкція, карта розкладки, чек-лист. → кількість сторінок."""
    opts = opts or {}; font_paths(); info = info or layout_info(data, opts.get("ecc", "m"), key)
    c = _rl.Canvas(path, pagesize=(210 * 72 / 25.4, 297 * 72 / 25.4)); c.setTitle(f"Розкладка табличок – дослід {data['exp']['id']}")
    pages = [("p", instruction_page(data, key, opts, info))] + [("l", pg) for pg in map_pages(data, key, opts)] + [("p", pg) for pg in checklist_pages(data, key, opts)]
    for kind, prims in pages:
        w_mm, h_mm = (297, 210) if kind == "l" else (210, 297)
        c.setPageSize((w_mm * 72 / 25.4, h_mm * 72 / 25.4)); draw_prims_pdf(c, prims, 0, 0, h_mm * 72 / 25.4); c.showPage()
    c.save(); return len(pages)


def layout_info(data, ecc="m", key="S"):
    import collections
    mx = max((p.get("qr_version", 0) for p in data["plots"]), default=0)
    return {"version": mx, "ecc": ecc, "module": module_mm(key, mx) if mx else 0.0, "tiers": dict(collections.Counter(p["tier"] for p in data["plots"]))}


def summary(data, key, ecc="m", style="descriptive"):
    """Зведення для користувача: версія QR, модуль, орієнтовна відстань сканування, кількість аркушів і попередження."""
    plots = data["plots"]; info = layout_info(data, ecc, key); P = LABEL_PRESETS[key]; warns = []
    if style == "descriptive":
        n_full = info["tiers"].get("full", 0)
        if n_full < len(plots):
            warns.append(f"У {len(plots) - n_full} із {len(plots)} табличок QR містить скорочений опис (щільність обмежена розміром таблички). Повний опис надруковано на самій табличці; "
                         f"для повного опису в QR оберіть більшу табличку.")
    else:
        warns.append("Компактний QR містить лише службовий рядок: звичайна камера покаже технічний текст; опис варіанта видно на самій табличці та у польовому журналі.")
    if len(data["variants"]) > len(PALETTE):
        warns.append(f"Варіантів більше, ніж кольорів ({len(PALETTE)}): кольори повторюються, орієнтуйтеся на номер варіанта.")
    if data.get("skipped"): warns.append(f"{data['skipped']} записів схеми без коректних факторів пропущено.")
    if info["module"] and info["module"] < MIN_MODULE_MM: warns.append(f"Модуль QR {info['module']:.2f} мм менший за {MIN_MODULE_MM} мм: оберіть більшу табличку.")
    info.update(scan_cm=scan_distance_cm(key, info["version"]) if info["version"] else 0.0, warnings=warns, per_sheet=P["cols"] * P["rows"],
                sheets=max(1, math.ceil(len(plots) / (P["cols"] * P["rows"]))))
    return info


# ═══════════════════════════════════════════════════════════════
# 8. Експорт: маніфест (для польового журналу), Excel, PNG, комплект
# ═══════════════════════════════════════════════════════════════
def manifest_dict(data, style="descriptive", ecc="m"):
    exp = data["exp"]
    return {"format": "SAD-plots", "version": 1, "qr_format": FORMAT,
            "experiment": {k: exp[k] for k in ("id", "name", "created", "model", "design", "origin", "signature", "n_plots", "n_variants", "n_reps", "factors", "trait")},
            "qr": {"style": style, "ecc": ecc.upper(), "check": "Luhn mod 36 over E+P(4)+R(3)+V(3)", "id_alphabet": ID_ALPHABET},
            "variants": [{"n": v["n"], "code": v["code"], "levels": v["levels"]} for v in data["variants"]],
            "plots": [{"plot": p["plot"], "code": p["code"], "rep": p["rep"], "variant": p["variant"], "variant_code": p["vcode"], "levels": p["levels"], "row": p["row"],
                       "segments": [list(s) for s in p["segments"]], "position_text": p["pos_text"], "left": p["left"], "right": p["right"], "payload": p["payload"]}
                      for p in data["plots"]]}


def export_manifest(path, data, style="descriptive", ecc="m"):
    with open(path, "w", encoding="utf-8") as f: json.dump(manifest_dict(data, style, ecc), f, ensure_ascii=False, indent=1)


def export_xlsx(path, data):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    exp = data["exp"]; wb = Workbook(); ws = wb.active; ws.title = "Розкладка"
    heads = ["№ ділянки", "Код", "Повторність", "Варіант", "Код варіанта"] + [f["name"] for f in exp["factors"]] + ["Ряд", "Місце", "Ліворуч", "Праворуч", "Встановлено", "Скановано"]
    ws.append(heads)
    for c in ws[1]: c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="D9E2F3"); c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin = Side(style="thin", color="999999")
    for p in data["plots"]:
        ws.append([p["plot"], p["code"], p["rep"], p["variant"], p["vcode"]] + list(p["levels"]) + [p["row"], p["pos_text"].split(", ", 1)[-1] if p["units"] == 1 else p["pos_text"],
                  (f"№{p['left']}" if p["left"] else "початок ряду"), (f"№{p['right']}" if p["right"] else "кінець ряду"), "", ""])
        ws.cell(ws.max_row, 4).fill = PatternFill("solid", fgColor=variant_color(p["variant"]).lstrip("#")); ws.cell(ws.max_row, 4).font = Font(bold=True, color=on_color(variant_color(p["variant"])).lstrip("#"))
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
        for c in row: c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for i, c in enumerate(ws[1], 1): ws.column_dimensions[c.column_letter].width = 14 if i > 5 else 11
    ws.column_dimensions["B"].width = 14; ws.freeze_panes = "A2"
    wv = wb.create_sheet("Варіанти"); wv.append(["№", "Код"] + [f["name"] for f in exp["factors"]])
    for c in wv[1]: c.font = Font(bold=True)
    for v in data["variants"]:
        wv.append([v["n"], v["code"]] + list(v["levels"])); wv.cell(wv.max_row, 1).fill = PatternFill("solid", fgColor=variant_color(v["n"]).lstrip("#")); wv.cell(wv.max_row, 1).font = Font(bold=True, color=on_color(variant_color(v["n"])).lstrip("#"))
    wi = wb.create_sheet("Дослід")
    for a, b in (("Код досліду", exp["id"]), ("Назва", exp["name"]), ("Створено", exp["created"]), ("Модель", exp["model"]), ("Спосіб рандомізації", exp["design"]), ("Ділянок", exp["n_plots"]),
                 ("Варіантів", exp["n_variants"]), ("Повторностей", exp["n_reps"]), ("Підпис схеми", exp["signature"]), ("Початок схеми", exp["origin"]), ("Формат QR", FORMAT)):
        wi.append([a, b]); wi.cell(wi.max_row, 1).font = Font(bold=True)
    wi.column_dimensions["A"].width = 22; wi.column_dimensions["B"].width = 50
    wb.save(path)


def save_png_labels(folder, data, key, opts=None, dpi=300):
    os.makedirs(folder, exist_ok=True); out = []
    for p in data["plots"]:
        fn = os.path.join(folder, f"{p['plot']:03d}_{p['code']}_П{p['rep']}_В{p['variant']}.png")
        render_label_png(data, p, key, opts, dpi).save(fn, dpi=(dpi, dpi)); out.append(fn)
    return out


def save_all(folder, data, key, opts=None, dpi=300):
    """Пише в теку весь комплект: PDF табличок, PDF розкладки, Excel, маніфест і PNG кожної таблички. → список файлів."""
    opts = opts or {}; os.makedirs(folder, exist_ok=True); E = data["exp"]["id"]; files = []
    f = os.path.join(folder, f"{E}_таблички.pdf"); render_labels_pdf(f, data, key, opts); files.append(f)
    f = os.path.join(folder, f"{E}_розкладка.pdf"); render_layout_pdf(f, data, key, opts); files.append(f)
    f = os.path.join(folder, f"{E}_розкладка.xlsx"); export_xlsx(f, data); files.append(f)
    f = os.path.join(folder, f"{E}_manifest.json"); export_manifest(f, data, opts.get("style", "descriptive"), opts.get("ecc", "m")); files.append(f)
    files += save_png_labels(os.path.join(folder, f"{E}_png"), data, key, opts, dpi)
    return files


def open_for_print(path):
    """Відкриває файл у програмі за замовчуванням (звідти його друкують)."""
    import sys, subprocess
    if sys.platform.startswith("win"): os.startfile(path)                      # noqa
    elif sys.platform == "darwin": subprocess.Popen(["open", path])
    else: subprocess.Popen(["xdg-open", path])
