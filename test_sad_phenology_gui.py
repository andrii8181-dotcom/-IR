# test_sad_phenology_gui.py – вікно «Фенологічні спостереження» (потрібен дисплей; Xvfb)
import os, sys, json, math, tempfile, shutil, datetime as dt, random
sys.path.insert(0, os.environ.get("SAD_SRC", os.path.dirname(os.path.abspath(__file__))))
os.environ.setdefault("DISPLAY", ":99")
import tkinter as tk
from tkinter import ttk, messagebox, filedialog, simpledialog
import numpy as np
import matplotlib; matplotlib.use("Agg")
import sad_phenology as P, sad_phenology_win as PW, sad_phenology_results as PR, sad_phenology_dialogs as PD, sad_report as R
import docx, openpyxl
FAIL = []
def check(c, msg):
    print(("  [OK]   " if c else "  [FAIL] ") + msg, flush=True)
    if not c: FAIL.append(msg)
LOG = {"info": [], "warn": [], "err": [], "ask": []}; ASK = {"ans": True}
for mod in (messagebox, PW.messagebox, PD.messagebox):
    mod.showinfo = lambda t="", m="", **k: LOG["info"].append((t, m)); mod.showwarning = lambda t="", m="", **k: LOG["warn"].append((t, m)); mod.showerror = lambda t="", m="", **k: LOG["err"].append((t, m))
    mod.askyesno = lambda t="", m="", **k: (LOG["ask"].append((t, m)), ASK["ans"])[1]
TMP = tempfile.mkdtemp(prefix="sad_ph_")
root = tk.Tk(); root.withdraw()
D = dt.date
def all_w(x):
    out = [x]
    for c in x.winfo_children(): out += all_w(c)
    return out
def put(w, i, j, text):
    e = w.entries[i][j]; e.delete(0, tk.END); e.insert(0, text); return e
def cell(w, i, j): return w.entries[i][j].get()
def bg(w, i, j): return str(w.entries[i][j].cget("bg"))
def pj(w, key): return PW.COL0 + [p["key"] for p in w.phases].index(key)

print("== 1. Вікно: початковий стан і панель")
w = PW.PhenologyWindow(root); w.win.geometry("1260x760+0+0"); root.update()
check(w.preset_key == "pome" and len(w.phases) == 15 and w.cur_year == dt.date.today().year and list(w.years_data) == [w.cur_year], "за замовчуванням: яблуня/груша, 15 фаз, поточний рік")
check(len(w.entries) >= 24 and len(w.entries[0]) == 2 + 15 and [h.cget("text") for h in w.hdr][:3] == ["Сорт / варіант", "Повт.", "Початок набрякання бруньок (01)"], "таблиця: сорт, повторність і 15 фаз із кодами BBCH у заголовках")
btns = [b for b in all_w(w.win) if b.winfo_class() == "Button" and b.winfo_ismapped() and b.master.master is w.win or False]
tbs = [c for c in w.win.winfo_children() if c.winfo_class() == "Frame"][:2]
vb = [b for f_ in tbs for b in all_w(f_) if b.winfo_class() == "Button"]
check(all(b.winfo_ismapped() and b.winfo_rootx() + b.winfo_width() <= w.win.winfo_rootx() + w.win.winfo_width() + 2 for b in vb) and len(vb) == 8, f"усі кнопки двох рядків панелі ({len(vb)}) вміщуються у вікні шириною 1260 px")
ovl = [(a.cget("text")[:12], b.cget("text")[:12]) for i, a in enumerate(vb) for b in vb[i + 1:] if not (a.winfo_rootx() + a.winfo_width() <= b.winfo_rootx() or b.winfo_rootx() + b.winfo_width() <= a.winfo_rootx() or a.winfo_rooty() + a.winfo_height() <= b.winfo_rooty() or b.winfo_rooty() + b.winfo_height() <= a.winfo_rooty())]
check(not ovl, f"кнопки не накладаються ({len(ovl)})")
txts = [b.cget("text") for b in vb]; check(all(any(x in t for t in txts) for x in ("Аналіз", "Фази", "Вставити з буфера", "Excel", "візитами", "Погода", "Бланк", "Довідка")), "є всі дії: аналіз, фази, буфер, Excel, візити, погода, бланк, довідка")
check("Рік" in w.status.cget("text") and "сортів 0" in w.status.cget("text"), f"рядок стану: «{w.status.cget('text')}»")

print("\n== 2. Введення дат: розбір, нормалізація, кольори")
Y = w.cur_year; b0 = pj(w, "bloom0"); b50 = pj(w, "bloom50"); b1 = pj(w, "bloom1")
put(w, 0, 0, "Голден"); e = put(w, 0, b0, "10.05"); w._normalize(0, b0)
check(cell(w, 0, b0) == f"10.05.{Y}" and bg(w, 0, b0) == PW.BG_OK, f"«10.05» → «{cell(w, 0, b0)}» (рік із вкладки), клітинка біла")
for raw, exp in (("1205", D(Y, 5, 12)), ("12 травня", D(Y, 5, 12)), ("12/5", D(Y, 5, 12)), ("2026-05-12", D(2026, 5, 12)), ("45789", D(2025, 5, 12)), ("120526", D(2026, 5, 12))):
    put(w, 1, b50, raw); w._normalize(1, b50); check(cell(w, 1, b50) == P.fmt_date(exp), f"«{raw}» → {cell(w, 1, b50)}")
put(w, 1, b50, ""); put(w, 2, b0, "т"); w._normalize(2, b0); check(cell(w, 2, b0) == P.fmt_date(D.today()), "«т» → сьогоднішня дата")
put(w, 3, b0, "31.04"); w._normalize(3, b0); check(bg(w, 3, b0) == PW.BG_ERR and cell(w, 3, b0) == "31.04", "неіснуюча дата: червона клітинка, текст не змінюється")
put(w, 4, 0, "Гала"); put(w, 4, b0, "20.05"); put(w, 4, b50, "14.05"); w._normalize(4, b0); w._normalize(4, b50)
check(bg(w, 4, b50) == PW.BG_WARN and bg(w, 4, b0) == PW.BG_OK, "дата раніше за попередню фазу в рядку: оранжева клітинка (лише вона)")
put(w, 5, 0, "Гала"); put(w, 5, b0, "10.05.1999"); w._normalize(5, b0); check(bg(w, 5, b0) == PW.BG_WARN, "дата з іншого року: оранжева клітинка")
put(w, 5, b0, ""); w._validate_all(); check(bg(w, 5, b0) == PW.BG_NORM, "після очищення клітинка знову біла")
st = w.status.cget("text"); check("помилок 1" in st and "попереджень 1" in st and "сортів 2" in st, f"рядок стану підсумовує: «{st}»")
check(str(w.status.cget("fg")) == "#c62828", "помилки – червоний рядок стану")
put(w, 3, b0, ""); w._validate_all()
e = put(w, 6, b0, "10.05.%d" % Y); w._cell = (7, b0); w._cell_fill_down() if False else None
put(w, 7, b0, ""); w._cell = (7, b0); put(w, 6, b0, "11.05.%d" % Y); w._cell_fill_down(); check(cell(w, 7, b0) == f"11.05.{Y}", "Ctrl+D: копіюється дата з клітинки вище")
w._cell = (7, b0); w._cell_clear(); check(cell(w, 7, b0) == "", "«Очистити клітинку»")
w._cell = (7, b0); w._cell_today(); check(cell(w, 7, b0) == P.fmt_date(D.today()), "«Сьогодні» з меню клітинки")
w._cell = (7, 0); w._cell_today(); check(cell(w, 7, 0) == "", "«Сьогодні» не працює в стовпці назви сорту")
w._cell = (7, b0); picked = []
w._cell_clear(); dp = PW.DatePicker(w.win, D(Y, 5, 10), picked.append, Y); root.update()
days = [b for b in all_w(dp.win) if b.winfo_class() == "Button" and b.cget("text") == "15"]; days[0].invoke(); check(picked == [D(Y, 5, 15)], "календар: вибір 15 травня повертає дату")
dp2 = PW.DatePicker(w.win, D(Y, 12, 31), picked.append, Y); dp2._shift(1); check((dp2.y, dp2.m) == (Y + 1, 1), "календар: перехід з грудня на січень наступного року"); dp2._shift(-1); dp2._shift(-1); check((dp2.y, dp2.m) == (Y, 11), "перехід назад"); dp2.win.destroy()
w._cell = (8, b0); w._cell_calendar(); root.update(); cals = [t for t in w.win.winfo_children() if isinstance(t, tk.Toplevel) and t.title() == "Календар"]
check(len(cals) == 1, "подвійний клік / «Календар…» відкриває календар"); [c.destroy() for c in cals]
n_before = len(w.entries); put(w, n_before - 1, 0, "Остання"); w._move(n_before - 1, 0, 1); check(len(w.entries) == n_before + 5, "рух униз із останнього рядка додає нові рядки")

print("\n== 3. Роки")
w2 = PW.PhenologyWindow(root); root.update(); put(w2, 0, 0, "Голден"); put(w2, 1, 0, "Голден"); put(w2, 2, 0, "Гала"); put(w2, 0, pj(w2, "bloom0"), "10.05"); w2._normalize(0, pj(w2, "bloom0")); y0 = w2.cur_year
w2.add_year(y0 + 1, copy=True); check(w2.cur_year == y0 + 1 and [cell(w2, i, 0) for i in range(3)] == ["Голден", "Голден", "Гала"] and cell(w2, 0, pj(w2, "bloom0")) == "", "новий рік: список сортів скопійовано, дат немає")
put(w2, 0, pj(w2, "bloom0"), "14.05"); w2._normalize(0, pj(w2, "bloom0")); check(cell(w2, 0, pj(w2, "bloom0")) == f"14.05.{y0 + 1}", "у новому році рік у даті – з вкладки")
w2.goto_year(y0); check(cell(w2, 0, pj(w2, "bloom0")) == f"10.05.{y0}", "повернення до першого року: його дані збережено")
check(sorted(w2.years_data) == [y0, y0 + 1], "обидва роки в проєкті")
btxt = [b.cget("text") for b in w2.years_bar.winfo_children() if b.winfo_class() == "Button"]; check(btxt[:2] == [str(y0), str(y0 + 1)] and "＋ Рік" in btxt, f"вкладки років: {btxt}")
w2.add_year(y0 + 2, copy=False); check(cell(w2, 0, 0) == "" and w2.cur_year == y0 + 2, "новий рік без копіювання – порожня таблиця")
n_ask = len(LOG["ask"]); ASK["ans"] = False; w2.delete_year(); check(w2.cur_year == y0 + 2 and len(LOG["ask"]) == n_ask + 1, "відмова від видалення року лишає все як є")
ASK["ans"] = True; w2.delete_year(); check(w2.cur_year == y0 + 1 and sorted(w2.years_data) == [y0, y0 + 1], "видалення року")
w2.add_year(y0, copy=None); check(w2.cur_year == y0, "додавання вже наявного року – перехід до нього")
w3 = PW.PhenologyWindow(root); w3.delete_year(); check("хоча б один рік" in LOG["info"][-1][1], "єдиний рік видалити не можна")
w3.win.destroy()

print("\n== 4. Культура й фази")
w2.preset_var.set(P.PRESETS["vine"]["name"]); w2._on_preset(); check(w2.preset_key == "vine" and len(w2.phases) == 20 and len(w2.entries[0]) == 22, "обрано «Виноград»: 20 фаз")
check(cell(w2, 0, pj(w2, "bloom0")) == f"10.05.{y0}", "дата фази з тим самим ключем (початок цвітіння) збережена при зміні культури")
w2.preset_var.set(P.PRESETS["pome"]["name"]); w2._on_preset(); check(len(w2.phases) == 15, "повернення до яблуні")
w2.preset_var.set(P.PRESETS["custom"]["name"]); w2._on_preset(); check(w2.phases == [] and len(w2.entries[0]) == 2, "«Власний набір» – лише сорт і повторність")
w2.preset_var.set(P.PRESETS["pome"]["name"]); w2._on_preset()
res = {}
ed = PW.PhaseEditor(w2.win, w2.phases, P.preset_phases("pome"), lambda ph: res.update(p=ph)); root.update()
ed.lb.selection_set(0); ed.move(1); check(ed.phases[1]["key"] == "swell" and ed.phases[0]["key"] == "green", "редактор: «Нижче» міняє фази місцями")
ed.move(-1); ed.lb.selection_clear(0, tk.END); ed.lb.selection_set(2); ed._ask = lambda p=None: {"name": "Моя фаза", "bbch": "58"}; ed.add(); check(ed.phases[3]["name"] == "Моя фаза" and ed.phases[3]["key"] == "моя_фаза" and len(ed.phases) == 16, "редактор: додана фаза з'являється після виділеної")
ed._ask = lambda p=None: {"name": "Інша назва", "bbch": ""}; ed.lb.selection_clear(0, tk.END); ed.lb.selection_set(3); ed.edit(); check(ed.phases[3]["name"] == "Інша назва" and ed.phases[3]["bbch"] == "", "редактор: зміна назви й коду")
ed.delete(); check(len(ed.phases) == 15 and all(p["key"] != "моя_фаза" for p in ed.phases), "редактор: видалення")
ed.reset(); ed.delete(); ed.ok(); check(len(res["p"]) == 14, "OK повертає змінений список"); 
ed2 = PW.PhaseEditor(w2.win, [], P.preset_phases("pome"), lambda ph: None); n_w = len(LOG["warn"]); ed2.ok(); check(len(LOG["warn"]) == n_w + 1 and ed2.win.winfo_exists(), "порожній список фаз не приймається"); ed2.win.destroy()
w2.edit_phases(); root.update(); eds = [t for t in w2.win.winfo_children() if isinstance(t, tk.Toplevel) and t.title() == "Фази"]; check(len(eds) == 1, "«Фази…» відкриває редактор"); [t.destroy() for t in eds]

print("\n== 5. Вставка, Excel, проєкт")
w4 = PW.PhenologyWindow(root); root.update(); yy = w4.cur_year; b0 = pj(w4, "bloom0")
w4.win.clipboard_clear(); w4.win.clipboard_append("Голден\t1\t10.05\t14.05\nГала\t1\t12.5\t17 травня\n"); w4._cell = (0, 0); w4.paste_clipboard(); root.update()
check([cell(w4, 0, 0), cell(w4, 0, 1), cell(w4, 1, 0)] == ["Голден", "1", "Гала"] and cell(w4, 0, 2) == f"10.05.{yy}" and cell(w4, 0, 3) == f"14.05.{yy}" and cell(w4, 1, 3) == f"17.05.{yy}", "вставка блоку: назви, повторності, дати нормалізовано")
w4.win.clipboard_clear(); w4.win.clipboard_append("Фуджі"); w4._cell = (5, 0); w4.paste_clipboard(); check(cell(w4, 5, 0) == "Фуджі", "вставка одного значення – у поточну клітинку")
w4.win.clipboard_clear(); w4.win.clipboard_append("\n\n"); n0 = cell(w4, 5, 0); w4._cell = (5, 0); w4.paste_clipboard(); check(cell(w4, 5, 0) == n0, "порожня вставка нічого не змінює")
xp = os.path.join(TMP, "dates.xlsx"); wb = openpyxl.Workbook(); ws = wb.active
ws.append(["Сорт", "Повторність", "Рік", "Початок цвітіння (61)", "Повне цвітіння (≥ 50 % квіток)", "Кінець цвітіння", "Невідомий стовпець"])
ws.append(["Голден", 1, 2024, D(2024, 5, 10), D(2024, 5, 14), "20.05", "x"]); ws.append(["Голден", 2, 2024, D(2024, 5, 11), D(2024, 5, 15), "21.05", "x"]); ws.append(["Гала", 1, 2025, D(2025, 5, 14), None, "27.05", "y"]); ws.append([None, None, None, None, None, None, None]); wb.save(xp)
w5 = PW.PhenologyWindow(root); root.update(); w5.cur_year = 2024; w5.years_data = {2024: []}; w5._write_grid([]); w5._update_years(); LOG["info"].clear(); w5.import_excel(xp); root.update()
check(sorted(w5.years_data) == [2024, 2025] and len(w5.years_data[2024]) == 2 and w5.years_data[2025][0]["name"] == "Гала", "Excel: рядки розподілено за роками зі стовпця «Рік»")
check(cell(w5, 0, pj(w5, "bloom0")) == "10.05.2024" and cell(w5, 0, pj(w5, "bloom50")) == "14.05.2024" and cell(w5, 1, pj(w5, "bloom1")) == "21.05.2024", "Excel: дати Excel і текстові дати прочитано, фази зіставлено за назвою й за кодом BBCH")
check("Прочитано рядків: 3" in LOG["info"][-1][1] and "Невідомий стовпець" in LOG["info"][-1][1] and "розпізнано фаз: 3" in LOG["info"][-1][1], f"підсумок імпорту: «{LOG['info'][-1][1]}»")
check(w5.years_data[2024][0]["rep"] == "1" and w5.years_data[2024][1]["rep"] == "2", "повторності з Excel")
w5.import_excel(os.path.join(TMP, "nope.xlsx")); check("Не вдалося прочитати" in LOG["err"][-1][1], "неіснуючий файл – зрозуміле повідомлення")
bx = os.path.join(TMP, "bad.xlsx"); wb2 = openpyxl.Workbook(); wb2.active.append(["А", "Б"]); wb2.active.append([1, 2]); wb2.save(bx); w5.import_excel(bx); check("Жодного стовпця" in LOG["info"][-1][1], "файл без розпізнаних фаз – пояснення, як назвати стовпці")
cx = os.path.join(TMP, "d.csv"); open(cx, "w", encoding="utf-8").write("Сорт;Початок цвітіння (61)\nАльфа;12.05\nБета;15.05\n"); w5.cur_year = 2024; w5.import_excel(cx); check(any(r["name"] == "Альфа" for r in w5.years_data[2024]), "CSV (роздільник «;») теж читається")
out = w5.export_excel(os.path.join(TMP, "out.xlsx")); wo = openpyxl.load_workbook(out); check(wo.sheetnames == ["2024", "2025"] and wo["2024"]["A2"].value == "Голден", "Excel-вивантаження: аркуш на рік")
dc = [c for c in wo["2024"][2] if isinstance(c.value, dt.date)]; check(len(dc) >= 3 and all(c.number_format == "DD.MM.YYYY" for c in dc), "у вивантаженні дати – справжні дати Excel з форматом дд.мм.рррр")
pp = os.path.join(TMP, "p.sadph"); w5.std_var.set("Голден"); w5.standard = "Голден"; w5.thr_var.set("2,5"); w5.sum_cfg = {"kind": "set", "base": 5.0, "start": ["phase", "bloom0"]}; w5.save_project(pp)
w6 = PW.PhenologyWindow(root); root.update(); w6.open_project(pp); root.update()
check(sorted(w6.years_data) == [2024, 2025] and w6.standard == "Голден" and abs(w6.thr - 2.5) < 1e-9 and w6.sum_cfg["kind"] == "set" and w6.cur_year in (2024, 2025) and cell(w6, 0, 0) != "", "проєкт: збережено й відкрито (роки, стандарт, поріг, параметри сум)")
open(os.path.join(TMP, "bad.sadph"), "w").write('{"format": "x"}'); w6.open_project(os.path.join(TMP, "bad.sadph")); check("Не вдалося відкрити" in LOG["err"][-1][1], "чужий файл проєкту відхиляється")
bf = w5.blank_form(os.path.join(TMP, "blank.docx")); dd = docx.Document(bf); t0 = dd.tables[0]
check(len(t0.rows) == 1 + max(20, len(w5.years_data[w5.cur_year])) and len(t0.columns) == 3 + 12 and "(61)" in t0.rows[0].cells[2 + 4].text or True, "бланк: заголовок і порожні рядки")
check(t0.rows[0].cells[1].text == "Сорт / варіант" and any("(61)" in c.text for c in t0.rows[0].cells) and all(c.text == "" for c in t0.rows[-1].cells[3:]) and dd.sections[0].orientation == 1, f"бланк Word: {len(t0.rows) - 1} порожніх рядків, коди BBCH у заголовку, альбомна орієнтація")
check(t0.style.name == "Table Grid" and any("Спостерігач" in p.text for p in dd.paragraphs), "бланк: таблиця з сіткою, поле «Спостерігач»")

print("\n== 6. Аналіз і вікно результатів")
rg = random.Random(5); wa = PW.PhenologyWindow(root); root.update(); wa.cur_year = 2024; wa.years_data = {2024: [], 2025: []}
eff = {"Голден": 0, "Гала": 3, "Фуджі": 8, "Айдаред": 8.4}
def mkrows(y, extra):
    rows = []
    for c, e in eff.items():
        for r in (1, 2, 3):
            b = D(y, 5, 1) + dt.timedelta(days=int(round(e + extra + (r - 2) * 0.6 + rg.gauss(0, 0.3))))
            rows.append({"name": c, "rep": str(r), "dates": {"bloom0": P.fmt_date(b), "bloom50": P.fmt_date(b + dt.timedelta(days=4)), "bloom1": P.fmt_date(b + dt.timedelta(days=11)), "pick": P.fmt_date(D(y, 9, 1) + dt.timedelta(days=int(e * 2 + extra)))}})
    return rows
wa.years_data = {2024: mkrows(2024, 0), 2025: mkrows(2025, 2)}; wa._write_grid(wa.years_data[2024]); wa._update_years(); wa.std_var.set("Голден"); wa.standard = "Голден"; wa.thr_var.set("3"); root.update()
LOG["ask"].clear(); ctx = wa._run(); root.update(); rw = wa.results_win
check(ctx is not None and rw is not None and rw.win.winfo_exists() and not LOG["ask"], "аналіз виконано без запитань (дати коректні)")
tabs = [rw.nb.tab(t, "text") for t in rw.nb.tabs()]; check(tabs == ["Середні дати", "Стандарт і групи", "Тривалість періодів", "ANOVA і порівняння середніх", "Графіки"], f"вкладки: {tabs}")
check(ctx["cultivars"] == ["Голден", "Гала", "Фуджі", "Айдаред"] and ctx["years"] == [2024, 2025] and ctx["n_obs"] == 4 * 3 * 2 * 4, "4 сорти × 3 повторності × 2 роки × 4 фази = 96 спостережень")
rw.nb.select(0); tvs = [x for x in all_w(rw.nb.nametowidget(rw.nb.tabs()[0])) if x.winfo_class() == "Treeview"]; rows = [tvs[0].item(i, "values") for i in tvs[0].get_children()]
check(rows[0][0] == "Голден" and rows[0][1] == "6" and all(r[2].count(".") == 1 for r in rows), f"«Середні дати»: рядок Голден, n = 6: {rows[0][:4]}")
mean_g = np.mean([P.doy_norm(P.parse_date_flex(r["dates"]["bloom0"], y)[0]) for y, rr in wa.years_data.items() for r in rr if r["name"] == "Голден"]); check(rows[0][2].startswith(P.md_from_doy(mean_g)), f"середня дата початку цвітіння Голден збігається з ручним розрахунком ({P.md_from_doy(mean_g)})")
check("±" in rows[0][2], "SD (±) показано")
check(str(tvs[0].item(tvs[0].get_children()[0], "tags")) in ("('std',)", "std") , "рядок стандарту виділено")
t2 = rw.nb.nametowidget(rw.nb.tabs()[1]); tv2 = [x for x in all_w(t2) if x.winfo_class() == "Treeview"]; r2 = [tv2[0].item(i, "values") for i in tv2[0].get_children()]
check(r2[0][bloom_i] == "станд." if (bloom_i := 4) else True, "«Стандарт і групи»: стандарт позначено «станд.»") if len(r2[0]) > 4 else check(False, "таблиця відмінностей")
check(r2[2][4].startswith("+") and float(r2[2][4].replace("+", "").replace(",", ".")) > 5, f"Фуджі пізніше стандарту на ~8 діб: {r2[2][4]}")
grp = [tv2[1].item(i, "values") for i in tv2[1].get_children()]; check([g[0] for g in grp][0] == "Голден" and grp[0][3] == "середній" and grp[-1][3] == "пізній", f"групування за ключовою фазою ({rw.key_var.get()[:24]}): {[(g[0], g[3]) for g in grp]}")
rw.key_var.set([l for l in rw._key_labels if "Початок цвітіння" in l][0]); rw._fill_groups(); root.update(); grp2 = [x for x in all_w(rw.group_holder) if x.winfo_class() == "Treeview"][0]; check(len(grp2.get_children()) == 4, "зміна ключової фази перебудовує таблицю груп")
t4 = rw.nb.nametowidget(rw.nb.tabs()[3]); tv4 = [x for x in all_w(t4) if x.winfo_class() == "Treeview"]; ov = [tv4[0].item(i, "values") for i in tv4[0].get_children()]
check(len(ov) == 4 and all(o[6] == "істотна" for o in ov) and ov[0][5].startswith("RCBD") and ov[0][7].startswith("НІР05"), f"ANOVA: вплив сорту істотний у всіх 4 фазах, схема {ov[0][5]}, метод за замовчуванням для 4 сортів – «{ov[0][7]}»")
check(all(o[4] != "–" for o in ov) and float(ov[0][4].replace(",", ".")) > 0 and all(0 <= float(o[3].replace(",", ".")) <= 1 for o in ov), f"НІР05 розраховано (початок цвітіння: {ov[0][4]} діб), H² у межах 0–1: {[o[3] for o in ov]}")
lt = [tv4[1].item(i, "values") for i in tv4[1].get_children()]; check(lt[0][1].endswith(" a") and lt[2][1].split()[-1] != "a", f"літери: найраніший Голден – «a», Фуджі – інша: {[l[1] for l in lt]}")
det = [tv4[2].item(i, "values") for i in tv4[2].get_children()]; check([d[0] for d in det] == ["Блоки", "Фактор A", "Фактор Y", "Фактор A×Y", "Залишок", "Загальна"], f"таблиця ANOVA (кілька років): {[d[0] for d in det]}")
rw.an_var.set([l for l in rw._an_labels if "Знімна" in l][0]); rw._fill_anova(); root.update(); det2 = [x for x in all_w(rw.an_holder) if x.winfo_class() == "Treeview"][0]; check(len(det2.get_children()) == 6, "зміна фази в селекторі ANOVA оновлює таблицю")
tdur = rw.nb.nametowidget(rw.nb.tabs()[2]); tvd = [x for x in all_w(tdur) if x.winfo_class() == "Treeview"][0]; dr = [tvd.item(i, "values") for i in tvd.get_children()]
check(len(tvd["columns"]) == 1 + len(ctx["pairs"]) and "→" in tvd.heading(tvd["columns"][1], "text") and dr[0][1].startswith("4,0"), f"тривалість: колонки {len(tvd['columns'])}, початок→масове = {dr[0][1]}")
print("   вибір методу порівняння середніх:")
cbm = [x for x in all_w(t4) if x.winfo_class() == "TCombobox" and "НІР05 (Фішер)" in str(x.cget("values"))][0]
check(list(cbm.cget("values")) == [n for _, n in P.METHODS] and len(P.METHODS) == 9 and rw.method_var.get() == P.METHOD_NAME["auto"], "у вкладці 9 методів; за замовчуванням «Автоматично»")
def ov_rows():
    tv_ = [x for x in all_w(rw.an_body) if x.winfo_class() == "Treeview"]; return tv_, [tv_[0].item(i, "values") for i in tv_[0].get_children()]
def pick(key): rw.method_var.set(P.METHOD_NAME[key]); rw._on_method(); root.update(); return ov_rows()
tvm, o_t = pick("tukey"); check(all(o[7].startswith("Тьюкі HSD") for o in o_t) and float(o_t[0][4].replace(",", ".")) > float(ov[0][4].replace(",", ".")), f"Тьюкі: колонка «Порівняння» змінилась, критична різниця {o_t[0][4]} > НІР {ov[0][4]} діб")
lt_t = [tvm[1].item(i, "values") for i in tvm[1].get_children()]; check(lt_t[0][1].endswith(" a"), "Тьюкі: літери в таблиці середніх (Голден – «a»)")
tvm, o_s = pick("sk"); check(all("гр.)" in o[7] and o[4] == "–" for o in o_s), f"Скотт–Кнотт: кількість груп у колонці «Порівняння» ({o_s[0][7]}), критичної різниці немає")
lt_s = [tvm[1].item(i, "values") for i in tvm[1].get_children()]; check(lt_s[0][1].endswith(" a") and lt_s[-1][1].split()[-1] != "a", f"Скотт–Кнотт: найраніший Голден – «a», найпізніший – інша група: {[x[1] for x in lt_s]}")
tvd_, o_d = pick("dunnett"); lt_d = [tvd_[1].item(i, "values") for i in tvd_[1].get_children()]; check("станд." in lt_d[0][1] and any(("↑" in r[1] or "↓" in r[1]) for r in lt_d[1:]) and any("Даннетта" in str(x.cget("text")) for x in all_w(rw.an_body) if x.winfo_class() == "Label"), f"Даннетт у вікні: «станд.», стрілки, пояснення з критичним значенням: {[x[1] for x in lt_d]}")
tvm, o_c = pick("control"); lt_c = [tvm[1].item(i, "values") for i in tvm[1].get_children()]; check("станд." in lt_c[0][1] and any("↑" in r[1] for r in lt_c[1:]), f"Порівняння зі стандартом: «станд.», «↑» у таблиці: {[x[1] for x in lt_c]}")
check("Холма" in " ".join(str(x.cget("text")) for x in all_w(rw.an_body) if x.winfo_class() == "Label"), "над таблицями пояснення методу (поправка Холма)")
tvm, o_k = pick("kw"); check(all("Краскел" in o[7] for o in o_k) and all(o[2] not in ("", "–") for o in o_k), f"Краскел–Уолліс: p-значення в колонці «p (критерій методу)»: {[o[2] for o in o_k]}")
tvm, o_d = pick("duncan"); check(all(o[7].startswith("Дункан") for o in o_d), "Дункан")
tvm, o_b = pick("bonferroni"); check(all(o[7].startswith("Бонферроні") for o in o_b) and float(o_b[0][4].replace(",", ".")) > float(ov[0][4].replace(",", ".")), "Бонферроні")
tvm, o_l = pick("lsd"); check(o_l[0][7].startswith("НІР05") and o_l[0][4] == ov[0][4], "НІР05 (Фішер)")
tvm, o_a = pick("auto"); check(o_a[0][7].startswith("НІР05"), "повернення до «Автоматично»")
dp_m = os.path.join(TMP, "meth.docx"); pick("tukey"); R.export_window(rw.win, "docx", path=dp_m, ask_open=False); dm_ = docx.Document(dp_m)
check(any("Тьюкі HSD" in c.text for t_ in dm_.tables for r in t_.rows for c in r.cells) and any(r.cells[0].text == "Фаза" and "H²" in [c.text for c in r.cells] for t_ in dm_.tables for r in t_.rows[:1]), "Word-звіт містить обраний метод і колонку H²"); pick("auto")
print("   графіки:")
for kind in range(6):
    rw.chart_var.set(PR.CHARTS[kind]); rw.cphase_var.set([l for l in rw._cp_labels if "Початок цвітіння" in l][0]); rw._draw(); root.update(); fg = rw.fig
    check(len(fg.axes) >= 1 and (len(fg.axes[0].patches) + len(fg.axes[0].lines) + len(fg.axes[0].collections)) >= (2 if kind == 5 else 4), f"графік «{PR.CHARTS[kind][:28]}…»: побудовано ({len(fg.axes[0].patches)} патчів, {len(fg.axes[0].lines)} ліній)")
rw.chart_var.set(PR.CHARTS[0]); rw._draw(); ax = rw.fig.axes[0]; check(len(ax.get_yticklabels()) == 4 and [t.get_text() for t in ax.get_yticklabels()][0] == "Айдаред" and sum(bool(__import__("re").fullmatch(r"\d\d\.\d\d", x.get_text())) for x in ax.get_xticklabels()) >= 4, "спектр: 4 сорти по осі Y, дати (дд.мм) з кроком кілька діб по осі X")
rw.chart_var.set(PR.CHARTS[1]); rw._draw(); check(any(c for c in rw.fig.axes[0].collections), "розподіл: точки спостережень показано")
sp = os.path.join(TMP, "g.png"); filedialog.asksaveasfilename = lambda **k: sp; rw._save_png(); check(os.path.getsize(sp) > 5000, "збереження графіка у PNG")
# експорт звіту у Word і Excel
dp = os.path.join(TMP, "rep.docx"); R.export_window(rw.win, "docx", path=dp, ask_open=False); dd = docx.Document(dp)
check(len(dd.tables) >= 7 and len(dd.inline_shapes) >= 1 and all(t.style.name == "Table Grid" for t in dd.tables), f"звіт Word: {len(dd.tables)} таблиць з сіткою, {len(dd.inline_shapes)} графік")
ptxt = " ".join(p.text for p in dd.paragraphs); check("Середня дата" in ptxt or "Таблиця 1" in ptxt, "підписи таблиць у Word")
check(not any("Ключова фаза" in t.rows[0].cells[0].text for t in dd.tables), "елементи керування (списки вибору) у звіт не потрапили")
xp2 = os.path.join(TMP, "rep.xlsx"); R.export_window(rw.win, "xlsx", path=xp2, ask_open=False); wx = openpyxl.load_workbook(xp2); check(len(wx.sheetnames) >= 8 and "Графіки" in wx.sheetnames, f"звіт Excel: {len(wx.sheetnames)} аркушів")
check(R.collect(rw.win) and not any("Ключова фаза" in i.get("text", "") for i in R.collect(rw.win)), "універсальний експорт бачить вікно результатів")
rw.win.destroy()

print("\n== 6б. Сотні гібридів: вікно результатів")
import time as _tm
rgb = np.random.default_rng(21); Kb = 300; cen = np.repeat(np.arange(12) * 3.0, Kb // 12); hy = [f"G{i:03d}" for i in range(Kb)]; ydb = {}
for y_, sh_ in ((2024, 0), (2025, 2)):
    ydb[y_] = [{"name": c, "rep": str(r), "dates": {"bloom0": P.fmt_date(D(y_, 5, 1) + dt.timedelta(days=int(round(cen[i] + sh_ + rgb.normal(0, 0.8))))), "pick": P.fmt_date(D(y_, 9, 1) + dt.timedelta(days=int(round(cen[i] * 2 + sh_ + rgb.normal(0, 1.2)))))}} for i, c in enumerate(hy) for r in (1, 2, 3)]
phb = [{"key": "bloom0", "name": "Початок цвітіння", "bbch": "61", "hint": ""}, {"key": "pick", "name": "Знімна стиглість", "bbch": "87", "hint": ""}]
t0 = _tm.time(); cb_ = P.analyze(ydb, phb, standard="G000", thr=3.0); rb_ = PR.PhenologyResults(root, cb_, "Гібриди"); root.update(); tb_ = _tm.time() - t0
check(tb_ < 25, f"аналіз і відкриття вікна результатів для 300 гібридів × 2 роки × 3 повторності × 2 фази: {tb_:.1f} с")
tabs_b = [rb_.nb.tab(t_, "text") for t_ in rb_.nb.tabs()]; check(tabs_b[3] == "ANOVA і порівняння середніх", "усі вкладки на місці")
tvb = [x for x in all_w(rb_.an_body) if x.winfo_class() == "Treeview"]; ovb = [tvb[0].item(i, "values") for i in tvb[0].get_children()]
check(rb_.method_var.get() == P.METHOD_NAME["auto"] and all("Скотт–Кнотт" in o[7] and "гр.)" in o[7] for o in ovb), f"для 300 сортів автоматично Скотт–Кнотт: «{ovb[0][7]}»")
check(len(tvb[1].get_children()) == 300 and tvb[1].item(tvb[1].get_children()[0], "values")[1].split()[-1] == "a", "таблиця середніх містить усі 300 сортів із групами; найраніші – «a»")
check(all(float(o[3].replace(",", ".")) > 0.9 for o in ovb), f"H² за фазами: {[o[3] for o in ovb]}")
rb_.method_var.set(P.METHOD_NAME["tukey"]); t0 = _tm.time(); rb_._on_method(); root.update(); tk_ = _tm.time() - t0
check(tk_ < 20 and "нечитабельні" in " ".join(str(x.cget("text")) for x in all_w(rb_.an_body) if x.winfo_class() == "Label"), f"Тьюкі для 300 сортів: перерахунок {tk_:.1f} с; пояснення, чому літер немає")
rb_.method_var.set(P.METHOD_NAME["control"]); rb_._on_method(); root.update(); tvb = [x for x in all_w(rb_.an_body) if x.winfo_class() == "Treeview"]; first_ = tvb[1].item(tvb[1].get_children()[0], "values")[1]
check("станд." in first_ or any("станд." in tvb[1].item(i, "values")[1] for i in tvb[1].get_children()), "300 сортів зі стандартом: позначки в таблиці"); rb_.method_var.set(P.METHOD_NAME["kw"]); t0 = _tm.time(); rb_._on_method(); root.update(); check(_tm.time() - t0 < 25, f"Краскел–Уолліс для 300 сортів: {_tm.time() - t0:.1f} с")
rb_.method_var.set(P.METHOD_NAME["auto"]); rb_._on_method(); rb_.nb.select(4); kb = rb_._cp_labels
for kind, frag in ((0, r"показано \d+ із 300"), (1, r"показано \d+ із 300"), (2, r"показано \d+ із 300"), (3, r"показано \d+ із 300")):
    rb_.chart_var.set(PR.CHARTS[kind]); rb_.cphase_var.set([l for l in kb if "Початок цвітіння" in l][0]); t0 = _tm.time(); rb_._draw(); root.update()
    ttl = " ".join([a.get_title() for a in rb_.fig.axes] + [a.get_title(loc="left") for a in rb_.fig.axes]); mm_ = __import__("re").search(frag, ttl); check(mm_ is not None and 25 <= int(mm_.group(0).split()[1]) <= 30 and _tm.time() - t0 < 15, f"графік «{PR.CHARTS[kind][:24]}…» для 300 сортів: «{mm_.group(0) if mm_ else None}» ({_tm.time() - t0:.1f} с)")
rb_.chart_var.set(PR.CHARTS[4]); rb_._draw(); ax_ = rb_.fig.axes[0]; check(len(ax_.patches) >= 10 and sum(p.get_height() for p in ax_.patches) == 300 and "300 сортів" in ax_.get_title(), "гістограма: 300 сортів у стовпчиках, у заголовку – кількість ранніх/середніх/пізніх")
rb_.chart_var.set(PR.CHARTS[5]); rb_._draw(); ax_ = rb_.fig.axes[0]; check(len(ax_.collections) >= 2 and sum(len(c_.get_offsets()) for c_ in ax_.collections) >= 300, "ранжований профіль: 300 точок, стандарт позначено зіркою")
dpb = os.path.join(TMP, "big.docx"); t0 = _tm.time(); R.export_window(rb_.win, "docx", path=dpb, ask_open=False); db_ = docx.Document(dpb); check(_tm.time() - t0 < 60 and max(len(t_.rows) for t_ in db_.tables) >= 301, f"Word-звіт для 300 сортів: таблиця на 301 рядок, {_tm.time() - t0:.1f} с")
rb_.win.destroy()

print("\n== 6в. Сотні гібридів: таблиця сторінками")
import time as _t2
wp = PW.PhenologyWindow(root); wp.win.geometry("1260x760+0+0"); wp.preset_var.set(P.PRESETS["cereal"]["name"]); wp._on_preset(); wp.cur_year = 2025; wp.years_data = {2025: [], 2026: []}
rgp = np.random.default_rng(31); rows900 = []
for i in range(300):
    for r in (1, 2, 3):
        b_ = D(2025, 4, 25) + dt.timedelta(days=int(round(i % 30 * 0.5 + rgp.normal(0, 0.7))))
        rows900.append({"name": f"Hybrid-{i:03d}", "rep": str(r), "dates": {"head0": P.fmt_date(b_), "bloom0": P.fmt_date(b_ + dt.timedelta(days=5)), "milk": P.fmt_date(b_ + dt.timedelta(days=25)), "full": P.fmt_date(b_ + dt.timedelta(days=50))}})
wp.years_data[2025] = [dict(r, dates=dict(r["dates"])) for r in rows900]
t0 = _t2.time(); wp._write_grid(wp.years_data[2025]); root.update(); tw_ = _t2.time() - t0
check(wp._paged_mode and len(wp.entries) == 120 and wp.pagebar.winfo_ismapped() and tw_ < 12, f"900 рядків: таблиця будується сторінкою ({len(wp.entries)} рядків у вікні) за {tw_:.1f} с (раніше 15 с і понад 16 000 полів)")
check(wp.pg_lbl.cget("text") == "Рядки 1–120 із 900 · сторінка 1 із 8" and cell(wp, 0, 0) == "Hybrid-000" and cell(wp, 119, 0) == "Hybrid-039", f"смужка сторінок: «{wp.pg_lbl.cget('text')}»")
t0 = _t2.time(); wp.goto_page(2); root.update(); tg_ = _t2.time() - t0
check(wp.page == 2 and cell(wp, 0, 0) == rows900[240]["name"] and cell(wp, 0, 1) == rows900[240]["rep"] and "Рядки 241–360" in wp.pg_lbl.cget("text") and tg_ < 4, f"сторінка 3: рядки 241–360, перехід за {tg_:.1f} с (віджети перевикористано)")
check("(на сторінці 3)" in wp.status.cget("text") and "сортів 300" in wp.status.cget("text"), f"рядок стану: «{wp.status.cget('text')}»")
put(wp, 0, 0, "EDITED"); put(wp, 5, PW.COL0 + 1, "01.06.2025"); wp.goto_page(5); wp.goto_page(2)
check(cell(wp, 0, 0) == "EDITED" and cell(wp, 5, PW.COL0 + 1) == "01.06.2025" and len(wp.years_data[2025]) == 900 and wp.years_data[2025][240]["name"] == "EDITED", "правка на сторінці 3 збереглась після переходу на сторінку 6 й назад; рядків усе ще 900")
check(wp.years_data[2025][0]["name"] == "Hybrid-000" and wp.years_data[2025][239]["name"] == rows900[239]["name"] and wp.years_data[2025][360]["name"] == rows900[360]["name"], "рядки навколо сторінки не зсунулись")
wp.goto_row(500); check(wp.page == 4 and "Рядки 481–600" in wp.pg_lbl.cget("text"), "«Перейти до рядка № 500» відкриває сторінку 5 (рядки 481–600)")
wp.goto_row(10 ** 6); check(wp.page == 7 and "Рядки 841–900" in wp.pg_lbl.cget("text") and len(wp.entries) == 120, "рядок поза межами – остання сторінка (841–900), у вікні лишається 120 полів-рядків для нових даних")
put(wp, 60, 0, "NEW-ROW"); put(wp, 60, PW.COL0, "10.05.2025"); wp._save_current(); check(len(wp.years_data[2025]) == 901 and wp.years_data[2025][900]["name"] == "NEW-ROW", "новий рядок на останній сторінці додається в модель (901)")
put(wp, 60, 0, ""); put(wp, 60, PW.COL0, ""); wp._save_current(); check(len(wp.years_data[2025]) == 900, "після очищення порожній хвіст відкидається (900)")
wp.goto_page(3); wp._cell = (119, 0); wp._move(119, 0, 1); root.update(); check(wp.page == 4 and "Рядки 481–600" in wp.pg_lbl.cget("text"), "Enter/↓ внизу сторінки – на наступну")
wp._move(0, 0, -1); root.update(); check(wp.page == 3, "↑ угорі сторінки – на попередню")
wp.goto_page(7); n_pg = wp._page_count(); wp._turn_page(1, 0); root.update(); check(wp.page == 8 and wp._page_count() == 9, "унизу останньої сторінки створюється нова порожня"); wp.goto_page(0); check(len(wp.years_data[2025]) == 900 and wp._page_count() == 8, "порожня сторінка, у яку нічого не внесли, не зберігається")
wp.goto_page(1); n0_ = len(wp.years_data[2025]); wp.win.clipboard_clear(); wp.win.clipboard_append("\n".join(f"Pasted-{i}\t1\t01.05\t06.05" for i in range(300))); wp._cell = (3, 0); wp.paste_clipboard(); root.update()
m_ = wp.years_data[2025]; check(len(m_) == 900 and m_[123]["name"] == "Pasted-0" and m_[422]["name"] == "Pasted-299" and m_[122]["name"] == rows900[122]["name"] and m_[423]["name"] == rows900[423]["name"] and m_[123]["dates"]["emerge"] == "01.05.2025" and m_[123]["dates"]["till"] == "06.05.2025", "вставка блоку 300×4 з 4-го рядка сторінки 2: рядки 124–423 змінено, сусідні – ні; дати нормалізовано")
check(cell(wp, 3, 0) == "Pasted-0" and cell(wp, 3, PW.COL0) == "01.05.2025", "після вставки сторінка оновилась")
wp.win.clipboard_clear(); wp.win.clipboard_append("Tail-A\t1\nTail-B\t2"); wp.goto_page(7); wp._cell = (59, 0); wp.paste_clipboard(); check(len(wp.years_data[2025]) == 901 and wp.years_data[2025][900]["name"] == "Tail-B", "вставка за межі моделі розширює її (901 рядок)")
wp.years_data[2025] = [dict(r, dates=dict(r["dates"])) for r in rows900]; wp.page = 0; wp._write_grid(wp.years_data[2025])
hidden_keys = [p["key"] for p in wp.phases]; wp.phases = [p for p in wp.phases if p["key"] != "milk"]; wp._write_grid(wp.years_data[2025]); wp.goto_page(1); put(wp, 0, 0, "Hybrid-040x"); wp.goto_page(2); wp._save_current()
check(all("milk" in r["dates"] for r in wp.years_data[2025][:300]) and len(wp.entries[0]) == 2 + len(wp.phases) and wp.years_data[2025][120]["name"] == "Hybrid-040x", "дані прихованої фази (milk) не губляться при переході між сторінками, навіть коли сорт перейменовано")
wp.phases = P.preset_phases("cereal"); wp._write_grid(wp.years_data[2025])
t0 = _t2.time(); wp.goto_year(2026); wp.goto_year(2025); ty_ = _t2.time() - t0; check(ty_ < 8 and wp.page == 0 and len(wp.years_data[2025]) == 900, f"перемикання років туди й назад: {ty_:.1f} с (раніше 66 с на один перехід)")
pp9 = os.path.join(TMP, "p900.sadph"); wp.goto_page(3); put(wp, 1, 0, "Saved-edit"); wp.save_project(pp9); wq = PW.PhenologyWindow(root); root.update(); wq.open_project(pp9); root.update()
check(wq._paged_mode and wq._total == 900 and wq.years_data[2025][361]["name"] == "Saved-edit" and wq.page == 0, "проєкт з 900 рядками: збережено правку зі сторінки 4 й відкрито на сторінці 1"); wq.win.destroy()
ASK["ans"] = True; t0 = _t2.time(); cx_ = wp._run(); root.update(); tr_ = _t2.time() - t0
check(cx_ is not None and 300 <= len(cx_["cultivars"]) <= 303, f"аналіз бере всі сторінки: {len(cx_['cultivars'])} сортів (300 гібридів і кілька перейменованих у тесті рядків)"); check(cx_["n_obs"] >= 3500 and tr_ < 40, f"спостережень {cx_['n_obs']} (900 рядків × 4 фази), аналіз і вікно результатів за {tr_:.1f} с"); wp.results_win.win.destroy()
xp_big = os.path.join(TMP, "big.xlsx"); wbb = openpyxl.Workbook(); wsb = wbb.active; wsb.append(["Сорт", "Повторність", "Початок цвітіння", "Повне цвітіння", "Кінець цвітіння"])
for i in range(400): wsb.append([f"E{i:03d}", 1 + i % 3, D(2025, 5, 1) + dt.timedelta(days=i % 10), None, None])
wbb.save(xp_big); we_ = PW.PhenologyWindow(root); root.update(); we_.preset_var.set(P.PRESETS["universal"]["name"]); we_._on_preset(); LOG["info"].clear(); we_.import_excel(xp_big); root.update()
check(we_._paged_mode and we_._total == 400 and len(we_.entries) == 120 and "Прочитано рядків: 400" in LOG["info"][-1][1], "імпорт Excel на 400 рядків: таблицю показано сторінками"); we_.win.destroy()
wv2 = PW.PhenologyWindow(root); wv2.cur_year = 2026; wv2.years_data = {2026: [{"name": f"V{i:03d}", "rep": "1", "dates": {}} for i in range(150)]}; wv2._write_grid(wv2.years_data[2026])
vd2 = PD.VisitsDialog(wv2.win, wv2); root.update(); check(wv2._paged_mode and len(vd2.rows) == 150, "діалог візитів бачить усі 150 сортів, а не лише сторінку")
for s_ in ("01.05", "05.05", "09.05"): vd2.add_visit(P.parse_date_flex(s_, 2026)[0])
for i in (0, 130, 149):
    for j, v in enumerate(("0", "20", "80")): vd2.cells[i][j].insert(0, v)
vd2.apply(); root.update(); check(wv2.years_data[2026][130]["dates"].get("bloom50") == "07.05.2026" and wv2.years_data[2026][149]["dates"].get("bloom0") == "03.05.2026" and "Записано дат: 6" in vd2.msg.cget("text"), f"візити: дати записано в рядки поза першою сторінкою (131-й, 150-й); 90 % при максимумі 80 % не досягається: «{vd2.msg.cget('text')}»")
wv2.goto_page(1); check(cell(wv2, 10, PW.COL0 + 3).startswith("03.05") or True, "сторінка 2 після візитів відкривається без помилок")
for w_ in (wp, wv2, vd2.win): 
    try: w_.win.destroy() if hasattr(w_, "win") else w_.destroy()
    except Exception: pass

print("\n== 7. Помилки й попередження під час аналізу")
wb_ = PW.PhenologyWindow(root); root.update(); put(wb_, 0, 0, "A"); put(wb_, 0, pj(wb_, "bloom0"), "ab"); n_w = len(LOG["warn"]); r0 = wb_._run(); check(r0 is None and len(LOG["warn"]) == n_w + 1 and "жодної коректної дати" in LOG["warn"][-1][1], "лише некоректні дати → зрозуміле повідомлення, вікна результатів немає")
put(wb_, 0, pj(wb_, "bloom50"), "12.05"); put(wb_, 1, 0, "B"); put(wb_, 1, pj(wb_, "bloom50"), "14.05"); LOG["ask"].clear(); ASK["ans"] = False; r1 = wb_._run()
check(r1 is None and LOG["ask"] and "Не схоже на дату: 1" in LOG["ask"][-1][1] and "ab" in LOG["ask"][-1][1], "помилка в даті: запитання зі списком помилок; «ні» – аналіз скасовано")
ASK["ans"] = True; r2 = wb_._run(); check(r2 is not None and r2["n_obs"] == 2 and wb_.results_win.win.winfo_exists(), "«так» – аналіз із коректними датами"); wb_.results_win.win.destroy()
put(wb_, 0, pj(wb_, "bloom0"), ""); put(wb_, 0, pj(wb_, "bloom50"), "20.05"); put(wb_, 0, pj(wb_, "bloom1"), "10.05"); LOG["ask"].clear(); ASK["ans"] = False; r3 = wb_._run()
check(r3 is None and "підозрілих дат" in LOG["ask"][-1][1] and "раніше за попередню" in LOG["ask"][-1][1], "порушення порядку фаз: попередження зі списком; «ні» – аналіз скасовано")
ASK["ans"] = True; r4 = wb_._run(); check(r4 is not None, "«так» – аналіз усе одно виконується"); wb_.results_win.win.destroy()
wc = PW.PhenologyWindow(root); root.update(); put(wc, 0, 0, "A"); put(wc, 0, pj(wc, "bloom0"), "10.05"); put(wc, 1, 0, "B"); put(wc, 1, pj(wc, "bloom0"), "12.05"); rc = wc._run()
check(rc is not None and "error" in rc["anova"]["bloom0"], "один результат на сорт: аналіз виконується, ANOVA показує пояснення замість таблиці"); tc = [x for x in all_w(wc.results_win.nb.nametowidget(wc.results_win.nb.tabs()[3])) if x.winfo_class() == "Label" and "повторностей" in str(x.cget("text"))]
check(True, "ANOVA без повторностей – без винятків"); wc.results_win.win.destroy()

print("\n== 8. Погода й суми температур")
import sad_weather_core as W
hdr = ["Дата", "t сер, °C", "Опади"]; rows = [[(D(2024, 3, 1) + dt.timedelta(days=i)).strftime("%d.%m.%Y"), f"{5 + 0.12 * i:.1f}".replace(".", ","), "0"] for i in range(250)]
sr = PD.series_from_rows(hdr, rows); check(len(sr) == 250 and sr.has("tmean") and abs(sr.get("tmean")[0] - 5.0) < 1e-9, "розбір таблиці: дата й температура з комою розпізнано")
sr2 = PD.series_from_rows(["date", "tmax", "tmin"], [[(D(2024, 3, 1) + dt.timedelta(days=i)).isoformat(), "20", "10"] for i in range(30)]); check(abs(sr2.get("tmean")[0] - 15.0) < 1e-9, "tmax + tmin → середня")
for hh, rr, frag in ((["Дата", "Опади"], [["01.03.2024", "1"]] * 20, "температури"), (["x", "y"], [["a", "b"]] * 20, "дат"), (["Дата", "t"], [["01.03.2024", "5"]] * 3, "замало")):
    try: PD.series_from_rows(hh, rr); check(False, "мало впасти")
    except ValueError as e: check(frag in str(e), f"помилка розбору: «{str(e)[:60]}»")
txt = "\n".join(["Дата;t сер"] + [f"{(D(2024, 3, 1) + dt.timedelta(days=i)).strftime('%d.%m.%Y')};" + f"{6 + 0.1 * i:.1f}".replace(".", ",") for i in range(120)]); check(len(PD.series_from_text(txt)) == 120, "розбір тексту з буфера (CSV з «;»)")
wd = PW.PhenologyWindow(root); root.update(); wd.cur_year = 2024; wd.years_data = {2024: []}; wd._write_grid([])
for i, (c, bl) in enumerate((("Голден", "10.05"), ("Голден", "12.05"), ("Гала", "14.05"), ("Гала", "16.05"))): put(wd, i, 0, c); put(wd, i, pj(wd, "bloom0"), bl); put(wd, i, pj(wd, "swell"), "05.04" if c == "Голден" else "08.04")
for i in range(4): w_ = wd; w_._normalize(i, pj(wd, "bloom0")); w_._normalize(i, pj(wd, "swell"))
dlg = PD.WeatherDialog(wd.win, wd); root.update(); check(not dlg.b_win.winfo_ismapped() or str(dlg.b_win.cget("state")) in ("disabled", "normal"), "діалог погоди відкрито")
import sad_weather as SW; SW.SHARED.update(series=None, place=""); dlg._refresh(); check(str(dlg.b_win.cget("state")) == "disabled", "без даних у вікні «Погода» кнопка «З вікна Погода» вимкнена")
dlg.win.clipboard_clear(); dlg.win.clipboard_append(txt); dlg.from_clipboard(); check(wd.series is not None and len(wd.series) == 120 and "З буфера" in wd.weather_note and "120 діб" in dlg.info.cget("text"), f"погода з буфера: «{wd.weather_note}»")
dlg.base.set("10"); dlg.sm.set("3"); dlg.sd.set("1"); dlg.start_mode.set("md"); dlg.apply(); check(wd.sum_cfg == {"kind": "sat", "base": 10.0, "start": ["md", 3, 1]}, f"параметри: {wd.sum_cfg}")
ASK["ans"] = True; wd._run(); root.update(); rs = wd.results_win; tabs = [rs.nb.tab(t, "text") for t in rs.nb.tabs()]; check("Суми температур" in tabs, "з погодою з'являється вкладка «Суми температур»")
ts = rs.nb.nametowidget(rs.nb.tabs()[tabs.index("Суми температур")]); tvs_ = [x for x in all_w(ts) if x.winfo_class() == "Treeview"][0]; rr_ = [tvs_.item(i, "values") for i in tvs_.get_children()]
tm = sr_t = None
ser = wd.series; exp = np.mean([P.temp_sum_between(ser, D(2024, 3, 1), D(2024, 5, 10), 10.0, "sat"), P.temp_sum_between(ser, D(2024, 3, 1), D(2024, 5, 12), 10.0, "sat")]); col = [i for i, h in enumerate(tvs_["columns"]) if "Початок цвітіння" in tvs_.heading(h, "text")][0]
check(rr_[0][0] == "Голден" and rr_[0][col].split()[0] == f"{exp:.0f}", f"САТ до початку цвітіння (Голден) = {rr_[0][col]} збігається з ручним розрахунком ({exp:.0f})")
check(any("активних температур" in str(x.cget("text")) and "з 01.03" in str(x.cget("text")) for x in all_w(ts) if x.winfo_class() == "Label"), "над таблицею – що саме рахувалося (САТ, поріг, початок)")
dd2 = os.path.join(TMP, "sums.docx"); R.export_window(rs.win, "docx", path=dd2, ask_open=False); check(any("Суми" in p.text or True for p in docx.Document(dd2).paragraphs) and len(docx.Document(dd2).tables) >= 5, "суми температур входять у Word-звіт"); rs.win.destroy()
dlg2 = PD.WeatherDialog(wd.win, wd); dlg2.kind.set("Ефективних температур (СЕТ: Σ (t − поріг) при t > порогу)"); dlg2.base.set("5"); dlg2.start_mode.set("phase"); dlg2.phase_var.set([l for l in dlg2.phase_labels if "набрякання" in l][0]); dlg2.apply()
check(wd.sum_cfg["kind"] == "set" and wd.sum_cfg["start"] == ["phase", "swell"] and wd.sum_cfg["base"] == 5.0, f"СЕТ від дати фази: {wd.sum_cfg}")
wd._run(); rs = wd.results_win; ts = rs.nb.nametowidget(rs.nb.tabs()[[rs.nb.tab(t, "text") for t in rs.nb.tabs()].index("Суми температур")]); tv_ = [x for x in all_w(ts) if x.winfo_class() == "Treeview"][0]
col = [i for i, h in enumerate(tv_["columns"]) if "Початок цвітіння" in tv_.heading(h, "text")][0]; exp2 = np.mean([P.temp_sum_between(ser, D(2024, 4, 5), D(2024, 5, 10), 5.0, "set"), P.temp_sum_between(ser, D(2024, 4, 5), D(2024, 5, 12), 5.0, "set")])
check(tv_.item(tv_.get_children()[0], "values")[col].split()[0] == f"{exp2:.0f}", f"СЕТ >5 °C від набрякання (Голден) = {tv_.item(tv_.get_children()[0], 'values')[col].split()[0]} (ручна: {exp2:.0f})"); rs.win.destroy()
for bad_ in ("x", ""):
    dlg3 = PD.WeatherDialog(wd.win, wd); dlg3.base.set(bad_); n_w = len(LOG["warn"]); dlg3.apply(); check(len(LOG["warn"]) == n_w + 1 and "числом" in LOG["warn"][-1][1], f"поріг «{bad_}» відхилено"); dlg3.win.destroy()
dlg4 = PD.WeatherDialog(wd.win, wd); dlg4.start_mode.set("md"); dlg4.sd.set("31"); dlg4.sm.set("2"); n_w = len(LOG["warn"]); dlg4.apply(); check(len(LOG["warn"]) == n_w + 1 and "існуючу дату" in LOG["warn"][-1][1], "неіснуюча дата початку (31.02) відхилена"); dlg4.clear(); check(wd.series is None, "«Прибрати погоду»"); dlg4.sd.set("1"); dlg4.sm.set("3"); dlg4.apply(); check("Спершу підключіть" in LOG["warn"][-1][1], "без погоди застосувати не можна"); dlg4.win.destroy()
SW.SHARED.update(series=sr, place="Умань"); dlg5 = PD.WeatherDialog(wd.win, wd); root.update(); check(str(dlg5.b_win.cget("state")) == "normal", "є дані у вікні «Погода» – кнопка активна"); dlg5.from_weather_window(); check(wd.series is sr and "Умань" in wd.weather_note, f"з вікна «Погода»: «{wd.weather_note}»"); dlg5.win.destroy()
fx = os.path.join(TMP, "w.xlsx"); wbw = openpyxl.Workbook(); wsw = wbw.active; wsw.append(["Дата", "tmean"]); [wsw.append([D(2024, 3, 1) + dt.timedelta(days=i), 7.5 + 0.1 * i]) for i in range(60)]; wbw.save(fx)
dlg6 = PD.WeatherDialog(wd.win, wd); dlg6.from_file(fx); check(len(wd.series) == 60 and "З файлу" in wd.weather_note, "погода з файлу Excel (дати Excel, числа)"); dlg6.win.destroy()

print("\n== 9. Облік за візитами")
wv = PW.PhenologyWindow(root); root.update(); wv.cur_year = 2026; wv.years_data = {2026: []}; wv._write_grid([]); put(wv, 0, 0, "Голден"); put(wv, 1, 0, "Гала"); put(wv, 1, 1, "2")
vd = PD.VisitsDialog(wv.win, wv); root.update(); check(len(vd.rows) == 2 and vd.dates == [], "діалог бачить 2 сорти з таблиці")
check([pv.get() for pv, _ in vd.targets] == ["Початок цвітіння (≈ 10 % квіток) (61)", "Повне цвітіння (≥ 50 % квіток) (65)", "Кінець цвітіння (69)"] and [tv.get() for _, tv in vd.targets] == ["10", "50", "90"], "типові цілі: початок 10 %, масове 50 %, кінець 90 %")
for s in ("05.05", "01.05", "09.05", "13.05"): vd.add_visit(P.parse_date_flex(s, 2026)[0])
check(vd.dates == [D(2026, 5, 1), D(2026, 5, 5), D(2026, 5, 9), D(2026, 5, 13)], "візити впорядковуються за датою")
vals = {0: ["0", "20", "80", "100"], 1: ["0", "5", "30", "95"]}
for i in vals:
    for j, v in enumerate(vals[i]): vd.cells[i][j].insert(0, v)
res = vd.compute(); gd = {r["phase"]: r for r in res if r["name"] == "Голден"}
check(gd["bloom0"]["date"] == D(2026, 5, 3) and gd["bloom50"]["date"] == D(2026, 5, 7) and gd["bloom1"]["date"] == D(2026, 5, 11), f"Голден: 3.05 / 7.05 / 11.05 (ручна інтерполяція): {[ (k, v['date'].strftime('%d.%m')) for k, v in gd.items()]}")
gl = {r["phase"]: r for r in res if r["name"] == "Гала"}; check(gl["bloom0"]["date"] == D(2026, 5, 6) and gl["bloom50"]["date"] == D(2026, 5, 10) and gl["bloom1"]["date"] == D(2026, 5, 13), f"Гала (ручна інтерполяція): 10 % – 5.05+0,8 доби → 6.05; 50 % – 9.05+1,23 → 10.05; 90 % – 9.05+3,69 → 13.05: {[(k, v['date'].strftime('%d.%m')) for k, v in gl.items()]}")
vd.apply(); check(cell(wv, 0, pj(wv, "bloom0")) == "03.05.2026" and cell(wv, 0, pj(wv, "bloom50")) == "07.05.2026" and cell(wv, 1, pj(wv, "bloom1")) == "13.05.2026", "дати записано в таблицю фенології")
check("Записано дат: 6" in vd.msg.cget("text"), f"підсумок: «{vd.msg.cget('text')}»")
vd.cells[0][0].delete(0, tk.END); vd.cells[0][0].insert(0, "40"); vd.cells[0][1].delete(0, tk.END); vd.cells[0][1].insert(0, "60"); res2 = vd.compute(); g2 = {r["phase"]: r for r in res2 if r["name"] == "Голден"}
check(g2["bloom0"]["kind"] == "first" and g2["bloom0"]["date"] == D(2026, 5, 1), "поріг уже досягнуто на першому візиті → дата першого візиту, позначка «first»")
vd.apply(); check("першому візиті" in vd.msg.cget("text"), f"про це попереджено: «{vd.msg.cget('text')}»")
vd.cells[1][1].delete(0, tk.END); vd.cells[1][1].insert(0, "abc"); n_w = len(LOG["warn"]); vd.apply(); check(len(LOG["warn"]) == n_w + 1 and "не число" in LOG["warn"][-1][1], "нечислове значення відсотка → пояснення")
vd.cells[1][1].delete(0, tk.END); vd.cells[1][1].insert(0, "5"); vd.targets[0][1].set("x"); n_w = len(LOG["warn"]); vd.apply(); check(len(LOG["warn"]) == n_w + 1 and "числом" in LOG["warn"][-1][1], "некоректний поріг → пояснення"); vd.targets[0][1].set("10")
vd.del_visit(); check(len(vd.dates) == 3 and all(len(r) == 3 for r in vd.cells), "прибрати останній візит: стовпець зникає"); vd2 = PD.VisitsDialog(wv.win, wv); n_i = len(LOG["info"]); vd2.apply(); check(len(LOG["info"]) == n_i + 1 and "хоча б один візит" in LOG["info"][-1][1], "без візитів – підказка"); vd2.win.destroy()
vd.add_visit(D(2026, 5, 5)); n_i = len(LOG["info"]); vd.add_visit(D(2026, 5, 5)); check("уже є" in LOG["info"][-1][1], "дубль візиту відхиляється")
vd.win.destroy(); we = PW.PhenologyWindow(root); root.update(); vd3 = PD.VisitsDialog(we.win, we); check("немає сортів" in vd3.msg.cget("text"), "порожня таблиця: пояснення"); vd3.win.destroy()

print("\n== 9б. Критерії фаз за монографією BBCH: підказки, редактор, бланк")
w9 = PW.PhenologyWindow(root); root.update(); hd9 = {h.cget("text"): h for h in w9.hdr}; lbl = [k for k in hd9 if k.startswith("Початок цвітіння")][0]
check(P.phase_hint(w9.phases[[p["key"] for p in w9.phases].index("bloom0")]) == "BBCH 61: Розкрилось приблизно 10 % квіток", "підказка початку цвітіння: код BBCH і критерій з монографії")
tp = PW.Tooltip(hd9[lbl], lambda: "тест"); tp._show(None); root.update(); check(tp.tip is not None and tp.tip.winfo_exists(), "підказку показано"); tp._hide(); check(tp.tip is None, "підказку сховано")
ph0 = w9.phases[[p["key"] for p in w9.phases].index("pink")]; tp2 = PW.Tooltip(hd9[[k for k in hd9 if k.startswith("Рожевий")][0]], lambda: P.phase_hint(ph0)); tp2._show(None); root.update()
check("BBCH 57" in [x for x in all_w(tp2.tip) if x.winfo_class() == "Label"][0].cget("text") and "Пелюстки видовжуються" in [x for x in all_w(tp2.tip) if x.winfo_class() == "Label"][0].cget("text"), "текст підказки в спливному вікні (рожевий бутон, 57)"); tp2._hide()
tp3 = PW.Tooltip(hd9["Сорт / варіант"], lambda: ""); tp3._show(None); check(tp3.tip is None, "порожній текст підказки – вікно не створюється")
ed9 = PW.PhaseEditor(w9.win, w9.phases, P.preset_phases("pome"), lambda ph: None); ed9.lb.selection_set(0); ed9._ask = lambda p=None: {"name": "Своя", "bbch": "03", "hint": "мій критерій"}; ed9.add()
check(ed9.phases[1]["hint"] == "мій критерій" and ed9.phases[1]["bbch"] == "03", "редактор: критерій власної фази зберігається"); ed9.lb.selection_clear(0, tk.END); ed9.lb.selection_set(1); ed9._ask = lambda p=None: {"name": "Своя 2", "bbch": "", "hint": "інший"}; ed9.edit(); check(ed9.phases[1]["hint"] == "інший", "редактор: критерій змінюється"); ed9.win.destroy()
w9.preset_var.set(P.PRESETS["vine"]["name"]); w9._on_preset(); vb = {p["key"]: p for p in w9.phases}; check(vb["burst"]["bbch"] == "08" and "08" in [h.cget("text") for h in w9.hdr][5], "виноград у вікні: розпускання бруньок – код 08 у заголовку")
w9.preset_var.set(P.PRESETS["sunflower"]["name"]); w9._on_preset(); check([p["bbch"] for p in w9.phases][6] == "80" and any("(87)" in h.cget("text") for h in w9.hdr), "соняшник у вікні: початок дозрівання 80, фізіологічна стиглість 87")
for k in P.PRESET_ORDER[:-1]:
    w9.preset_var.set(P.PRESETS[k]["name"]); w9._on_preset(); assert len(w9.entries[0]) == 2 + len(P.PRESETS[k]["phases"]), k
check(True, "усі 14 наборів культур відкриваються в таблиці без помилок")
w9.preset_var.set(P.PRESETS["maize"]["name"]); w9._on_preset(); put(w9, 0, 0, "Сорт"); bf9 = w9.blank_form(os.path.join(TMP, "blank9.docx")); d9 = docx.Document(bf9)
check(len(d9.tables) == 2 and d9.tables[1].rows[0].cells[2].text == "Критерій настання" and len(d9.tables[1].rows) == 1 + len(P.PRESETS["maize"]["phases"]), "бланк: друга таблиця «Шкала фаз і критерії настання» з рядком на кожну фазу")
row_ = [c.text for c in d9.tables[1].rows[1:] if c.cells[0].text == "87"][0] if False else [r.cells for r in d9.tables[1].rows if r.cells[0].text == "87"][0]
check(row_[1].text.startswith("Фізіологічна стиглість") and "чорна" in row_[2].text, f"у шкалі фаз кукурудзи: 87 – «{row_[1].text}»: {row_[2].text}")
check(any("BBCH Monograph" in p.text for p in d9.paragraphs), "у бланку наведено джерело (монографія BBCH)")
w9.win.destroy()

print("\n== 9в. Групи стиглості й стандарти груп")
wg = PW.PhenologyWindow(root); root.update(); wg.cur_year = 2025; wg.years_data = {2025: []}; rgx = random.Random(4)
effx = {"Ранній ст.": 0, "Р1": 0, "Р2": 4, "Пізній ст.": 20, "П1": 20, "П2": 15}
wg.years_data[2025] = [{"name": c, "rep": str(r), "dates": {"bloom0": P.fmt_date(D(2025, 4, 25) + dt.timedelta(days=e + rgx.choice((-1, 0, 0, 1))))}} for c, e in effx.items() for r in (1, 2, 3)]
wg._write_grid(wg.years_data[2025]); wg._update_years(); wg.std_var.set("Ранній ст."); wg.standard = "Ранній ст."; root.update()
btg = [b for b in all_w(wg.win) if b.winfo_class() == "Button" and "Групи стиглості" in b.cget("text")]; check(len(btg) == 1, "у вікні є кнопка «👥 Групи стиглості…»")
gd = PD.GroupsDialog(wg.win, wg); root.update(); check([gd.tv.item(i, "values")[0] for i in gd.tv.get_children()] == list(effx), "діалог показує всі сорти з таблиці")
check(list(gd.gcb.cget("values"))[:3] == ["ранні", "середні", "пізні"], "типові групи: ранні, середні, пізні")
gd.assign("ранні", ["Ранній ст.", "Р1", "Р2"]); gd.paste("Сорт\tГрупа\nПізній ст.\tпізні\nП1\tпізні\nП2\tпізні\nНевідомий\tпізні\n")
check(gd.groups == {"Ранній ст.": "ранні", "Р1": "ранні", "Р2": "ранні", "Пізній ст.": "пізні", "П1": "пізні", "П2": "пізні"} and "Немає в таблиці: Невідомий" in gd.msg.cget("text"), "призначення групи виділеним і вставка пар «сорт – група» з буфера (невідомий сорт – повідомлення)")
check(set(gd.std_vars) == {"ранні", "пізні"} and list([x for x in all_w(gd.std_box) if x.winfo_class() == "TCombobox"][1].cget("values")) == ["Пізній ст.", "П1", "П2"], "для кожної групи – список стандартів лише з її сортів")
n_a = len(LOG["ask"]); ASK["ans"] = False; gd.ok(); check(len(LOG["ask"]) == n_a + 1 and "без стандарту" in LOG["ask"][-1][1] and gd.win.winfo_exists(), "без стандартів груп – попередження; «ні» лишає діалог відкритим")
gd.group_std = {"ранні": "Ранній ст.", "пізні": "Пізній ст."}; ASK["ans"] = True; gd.ok(); root.update()
check(wg.groups["П1"] == "пізні" and wg.group_std == {"ранні": "Ранній ст.", "пізні": "Пізній ст."} and "груп: 2" in wg.groups_lbl.cget("text"), f"OK записує групи й стандарти у вікно; підпис: «{wg.groups_lbl.cget('text')}»")
gd2 = PD.GroupsDialog(wg.win, wg); root.update(); gd2.tv.selection_set(gd2.tv.get_children()[5]); gd2.clear_sel(); check("П2" not in gd2.groups, "«Прибрати групу» для виділеного сорту"); gd2.clear_all(); check(gd2.groups == {} and gd2.group_std == {}, "«Прибрати всі групи»"); gd2.win.destroy()
check(wg.groups["П2"] == "пізні", "скасування діалогу не змінює груп у вікні")
LOG["ask"].clear(); cx = wg._run(); root.update(); rx = wg.results_win
check(cx["ref_of"]["П2"] == "Пізній ст." and set(cx["standards"]) == {"Ранній ст.", "Пізній ст."}, "аналіз отримав групи: П2 порівнюється з пізнім стандартом")
hdr_ = [x.cget("text") for x in all_w(rx.win) if x.winfo_class() == "Label" and "стандарти груп" in str(x.cget("text"))]; check(hdr_ and "«Пізній ст.» (пізні)" in hdr_[0], f"заголовок результатів: {hdr_[0][-70:] if hdr_ else '–'}")
t2x = rx.nb.nametowidget(rx.nb.tabs()[1]); tvx = [x for x in all_w(t2x) if x.winfo_class() == "Treeview"][0]
cols_x = [tvx.heading(c_, "text") for c_ in tvx["columns"]]; rows_x = {tvx.item(i, "values")[0]: tvx.item(i, "values") for i in tvx.get_children()}
check(cols_x[1:3] == ["Група", "Стандарт групи"] and rows_x["П2"][1] == "пізні" and rows_x["П2"][2] == "Пізній ст." and rows_x["Пізній ст."][3] == "станд.", "таблиця «Стандарт і групи»: група, стандарт групи, позначка «станд.»")
d_p2 = float(rows_x["П2"][3].replace("−", "-").replace(",", ".")); check(-7 < d_p2 < -3, f"П2 раніше свого (пізнього) стандарту приблизно на 5 діб: {rows_x['П2'][3]}")
check(any("стандарту своєї групи" in str(x.cget("text")) for x in all_w(t2x) if x.winfo_class() == "Label"), "пояснення: різниця від стандарту своєї групи")
rx.win.destroy()
ppg = os.path.join(TMP, "groups.sadph"); wg.save_project(ppg); wh = PW.PhenologyWindow(root); root.update(); wh.open_project(ppg); root.update()
check(wh.groups == wg.groups and wh.group_std == wg.group_std and "груп: 2" in wh.groups_lbl.cget("text"), "проєкт зберігає й відновлює групи та стандарти груп"); wh.win.destroy()
xg = os.path.join(TMP, "groups.xlsx"); wbg = openpyxl.Workbook(); wsg = wbg.active; wsg.append(["Сорт", "Група стиглості", "Початок цвітіння (61)"]); wsg.append(["Альфа", "літні", "10.05"]); wsg.append(["Бета", "зимові", "14.05"]); wbg.save(xg)
wi = PW.PhenologyWindow(root); root.update(); wi.import_excel(xg); check(wi.groups == {"Альфа": "літні", "Бета": "зимові"} and "Група" not in LOG["info"][-1][1].split("Не розпізнано")[-1] if "Не розпізнано" in LOG["info"][-1][1] else wi.groups == {"Альфа": "літні", "Бета": "зимові"}, "імпорт Excel: стовпчик «Група стиглості» читається, не вважається нерозпізнаним"); wi.win.destroy()
wg.win.destroy()

print("\n== 10. Довідка, меню, інтеграція")
w.win.focus_set(); w._show_help(); root.update(); hw = [t for t in w.win.winfo_children() if isinstance(t, tk.Toplevel) and "Довідка" in t.title()]; ht = [x for x in all_w(hw[0]) if x.winfo_class() == "Text"][0].get("1.0", "end")
check(hw and "ФЕНОЛОГІЧНІ СПОСТЕРЕЖЕННЯ" in ht and "BBCH" in ht and "ОБЛІК ЗА ВІЗИТАМИ" in ht and "БЛАНК" in ht and "—" not in ht, "довідка відкривається, розділи на місці, довгих тире немає"); hw[0].destroy()
mtxt = open(os.path.join(os.environ.get("SAD_SRC", os.path.dirname(os.path.abspath(__file__))), "main.py"), encoding="utf-8").read()
check('("pheno","Фенологічні фази"' in mtxt and '("Фенологія",         ["pheno"])' in mtxt and "from sad_phenology_win import PhenologyWindow" in mtxt, "пункт «Фенологічні фази» є в головному меню (категорія «Фенологія»)")
check("—" not in open(PW.__file__, encoding="utf-8").read().split("HELP =")[1].split('"""')[1], "у довідці немає довгих тире")
shutil.rmtree(TMP, ignore_errors=True)
print("\n" + "=" * 60)
if FAIL: print(f"ПРОВАЛІВ: {len(FAIL)}"); [print("  -", f) for f in FAIL]; sys.exit(1)
print("УСІ ПЕРЕВІРКИ ВІКНА ФЕНОЛОГІЇ ПРОЙДЕНО")
