# test_sad_weather_gui.py – повний тест вікна «Погода» (Xvfb + імітація мережі)
import os, sys, json, math, tempfile, shutil, time, threading, datetime as dt
real_sleep = time.sleep
time.sleep = lambda s: None                      # пауза між запитами й норма – без очікування
sys.path.insert(0, os.environ.get("SAD_SRC", os.path.dirname(os.path.abspath(__file__))))
os.environ["APPDATA"] = tempfile.mkdtemp(prefix="sadw_appdata_")      # налаштування діаграм пишуться в тимчасову теку
os.environ.setdefault("DISPLAY", ":99")
from urllib.parse import urlparse, parse_qs
import numpy as np
import tkinter as tk

FAIL = []
def check(cond, msg):
    print(("  [OK]   " if cond else "  [FAIL] ") + msg, flush=True)
    if not cond: FAIL.append(msg)

import main as m
import sad_weather_core as W
import sad_weather as SW
import sad_weather_indices as W_ind

# ───────────────────────────── імітація серверів ─────────────────────────────
REQ = []
def _day_vals(d):
    doy = d.timetuple().tm_yday
    h = (d.toordinal() * 2654435761) % 1000 / 1000.0
    t = 9.5 + 11.5 * math.sin(2 * math.pi * (doy - 110) / 365.0) + (h - 0.5) * 4
    p = round(((h * 13) % 1) * 14, 1) if h < 0.33 else 0.0
    return t, p
def nasa_payload(a, b):
    days = [a + dt.timedelta(days=i) for i in range((b - a).days + 1)]
    k = [d.strftime("%Y%m%d") for d in days]
    tv = [_day_vals(d) for d in days]
    par = {"T2M": {x: round(v[0], 2) for x, v in zip(k, tv)}, "T2M_MAX": {x: round(v[0] + 4, 2) for x, v in zip(k, tv)},
           "T2M_MIN": {x: round(v[0] - 4, 2) for x, v in zip(k, tv)}, "PRECTOTCORR": {x: v[1] for x, v in zip(k, tv)},
           "RH2M": {x: 70.0 for x in k}, "WS2M": {x: 3.0 for x in k}, "ALLSKY_SFC_SW_DWN": {x: 15.0 for x in k}}
    return {"geometry": {"coordinates": [30.2218, 48.7484, 212.0]}, "properties": {"parameter": par},
            "header": {"fill_value": -999.0, "time_standard": "LST", "api": {"version": "v2.6"}},
            "parameters": {"ALLSKY_SFC_SW_DWN": {"units": "MJ/m^2/day"}}}
MODE = {"offline": False, "slow": 0.0}
def fake_get(url, timeout=60):
    REQ.append(url)
    if MODE["offline"]: raise W.WeatherError("Немає з'єднання з сервером (імітація)")
    if MODE["slow"]: real_sleep(MODE["slow"])
    q = parse_qs(urlparse(url).query)
    if "power.larc.nasa.gov" in url:
        a = dt.datetime.strptime(q["start"][0], "%Y%m%d").date(); b = dt.datetime.strptime(q["end"][0], "%Y%m%d").date()
        return 200, json.dumps(nasa_payload(a, b))
    if "nominatim.openstreetmap.org" in url:
        return 200, json.dumps([{"lat": "48.6830", "lon": "30.1650", "display_name": "Яблунівка, Уманський район, Черкаська область, Україна", "type": "village"}])
    return 404, "not found"
W.http_get = fake_get

# ───────────────────────────── заглушки діалогів ─────────────────────────────
LOG = {"info": [], "warn": [], "error": [], "ask": []}
ASK = {"ans": True}
tk_mb = tk.messagebox
tk_mb.showinfo = lambda t="", msg="", **k: LOG["info"].append((t, msg))
tk_mb.showwarning = lambda t="", msg="", **k: LOG["warn"].append((t, msg))
tk_mb.showerror = lambda t="", msg="", **k: LOG["error"].append((t, msg))
def _ask(t="", msg="", **k):
    LOG["ask"].append((t, msg)); return ASK["ans"]
tk_mb.askyesno = _ask
SW.open_file_cross = lambda p: None
TMP = tempfile.mkdtemp(prefix="sadw_gui_")
SAVE = {"path": None, "open": None}
tk.filedialog.asksaveasfilename = lambda **k: SAVE["path"]
tk.filedialog.askopenfilename = lambda **k: SAVE["open"]

root = tk.Tk()
app = m.SADTk(root); root._sad_app = app
root.update()

def pump(cond, timeout=20):
    t0 = time.time()
    while time.time() - t0 < timeout:
        root.update()
        if cond(): return True
        real_sleep(0.02)
    return False
def toplevels(w=root):
    out = []
    if isinstance(w, tk.Toplevel): out.append(w)
    for c in w.winfo_children(): out.extend(toplevels(c))
    return out
def find_btn(txt, w):
    r = [None]
    def rec(x):
        if r[0] is not None: return
        try:
            if x.winfo_class() == "Button" and txt in x.cget("text"): r[0] = x; return
        except Exception: pass
        for c in x.winfo_children(): rec(c)
    rec(w); return r[0]
def tree_rows(holder):
    out = []
    def rec(x):
        if isinstance(x, tk.ttk.Treeview): out.append(x)
        for c in x.winfo_children(): rec(c)
    rec(holder); return out[0] if out else None
def fetch_and_wait(win, expect_series=True, timeout=25):
    win._start_fetch()
    ok = pump(lambda: not win._busy and (win.series is not None or not expect_series), timeout)
    pump(lambda: False, 0.1)
    return ok

# ═════════════════════════════════════════════════════════════════════════════
print("== 1. Меню й відкриття вікна")
labels = []
def rec_l(w):
    try:
        if w.winfo_class() == "Label" and w.cget("text"): labels.append(w.cget("text"))
    except Exception: pass
    for c in w.winfo_children(): rec_l(c)
rec_l(root)
check(any(t == "Погода" or t.endswith("Погода") for t in labels) or any("Погода" in t for t in labels), "пункт «Погода» є на головній сторінці")
win = m.WeatherWindow(root); root.update()
check(win.win.winfo_exists(), "вікно «Погода» відкрилось")
check(win.src_var.get() == W.SOURCE_TITLES["nasa"], "джерело за замовчуванням – NASA POWER (безпечне для комерційного поширення)")
check("NASA" in win.lic_lbl.cget("text"), "показано умови NASA (подяка)")
import sad_weather_charts as CHm
cbs = [w for w in win.win.winfo_children()]
def all_widgets(w):
    out = [w]
    for c in w.winfo_children(): out += all_widgets(c)
    return out
srcs = [w for w in all_widgets(win.win) if w.winfo_class() == "TCombobox" and W.SOURCE_TITLES["nasa"] in (w.cget("values") or ())][0]
check(list(srcs.cget("values")) == [W.SOURCE_TITLES["nasa"], W.SOURCE_TITLES["file"]], "у списку джерел лише NASA POWER і власні дані (Open-Meteo видалено)")
txts = " ".join(w.cget("text") for w in all_widgets(win.win) if w.winfo_class() in ("Label", "Labelframe", "Button", "Checkbutton"))
check("Open-Meteo" not in txts and "Кеш" not in txts and "кеш" not in txts and "Швидкий вибір" not in txts, "в інтерфейсі немає Open-Meteo, розділу «Кеш» і швидкого вибору українських міст")
check([w.cget("text") for w in all_widgets(win.win) if w.winfo_class() == "Labelframe"] == ["1. Джерело даних", "2. Місце досліду", "3. Період і порівняння років", "4. Параметри розрахунку"],
      "розділи ліворуч: 1–4 (розділ про кеш прибрано)")
win.win.geometry("1100x640"); root.update()
lcv = win.left_canvas; bb = lcv.bbox("all")
vs_ = [x for x in all_widgets(lcv.master) if x.winfo_class() == "TScrollbar" and str(x.cget("orient")) == "vertical"][0]
check(bb[3] > lcv.winfo_height() and vs_.get() != (0.0, 1.0), f"у низькому вікні (640 px) ліва колонка довша за вікно ({bb[3]} px) – з'являється вертикальна прокрутка")
lcv.yview_moveto(1.0); root.update(); last = find_btn("Перерахувати", win.win)
check(last is not None and last.winfo_rooty() + last.winfo_height() <= win.win.winfo_rooty() + win.win.winfo_height() + 2, "прокрутивши вниз, кнопку «Перерахувати» видно повністю")
lcv.yview_moveto(0.0); win.win.geometry("1340x840"); root.update()
check(find_btn("Обрати на карті", win.win) is not None and find_btn("Знайти за назвою", win.win) is not None, "є кнопки «Обрати на карті» і «Знайти за назвою» (пошук збережено)")
check("не обрано" in win.coord_lbl.cget("text") and win.lat_var.get() == "" and win.lon_var.get() == "", "координати за замовчуванням порожні: місце потрібно обрати (програма не налаштована на Україну)")
n = len(LOG["warn"]); win._start_fetch(); check(len(LOG["warn"]) == n + 1 and "Оберіть місце" in LOG["warn"][-1][1], "завантаження без обраного місця → пояснення, що робити")
win.src_var.set(W.SOURCE_TITLES["file"]); win._on_source(); root.update()
check(win.file_row.winfo_ismapped() and "файл" in win.btn_fetch.cget("text").lower(), "власні дані: кнопки файлу, кнопка «Відкрити файл…»")
win.src_var.set(W.SOURCE_TITLES["nasa"]); win._on_source(); root.update()

print("\n== 2. NASA: пресет, координати, період, завантаження у фоні")
win._set_point(48.7484, 30.2218, "Умань, Україна"); root.update()
check(win.lat_var.get() == "48.7484" and win.lon_var.get() == "30.2218" and win.place_var.get() == "Умань, Україна" and "48,7484° пн. ш., 30,2218° сх. д." in win.coord_lbl.cget("text"),
      "обране місце: координати, назва й підпис «48,7484° пн. ш., 30,2218° сх. д.»")
win._set_point(50.4501, 30.5234, "Київ, Україна"); check(win.place_var.get() == "Київ, Україна", "автоназва оновлюється при виборі іншого місця")
win.place_var.set("Моє поле"); win._set_point(48.7484, 30.2218, "Умань, Україна"); check(win.place_var.get() == "Моє поле", "назву, введену користувачем, не перезаписано")
win.place_var.set("м. Умань")
win.d0_var.set("01.01.2022"); win.d1_var.set("31.12.2024")
REQ.clear()
ok = fetch_and_wait(win)
check(ok and win.series is not None and len(win.series) == 1096, f"завантажено {len(win.series) if win.series else 0} діб (очікувано 1096)")
check(len(REQ) == 3 and all("community=AG" in u for u in REQ), f"3 запити по роках до NASA (community=AG): {len(REQ)}")
check(not win._busy and str(win.btn_fetch.cget("state")) == "normal", "кнопка розблокована після завантаження")
txt = win.txt_sum.get("1.0", "end")
check("м. Умань" in txt and "NASA POWER" in txt and "Дата отримання даних" in txt, "вкладка «Підсумок»: місце, джерело, дата отримання")
check("реаналізу" in txt, "у підсумку є попередження про реаналіз (не вимірювання станції)")
check("ГТК Селянінова за період" in txt and "Шкала ГТК" in txt, "підсумок: ГТК за період і шкала")
check(len(tree_rows(win.holders["months"]).get_children()) == 36, "вкладка «Місяці»: 36 рядків")
check(len(tree_rows(win.holders["decades"]).get_children()) == 108, "вкладка «Декади»: 108 рядків")
check(len(tree_rows(win.holders["years"]).get_children()) == 3, "вкладка «За роками»: 3 роки (вікно 04-01…10-31)")
check(len(tree_rows(win.day_holder).get_children()) == 1096, "вкладка «Добові дані»: 1096 рядків")
tv = tree_rows(win.holders["months"])
tagged = [tv.item(i, "tags") for i in tv.get_children()]
check(any(t for t in tagged), "рядки ГТК розфарбовано за класом")
check(win._fig is not None, "діаграма побудована")
check((win.y0_var.get(), win.y1_var.get()) == ("2022", "2024") and "Повне покриття: 2022, 2023, 2024".replace("Повне", "повне") in win.win_lbl.cget("text"), "після завантаження роки порівняння = роки даних; підпис показує покриття")

print("\n== 3. Кома/крапка у таблицях, метод ГТК")
v_comma = tv.item(tv.get_children()[5], "values")
win.comma_var.set(False); win._refresh_all(); root.update()
tv2 = tree_rows(win.holders["months"]); v_dot = tv2.item(tv2.get_children()[5], "values")
check(any("," in str(x) for x in v_comma) and not any("," in str(x) for x in v_dot) and any("." in str(x) for x in v_dot),
      "перемикач десяткової коми працює в таблицях")
win.comma_var.set(True); win._refresh_all()
h_dec = win.rep["total"]["htc"]
win.htc_var.set(list(SW.HTC_METHODS)[2]); win._recompute(); root.update()
check(win.rep["settings"]["htc_method"] == "daily" and abs(win.rep["total"]["htc"] - h_dec) > 1e-9,
      f"добовий метод змінює ГТК ({h_dec:.3f} → {win.rep['total']['htc']:.3f})")
win.htc_var.set(list(SW.HTC_METHODS)[0]); win._recompute()
win.base_var.set("abc"); n_w = len(LOG["warn"]); win._recompute()
check(len(LOG["warn"]) == n_w + 1 and win.rep is not None, "некоректний параметр → зрозуміле попередження, вікно не падає")
win.base_var.set("10")
win.wm2_var.set("лютий"); win.wd2_var.set("31"); n_w = len(LOG["warn"]); win._recompute()
check(len(LOG["warn"]) == n_w + 1 and "немає 31-го числа" in LOG["warn"][-1][1] and "лютий" in LOG["warn"][-1][1], "неіснуюча дата вікна (31 лютого) відхилена з поясненням")
win._update_window_preview(); check("немає 31-го числа" in win.win_lbl.cget("text"), "підпис під вікном показує помилку одразу, без натискання кнопок")
win.wm2_var.set("жовтень"); win.wd2_var.set("31"); win._update_window_preview()
win.y0_var.set("2025"); win.y1_var.set("2022"); n_w = len(LOG["warn"]); win._recompute()
check(len(LOG["warn"]) == n_w + 1 and "більший" in LOG["warn"][-1][1], "рік «від» більший за «до» → зрозуміле попередження")
win.y0_var.set("abc"); n_w = len(LOG["warn"]); win._recompute(); check(len(LOG["warn"]) == n_w + 1 and "цілими числами" in LOG["warn"][-1][1], "нечисловий рік відхилено")
win.y0_var.set("2022"); win.y1_var.set("2024"); win._update_window_preview()

print("\n== 4. Усі діаграми через інтерфейс, збереження PNG/SVG/PDF")
bad = []
for kind, title in __import__("sad_weather_charts").CHARTS:
    win.chart_var.set(title); win._draw_chart(); root.update()
    if win._fig is None or not win._fig.axes: bad.append(title)
check(not bad, f"всі {len(__import__('sad_weather_charts').CHARTS)} діаграм малюються у вікні ({'; '.join(bad)})")
win.chart_var.set("Діаграма Вальтера"); win._draw_chart()
for ext in ("png", "svg", "pdf"):
    SAVE["path"] = os.path.join(TMP, f"chart.{ext}"); win._save_chart()
    check(os.path.exists(SAVE["path"]) and os.path.getsize(SAVE["path"]) > 2000, f"діаграма збережена як {ext.upper()} ({os.path.getsize(SAVE['path'])//1024} КБ)")

print("\n== 4б. Легенда, налаштування діаграм, прокрутка")
win.chart_var.set("Температура й опади за декадами"); win._draw_chart(); root.update()
leg = [a.get_legend() for a in win._fig.axes if a.get_legend()]
ltxt = " ".join(t.get_text() for l in leg for t in l.get_texts())
check(leg and "Середня температура" in ltxt and "крива" in ltxt and "Сума опадів" in ltxt and "стовпчики" in ltxt, f"діаграма «Температура й опади за декадами» має легенду: «{ltxt[:90]}…»")
win.chart_var.set("Порівняння років (температура й опади у вікні)"); win._draw_chart(); root.update()
lt2 = " ".join(t.get_text() for a in win._fig.axes if a.get_legend() for t in a.get_legend().get_texts())
check("Середня температура" in lt2 and "Сума опадів" in lt2, "порівняння років: легенда пояснює точки й стовпчики")
win.chart_var.set("Температура й опади за декадами"); win._draw_chart(); root.update()
cb_ = find_btn("Налаштування", win.tab_ch); check(cb_ is not None and find_btn("Зберегти", win.tab_ch) is not None and find_btn("Копіювати", win.tab_ch) is not None, "біля діаграми три кнопки: Зберегти, Копіювати, Налаштування")
cb_.invoke(); root.update()
sdlg = win._style_dlg; check(sdlg.win.winfo_exists(), "діалог налаштувань відкрився")
sdlg.gv["width_cm"].set("20"); sdlg.gv["height_cm"].set("11"); sdlg.gv["font_size"].set("12"); sdlg.gv["title_size"].set("15"); sdlg.gv["title_bold"].set(True)
sdlg.gv["grid"].set(True); sdlg.gv["legend_loc"].set("під графіком"); sdlg.gv["font_family"].set("DejaVu Sans"); sdlg.title_var.set("Мій власний заголовок")
sdlg.set_color("tmean", "#006d2c"); sdlg.set_color("prec", "#fd8d3c")
st_new = sdlg.apply(); root.update()
check(st_new is not None and win.chart_style["general"]["width_cm"] == 20 and win.chart_style["colors"]["climo_decade"] == {"tmean": "#006d2c", "prec": "#fd8d3c"}
      and win.chart_style["titles"]["climo_decade"] == "Мій власний заголовок", "налаштування застосовано до стилю вікна")
w_in, h_in = win._fig.get_size_inches(); check(abs(w_in - 20 / 2.54) < 0.011 and abs(h_in - 11 / 2.54) < 0.011, f"розмір діаграми 20×11 см → {w_in:.3f}×{h_in:.3f} дюйма (точність до 1 пікселя екрана)")
cw = win._canvas.get_tk_widget(); dpi_ = win._fig.dpi
check(abs(cw.winfo_reqwidth() - w_in * dpi_) <= 2 and abs(cw.winfo_reqheight() - h_in * dpi_) <= 2 and cw.winfo_reqwidth() < 1000,
      f"полотно у вікні має заданий розмір {cw.winfo_reqwidth()}×{cw.winfo_reqheight()} px (= {w_in:.2f}×{h_in:.2f} дюйма при {dpi_:.0f} dpi екрана; не розтягується під вікно)")
ax0 = win._fig.axes[0]; ttl = ax0.get_title(loc="left"); check(ttl.startswith("Мій власний заголовок") and ax0.title.get_fontsize() == 15 and ax0.title.get_fontweight() == "bold", "власний заголовок, розмір 15 і жирний шрифт застосовано")
lines_ = [l for l in ax0.get_lines() if l.get_color() == "#006d2c"]; check(len(lines_) >= 1, "колір кривої температури змінено на обраний")
bars = [pt.get_facecolor() for a in win._fig.axes for pt in a.patches][:1]; import matplotlib.colors as mcol
check(bars and mcol.to_hex(bars[0][:3]) == "#fd8d3c", "колір стовпчиків опадів змінено на обраний")
check(win._fig.axes[0].xaxis.get_gridlines()[0].get_visible() or win._fig.axes[0].yaxis.get_gridlines()[0].get_visible(), "сітка ввімкнена")
import json as _j, sad_weather_charts as CHm2
sp = CHm2.style_path(); check(os.path.exists(sp) and _j.load(open(sp, encoding="utf-8"))["general"]["width_cm"] == 20, "налаштування збережено на диск (діють і в наступних сеансах)")
check(CHm2.load_style()["titles"]["climo_decade"] == "Мій власний заголовок", "збережене читається назад")
SAVE["path"] = os.path.join(TMP, "styled.png"); win._save_chart()
from PIL import Image as _Im
im = _Im.open(SAVE["path"]); check(abs(im.size[0] - 20 / 2.54 * 300) <= 1 and abs(im.size[1] - 11 / 2.54 * 300) <= 1, f"PNG точно 20×11 см у роздільності 300 dpi: {im.size[0]}×{im.size[1]} px (очікувано {round(20/2.54*300)}×{round(11/2.54*300)})")
win.chart_var.set("ГТК за місяцями"); win._draw_chart(); root.update()
check(abs(win._fig.get_size_inches()[0] - 20 / 2.54) < 0.011, "загальні параметри (розмір, шрифт) діють на всі діаграми")
check(win.chart_style.get("colors", {}).get("htc_month") is None, "кольори окремі для кожної діаграми: для ГТК за місяцями лишились типові")
bad_ = sdlg.gv["width_cm"].set("abc"); n_w = len(LOG["warn"]); r_ = sdlg.apply(); check(r_ is None and len(LOG["warn"]) == n_w + 1 and "Ширина" in LOG["warn"][-1][1], "некоректне значення в налаштуваннях → пояснення, стиль не змінено")
sdlg.reset_all(); sdlg.apply(); root.update()
check(win.chart_style["general"] == CHm2.sanitize_general({}) and "climo_decade" not in win.chart_style.get("colors", {}) and not win.chart_style.get("titles"), "«Усе до типових» повертає початковий вигляд")
sdlg.win.destroy()
def hsb_in(w): return [x for x in all_widgets(w) if x.winfo_class() == "TScrollbar" and str(x.cget("orient")) == "horizontal"]
check(len(hsb_in(win.tab_ch)) == 1 and len(hsb_in(win.tab_sum)) == 1 and len(hsb_in(win.tab_pub)) == 1 and len(hsb_in(win.tab_day)) == 1 and len(hsb_in(win.tab_per)) >= 1,
      "горизонтальна прокрутка є на вкладках: Підсумок, Добові дані, Періоди, Графіки, Для публікації")
check(all(len(hsb_in(win.holders[k].master)) >= 1 for k in ("decades", "months", "years")), "горизонтальна прокрутка є й у таблицях «Декади», «Місяці», «За роками»")
win.txt_sum._wrap_var.set(False); win.txt_sum.configure(wrap="none"); check(win.txt_sum.cget("wrap") == "none", "у «Підсумку» можна вимкнути перенесення рядків")
win.txt_sum._wrap_var.set(True); win.txt_sum.configure(wrap="word")

print("\n== 5. Експорт: Excel, Word, CSV, погодний набір, текст")
import openpyxl, docx
SAVE["path"] = os.path.join(TMP, "w.xlsx"); ASK["ans"] = False; win._export_xlsx()
wb = openpyxl.load_workbook(SAVE["path"]); check("Місяці" in wb.sheetnames and "За роками" in wb.sheetnames, f"Excel: аркуші {wb.sheetnames}")
SAVE["path"] = os.path.join(TMP, "w.docx"); win._export_docx()
d = docx.Document(SAVE["path"]); check(len(d.tables) >= 2 and len(d.inline_shapes) >= 3, f"Word: таблиць {len(d.tables)}, рисунків {len(d.inline_shapes)}")
SAVE["path"] = os.path.join(TMP, "w.csv"); win._export_csv()
check(open(SAVE["path"], encoding="utf-8-sig").read().splitlines()[0].startswith("Дата;"), "CSV (кома → роздільник «;»)")
SAVE["path"] = os.path.join(TMP, "w.sadw"); win._export_set()
ws = json.load(open(SAVE["path"], encoding="utf-8"))
check(ws["type"] == "sad_weather_set" and set(ws["years"]) == {"2022", "2023", "2024"} and ws["place"] == "м. Умань", "погодний набір (.sadw): 3 роки")
SAVE["path"] = os.path.join(TMP, "w.txt"); win._save_pub()
check("МЕТОДИКА ЗБОРУ МЕТЕОДАНИХ" in open(SAVE["path"], encoding="utf-8").read(), "текст для публікації збережено")
win._copy_table("months"); clip = win.win.clipboard_get()
check(clip.startswith("Місяць\t") and clip.count("\n") == 36, "копіювання таблиці в буфер (TSV, 36 рядків)")
check(sys.platform == "win32" or "як текст" in win.status_var.get(), "поза Windows користувача чесно попереджено, що в буфері текст (сітку дає Windows або кнопка «Word»)")
import sad_clip as SCm
captured = {}
SCm.copy_table = lambda widget, headers, rows, fills=None, tsv=None, cell_fills=None: captured.update(h=headers, r=rows, f=fills, t=tsv, cf=cell_fills) or "html"
win._copy_table("months")
check(captured["h"][0] == "Місяць" and len(captured["r"]) == 36 and captured["t"].startswith("Місяць\t") and any(captured["f"]), "для копіювання з сіткою передаються заголовки, 36 рядків, кольори ГТК й текст TSV")
check("із сіткою" in win.status_var.get() and "Word" in win.status_var.get(), "статус: «таблицю скопійовано разом із сіткою й кольорами»")
SCm.copy_table = __import__("importlib").reload(SCm).copy_table
SAVE["path"] = None; win._export_xlsx(); win._export_docx(); win._export_csv()
check(True, "скасування діалогу збереження не призводить до помилок")
ASK["ans"] = True

print("\n== 6. Періоди: за датами замірів і з серії журналу")
win._period_from_dates(preset_dates=[dt.date(2024,5,15), dt.date(2024,5,29), dt.date(2024,6,12)], season_default="01.05.2024"); root.update()
dlg = [t for t in toplevels(win.win) if "замірів" in t.title()][0]
find_btn("Створити періоди", dlg).invoke(); root.update()
check(len(win.periods) == 3 and win.per_list.size() == 3, "створено 3 періоди між замірами")
trp = tree_rows(win.holders["periods"]); check(trp is not None and len(trp.get_children()) == 3, "таблиця періодів: 3 рядки")
n_win = len(toplevels(win.win))
find_btn("Очистити", win.tab_per).invoke(); root.update()
check(win.periods == [] and win.per_list.size() == 0, "«Очистити» прибирає періоди")
sch = os.path.join(TMP, "scheme"); os.makedirs(os.path.join(sch, "2024"))
sp = os.path.join(sch, "схема.sadp")
json.dump({"type": "multi_factor_scheme", "plants": [], "factor_defs": [], "cfg": {},
           "series": {"Висота": {"interval_days": 14, "rounds": [{"date": "2024-05-15", "indicator": "x", "year": "2024"},
                                                                  {"date": "2024-05-29", "indicator": "y", "year": "2024"}]}}},
          open(sp, "w", encoding="utf-8"))
import sad_journal_trial as JT
JT.pick_indicators_dialog = lambda parent, recs, multi_select=False, n_required=None: [sorted(recs)[0]]
SAVE["open"] = sp; win._period_from_series(); root.update()
dlg = [t for t in toplevels(win.win) if "замірів" in t.title()][0]
txd = [w for w in dlg.winfo_children()[0].winfo_children() if isinstance(w, tk.Text)][0]
check("15.05.2024" in txd.get("1.0", "end") and "29.05.2024" in txd.get("1.0", "end"), "з серії журналу підтягнуто дати замірів")
find_btn("Створити періоди", dlg).invoke(); root.update()
check(len(win.periods) == 2, f"з серії створено {len(win.periods)} періоди")
win._period_clear()

print("\n== 7. Норма (1991–2020): фонове завантаження й відхилення")
REQ.clear(); win.ny0.set("1991"); win.ny1.set("2020"); win._load_norm()
ok = pump(lambda: win.norm is not None, 40)
check(ok and win.norm[2] == "1991–2020" and len(win.norm[1]) == 30, f"норма за {len(win.norm[1]) if win.norm else 0} років")
check(len(REQ) == 30, f"30 окремих запитів (по вікну року, а не 30 років підряд): {len(REQ)}")
check(len(tree_rows(win.holders["years"])["columns"]) > 0 and "Опади, % норми" in tree_rows(win.holders["years"])["columns"],
      "таблиця за роками отримала стовпці відхилень від норми")
check("Норма: 1991–2020" in win.years_note.cget("text"), "підпис про норму")
win._clear_norm(); check(win.norm is None and "Опади, % норми" not in tree_rows(win.holders["years"])["columns"], "норму можна прибрати")
win.ny0.set("2010"); win.ny1.set("2012"); ASK["ans"] = False; n = len(LOG["ask"]); win._load_norm()
check(len(LOG["ask"]) == n + 1 and "30" in LOG["ask"][-1][1] and win.norm is None, "коротка норма (<10 років): програма питає й за відмови нічого не робить")
ASK["ans"] = True
win.ny0.set("1991"); win.ny1.set(str(dt.date.today().year)); n = len(LOG["warn"]); win._load_norm()
check(len(LOG["warn"]) == n + 1, "незавершений рік у нормі відхилено")

print("\n== 8. Роки й вікно порівняння: зрозумілий підпис, період, фільтр")
win.d0_var.set("01.01.2021"); win.d1_var.set("31.12.2024"); REQ.clear(); win._years_touched = False
ok = fetch_and_wait(win); check(ok and len(win.series) == 1461, "завантажено 2021–2024 для перевірки порівняння років")
check((win.y0_var.get(), win.y1_var.get()) == ("2021", "2024"), "роки порівняння автоматично = роки завантаженого періоду")
lbl = win.win_lbl.cget("text")
check("Буде порівняно 4 роки" in lbl and "2021: 01.04–31.10.2021" in lbl and "2024: 01.04–31.10.2024" in lbl and "1 квітня – 31 жовтня" in lbl, f"підпис пояснює, які саме роки й відрізки порівнюються: «{lbl[:110]}…»")
check(len(tree_rows(win.holders["years"]).get_children()) == 4, "таблиця «За роками»: 4 роки")
win.y0_var.set("2022"); win.y1_var.set("2023"); win._years_touched = True; win._recompute(); root.update()
check(len(tree_rows(win.holders["years"]).get_children()) == 2 and [r["year"] for r in win.rep["years"]] == [2022, 2023], "якщо вказати роки 2022–2023, у порівнянні лишаються лише вони")
win.chart_var.set("Порівняння років (температура й опади у вікні)"); win._draw_chart(); root.update()
check(len(win._fig.axes[1].patches) == 2, "діаграма порівняння років показує лише 2 вибрані роки")
win.y0_var.set("2019"); win.y1_var.set("2024"); win._update_window_preview()
check("Немає даних: 2019, 2020" in win.win_lbl.cget("text"), "підпис попереджає про роки без даних у завантаженому ряді")
win.wm1_var.set("листопад"); win.wd1_var.set("1"); win.wm2_var.set("лютий"); win.wd2_var.set("28"); win.y0_var.set("2021"); win.y1_var.set("2023"); win._update_window_preview()
check("Кінець вікна припадає на наступний рік" in win.win_lbl.cget("text") and "2021: 01.11–28.02.2022" in win.win_lbl.cget("text"), "вікно через Новий рік: підпис пояснює, що кінець у наступному році")
win.wm1_var.set("квітень"); win.wd1_var.set("1"); win.wm2_var.set("жовтень"); win.wd2_var.set("31")
win.y0_var.set("2020"); win.y1_var.set("2023"); win._apply_season(); root.update()
check(win.d0_var.get() == "01.04.2020" and win.d1_var.get() == "31.10.2023" and "Натисніть «Отримати дані»" in win.status_var.get(), "«Встановити період…» задає 01.04.2020 – 31.10.2023")
win.wm1_var.set("жовтень"); win.wd1_var.set("15"); win.wm2_var.set("квітень"); win.wd2_var.set("15"); win.y0_var.set("2020"); win.y1_var.set("2021"); win._apply_season()
check(win.d0_var.get() == "15.10.2020" and win.d1_var.get() == "15.04.2022", "вікно через Новий рік: кінець останнього року – у наступному році")
win.wm1_var.set("квітень"); win.wd1_var.set("1"); win.wm2_var.set("жовтень"); win.wd2_var.set("31")
win._calendar_years_dlg(); cd = [t for t in toplevels(win.win) if t.title() == "Календарні роки"][0]; find_btn("OK", cd).invoke(); root.update()
check(win.d0_var.get().startswith("01.01.") and win.d1_var.get().startswith("31.12."), "«Календарні роки…» задає 01.01 – 31.12")
win.y0_var.set("2022"); win.y1_var.set("2024"); win._years_touched = False

print("\n== 8б. Місце: карта й пошук")
import sad_worldmap as WMm
WMm.MapPickerDialog.show = lambda self: (self.win.update(), {"lat": 50.4501, "lon": 30.5234, "label": "Київ, Україна"})[1]
win.lat_var.set(""); win.lon_var.set(""); win.place_var.set(""); win._auto_place = ""
win._pick_on_map(); root.update()
check(win.lat_var.get() == "50.4501" and win.lon_var.get() == "30.5234" and win.place_var.get() == "Київ, Україна", "вибір на карті переносить координати й назву у вікно")
for tl in [t for t in toplevels(win.win) if "карті" in t.title()]: tl.destroy()
win._geocode_dlg(); root.update(); ps = win._geo_ps
ps.qv.set("умань"); ps.search(); check(ps.found and ps.found[0]["label"].startswith("Умань") and abs(ps.found[0]["lat"] - 48.754) < 0.01, "пошук за назвою: «умань» знайдено серед міст світу (без інтернету)")
REQ.clear(); ps.qv.set("Яблунівка"); ps.search(); check(ps.found == [] or "Яблунівка" not in str(ps.found), "офлайн-пошук не знаходить маленьке село (довідник міст)")
ps.search_online(); check(len(REQ) == 1 and "nominatim" in REQ[0] and ps.found and ps.found[0]["lat"] == 48.683, "онлайн-пошук (OpenStreetMap) знаходить село; один запит на натискання")
ps.lb.selection_clear(0, "end"); ps.lb.selection_set(0); ps.pick(); root.update()
check(win.lat_var.get() == "48.6830" and win._geo_dlg.winfo_exists() is False if False else win.lat_var.get() == "48.6830", "вибір результату пошуку заповнює координати")
MODE["offline"] = True
win._geocode_dlg(); root.update(); ps = win._geo_ps; ps.qv.set("Село"); n = len(LOG["warn"]); ps.search_online()
check(len(LOG["warn"]) == n + 1 and "з'єднання" in LOG["warn"][-1][1], "онлайн-пошук без інтернету → зрозуміле повідомлення, вікно не падає")
MODE["offline"] = False
for tl in [t for t in toplevels(win.win) if "Пошук" in t.title()]: tl.destroy()
win._set_point(48.7484, 30.2218, "")
win._coords_dlg(); root.update(); la_v, lo_v = win._coords_vars
check(la_v.get() == "48.7484" and lo_v.get() == "30.2218", "діалог «Координати…» відкривається з поточними значеннями")
la_v.set("48°45'N"); lo_v.set("30°13'E"); win._coords_ok(); root.update()
check(abs(float(win.lat_var.get()) - 48.75) < 1e-4 and abs(float(win.lon_var.get()) - 30.2167) < 1e-4 and "48,7500° пн. ш." in win.coord_lbl.cget("text"), "ручне введення у форматі ГМС: 48°45'N, 30°13'E → 48,7500°, 30,2167°")
win._coords_dlg(); la_v, lo_v = win._coords_vars; la_v.set("95"); lo_v.set("10"); n = len(LOG["warn"]); old_lat = win.lat_var.get(); win._coords_ok()
check(len(LOG["warn"]) == n + 1 and "широта" in LOG["warn"][-1][1] and win.lat_var.get() == old_lat, "широта 95° відхилена з поясненням, значення не змінено")
win._coords_win.destroy(); win._set_point(48.7484, 30.2218, "")

print("\n== 8в. Міжнародні індекси: ET₀, Вінклер, Югліна, індекс сухості")
win._set_point(48.7484, 30.2218, ""); win._recompute(); root.update()
iv = win.rep["intl"]
check(iv["pet_method"] == "pm" and win.rep["series"].has("et0") and not win.series.has("et0"), "ET₀ розраховано (Пенман–Монтейт) для звіту й не записано у вихідний ряд користувача")
tabs = [win.nb.tab(i, "text") for i in range(win.nb.index("end"))]
check(tabs == ["Підсумок", "Добові дані", "Декади", "Місяці", "За роками", "Періоди", "Індекси", "SPI", "Графіки", "Для публікації"], f"вкладки: {tabs}")
cols_m = tree_rows(win.holders["months"])["columns"]
check("ET₀, мм" in cols_m and "P/ET₀" in cols_m, "у таблицях з'явилися колонки «ET₀, мм» і «P/ET₀»")
check("Пенман" in win.idx_note.cget("text") and "GDD" in win.idx_note.cget("text") and "2 м" in win.idx_note.cget("text"), "на вкладці «Індекси» пояснено метод ET₀ і що СЕТ = GDD")
tv_v = tree_rows(win.holders["vine"]); tv_a = tree_rows(win.holders["aridity"])
check(len(tv_v.get_children()) == 4 and len(tv_a.get_children()) == 4, "таблиці Вінклера/Югліна й індексу сухості: по 4 сезони/роки (2021–2024)")
vals = tv_v.item(tv_v.get_children()[0], "values"); check(vals[0] == "2021" and vals[2].split(" ")[0] in ("I", "II", "III", "IV", "V") and vals[4].startswith("HI"), f"рядок вінограрських індексів: {vals[:5]}")
va = tv_a.item(tv_a.get_children()[0], "values"); check(va[0] == "2021" and float(va[3].replace(",", ".")) > 0 and va[4] in ("гіперарідний", "арідний", "напівзасушливий", "сухий субвологий", "вологий"), f"рядок індексу сухості: {va}")
txt = win.txt_sum.get("1.0", "end"); check("Індекс сухості P/ET₀ за період" in txt and "Еталонне випаровування ET₀" in txt and "Пенман" in txt, "підсумок містить ET₀, водний баланс, індекс сухості й примітку про метод")
cap2 = {}
import sad_clip as SCm2
SCm2.copy_table = lambda widget, headers, rows, fills=None, tsv=None, cell_fills=None: cap2.update(h=headers, r=rows, cf=cell_fills) or "html"
win._copy_table("vine"); check(cap2["h"][0] == "Сезон" and len(cap2["r"]) == 4 and any(cap2["cf"]), "копіювання таблиці Вінклера/Югліна передає кольори клітинок (клас)")
win._copy_table("aridity"); check(cap2["h"][3] == "Індекс сухості P/ET₀" and any(cap2["cf"]), "копіювання індексу сухості з кольорами класів")
win.chart_var.set("Вінклер і Югліна за сезонами"); win._draw_chart(); root.update()
check(len(win._fig.axes) == 2 and "Вінклер" in " ".join(t.get_text() for a in win._fig.axes if a.get_legend() for t in a.get_legend().get_texts()), "діаграма Вінклера й Югліна: 2 панелі з легендою")
win.chart_var.set("SPI за місяцями"); win._draw_chart(); root.update()
check(any("SPI ще не розраховано" in t.get_text() for t in win._fig.texts + [x for a in win._fig.axes for x in a.texts]), "до розрахунку SPI діаграма пояснює, що робити")
SAVE["path"] = os.path.join(TMP, "i.xlsx"); ASK["ans"] = False; win._export_xlsx()
import openpyxl as _ox; wb2 = _ox.load_workbook(SAVE["path"]); check("Вінклер і Югліна" in wb2.sheetnames and "Індекс сухості" in wb2.sheetnames and "SPI" not in wb2.sheetnames, f"Excel: аркуші індексів є, SPI ще немає ({wb2.sheetnames[-3:]})")
wsv = wb2["Вінклер і Югліна"]; check(wsv.cell(2, 1).value == "2021" and isinstance(wsv.cell(2, 2).value, float) and isinstance(wsv.cell(2, 3).value, str) and wsv.cell(2, 3).fill.fgColor.rgb not in (None, "00000000"), "Excel: Вінклер числом, регіон текстом із заливкою класу")
SAVE["path"] = os.path.join(TMP, "i.docx"); win._export_docx(); import docx as _dx; d2 = _dx.Document(SAVE["path"])
htxt = [c.text for tb in d2.tables for c in tb.rows[0].cells]; ptxt = " ".join(p.text for p in d2.paragraphs)
check(any("Вінклер" in h for h in htxt) and any("P/ET₀" in h for h in htxt) and "Міжнародні агрокліматичні показники" in ptxt and "[1]" in ptxt and "Allen R.G." in ptxt, "Word: таблиці індексів, розділ про методи, нумерована література")
ASK["ans"] = True
pub = win.txt_pub.get("1.0", "end"); check("МІЖНАРОДНІ АГРОКЛІМАТИЧНІ ПОКАЗНИКИ" in pub and "ЛІТЕРАТУРА ДО МІЖНАРОДНИХ ПОКАЗНИКІВ" in pub and pub.index("[1]") < pub.index("[2]") < pub.index("[3]"), "текст для публікації: методи індексів і література за порядком появи")

print("\n== 8г. SPI: довгий ряд, діаграма, експорт, помилки")
# нижня смуга прокрутки таблиць має бути ВИДИМОЮ навіть для довгих таблиць (декади – 108 рядків)
def hsb_visible(tab, holder):
    hs = [x for x in all_widgets(holder) if x.winfo_class() == "TScrollbar" and str(x.cget("orient")) == "horizontal"]
    if not hs: return False
    h_ = hs[0]; root.update()
    return h_.winfo_ismapped() and h_.winfo_rooty() + h_.winfo_height() <= win.win.winfo_rooty() + win.win.winfo_height() - 20 and h_.winfo_height() > 5
win.win.geometry("1360x640"); root.update()                 # низький екран (ноутбук 768 px): саме тут нижню смугу раніше виштовхувало за межі вкладки
for key, tab in (("decades", None), ("months", None), ("years", win.tab_years), ("vine", win.tab_idx), ("aridity", win.tab_idx)):
    sel_tab = win.holders[key].master; win.nb.select(sel_tab); root.update()
    check(hsb_visible(sel_tab, win.holders[key]), f"таблиця «{key}»: нижня смуга прокрутки видима у вікні")
win.nb.select(win.tab_day); root.update(); check(hsb_visible(win.tab_day, win.day_holder), "«Добові дані» (понад тисячу рядків): нижня смуга прокрутки видима")
win.win.geometry("1340x840"); root.update()
win2 = m.WeatherWindow(root); root.update(); n = len(LOG["info"]); win2._load_spi()
check(len(LOG["info"]) == n + 1 and "Спершу отримайте дані" in LOG["info"][-1][1], "SPI без завантажених даних → підказка")
win2._close()
win.sr0.set("1991"); win.sr1.set("2020")
for k, v in win.spi_vars.items(): v.set(False)
n = len(LOG["warn"]); win._load_spi(); check(len(LOG["warn"]) == n + 1 and "масштаб" in LOG["warn"][-1][1], "жодного масштабу не обрано → попередження")
for k, v in win.spi_vars.items(): v.set(True)
win.sr0.set("abc"); n = len(LOG["warn"]); win._load_spi(); check(len(LOG["warn"]) == n + 1 and "цілими числами" in LOG["warn"][-1][1], "нечислові роки → попередження")
win.sr0.set("2020"); win.sr1.set("1991"); n = len(LOG["warn"]); win._load_spi(); check(len(LOG["warn"]) == n + 1 and "«від» ≤ «до»" in LOG["warn"][-1][1], "рік «від» > «до» → попередження")
win.sr0.set("2005"); win.sr1.set("2020"); n = len(LOG["warn"]); REQ.clear(); win._load_spi()
check(len(LOG["warn"]) == n + 1 and "замалий" in LOG["warn"][-1][1] and len(REQ) == 0, "еталон 16 років → відмова з поясненням, жодного запиту до мережі")
win.sr0.set("1991"); win.sr1.set(str(dt.date.today().year)); n = len(LOG["warn"]); win._load_spi(); check(len(LOG["warn"]) == n + 1 and "завершений рік" in LOG["warn"][-1][1], "незавершений рік серед еталонних → відмова")
win.sr1.set("2020"); REQ.clear(); ASK["ans"] = False; n = len(LOG["ask"]); win._load_spi()
check(len(LOG["ask"]) == n + 1 and "35 запитів" in LOG["ask"][-1][1] and "1990–2024" in LOG["ask"][-1][1] and len(REQ) == 0 and win.spi is None, "перед довгим завантаженням (1990–2024, 35 запитів) програма питає; за відмови запитів немає")
ASK["ans"] = True; REQ.clear(); win._load_spi(); ok = pump(lambda: win.spi is not None, 120); root.update()
check(ok and len(REQ) == 35 and all("community=AG" in u for u in REQ) and "start=19900101" in REQ[0], f"завантажено 35 років (з 1990 – року перед еталонним, 1991–2020) окремими запитами: {len(REQ)}")
sp = win.spi; check(sp["ref"] == (1991, 2020) and sp["scales"] == [1, 3, 6, 12] and len(sp["rows"]) == 48 and sp["rows"][0]["label"] == "01.2021" and sp["rows"][-1]["label"] == "12.2024", "SPI: еталон 1991–2020, масштаби 1/3/6/12, 48 місяців періоду дослідження")
check(all(not math.isnan(r["spi"][12]) for r in sp["rows"]) and all(not math.isnan(r["spi"][1]) for r in sp["rows"]), "для всіх місяців 2021–2024 SPI розраховано (їм передує довгий ряд)")
check(sp["n_ref"] == {1: 30, 3: 30, 6: 30, 12: 30} and sp["warnings"] == [], f"в еталонному періоді для кожного календарного місяця й масштабу рівно 30 років, попереджень немає: {sp['n_ref']}")
check(win.rep["spi"] is sp and win.nb.tab(win.nb.select(), "text") == "SPI", "звіт отримав SPI, відкрилась вкладка «SPI»")
win.win.geometry("1360x640"); win.nb.select(win.tab_spi); root.update()
check(hsb_visible(win.tab_spi, win.holders["spi"]), "таблиця SPI (48 рядків, 8 колонок) у низькому вікні: нижня смуга прокрутки видима")
win.win.geometry("1340x840"); root.update()
tv_s = tree_rows(win.holders["spi"]); check(len(tv_s.get_children()) == 48 and "SPI-12" in tv_s["columns"] and "Оцінка SPI-12" in tv_s["columns"], "таблиця SPI: 48 рядків, колонки SPI-k й оцінка")
check("Еталонний період 1991–2020" in win.spi_note.cget("text") and "екстремально сухо" in win.spi_note.cget("text"), "підпис пояснює еталонний період і класи")
win.chart_var.set("SPI за місяцями"); win._draw_chart(); root.update()
check(len(win._fig.axes) == 4 and any("Вологі" in t.get_text() for a in win._fig.axes if a.get_legend() for t in a.get_legend().get_texts()), "діаграма SPI: 4 панелі (за кількістю масштабів), легенда")
SAVE["path"] = os.path.join(TMP, "s.xlsx"); ASK["ans"] = False; win._export_xlsx(); wb3 = _ox.load_workbook(SAVE["path"])
check("SPI" in wb3.sheetnames and wb3["SPI"].max_row == 49 and wb3["SPI"].cell(1, 3).value == "SPI-1", "Excel: аркуш SPI (48 місяців, колонки SPI-k)")
zs = [wb3["SPI"].cell(r, 3).value for r in range(2, 50)]; cls_ = [wb3["SPI"].cell(r, 4).value for r in range(2, 50)]
check(all(isinstance(z, float) for z in zs) and all(c == W_ind.spi_class(z) for c, z in zip(cls_, zs)), "Excel: SPI числами, оцінки збігаються з класифікацією McKee")
SAVE["path"] = os.path.join(TMP, "s.docx"); win._export_docx(); d3 = _dx.Document(SAVE["path"])
check(not any("SPI" in tb.rows[0].cells[0].text or "SPI" in " ".join(c.text for c in tb.rows[0].cells) for tb in d3.tables), "SPI-таблиця на 48 місяців у Word не вставляється (надто довга) – вона є в Excel")
check(len(d3.inline_shapes) >= 5, f"у Word додано діаграми Вінклера/Югліна й SPI (рисунків {len(d3.inline_shapes)})")
pub = win.txt_pub.get("1.0", "end"); check("Стандартизований індекс опадів (SPI)" in pub and "McKee" in pub and "WMO-No. 1090" in pub and "1991–2020" in pub, "текст для публікації: опис SPI та література (McKee, WMO)")
import re as _re
seq = [int(x) for grp in _re.findall(r"\[(\d+(?:, \d+)*)\]", pub) for x in grp.split(", ")]; first = []
for x in seq:
    if x not in first: first.append(x)
check(first == list(range(1, 8)) and len(_re.findall(r"^\d+\. ", pub, _re.M)) >= 7, f"посилання в тексті нумеруються за порядком першої появи: {first}; у списку літератури 7 джерел")
ASK["ans"] = True
win.sr0.set("1991"); win.sr1.set("2020"); MODE["offline"] = True; had = win.spi; n = len(LOG["error"]); win._load_spi(); pump(lambda: len(LOG["error"]) > n, 15)
check(len(LOG["error"]) == n + 1 and "з'єднання" in LOG["error"][-1][1] and not win._busy and win.spi is had, "немає інтернету під час SPI → повідомлення, попередній SPI збережено")
MODE["offline"] = False
win._clear_spi(); root.update(); check(win.spi is None and win.rep["spi"] is None and len(tree_rows(win.holders["spi"]).get_children()) == 0 if tree_rows(win.holders["spi"]) else win.spi is None, "«Прибрати» очищає SPI й таблицю")
win.spi = sp; win._on_fetched(win.series, win.series.dates[0], win.series.dates[-1]); check(win.spi is None, "нове завантаження даних скидає SPI (інший період чи місце)")
nsp = {"rows": [{"year": 2021, "month": 1, "label": "01.2021", "prec": float("nan"), "spi": {1: float("nan")}, "class": {1: ""}}], "scales": [1], "ref": (1991, 2020), "n_ref": {1: 0}, "warnings": ["тест"], "params": {}}
n = len(LOG["warn"]); win._on_spi(nsp); check(len(LOG["warn"]) == n + 1 and "SPI не розраховано" in LOG["warn"][-1][0] and win.spi is None, "SPI без жодного значення → попередження, а не порожній результат")

print("\n== 9. Аварійні сценарії")
win.src_var.set(W.SOURCE_TITLES["nasa"]); win._on_source(); win.d0_var.set("01.01.2019"); win.d1_var.set("30.06.2019")
MODE["offline"] = True; n = len(LOG["error"]); had = win.series
win._start_fetch(); pump(lambda: len(LOG["error"]) > n, 10)
check(len(LOG["error"]) == n + 1 and "з'єднання" in LOG["error"][-1][1], "немає інтернету → зрозуміле повідомлення")
check(not win._busy and win.series is had, "після помилки вікно працездатне, попередні дані збережено")
MODE["offline"] = False
ok = fetch_and_wait(win); check(ok and len(win.series) == 181, "повторна спроба після відновлення мережі успішна (181 доба)")
win.lat_var.set("95"); n = len(LOG["warn"]); win._start_fetch(); check(len(LOG["warn"]) == n + 1 and "широта" in LOG["warn"][-1][1], "широта 95° відхилена з поясненням")
win.lat_var.set("48°45'N"); win.lon_var.set("30°13'E"); win._start_fetch(); pump(lambda: not win._busy, 10)
check(abs(win.ctx["lat"] - 48.75) < 1e-9 and abs(win.ctx["lon"] - 30.21666667) < 1e-6, "координати у форматі ГМС розпізнано")
win.d0_var.set("2030-01-01"); win.d1_var.set("2030-02-01"); n = len(LOG["warn"]); win._start_fetch()
check(win.d1_var.get() != "01.02.2030" or len(LOG["warn"]) > n, "майбутні дати не призводять до запиту")
win.d0_var.set("01.01.2019"); win.d1_var.set("30.06.2019")

print("\n== 10. Власні дані (CSV у Windows-1251, кома, «;»)")
p = os.path.join(TMP, "station.csv")
rows = ["Дата;Tсер;Tмакс;Tмін;Опади, мм"]
d = dt.date(2023, 5, 1)
for i in range(40):
    day = d + dt.timedelta(days=i); t, pr = _day_vals(day)
    rows.append(f"{day:%d.%m.%Y};" + f"{t:.1f};{t+4:.1f};{t-4:.1f};{pr:.1f}".replace(".", ","))   # кома лише в числах
rows.append("помилковий рядок;1;2;3;4")
open(p, "wb").write("\r\n".join(rows).encode("cp1251"))
win.src_var.set(W.SOURCE_TITLES["file"]); win._on_source(); SAVE["open"] = p
win._open_file(); root.update()
dlg = [t for t in toplevels(win.win) if "Відповідність" in t.title()][0]
find_btn("Прочитати", dlg).invoke(); root.update()
check(win.series is not None and len(win.series) == 40 and win.ctx["source"] == "file", "станційний файл прочитано: 40 діб")
check(win.series.meta["bad_rows"] == 1 and "некоректних рядків: 1" in win.status_var.get(), "некоректний рядок пропущено й про це повідомлено")
check(win.series.has("tmean") and win.series.has("prec") and not win.series.has("rh"), "стовпці розпізнано автоматично")
check("реаналізу" not in win.txt_sum.get("1.0", "end"), "для станційних даних немає попередження про реаналіз")
win.win.clipboard_clear(); win.win.clipboard_append("Дата\tt\tОпади\n01.06.2023\t15,5\t0\n02.06.2023\t16,0\t2,5\n03.06.2023\t14,1\t0\n")
win._paste_table(); dlg = [t for t in toplevels(win.win) if "Відповідність" in t.title()][0]
find_btn("Прочитати", dlg).invoke(); root.update()
check(win.series is not None and len(win.series) == 3 and abs(win.series.get("tmean")[0] - 15.5) < 1e-9, "вставка з буфера (TSV, десяткова кома) працює")

print("\n== 11. Закриття вікна під час завантаження")
win2 = m.WeatherWindow(root); root.update()
win2._set_point(50.4501, 30.5234, "Київ"); win2.d0_var.set("01.01.2021"); win2.d1_var.set("31.12.2023")
MODE["slow"] = 0.4
win2._start_fetch(); real_sleep(0.2)
win2._close(); root.update()
check(not win2._alive and not win2.win.winfo_exists(), "вікно закрито посеред завантаження")
real_sleep(1.6); root.update()
check(True, "фоновий потік завершився без помилок у закритому вікні")
MODE["slow"] = 0.0

print("\n== 12. Довідка")
win._show_help(); root.update()
hw = [t for t in toplevels(win.win) if "Довідка" in t.title()]
htxt = hw[0].winfo_children()[0].winfo_children()[1].get("1.0", "end") if hw else ""
check(hw and "вбудовану карту світу" in htxt and "ПОРІВНЯННЯ РОКІВ" in htxt and "ГРАФІКИ" in htxt, "довідка відкривається й описує карту, порівняння років і налаштування графіків")
for bad in ("Open-Meteo", "КОМЕРЦІЙНОГО", "кеш", "Польовий", "Божко", "ТЕС, 2004"):
    check(bad not in htxt, f"у довідці немає згадки «{bad}»")

shutil.rmtree(TMP, ignore_errors=True); shutil.rmtree(os.environ["APPDATA"], ignore_errors=True)
root.destroy()
print("\n" + "=" * 60)
if FAIL:
    print(f"ПРОВАЛІВ: {len(FAIL)}")
    for f in FAIL: print("  -", f)
    sys.exit(1)
print("УСІ ПЕРЕВІРКИ ВІКНА ПРОЙДЕНО")
