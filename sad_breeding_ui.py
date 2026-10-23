# sad_breeding_ui.py — S.A.D. v2.1: переглядач результатів, експорт і графіки селекційних аналізів
# -*- coding: utf-8 -*-
from sad_common import *
import io, math
import matplotlib
from matplotlib.figure import Figure
import sad_gxe_core as GX

RF = ("Times New Roman", 11)
RB = ("Times New Roman", 11, "bold")
TAG_COLORS = {"good": "#e2efda", "warn": "#fff2cc", "bad": "#f4b6b6", "info": "#ddebf7"}
XL_FILL = {"good": "E2EFDA", "warn": "FFF2CC", "bad": "F4B6B6", "info": "DDEBF7"}
_RC = {"font.family": "serif", "font.serif": ["Times New Roman", "Times", "DejaVu Serif"],
       "font.size": 10, "axes.titlesize": 11, "axes.labelsize": 10,
       "axes.spines.top": False, "axes.spines.right": False, "axes.unicode_minus": False}
PALETTE = ["#1f77b4", "#ff7f0e", "#2ca02c", "#d62728", "#9467bd", "#8c564b", "#e377c2", "#7f7f7f",
           "#bcbd22", "#17becf"]


# ═══════════════════════════════════════════════════════════════
# Розділи результатів
# ═══════════════════════════════════════════════════════════════
def fcell(v, nd, comma):
    """Форматування комірки: None/NaN → «—»; рядки без змін; числа з nd знаками (кома за потреби)."""
    if v is None or (isinstance(v, float) and math.isnan(v)): return "—"
    if isinstance(v, tuple):                      # (значення, суфікс) — напр. (1.23, "**")
        return fcell(v[0], nd, comma) + v[1]
    if isinstance(v, str): return v
    if nd is None: return str(v)
    if nd == "p":                                 # p-значення: «<0,0001» замість «0,0000»
        s = "<0.0001" if v < 0.0001 else f"{v:.4f}"
        return s.replace(".", ",") if comma else s
    s = f"{v:.{nd}f}"
    return s.replace(".", ",") if comma else s


def sec_table(title, headers, rows, nd=None, notes=(), tags=None, hint="", key=None):
    """Таблиця: rows — сирі значення; nd — список знаків після коми по стовпцях (None — як є)."""
    nd = nd if nd is not None else [None] * len(headers)
    return {"kind": "table", "title": title, "headers": list(headers), "rows": [list(r) for r in rows],
            "nd": list(nd), "notes": list(notes), "tags": tags, "hint": hint, "key": key or title}


def sec_text(title, text, hint="", key=None):
    return {"kind": "text", "title": title, "text": text, "hint": hint, "key": key or title}


def sec_fig(title, build, hint="", size=(8.8, 5.8), key=None):
    return {"kind": "figure", "title": title, "build": build, "hint": hint, "size": size, "key": key or title}


def table_strings(sec, comma):
    rows = [[fcell(v, n, comma) for v, n in zip(r, sec["nd"])] for r in sec["rows"]]
    return sec["headers"], rows


def fig_to_png(sec, dpi=200):
    fig = Figure(figsize=sec["size"], dpi=dpi)
    with matplotlib.rc_context(_RC):
        res = sec["build"](fig)
    if res: return None
    buf = io.BytesIO(); fig.savefig(buf, format="png", dpi=dpi, bbox_inches="tight"); return buf.getvalue()


def _sheet_name(title, used):
    bad = '[]:*?/\\'
    nm = "".join("_" if c in bad else c for c in title)[:31] or "Аркуш"
    base, k = nm, 2
    while nm in used:
        nm = f"{base[:28]}_{k}"; k += 1
    used.add(nm); return nm


def export_xlsx(path, sections, title="Результати"):
    """Усі таблиці — на окремі аркуші (числа лишаються числами); текстові розділи — на аркуш «Примітки»."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    hf = Font(bold=True, name="Times New Roman", size=11); hfill = PatternFill("solid", fgColor="D9E2F3")
    bf = Font(name="Times New Roman", size=11)
    thin = Side(style="thin", color="BBBBBB"); bd = Border(left=thin, right=thin, top=thin, bottom=thin)
    used = set()
    notes = []
    for sec in sections:
        if sec["kind"] == "text":
            notes.append((sec["title"], sec["text"])); continue
        if sec["kind"] != "table": continue
        ws = wb.create_sheet(_sheet_name(sec["title"], used))
        ws.append(sec["headers"])
        for j in range(1, len(sec["headers"]) + 1):
            c = ws.cell(row=1, column=j); c.font = hf; c.fill = hfill; c.border = bd
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.column_dimensions[get_column_letter(j)].width = max(11, min(26, len(str(sec["headers"][j - 1])) + 3))
        ws.row_dimensions[1].height = 34; ws.freeze_panes = "A2"
        for ri, r in enumerate(sec["rows"]):
            ws.append([(fcell(v[0], 3, True) + v[1] if isinstance(v, tuple) else
                        (None if (isinstance(v, float) and math.isnan(v)) else v)) for v in r])
            tag = (sec["tags"][ri] if sec.get("tags") else None)
            for c, n in zip(ws[ws.max_row], sec["nd"]):
                c.font = bf; c.border = bd
                if isinstance(c.value, float) and n is not None:
                    c.number_format = "0.0000" if n == "p" else ("0" if n == 0 else "0." + "0" * n)
                if tag in XL_FILL: c.fill = PatternFill("solid", fgColor=XL_FILL[tag])
        if sec.get("notes"):
            ws.append([])
            for nline in sec["notes"]: ws.append([nline])
    if notes:
        ws = wb.create_sheet("Примітки")
        ws.column_dimensions["A"].width = 120
        for t_, txt in notes:
            ws.append([t_]); ws.cell(row=ws.max_row, column=1).font = hf
            for line in txt.split("\n"): ws.append([line])
            ws.append([])
    if not wb.sheetnames: wb.create_sheet("Порожньо")
    wb.save(path)


def export_docx(path, title, sections, comma=True):
    try:
        from docx import Document
        from docx.shared import Pt, Cm, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.table import WD_TABLE_ALIGNMENT
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
    except ImportError:
        raise RuntimeError("Для експорту у Word встановіть python-docx:\n  pip install python-docx")
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    for sec_ in doc.sections:
        sec_.top_margin = Cm(2); sec_.bottom_margin = Cm(2); sec_.left_margin = Cm(2.5); sec_.right_margin = Cm(1.5)

    def shade(cell, hex_):
        tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_); tcPr.append(shd)

    def para(text, bold=False, italic=False, size=None, align=None, after=6):
        p = doc.add_paragraph(); r = p.add_run(text); r.bold = bold; r.italic = italic; r.font.name = "Times New Roman"
        if size: r.font.size = Pt(size)
        if align == "center": p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(after); return p

    h = doc.add_heading(title, level=1)
    for r in h.runs: r.font.name = "Times New Roman"; r.font.color.rgb = RGBColor(0, 0, 0)
    n_tab = n_fig = 0
    for sec in sections:
        if sec["kind"] == "text":
            para(sec["title"], bold=True, size=12, after=2)
            for line in sec["text"].split("\n"):
                if line.strip(): para(line, size=11, after=2)
        elif sec["kind"] == "table":
            headers, rows = table_strings(sec, comma)
            if not rows: continue
            n_tab += 1
            para(f"Таблиця {n_tab} — {sec['title']}", italic=True, after=3)
            tbl = doc.add_table(rows=1 + len(rows), cols=len(headers)); tbl.style = "Table Grid"
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            for j, hd in enumerate(headers):
                c = tbl.rows[0].cells[j]; c.text = ""; run = c.paragraphs[0].add_run(str(hd))
                run.bold = True; run.font.size = Pt(9 if len(headers) > 9 else 10); run.font.name = "Times New Roman"
                c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER; shade(c, "D9E2F3")
            for i, row in enumerate(rows, 1):
                tag = sec["tags"][i - 1] if sec.get("tags") else None
                for j, v in enumerate(row):
                    c = tbl.rows[i].cells[j]; c.text = ""; run = c.paragraphs[0].add_run(v)
                    run.font.size = Pt(9 if len(headers) > 9 else 10); run.font.name = "Times New Roman"
                    c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER if j else WD_ALIGN_PARAGRAPH.LEFT
                    if tag in XL_FILL: shade(c, XL_FILL[tag])
            for nline in sec.get("notes", []): para(nline, italic=True, size=10, after=2)
            para("", after=4)
        else:
            png = fig_to_png(sec)
            if png is None: continue
            n_fig += 1
            doc.add_picture(io.BytesIO(png), width=Cm(15.5))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            para(f"Рис. {n_fig} — {sec['title']}", align="center", size=11)
    doc.save(path)



def scrollable_column(parent, bg="#2c3e50", width=210):
    """Прокручувана колонка (Canvas + Scrollbar + колесо миші). Повертає (outer, inner):
    outer пакуємо у батьківський контейнер, у inner додаємо кнопки. Потрібна, коли пунктів більше,
    ніж вміщає вікно — інакше нижні розділи недосяжні."""
    outer = tk.Frame(parent, bg=bg)
    canvas = tk.Canvas(outer, bg=bg, highlightthickness=0, width=width)
    sb = ttk.Scrollbar(outer, orient="vertical", command=canvas.yview)
    inner = tk.Frame(canvas, bg=bg)
    wid = canvas.create_window((0, 0), window=inner, anchor="nw")
    inner.bind("<Configure>", lambda e: canvas.configure(scrollregion=canvas.bbox("all")))
    canvas.bind("<Configure>", lambda e: canvas.itemconfig(wid, width=e.width))
    canvas.configure(yscrollcommand=sb.set)
    canvas.pack(side=tk.LEFT, fill=tk.BOTH, expand=True); sb.pack(side=tk.RIGHT, fill=tk.Y)

    def wheel(e):
        step = int(-1 * (e.delta / 120)) if getattr(e, "delta", 0) else (1 if e.num == 5 else -1)
        canvas.yview_scroll(step, "units")
    outer.bind("<Enter>", lambda e: (outer.bind_all("<MouseWheel>", wheel), outer.bind_all("<Button-4>", wheel),
                                     outer.bind_all("<Button-5>", wheel)))
    outer.bind("<Leave>", lambda e: (outer.unbind_all("<MouseWheel>"), outer.unbind_all("<Button-4>"),
                                     outer.unbind_all("<Button-5>")))
    outer._canvas = canvas
    return outer, inner


def sidebar_button(parent, text, command=None):
    """Компактна кнопка бічного меню (однакова за виглядом із наявними, але нижча)."""
    return tk.Button(parent, text="  " + text, bg="#2c3e50", fg="#bdc3c7", font=("Times New Roman", 11), relief=tk.FLAT,
                     anchor="w", padx=10, pady=4, wraplength=180, justify="left", activebackground="#c62828",
                     activeforeground="white", command=command)

# ═══════════════════════════════════════════════════════════════
# Відображення розділу у фреймі
# ═══════════════════════════════════════════════════════════════
def render_section(frame, sec, comma_var, owner_win):
    for w in frame.winfo_children(): w.destroy()
    kind = sec["kind"]
    if kind == "text":
        txt = ScrolledText(frame, wrap="word", font=("Times New Roman", 12), relief=tk.FLAT)
        txt.pack(fill=tk.BOTH, expand=True, padx=8, pady=8)
        txt.insert("1.0", sec["text"]); txt.configure(state="disabled")
    elif kind == "table":
        bar = tk.Frame(frame, padx=6, pady=4); bar.pack(fill=tk.X)
        def copy():
            headers, rows = table_strings(sec, comma_var.get())
            owner_win.clipboard_clear()
            owner_win.clipboard_append("\n".join("\t".join(map(str, r)) for r in [headers] + rows))
            messagebox.showinfo("Скопійовано", "Таблицю скопійовано. Вставте у Word/Excel (Ctrl+V).", parent=owner_win)
        tk.Button(bar, text="📋 Копіювати таблицю", font=RF, command=copy).pack(side=tk.LEFT)
        if sec.get("hint"):
            tk.Label(bar, text=sec["hint"], font=("Times New Roman", 9), fg="#666", justify="left",
                     wraplength=760).pack(side=tk.LEFT, padx=10)
        headers, rows = table_strings(sec, comma_var.get())
        if rows:
            frm, tv = make_tv(frame, headers, rows, min_col=40)
            frm.pack(fill=tk.BOTH, expand=True, padx=8, pady=4)
            fnt = tkfont.Font(family="Times New Roman", size=11)           # ущільнення: ширина за вмістом
            for i_, h_ in enumerate(headers):
                wmax = max([fnt.measure(str(h_)) + 16] + [fnt.measure(str(r_[i_])) + 14 for r_ in rows[:300] if i_ < len(r_)])
                tv.column(h_, width=max(44, wmax), minwidth=40, stretch=(i_ == len(headers) - 1))
            for k, bg in TAG_COLORS.items(): tv.tag_configure(k, background=bg)
            if sec.get("tags"):
                for iid, tg in zip(tv.get_children(), sec["tags"]):
                    if tg in TAG_COLORS: tv.item(iid, tags=(tg,))
        if sec.get("notes"):
            tk.Label(frame, text="\n".join(sec["notes"]), font=("Times New Roman", 10), fg="#555", justify="left",
                     wraplength=1000, anchor="w").pack(fill=tk.X, padx=10, pady=(2, 8))
    else:
        if not HAS_MPL:
            tk.Label(frame, text="Для графіків потрібен matplotlib (pip install matplotlib).").pack(pady=30); return
        bar = tk.Frame(frame, padx=6, pady=4); bar.pack(fill=tk.X)
        fig = Figure(figsize=sec["size"], dpi=100)
        with matplotlib.rc_context(_RC):
            msg = sec["build"](fig)
        canvas = FigureCanvasTkAgg(fig, master=frame)
        def save():
            p = filedialog.asksaveasfilename(parent=owner_win, defaultextension=".png",
                                             filetypes=[("PNG", "*.png"), ("SVG", "*.svg"), ("PDF", "*.pdf")])
            if p:
                cur = fig.get_size_inches().copy()
                try:
                    fig.set_size_inches(*sec["size"])               # зберігаємо в задуманому розмірі, а не в розмірі вікна
                    fig.savefig(p, dpi=300, bbox_inches="tight")
                except Exception as e: messagebox.showerror("", str(e), parent=owner_win)
                finally:
                    fig.set_size_inches(*cur); canvas.draw_idle()
        def copy():
            ok, m = _copy_fig_to_clipboard(fig)
            (messagebox.showinfo if ok else messagebox.showwarning)("", "Графік скопійовано." if ok else f"Помилка: {m}",
                                                                      parent=owner_win)
        tk.Button(bar, text="💾 Зберегти (PNG/SVG/PDF)…", font=RF, command=save).pack(side=tk.LEFT)
        tk.Button(bar, text="📋 Копіювати", font=RF, command=copy).pack(side=tk.LEFT, padx=4)
        if sec.get("hint"):
            tk.Label(bar, text=sec["hint"], font=("Times New Roman", 9), fg="#666", justify="left",
                     wraplength=700).pack(side=tk.LEFT, padx=10)
        cw = canvas.get_tk_widget(); cw.configure(width=300, height=240)        # мале «бажане» значення: розмір визначає вікно,
        cw.pack(fill=tk.BOTH, expand=True); canvas.draw_idle()                 # а діаграма підлаштовується (інакше її підрізає)
        frame._fig = fig


def do_export(win, sections, title, kind, comma):
    """Спільний обробник експорту (Excel/Word) для переглядача та вікна стабільності."""
    if kind == "xlsx":
        if not HAS_OPENPYXL:
            messagebox.showerror("", "Потрібен openpyxl: pip install openpyxl", parent=win); return
        p = filedialog.asksaveasfilename(parent=win, defaultextension=".xlsx", filetypes=[("Excel", "*.xlsx")])
        if not p: return
        try: export_xlsx(p, sections, title)
        except Exception as e: messagebox.showerror("Помилка експорту", str(e), parent=win); return
    else:
        p = filedialog.asksaveasfilename(parent=win, defaultextension=".docx", filetypes=[("Word", "*.docx")])
        if not p: return
        win.configure(cursor="watch"); win.update_idletasks()
        try: export_docx(p, title, sections, comma)
        except Exception as e:
            win.configure(cursor=""); messagebox.showerror("Помилка експорту", str(e), parent=win); return
        win.configure(cursor="")
    if messagebox.askyesno("", f"Файл збережено:\n{p}\n\nВідкрити?", parent=win): open_file_cross(p)


# ═══════════════════════════════════════════════════════════════
# Переглядач результатів
# ═══════════════════════════════════════════════════════════════
class ResultsViewer:
    def __init__(self, parent, title, sections, banner=None, geometry="1320x780"):
        self.sections = sections; self.title = title
        self.win = tk.Toplevel(parent); self.win.title(title); self.win.geometry(geometry); set_icon(self.win)
        self.comma_var = tk.BooleanVar(value=True)
        top = tk.Frame(self.win, padx=8, pady=5); top.pack(fill=tk.X)
        tk.Button(top, text="📊 Excel", font=RF, command=lambda: do_export(
            self.win, self.sections, self.title, "xlsx", self.comma_var.get())).pack(side=tk.LEFT, padx=3)
        tk.Button(top, text="📄 Word", font=RF, command=lambda: do_export(
            self.win, self.sections, self.title, "docx", self.comma_var.get())).pack(side=tk.LEFT, padx=3)
        tk.Checkbutton(top, text="Десяткова кома", variable=self.comma_var, font=("Times New Roman", 10),
                       command=self._rerender).pack(side=tk.RIGHT)
        if banner:
            bf = tk.Frame(self.win, bg="#fff4d6", padx=8, pady=5); bf.pack(fill=tk.X, padx=8)
            tk.Label(bf, text="⚠ " + "\n⚠ ".join(banner), bg="#fff4d6", fg="#7a4b00", justify="left",
                     font=("Times New Roman", 10), wraplength=1080, anchor="w").pack(fill=tk.X)
        main = tk.Frame(self.win); main.pack(fill=tk.BOTH, expand=True, pady=(4, 0))
        outer, self.sidebar = scrollable_column(main, "#2c3e50", 230)
        outer.pack(side=tk.LEFT, fill=tk.Y)
        self.content = tk.Frame(main); self.content.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        self.frames, self.btns, self.rendered = [], [], []
        for i, sec in enumerate(sections):
            b = sidebar_button(self.sidebar, sec["title"], lambda k=i: self.show(k))
            b.pack(fill=tk.X); tk.Frame(self.sidebar, bg="#3d5166", height=1).pack(fill=tk.X)
            self.btns.append(b); self.frames.append(tk.Frame(self.content)); self.rendered.append(False)
        self.cur = None
        if sections: self.show(0)

    def show(self, k):
        if self.cur is not None:
            self.frames[self.cur].pack_forget(); self.btns[self.cur].configure(bg="#2c3e50", fg="#bdc3c7")
        if not self.rendered[k]:
            render_section(self.frames[k], self.sections[k], self.comma_var, self.win); self.rendered[k] = True
        self.frames[k].pack(fill=tk.BOTH, expand=True); self.btns[k].configure(bg="#c62828", fg="white"); self.cur = k

    def _rerender(self):
        self.rendered = [False] * len(self.sections)
        if self.cur is not None: self.show(self.cur)


# ═══════════════════════════════════════════════════════════════
# Графіки
# ═══════════════════════════════════════════════════════════════
def _label_points(ax, xs, ys, names, color="#222", fs=8):
    for x, y, n in zip(xs, ys, names):
        ax.annotate(str(n), (x, y), xytext=(4, 4), textcoords="offset points", fontsize=fs, color=color)



def _square(ax, arrays, pad=1.18):
    """Симетричні квадратні межі навколо початку координат + рівний масштаб осей."""
    m = max(float(np.max(np.abs(a))) for a in arrays if np.size(a)) * pad or 1.0
    ax.set_xlim(-m, m); ax.set_ylim(-m, m); ax.set_aspect("equal", adjustable="box")

def _pc_labels(ax, pct, i=0, j=1, prefix="PC"):
    ax.set_xlabel(f"{prefix}{i + 1} ({pct[i]:.1f} %)"); ax.set_ylabel(f"{prefix}{j + 1} ({pct[j]:.1f} %)" if len(pct) > j else "")


def fig_ammi1(fig, am, gnames, enames):
    ax = fig.add_subplot(111)
    gs, es = am["gen_scores"][:, 0], am["env_scores"][:, 0]
    ax.scatter(am["means"], gs, s=60, c="#1f77b4", zorder=3, label="Генотипи")
    ax.scatter(am["env_means"], es, s=70, c="#d62728", marker="^", zorder=3, label="Середовища")
    _label_points(ax, list(am["means"]), list(gs), gnames, "#1f4e79")
    _label_points(ax, list(am["env_means"]), list(es), enames, "#9b1c1c")
    ax.axhline(0, color="#777", lw=0.8); ax.axvline(am["grand_mean"], color="#777", lw=0.8, ls="--")
    pc1 = am["axes"][0]["pct"]
    ax.set_xlabel("Середнє значення показника"); ax.set_ylabel(f"IPCA1 ({pc1:.1f} % взаємодії)")
    ax.set_title("AMMI1-біплот: продуктивність і стабільність\n(ближче до горизонталі 0 — стабільніший)", loc="left", fontsize=10.5)
    ax.legend(frameon=False, fontsize=8, loc="best"); fig.tight_layout()


def fig_ammi2(fig, am, gnames, enames):
    if am["k_all"] < 2: return "Для AMMI2-біплота потрібно щонайменше 2 осі взаємодії."
    ax = fig.add_subplot(111)
    g1, g2 = am["gen_scores"][:, 0], am["gen_scores"][:, 1]; e1, e2 = am["env_scores"][:, 0], am["env_scores"][:, 1]
    ax.scatter(g1, g2, s=60, c="#1f77b4", zorder=3)
    for x, y in zip(e1, e2): ax.annotate("", xy=(x, y), xytext=(0, 0), arrowprops=dict(arrowstyle="->", color="#d62728", lw=1.2))
    _label_points(ax, list(g1), list(g2), gnames, "#1f4e79"); _label_points(ax, list(e1), list(e2), enames, "#9b1c1c")
    ax.axhline(0, color="#777", lw=0.8); ax.axvline(0, color="#777", lw=0.8)
    ax.set_xlabel(f"IPCA1 ({am['axes'][0]['pct']:.1f} %)"); ax.set_ylabel(f"IPCA2 ({am['axes'][1]['pct']:.1f} %)")
    ax.set_title("AMMI2-біплот: взаємодія генотип × середовище\n(генотип поблизу стрілки середовища — добре до нього адаптований)",
                 loc="left", fontsize=10.5)
    _square(ax, [am["gen_scores"][:, :2], am["env_scores"][:, :2]]); fig.tight_layout()


def fig_gge_www(fig, gg, ww, gnames, enames):
    G, E = gg["G"], gg["E"]
    if G.shape[1] < 2: return "Для GGE-біплота потрібно щонайменше 2 головні компоненти."
    ax = fig.add_subplot(111)
    ax.scatter(G[:, 0], G[:, 1], s=55, c="#1f77b4", zorder=3)
    hull = ww["hull"]
    if len(hull) >= 3:
        pts = G[hull][:, :2]
        order = np.argsort(np.arctan2(pts[:, 1] - pts[:, 1].mean(), pts[:, 0] - pts[:, 0].mean()))
        hp = pts[order]
        ax.plot(np.append(hp[:, 0], hp[0, 0]), np.append(hp[:, 1], hp[0, 1]), color="#2c3e50", lw=1.4, zorder=2)
        R = float(np.max(np.abs(G[:, :2]))) * 1.35 + 1e-9
        for k in range(len(hp)):
            a, b = hp[k], hp[(k + 1) % len(hp)]
            d = b - a; n = np.array([-d[1], d[0]])
            if n @ a < 0: n = -n
            nn = np.linalg.norm(n)
            if nn > 0: ax.plot([0, R * n[0] / nn], [0, R * n[1] / nn], color="#7f8c8d", lw=0.9, ls="--", zorder=1)
    scale = (np.max(np.abs(G[:, :2])) / (np.max(np.abs(E[:, :2])) + 1e-12)) * 0.85
    for j in range(E.shape[0]):
        ax.annotate("", xy=(E[j, 0] * scale, E[j, 1] * scale), xytext=(0, 0), arrowprops=dict(arrowstyle="->", color="#d62728", lw=1.1))
        ax.annotate(enames[j], (E[j, 0] * scale, E[j, 1] * scale), fontsize=8, color="#9b1c1c", xytext=(3, 3), textcoords="offset points")
    _label_points(ax, list(G[:, 0]), list(G[:, 1]), gnames, "#1f4e79")
    ax.axhline(0, color="#aaa", lw=0.7); ax.axvline(0, color="#aaa", lw=0.7)
    _pc_labels(ax, gg["pct"])
    ax.set_title("GGE-біплот «хто де переміг» (which-won-where)\nГенотипи на вершинах багатокутника — переможці в своїх секторах",
                 loc="left", fontsize=10.5)
    _square(ax, [G[:, :2], E[:, :2] * scale]); fig.tight_layout()


def fig_gge_aec(fig, gg, ae, gnames, enames):
    G, E = gg["G"], gg["E"]
    if G.shape[1] < 2: return "Для GGE-біплота потрібно щонайменше 2 головні компоненти."
    ax = fig.add_subplot(111)
    u = ae["axis"]; perp = np.array([-u[1], u[0]])
    R = float(np.max(np.abs(G[:, :2]))) * 1.3 + 1e-9
    ax.annotate("", xy=(u[0] * R, u[1] * R), xytext=(-u[0] * R, -u[1] * R), arrowprops=dict(arrowstyle="->", color="#2c3e50", lw=1.6))
    ax.plot([-perp[0] * R, perp[0] * R], [-perp[1] * R, perp[1] * R], color="#2c3e50", lw=1.0, ls="--")
    ax.scatter(G[:, 0], G[:, 1], s=55, c="#1f77b4", zorder=3)
    for i in range(G.shape[0]):
        proj = (G[i, :2] @ u) * u
        ax.plot([G[i, 0], proj[0]], [G[i, 1], proj[1]], color="#bbb", lw=0.7, zorder=1)
    scale = (np.max(np.abs(G[:, :2])) / (np.max(np.abs(E[:, :2])) + 1e-12)) * 0.85
    ax.scatter(E[:, 0] * scale, E[:, 1] * scale, s=45, c="#d62728", marker="^", zorder=3)
    _label_points(ax, list(G[:, 0]), list(G[:, 1]), gnames, "#1f4e79")
    _label_points(ax, list(E[:, 0] * scale), list(E[:, 1] * scale), enames, "#9b1c1c")
    ax.scatter([ae["env_point"][0] * scale], [ae["env_point"][1] * scale], s=130, facecolors="none", edgecolors="#d62728", lw=1.8, zorder=4)
    ax.axhline(0, color="#ddd", lw=0.5); ax.axvline(0, color="#ddd", lw=0.5)
    _pc_labels(ax, gg["pct"])
    ax.set_title("GGE-біплот: продуктивність і стабільність\nСтрілка — вісь середнього середовища (праворуч — вища продуктивність);\n"
                 "чим довша проекція перпендикуляра — тим нестабільніший генотип", loc="left", fontsize=10)
    _square(ax, [G[:, :2], E[:, :2] * scale]); fig.tight_layout()


def fig_bars(fig, labels, series, title, ylabel):
    """Групові стовпчики: series — {назва: список значень}."""
    ax = fig.add_subplot(111)
    n = len(series); w = 0.8 / max(n, 1); x = np.arange(len(labels))
    for k, (nm, vals) in enumerate(series.items()):
        ax.bar(x + (k - (n - 1) / 2) * w, [0 if v is None or (isinstance(v, float) and math.isnan(v)) else v for v in vals],
               w, label=nm, color=PALETTE[k % 10])
    ax.set_xticks(x); ax.set_xticklabels(labels, rotation=30 if len(labels) > 4 else 0, ha="right" if len(labels) > 4 else "center")
    ax.set_ylabel(ylabel); ax.set_title(title, loc="left", fontsize=10.5); ax.legend(frameon=False, fontsize=8); fig.tight_layout()


def fig_effects(fig, groups, title, ylabel):
    """Стовпчики ефектів з ±SE: groups — [(підпис групи, імена, значення, SE)]."""
    n = len(groups)
    axes = fig.subplots(1, n, squeeze=False)[0]
    for ax, (gl, names, vals, se) in zip(axes, groups):
        cols = ["#2ca02c" if v >= 0 else "#d62728" for v in vals]
        ax.bar(range(len(vals)), vals, yerr=se, color=cols, capsize=3, ecolor="#555")
        ax.axhline(0, color="#333", lw=0.8)
        ax.set_xticks(range(len(vals))); ax.set_xticklabels(names, rotation=40, ha="right", fontsize=8)
        ax.set_title(gl, fontsize=10, loc="left"); ax.set_ylabel(ylabel)
    fig.suptitle(title, x=0.01, ha="left", fontsize=10.5); fig.tight_layout()


def fig_heat(fig, M, rown, coln, title, cbar):
    ax = fig.add_subplot(111)
    M = np.asarray(M, float); v = np.nanmax(np.abs(M)) or 1.0
    im = ax.imshow(M, cmap="RdYlGn", vmin=-v, vmax=v, aspect="auto")
    ax.set_xticks(range(len(coln))); ax.set_xticklabels(coln, rotation=40, ha="right", fontsize=8)
    ax.set_yticks(range(len(rown))); ax.set_yticklabels(rown, fontsize=8)
    for i in range(M.shape[0]):
        for j in range(M.shape[1]):
            if not np.isnan(M[i, j]): ax.text(j, i, f"{M[i, j]:.2f}".replace(".", ","), ha="center", va="center", fontsize=7)
    fig.colorbar(im, ax=ax, label=cbar, shrink=0.8)
    for sp in ax.spines.values(): sp.set_visible(False)
    ax.set_title(title, loc="left", fontsize=10.5); fig.tight_layout()


def fig_stress(fig, st, names):
    ax = fig.add_subplot(111)
    cols = {"A": "#2ca02c", "B": "#1f77b4", "C": "#ff7f0e", "D": "#7f7f7f"}
    for grp in "ABCD":
        idx = [i for i, g in enumerate(st["group"]) if g == grp]
        if idx: ax.scatter(st["Yp"][idx], st["Ys"][idx], s=60, c=cols[grp], label=f"Група {grp}", zorder=3)
    for x, y, n in zip(st["Yp"], st["Ys"], names): ax.annotate(str(n), (x, y), xytext=(4, 4), textcoords="offset points", fontsize=8)
    ax.axvline(st["mean_Yp"], color="#777", ls="--", lw=0.9); ax.axhline(st["mean_Ys"], color="#777", ls="--", lw=0.9)
    lim = [min(st["Yp"].min(), st["Ys"].min()) * 0.95, st["Yp"].max() * 1.05]
    ax.plot(lim, lim, color="#bbb", lw=0.8, ls=":")
    ax.set_xlabel("Yp — урожай у сприятливих умовах"); ax.set_ylabel("Ys — урожай у стресових умовах")
    ax.set_title("Групи Фернандеса за Yp і Ys\nA — високі в обох умовах (найцінніші); B — лише в сприятливих; "
                 "C — лише в стресових; D — низькі в обох", loc="left", fontsize=10)
    ax.legend(frameon=False, fontsize=8); fig.tight_layout()


def fig_path_bars(fig, pa):
    ax = fig.add_subplot(111); k = pa["k"]; x = np.arange(k)
    ax.bar(x - 0.2, pa["rxy"], 0.4, label="Кореляція r(X, Y)", color="#1f77b4")
    ax.bar(x + 0.2, pa["direct"], 0.4, label="Прямий ефект p", color="#ff7f0e")
    ax.axhline(0, color="#333", lw=0.8); ax.set_xticks(x); ax.set_xticklabels(pa["names"], rotation=30, ha="right")
    ax.set_title(f"Кореляція та прямий ефект на «{pa['y_name']}»\n(розбіжність знаків — ознака колінеарності чи придушення)",
                 loc="left", fontsize=10.5); ax.legend(frameon=False, fontsize=8); fig.tight_layout()


def fig_path_diagram(fig, pa):
    ax = fig.add_subplot(111); ax.axis("off"); k = pa["k"]
    ys = np.linspace(0.9, 0.1, k)
    ax.text(0.88, 0.5, pa["y_name"], ha="center", va="center", fontsize=11, bbox=dict(boxstyle="round", fc="#ddebf7", ec="#1f4e79"))
    for i in range(k):
        ax.text(0.12, ys[i], pa["names"][i], ha="center", va="center", fontsize=10, bbox=dict(boxstyle="round", fc="#e2efda", ec="#2e7d32"))
        d = pa["direct"][i]; ty = 0.5 + (ys[i] - 0.5) * 0.25
        ax.annotate("", xy=(0.80, ty), xytext=(0.20, ys[i]),
                    arrowprops=dict(arrowstyle="->", color="#c62828" if d < 0 else "#2e7d32", lw=1 + 3 * min(abs(d), 1)))
        ax.text(0.5, (ys[i] + ty) / 2 + 0.025, f"{d:.3f}".replace(".", ",").replace("-", "−"), fontsize=9, ha="center", color="#222")
    ax.annotate("", xy=(0.88, 0.42), xytext=(0.88, 0.12), arrowprops=dict(arrowstyle="->", color="#7f7f7f", lw=1.2))
    ax.text(0.88, 0.07, f"залишок = {pa['residual']:.3f}".replace(".", ","), ha="center", fontsize=9, color="#555")
    ax.set_title(f"Діаграма шляхового аналізу (R² = {pa['R2']:.3f})".replace(".", ",") +
                 "\nТовщина стрілки — сила прямого ефекту; червоні — від'ємні", loc="left", fontsize=10.5)


def fig_forest(fig, res, title, unit=""):
    """Різниця «номер − стандарт» із одночасним інтервалом Дуннетта; значущі — кольором."""
    rows = sorted(res["rows"], key=lambda r: r["diff"])
    if not rows: return "Немає номерів для порівняння."
    ax = fig.add_subplot(111); y = np.arange(len(rows))
    for i, r in enumerate(rows):
        col = "#2ca02c" if (r["sig"] and r["diff"] > 0) else "#d62728" if (r["sig"] and r["diff"] < 0) else "#7f7f7f"
        lo = r["ci_lo"] if np.isfinite(r["ci_lo"]) else r["diff"] - 4 * r["se_diff"]
        hi = r["ci_hi"] if np.isfinite(r["ci_hi"]) else r["diff"] + 4 * r["se_diff"]
        ax.plot([lo, hi], [i, i], color=col, lw=2.2, solid_capstyle="round"); ax.plot(r["diff"], i, "o", color=col, ms=6)
    ax.axvline(0, color="#333", lw=1)
    ax.set_yticks(y); ax.set_yticklabels([r["entry"] for r in rows], fontsize=8)
    ax.set_xlabel(f"Різниця від стандарту{', ' + unit if unit else ''}")
    ax.set_title(title + "\nІнтервал Дуннетта (одночасний); зелене — достовірно вище, червоне — нижче, сіре — не відрізняється",
                 loc="left", fontsize=10)
    fig.tight_layout()


def fig_caterpillar(fig, bt, trait_name):
    """Впорядковані BLUP із 95 % інтервалом (±1,96·SE) — «гусениця»."""
    ax = fig.add_subplot(111)
    order = np.argsort(-bt["blup"]); y = np.arange(len(order))[::-1]
    v = bt["blup"][order]; se = bt["se"][order]
    ax.errorbar(v, y, xerr=1.96 * se, fmt="o", ms=4, color="#1f77b4", ecolor="#9ecae1", elinewidth=1.4, capsize=0)
    ax.axvline(0, color="#333", lw=0.9)
    ax.set_yticks(y); ax.set_yticklabels([bt["levels"][i] for i in order], fontsize=7)
    ax.set_xlabel("BLUP (відхилення від середнього)")
    ax.set_title(f"BLUP генотипів: {trait_name}\n(інтервали ±1,96·SE; ті, що не перетинають нуль, достовірно відрізняються від середнього)",
                 loc="left", fontsize=10)
    fig.tight_layout()


def fig_score_dist(fig, names, classes, percents, title, order_note=""):
    """Горизонтальні стовпці розподілу балів (у %) для кожного генотипу; classes — спільний список балів."""
    ax = fig.add_subplot(111); y = np.arange(len(names))[::-1]; left = np.zeros(len(names))
    cmap = [(0.20 + 0.75 * i / max(len(classes) - 1, 1),) for i in range(len(classes))]
    for ci, c in enumerate(classes):
        vals = np.array([percents[g][ci] for g in range(len(names))])
        shade = 0.92 - 0.82 * ci / max(len(classes) - 1, 1)
        ax.barh(y, vals, left=left, color=(shade, 0.35 + 0.4 * (1 - shade), 1 - 0.6 * (1 - shade)), edgecolor="white", linewidth=0.5, label=f"{c:g}")
        left += vals
    ax.set_yticks(y); ax.set_yticklabels(names, fontsize=8); ax.set_xlim(0, 100); ax.set_xlabel("Частка рослин (ділянок), %")
    ax.legend(title="Бал", fontsize=7, ncol=min(len(classes), 9), loc="upper center", bbox_to_anchor=(0.5, -0.12), frameon=False)
    ax.set_title(title + (("\n" + order_note) if order_note else ""), loc="left", fontsize=10.5); fig.tight_layout()


def fig_curves(fig, times, curves, title, ylabel):
    """Середні криві розвитку хвороби за генотипами."""
    ax = fig.add_subplot(111)
    for k, (nm, ys) in enumerate(curves.items()):
        ax.plot(times, ys, marker="o", ms=4, lw=1.5, color=PALETTE[k % 10], label=nm)
    ax.set_xlabel("Час, діб від першого обліку"); ax.set_ylabel(ylabel); ax.set_title(title, loc="left", fontsize=10.5)
    ax.legend(fontsize=8, frameon=False, ncol=2 if len(curves) > 6 else 1); fig.tight_layout()


def fig_groups(fig, names, groups, ylabel, title):
    """Ящики з розмахом, окремі точки (з розсіюванням) і середні (ромб) – щоб розподіл у групах було видно, а не лише підсумок."""
    ax = fig.add_subplot(111)
    ax.boxplot(groups, positions=range(1, len(groups) + 1), widths=0.5, showfliers=False, medianprops=dict(color="#c62828", lw=1.6),
               boxprops=dict(color="#555"), whiskerprops=dict(color="#555"), capprops=dict(color="#555"))
    rg = np.random.default_rng(1)
    for i, g in enumerate(groups, 1):
        ax.scatter(i + rg.uniform(-0.14, 0.14, len(g)), g, s=16, color=PALETTE[(i - 1) % 10], alpha=0.65, zorder=3)
        ax.scatter([i], [np.mean(g)], marker="D", s=46, color="white", edgecolor="#222", zorder=4)
    ax.set_xticks(range(1, len(groups) + 1)); ax.set_xticklabels([f"{n}\n(n = {len(g)})" for n, g in zip(names, groups)], fontsize=9)
    ax.set_ylabel(ylabel); ax.set_title(title + "\nлінія – медіана, ромб – середнє, точки – окремі спостереження", loc="left", fontsize=10)
    fig.tight_layout()


def fig_stacked(fig, rows, cols, T, title):
    """Частки стовпців усередині кожного рядка таблиці (100 % стовпчики) із числом спостережень."""
    ax = fig.add_subplot(111); T = np.asarray(T, float); share = T / T.sum(1, keepdims=True); left = np.zeros(len(rows))
    for j, cn in enumerate(cols):
        ax.barh(range(len(rows)), 100 * share[:, j], left=left, color=PALETTE[j % 10], label=str(cn), edgecolor="white")
        for i in range(len(rows)):
            if share[i, j] > 0.06: ax.text(left[i] + 50 * share[i, j], i, f"{int(T[i, j])}", ha="center", va="center", fontsize=8, color="white")
        left += 100 * share[:, j]
    ax.set_yticks(range(len(rows))); ax.set_yticklabels([f"{r} (n = {int(T[i].sum())})" for i, r in enumerate(rows)]); ax.invert_yaxis()
    ax.set_xlabel("Частка в рядку, %"); ax.set_xlim(0, 100)
    ax.set_title(title, loc="left", fontsize=10.5); ax.legend(frameon=False, fontsize=8, bbox_to_anchor=(1.0, 1.0), loc="upper left"); fig.tight_layout()
