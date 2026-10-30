# sad_report.py – S.A.D. v2.1: універсальний експорт звіту (Word, Excel) із будь-якого вікна результатів
# -*- coding: utf-8 -*-
"""
Експортер обходить дерево віджетів готового вікна результатів і збирає:
  • заголовки й абзаци (Label, Text), розділи (LabelFrame, вкладки Notebook);
  • таблиці (ttk.Treeview) – з кольоровими заливками рядків (теги);
  • графіки (віджети, створені embed_figure, мають атрибут _sad_fig).
Панелі інструментів (кнопки, поля вводу) пропускаються. Із зібраного формується:
  • Word: таблиці з сіткою, заголовок таблиці, що повторюється на кожній сторінці, підписи «Таблиця N» / «Рисунок N», графіки;
  • Excel: кожна таблиця на власному аркуші (числа – числами, заливки збережено), текст – на аркуші «Звіт», графіки – на аркуші «Графіки».
"""
import os, re, io, datetime as _dt
import tkinter as tk
from tkinter import ttk, messagebox, filedialog
import tkinter.font as tkfont

SKIP_CLASSES = {"Button", "TButton", "Checkbutton", "TCheckbutton", "Radiobutton", "TRadiobutton", "Entry", "TEntry", "Spinbox", "TSpinbox",
                "Combobox", "TCombobox", "Scale", "TScale", "Menubutton", "TMenubutton", "Menu", "Scrollbar", "TScrollbar", "TSeparator", "Listbox",
                "Progressbar", "TProgressbar"}
INTERACTIVE = SKIP_CLASSES - {"TSeparator", "Menu", "Scrollbar", "TScrollbar"}


# ═══════════════════════════════════════════════════════════════
# 1. Збирання вмісту з вікна
# ═══════════════════════════════════════════════════════════════
def _has_content(w):
    if hasattr(w, "_sad_fig") or w.winfo_class() in ("Treeview", "Text"): return True
    return any(_has_content(c) for c in w.winfo_children())


def _count_interactive(w):
    return (1 if w.winfo_class() in INTERACTIVE else 0) + sum(_count_interactive(c) for c in w.winfo_children())


def _is_toolbar(w):
    """Панель інструментів чи навігації: серед прямих нащадків є кнопки/поля, або всередині кілька елементів керування, а таблиць, тексту
    й графіків усередині немає. (Бокова панель із кнопками й підписами – теж навігація, а не вміст звіту.)"""
    if _has_content(w): return False
    return any(c.winfo_class() in INTERACTIVE for c in w.winfo_children()) or _count_interactive(w) >= 2


def _label_text(w):
    try:
        var = str(w.cget("textvariable"))
        if var: return str(w.tk.globalgetvar(var))
    except Exception: pass
    try: return str(w.cget("text"))
    except Exception: return ""


def _font_info(w):
    try:
        f = tkfont.Font(font=w.cget("font")).actual(); return f.get("weight") == "bold", abs(int(f.get("size", 11)))
    except Exception: return False, 11


def _css(w, color):
    """Колір Tk (ім'я чи #hex) → 'RRGGBB' або None для порожнього/білого."""
    if not color: return None
    try: r, g, b = (v // 256 for v in w.winfo_rgb(color))
    except Exception: return None
    h = f"{r:02X}{g:02X}{b:02X}"
    return None if h in ("FFFFFF", "D9D9D9", "F0F0F0") else h


def _tree_table(tv):
    cols = list(tv["columns"]); disp = tv["displaycolumns"]
    if disp and disp != ("#all",) and disp != "#all": cols = [c for c in disp]
    headers = [str(tv.heading(c, "text")) for c in cols]
    widths = [max(30, int(tv.column(c, "width"))) for c in cols]
    rows, fills = [], []; tag_bg = {}
    def walk(item):
        for it in tv.get_children(item):
            vals = tv.item(it, "values"); allcols = list(tv["columns"])
            row = [str(vals[allcols.index(c)]) if allcols.index(c) < len(vals) else "" for c in cols]
            rows.append(row)
            f = None
            for tg in tv.item(it, "tags") or ():
                if tg not in tag_bg:
                    try: tag_bg[tg] = _css(tv, tv.tag_configure(tg, "background"))
                    except Exception: tag_bg[tg] = None
                f = f or tag_bg[tg]
            fills.append(f); walk(it)
    walk("")
    return headers, rows, fills, widths


def collect(root, title=None, roots=None):
    """→ список елементів: {'kind': 'heading'|'text'|'table'|'figure', ...}. Порядок – як у вікні (за створенням віджетів).
    roots – необов'язковий список (або функція, що його повертає) віджетів, з яких збирати вміст: коли результати лежать у тому самому
    вікні, що й введення даних (регресія, t-тест), експортується лише область результатів."""
    items = []; seen_fig = set()
    def add_text(txt, w):
        txt = txt.strip()
        if not txt: return
        bold, size = _font_info(w)
        if bold and len(txt) <= 140 and "\n" not in txt:
            items.append({"kind": "heading", "text": txt, "level": 2 if size >= 13 else 3})
        else:
            items.append({"kind": "text", "text": txt, "warn": _css(w, w.cget("fg")) in ("C62828", "B71C1C", "8C1A1A", "A01010", "D32F2F") if "fg" in w.keys() else False})
    def walk(w):
        cls = w.winfo_class()
        if getattr(w, "_sad_noexport", False) or cls in SKIP_CLASSES: return
        if isinstance(w, tk.Toplevel) and w is not root: return
        fig = getattr(w, "_sad_fig", None)
        if fig is not None:
            if id(fig) not in seen_fig: seen_fig.add(id(fig)); items.append({"kind": "figure", "fig": fig})
            return
        if cls == "Treeview":
            h, r, f, wd = _tree_table(w)
            if h and r: items.append({"kind": "table", "headers": h, "rows": r, "fills": f, "widths": wd})
            return
        if cls == "Text":
            t = w.get("1.0", "end").rstrip()
            if t.strip(): items.extend(_text_blocks(t))
            return
        if cls in ("Label", "TLabel"):
            if getattr(w, "_sad_image_only", False): return
            if w.master.winfo_class() == "Canvas" and not hasattr(w.master, "_sad_scroll_body"): return      # підписи осей усередині полотна графіка – оздоблення
            add_text(_label_text(w), w); return
        if cls == "TNotebook":
            for tab in w.tabs():
                try: label = w.tab(tab, "text")
                except Exception: label = ""
                if label.strip() and _has_content(w.nametowidget(tab)): items.append({"kind": "heading", "text": label.strip(), "level": 1})
                walk(w.nametowidget(tab))
            return
        if cls in ("Labelframe", "TLabelframe"):
            try:
                t = str(w.cget("text")).strip()
                if t and _has_content(w): items.append({"kind": "heading", "text": t, "level": 3})
            except Exception: pass
        if _is_toolbar(w): return
        for c in w.winfo_children(): walk(c)
    if callable(roots): roots = roots()
    if roots:
        for r_ in roots: walk(r_)
    else:
        for c in root.winfo_children(): walk(c)
    return _tidy(items)


def _text_blocks(t):
    """Вміст Text-віджета: абзаци; послідовність рядків зі стовпчиками (розділеними 2+ пробілами/табуляцією) → таблиця."""
    lines = ["" if (l.strip() and set(l.strip()) <= set("-=_─━ ")) else l for l in t.split("\n")]; out = []; i = 0      # лінії-розділювачі розділяють абзаци
    def cells(l): return [c for c in re.split(r"\t+|\s{2,}", l.strip()) if c != ""]
    while i < len(lines):
        if not lines[i].strip(): i += 1; continue
        c0 = cells(lines[i]); j = i
        if len(c0) >= 2:
            while j + 1 < len(lines) and len(cells(lines[j + 1])) == len(c0) and not set(lines[j + 1].strip()) <= set("-=_─ "): j += 1
        kv = sum(1 for l in lines[i:j + 1] for c in cells(l) if " = " in c or c.rstrip().endswith(":")) >= 0.5 * sum(len(cells(l)) for l in lines[i:j + 1])
        if len(c0) >= 2 and j - i >= 2 and not kv:                       # «ключ = значення» – це не таблиця
            rows = [cells(l) for l in lines[i:j + 1] if not set(l.strip()) <= set("-=_─ ")]
            out.append({"kind": "table", "headers": rows[0], "rows": rows[1:], "fills": [None] * (len(rows) - 1), "widths": [100] * len(c0)}); i = j + 1; continue
        k = i; para = []
        while k < len(lines) and lines[k].strip(): para.append(lines[k].strip()); k += 1
        txt = "\n".join(para)
        if not set(txt) <= set("-=_─ "): out.append({"kind": "text", "text": txt, "warn": False})
        i = k
    return out


_CAP = re.compile(r"^\s*(ТАБЛИЦЯ|Таблиця)\s*(\d+)\s*[.:–-]?\s*(.*)$", re.S)
_SPACED = re.compile(r"^(?:\S ){2,}\S(?:\s{2,}(?:\S ){1,}\S)*$")


def _despace(t):
    """«З В І Т   А Н А Л І З У» → «ЗВІТ АНАЛІЗУ»."""
    if not _SPACED.match(t.strip()): return t
    return " ".join(w.replace(" ", "") for w in re.split(r"\s{2,}", t.strip()))


def _merge_pairs(items):
    """Підряд ≥ 3 пар «жирна мітка із двокрапкою + значення» (типово: Модель: …, R²: …, F: …) → таблиця «Показник | Значення»."""
    def is_label(it): return it["kind"] == "heading" and it["level"] >= 3 and it["text"].rstrip().endswith(":") and len(it["text"]) <= 48
    def is_value(it): return it["kind"] == "text" and "\n" not in it["text"] and len(it["text"]) <= 140
    out = []; i = 0
    while i < len(items):
        j = i; pairs = []
        while j + 1 < len(items) and is_label(items[j]) and is_value(items[j + 1]): pairs.append([items[j]["text"].rstrip().rstrip(":"), items[j + 1]["text"]]); j += 2
        if len(pairs) >= 3:
            out.append({"kind": "table", "headers": ["Показник", "Значення"], "rows": pairs, "fills": [None] * len(pairs), "widths": [140, 260]}); i = j
        else:
            out.append(items[i]); i += 1
    return out


def _short_dash(t): return t.replace("\u2014", "\u2013")           # довге тире → коротке (правило оформлення)


def _tidy(items):
    """Прибирає дублікати підряд і порожні розділи; заголовки «ТАБЛИЦЯ N. …» перетворює на підписи таблиць; розряджені назви стискає;
    довге тире замінює коротким."""
    for it in items:
        if it["kind"] in ("heading", "text"): it["text"] = _short_dash(_despace(it["text"]))
        elif it["kind"] == "table":
            it["headers"] = [_short_dash(str(h)) for h in it["headers"]]; it["rows"] = [[_short_dash(str(v)) for v in r] for r in it["rows"]]
    out = []
    for it in items:
        if out and it["kind"] in ("heading", "text") and out[-1]["kind"] == it["kind"] and out[-1]["text"] == it["text"]: continue
        out.append(it)
    out = _merge_pairs(out)
    res = []
    for i, it in enumerate(out):
        if it["kind"] == "heading":
            nxt = next((o for o in out[i + 1:]), None)
            if nxt is None or (nxt["kind"] == "heading" and nxt["level"] <= it["level"]): continue
            m = _CAP.match(it["text"])
            if m and nxt["kind"] == "table":                      # власна нумерація програми → підпис таблиці (нумерація Word-звіту наскрізна)
                res.append({"kind": "caption", "text": m.group(3).strip().rstrip(".") or "Таблиця"}); continue
        res.append(it)
    return res


# ═══════════════════════════════════════════════════════════════
# 2. Допоміжне: числа, графіки
# ═══════════════════════════════════════════════════════════════
_NUM = re.compile(r"^[+-]?\d+(\.\d+)?([eE][+-]?\d+)?$")


def parse_number(s):
    """'0,0123' → 0.0123; '−1 234,5' → −1234.5; текст, «<0,0001», «–», коди з ведучими нулями лишаються текстом."""
    if not isinstance(s, str): return None
    t = s.strip().replace("\u00a0", "").replace(" ", "").replace("−", "-").replace("–", "-").replace(",", ".")
    if not t or t in ("-", "+") or not _NUM.match(t): return None
    if re.match(r"^[+-]?0\d", t) and "." not in t: return None
    try: v = float(t)
    except ValueError: return None
    return int(v) if re.fullmatch(r"[+-]?\d+", t) and abs(v) < 1e15 else v


def figure_png(fig, dpi=200):
    buf = io.BytesIO(); fig.savefig(buf, format="png", dpi=dpi, bbox_inches="tight", facecolor="white"); return buf.getvalue()


def figure_caption(fig):
    try:
        st = getattr(fig, "_suptitle", None)
        if st is not None and st.get_text().strip(): return st.get_text().strip().replace("\n", " ")
        for ax in fig.axes:
            t = ax.get_title().strip()
            if t: return t.replace("\n", " ")
    except Exception: pass
    return ""


def _stats(items):
    t = [i for i in items if i["kind"] == "table"]; return {"tables": len(t), "figures": sum(1 for i in items if i["kind"] == "figure"),
                                                          "max_cols": max((len(i["headers"]) for i in t), default=0)}


# ═══════════════════════════════════════════════════════════════
# 3. Word
# ═══════════════════════════════════════════════════════════════
def export_docx(path, items, title="Звіт", subtitle=""):
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    FONT = "Times New Roman"; st = _stats(items); landscape = st["max_cols"] > 9
    doc = Document(); n = doc.styles["Normal"]; n.font.name = FONT; n.font.size = Pt(11); n.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    for lvl, sz in ((1, 15), (2, 13), (3, 12)):
        h = doc.styles[f"Heading {lvl}"]; h.font.name = FONT; h.font.size = Pt(sz); h.font.bold = True; h.font.color.rgb = RGBColor(0, 0, 0); h.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        rf = h.element.rPr.rFonts; rf.set(qn("w:ascii"), FONT); rf.set(qn("w:hAnsi"), FONT); rf.set(qn("w:cs"), FONT)
        for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):          # теми шаблона Word перекривають явний шрифт
            if qn(a) in rf.attrib: del rf.attrib[qn(a)]
    sec = doc.sections[0]
    if landscape: sec.orientation = WD_ORIENT.LANDSCAPE; sec.page_width, sec.page_height = sec.page_height, sec.page_width
    sec.top_margin = sec.bottom_margin = Cm(2); sec.left_margin = Cm(2.2 if not landscape else 1.8); sec.right_margin = Cm(1.8 if not landscape else 1.8)
    avail = (sec.page_width - sec.left_margin - sec.right_margin) / 360000.0                      # доступна ширина, см
    fp = sec.footer.paragraphs[0]; fp.text = f"S.A.D. v2.1 · {_dt.datetime.now():%d.%m.%Y %H:%M}"; fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for r in fp.runs: r.font.size = Pt(8); r.font.name = FONT
    title = _short_dash(title); subtitle = _short_dash(subtitle)
    h0 = doc.add_heading(title, level=1); h0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if subtitle:
        p = doc.add_paragraph(subtitle); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs: r.font.size = Pt(10)
    def run_fmt(run, size=None, bold=None, color=None):
        run.font.name = FONT; run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        if size: run.font.size = Pt(size)
        if bold is not None: run.bold = bold
        if color: run.font.color.rgb = RGBColor.from_string(color)
    def shade(cell, hexc):
        tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hexc); tcPr.append(shd)
    def row_flag(row, tag):
        trPr = row._tr.get_or_add_trPr(); el = OxmlElement(tag); el.set(qn("w:val"), "true"); trPr.append(el)
    ctx = ""; tno = fno = 0
    for it in items:
        k = it["kind"]
        if k == "caption":
            ctx = it["text"]
        elif k == "heading":
            doc.add_heading(it["text"], level=min(3, max(1, it["level"]))); ctx = it["text"]
        elif k == "text":
            p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(3)
            for i, line in enumerate(it["text"].split("\n")):
                if i: p.add_run().add_break()
                run_fmt(p.add_run(line), 11, None, "C62828" if it.get("warn") else None)
        elif k == "table":
            tno += 1; hd, rows = it["headers"], it["rows"]; ncol = len(hd)
            cap = doc.add_paragraph(); cap.paragraph_format.keep_with_next = True; cap.paragraph_format.space_before = Pt(6); cap.paragraph_format.space_after = Pt(2)
            run_fmt(cap.add_run(f"Таблиця {tno}" + (f" – {ctx}" if ctx else "")), 10.5, True)
            tbl = doc.add_table(rows=1 + len(rows), cols=ncol); tbl.style = "Table Grid"; tbl.alignment = WD_TABLE_ALIGNMENT.CENTER; tbl.autofit = False
            fsz = 10 if ncol <= 7 else (9 if ncol <= 10 else (8 if ncol <= 14 else 7))
            wts = [max(30, w) for w in it.get("widths") or [100] * ncol]; tot = float(sum(wts)); widths = [max(1.2, avail * w / tot) for w in wts]
            if sum(widths) > avail: widths = [w * avail / sum(widths) for w in widths]
            text_col = [sum(len(str(r[j])) for r in rows[:30] if j < len(r)) / max(1, min(30, len(rows))) > 18 for j in range(ncol)]
            for j, h in enumerate(hd):
                c = tbl.rows[0].cells[j]; c.text = ""; p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER; run_fmt(p.add_run(str(h)), fsz, True); shade(c, "D9E2F3"); c.width = Cm(widths[j])
            row_flag(tbl.rows[0], "w:tblHeader")
            for i, row in enumerate(rows):
                fill = it["fills"][i] if i < len(it["fills"]) else None; tr = tbl.rows[i + 1]; row_flag(tr, "w:cantSplit")
                for j in range(ncol):
                    c = tr.cells[j]; c.text = ""; p = c.paragraphs[0]; v = str(row[j]) if j < len(row) and row[j] is not None else ""
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT if (j == 0 or text_col[j]) else WD_ALIGN_PARAGRAPH.CENTER; run_fmt(p.add_run(v), fsz); c.width = Cm(widths[j])
                    if fill: shade(c, fill)
            for r_ in tbl.rows:
                for c in r_.cells:
                    for p in c.paragraphs: p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(0)
            sp = doc.add_paragraph(); sp.paragraph_format.space_after = Pt(4)
        elif k == "figure":
            fno += 1
            try: png = figure_png(it["fig"], 200)
            except Exception as e:
                run_fmt(doc.add_paragraph().add_run(f"[Рисунок {fno}: не вдалося побудувати зображення – {e}]"), 10, None, "C62828"); continue
            from PIL import Image
            wpx, hpx = Image.open(io.BytesIO(png)).size; w_cm = min(avail, wpx / 200 * 2.54, avail); h_cm = w_cm * hpx / wpx
            max_h = (sec.page_height - sec.top_margin - sec.bottom_margin) / 360000.0 - 4.0       # місце під заголовок/підпис, щоб рисунок не «вилітав» на порожню сторінку
            if h_cm > max_h: w_cm *= max_h / h_cm
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.keep_with_next = True; p.add_run().add_picture(io.BytesIO(png), width=Cm(w_cm))
            cap = doc.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER; ct = figure_caption(it["fig"]) or ctx
            run_fmt(cap.add_run(f"Рисунок {fno}" + (f" – {ct}" if ct else "")), 10.5, True)
    doc.save(path); return _stats(items)


# ═══════════════════════════════════════════════════════════════
# 4. Excel
# ═══════════════════════════════════════════════════════════════
def export_xlsx(path, items, title="Звіт", subtitle=""):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.drawing.image import Image as XImage
    from openpyxl.utils import get_column_letter
    title = _short_dash(title); subtitle = _short_dash(subtitle)
    wb = Workbook(); ws = wb.active; ws.title = "Звіт"; thin = Side(style="thin", color="999999")
    ws.append([title]); ws.cell(1, 1).font = Font(bold=True, size=14)
    if subtitle: ws.append([subtitle])
    ws.append([f"S.A.D. v2.1 · {_dt.datetime.now():%d.%m.%Y %H:%M}"]); ws.append([])
    names = set(["Звіт"]); tno = 0; ctx = ""; figs = []
    def sheet_name(base):
        base = re.sub(r"[\\/*?:\[\]]", " ", base).strip()[:28] or "Таблиця"; nm = base; k = 2
        while nm in names: nm = f"{base[:26]} {k}"; k += 1
        names.add(nm); return nm
    for it in items:
        k = it["kind"]
        if k == "caption":
            ctx = it["text"]
        elif k == "heading":
            ws.append([it["text"]]); ws.cell(ws.max_row, 1).font = Font(bold=True, size=13 if it["level"] <= 2 else 12); ctx = it["text"]
        elif k == "text":
            ws.append([it["text"]]); c = ws.cell(ws.max_row, 1); c.alignment = Alignment(wrap_text=True, vertical="top")
            if it.get("warn"): c.font = Font(color="C62828")
        elif k == "table":
            tno += 1; nm = sheet_name(f"Т{tno} {ctx}"); wt = wb.create_sheet(nm); wt.append(it["headers"])
            ws.append([f"Таблиця {tno}" + (f" – {ctx}" if ctx else "") + f"  →  аркуш «{nm}»"]); ws.cell(ws.max_row, 1).font = Font(italic=True)
            for c in wt[1]: c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="D9E2F3"); c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            for i, row in enumerate(it["rows"]):
                vals = []
                for v in row:
                    num = parse_number(v); vals.append(num if num is not None else v)
                wt.append(vals); fill = it["fills"][i] if i < len(it["fills"]) else None
                for j, c in enumerate(wt[wt.max_row]):
                    c.border = Border(left=thin, right=thin, top=thin, bottom=thin); c.alignment = Alignment(horizontal="left" if j == 0 or isinstance(c.value, str) and len(c.value) > 18 else "center")
                    if fill: c.fill = PatternFill("solid", fgColor=fill)
                    if isinstance(c.value, float): c.number_format = "0.0000" if abs(c.value) < 1 and c.value != 0 else "0.00"
            for j, h in enumerate(it["headers"], 1):
                mx = max([len(str(h))] + [len(str(r[j - 1])) for r in it["rows"] if j - 1 < len(r)]); wt.column_dimensions[get_column_letter(j)].width = min(48, max(9, mx + 2))
            wt.freeze_panes = "A2"
        elif k == "figure":
            figs.append(it)
    ws.column_dimensions["A"].width = 110
    if figs:
        wf = wb.create_sheet("Графіки"); row = 1
        for i, it in enumerate(figs, 1):
            try: png = figure_png(it["fig"], 130)
            except Exception: continue
            wf.cell(row, 1, f"Рисунок {i}" + (f" – {figure_caption(it['fig'])}" if figure_caption(it["fig"]) else "")).font = Font(bold=True)
            img = XImage(io.BytesIO(png)); sc = min(1.0, 900 / img.width); img.width, img.height = int(img.width * sc), int(img.height * sc); wf.add_image(img, f"A{row + 1}"); row += int(img.height / 20) + 4
    wb.save(path); return _stats(items)


# ═══════════════════════════════════════════════════════════════
# 5. Кнопки експорту у вікні результатів
# ═══════════════════════════════════════════════════════════════
def _safe_name(s):
    return re.sub(r"[^\w\-. ]+", "_", s, flags=re.U).strip()[:60] or "звіт"


def _nav_buttons(win):
    """Кнопки бокового меню панелей (темна бічна панель: плоскі кнопки з текстом ліворуч, активна – червона)."""
    out = []
    def walk(w):
        for c in w.winfo_children():
            if isinstance(c, tk.Toplevel): continue
            if c.winfo_class() == "Button":
                try:
                    if str(c.cget("anchor")) == "w" and str(c.cget("relief")) == "flat" and str(c.cget("activebackground")).lower() == "#c62828": out.append(c)
                except tk.TclError: pass
            walk(c)
    walk(win); return out


def visit_panels(win):
    """Деякі вікна будують панелі (графіки) лише при першому відкритті. Відкриваємо кожну панель бокового меню один раз, як це зробив би
    користувач, і повертаємо початкову – тоді експорт бачить усе."""
    btns = _nav_buttons(win)
    if len(btns) < 2: return 0
    active = next((b for b in btns if str(b.cget("bg")).lower() == "#c62828"), btns[0])
    for b in btns:
        if b is not active:
            try: b.invoke(); win.update_idletasks()
            except Exception: pass
    try: active.invoke(); win.update_idletasks()
    except Exception: pass
    return len(btns)


def export_window(win, kind, title=None, path=None, ask_open=True, roots=None):
    """kind: 'docx' | 'xlsx'. Збирає вміст вікна й зберігає. path=None – запитати у користувача."""
    title = title or win.title()
    if roots is None: visit_panels(win)
    items = collect(win, roots=roots)
    st = _stats(items)
    if not items:
        messagebox.showinfo("Звіт", "У цьому вікні немає результатів для експорту.", parent=win); return None
    ext = ".docx" if kind == "docx" else ".xlsx"
    if path is None:
        path = filedialog.asksaveasfilename(parent=win, defaultextension=ext, initialfile=_safe_name(title) + ext,
                                            filetypes=[("Word", "*.docx")] if kind == "docx" else [("Excel", "*.xlsx")], title="Зберегти звіт")
        if not path: return None
    try:
        (export_docx if kind == "docx" else export_xlsx)(path, items, title, f"Створено {_dt.datetime.now():%d.%m.%Y %H:%M}")
    except Exception as e:
        messagebox.showerror("Звіт", f"Не вдалося зберегти: {e}", parent=win); return None
    msg = f"Збережено: {path}\nТаблиць: {st['tables']}, графіків: {st['figures']}."
    if ask_open and messagebox.askyesno("Звіт збережено", msg + "\n\nВідкрити файл?", parent=win):
        try: os.startfile(path)                                                      # noqa (Windows)
        except Exception:
            import subprocess, sys; subprocess.Popen(["open" if sys.platform == "darwin" else "xdg-open", path])
    return path


def attach_export_bar(win, title=None, text="Звіт повністю (таблиці з сіткою, графіки):", roots=None):
    """Додає у верхню частину вікна смужку «📄 Word · 📊 Excel». Повертає смужку (або None, якщо вже є)."""
    if getattr(win, "_sad_export_bar", None) is not None: return None
    bar = tk.Frame(win, bg="#eef3f8", padx=6, pady=3); bar._sad_noexport = True
    tk.Label(bar, text=text, bg="#eef3f8", fg="#1a4b8c", font=("Times New Roman", 10, "bold")).pack(side=tk.LEFT, padx=(2, 8))
    tk.Button(bar, text="📄 Word", font=("Times New Roman", 10, "bold"), bg="#1a4b8c", fg="white", relief=tk.FLAT, padx=10, pady=2, cursor="hand2",
              command=lambda: export_window(win, "docx", title, roots=roots)).pack(side=tk.LEFT, padx=3)
    tk.Button(bar, text="📊 Excel", font=("Times New Roman", 10, "bold"), bg="#1a6b3a", fg="white", relief=tk.FLAT, padx=10, pady=2, cursor="hand2",
              command=lambda: export_window(win, "xlsx", title, roots=roots)).pack(side=tk.LEFT, padx=3)
    try:
        first = win.pack_slaves()[0] if win.pack_slaves() else None
        if win.grid_slaves(): bar.place(relx=1.0, rely=0.0, anchor="ne")                  # вікно на grid: смужка поверх правого верхнього кута
        elif first is not None: bar.pack(side=tk.TOP, fill=tk.X, before=first)
        else: bar.pack(side=tk.TOP, fill=tk.X)
    except tk.TclError: bar.place(relx=1.0, rely=0.0, anchor="ne")
    win._sad_export_bar = bar
    return bar
