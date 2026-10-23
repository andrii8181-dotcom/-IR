# test_sad_gencorr_sections.py — розділи результатів генотипних кореляцій, логіка висновків, експорт
import os, sys, math, tempfile, shutil
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import numpy as np
import sad_gencorr_core as GC, sad_gencorr_sections as GS, sad_breeding_ui as U
from sad_gxe_core import BreedingError
FAIL = []
def check(c, msg):
    print(("  [OK]   " if c else "  [FAIL] ") + msg, flush=True)
    if not c: FAIL.append(msg)
def fmt(v): return f"{v:.4f}".replace(".", ",")
def make_rows(g, r, rG, h2m, seed, k=3):
    rg = np.random.default_rng(seed)
    sG = np.full((k, k), rG) + np.eye(k) * (1 - rG); vE = (1 - h2m) / h2m * r
    ge = rg.multivariate_normal(np.zeros(k), sG, size=g); rows = []
    for i in range(g):
        for j in range(r):
            y = 10 + ge[i] + rg.normal(0, math.sqrt(vE), k); rows.append([f"G{i + 1}", f"R{j + 1}"] + [fmt(v) for v in y] + [""] * (4 - k))
    return rows
HEAD = ["Генотип", "Повторність", "Висота", "Маса", "Урожай", "Порожня"]
def struct_ok(S):
    return not [s["title"] for s in S if s["kind"] == "table" and (any(len(x) != len(s["headers"]) for x in s["rows"]) or len(s["nd"]) != len(s["headers"]))]

print("== 1. Розділи на змодельованих даних (g = 60, сильна rG = 0,9, h² = 0,9)")
rows = make_rows(60, 3, 0.9, 0.9, 11)
S, warns, res = GS.sections_gencorr(rows, HEAD, 0.05, 500)
check([s["key"] for s in S] == ["pairs", "matrix", "comp", "heat", "bars", "help"], "розділи: пари, матриця, компоненти, два графіки, пояснення")
check(struct_ok(S), "структура таблиць коректна")
from matplotlib.figure import Figure
import matplotlib
bad = []
for s in S:
    if s["kind"] == "figure":
        f = Figure(figsize=s["size"])
        with matplotlib.rc_context(U._RC): m_ = s["build"](f)
        if m_: bad.append(m_)
check(not bad, "обидва графіки будуються")
pt = S[0]; vi = pt["headers"].index("Висновок про rG")
check(len(pt["rows"]) == 3 and all(r_[vi] == "істотна додатна" for r_ in pt["rows"]) and pt["tags"] == ["good"] * 3,
      "сильна rG = 0,9 при g = 60: усі три пари «істотна додатна», зелені")
check(warns == [], "для g = 60 попереджень немає")
mt = S[1]; check(isinstance(mt["rows"][0][1], tuple) and mt["rows"][0][1][0] == res["h2"][0] and mt["rows"][0][2][1] in ("*", "**"),
                 "матриця: на діагоналі h², над діагоналлю rG із зірочкою за інтервалом")
check(isinstance(mt["rows"][1][1], tuple) and mt["rows"][1][1][0] == res["rP_mean"][1, 0], "під діагоналлю — rP на середніх")
cp = S[2]; check(len(cp["rows"]) == 3 and all(r_[5] > 0.5 for r_ in cp["rows"]), "компоненти: h² > 0,5 за побудовою")

print("\n== 2. Логіка висновків на заданих результатах (детерміновано)")
def fake(pairs, g=40):
    k = 3; z = np.zeros((k, k))
    return {"g": g, "r": 3, "k": k, "traits": ["A", "B", "C"], "gen_names": [f"G{i}" for i in range(g)], "G": np.eye(k), "E": np.eye(k), "P_plot": np.eye(k) * 2,
            "P_mean": np.eye(k), "rG": z, "rP": z, "rP_mean": z, "rE": z, "h2": np.array([0.5, 0.5, 0.5]), "pairs": pairs, "alpha": 0.05, "n_boot": 100,
            "means": np.zeros((g, k)), "warnings": []}
def pair(i, j, rG, lo, hi, oor=False, bv=100):
    return {"i": i, "j": j, "a": "ABC"[i], "b": "ABC"[j], "rG": rG, "rG_lo": lo, "rG_hi": hi, "rG_lo99": lo - .1 if not math.isnan(lo) else lo, "rG_hi99": hi + .1 if not math.isnan(hi) else hi,
            "rG_p_trad": 0.01, "rP": 0.3, "rP_mean": 0.3, "rP_mean_p": 0.04, "rE": 0.1, "boot_valid": bv, "out_of_range": oor}
real = GC.genotypic_correlations
try:
    GC.genotypic_correlations = lambda *a, **k: fake([pair(0, 1, 0.6, 0.3, 0.8), pair(0, 2, -0.5, -0.8, -0.1), pair(1, 2, 0.2, -0.3, 0.6),
                                                      pair(0, 1 + 0, float("nan"), float("nan"), float("nan")), pair(1, 2, 0.4, float("nan"), float("nan")),
                                                      pair(0, 2, 1.2, 0.8, 1.0, oor=True)])
    S2, w2, _ = GS.sections_gencorr(rows, HEAD, 0.05, 100)
finally:
    GC.genotypic_correlations = real
vd = {(r_[0], r_[1], round(r_[2], 3) if r_[2] == r_[2] else "nan"): (r_[vi], t_) for r_, t_ in zip(S2[0]["rows"], S2[0]["tags"])}
check(vd[("A", "B", 0.6)] == ("істотна додатна", "good"), "інтервал [0,3; 0,8] без нуля, rG > 0 → «істотна додатна», зелене")
check(vd[("A", "C", -0.5)] == ("істотна від'ємна", "bad"), "інтервал [−0,8; −0,1] → «істотна від'ємна», червоне")
check(vd[("B", "C", 0.2)] == ("не істотна", None), "інтервал [−0,3; 0,6] містить 0 → «не істотна»")
check(vd[("A", "B", "nan")][0].startswith("не визначено") and vd[("A", "B", "nan")][1] == "bad", "rG = NaN → «не визначено», червоне")
check(vd[("B", "C", 0.4)] == ("інтервал ненадійний", "warn"), "немає валідного інтервалу → «інтервал ненадійний», жовте")
oor = [v for k_, v in vd.items() if k_[2] == 1.2][0]; check("|rG| > 1" in oor[0] and oor[1] == "warn", "|rG| > 1 → позначено як похибку оцінки, жовте")

print("\n== 3. Помилки введення")
for fn, frag in [(lambda: GS.sections_gencorr([], HEAD), "порожня"),
                 (lambda: GS.sections_gencorr([[""] + r_[1:] for r_ in rows], HEAD), "назви генотипу"),
                 (lambda: GS.sections_gencorr([r_[:1] + [""] + r_[2:] for r_ in rows], HEAD), "номера повторності"),
                 (lambda: GS.sections_gencorr([r_[:3] + [""] * 4 for r_ in rows], HEAD), "щонайменше 2"),
                 (lambda: GS.sections_gencorr([r_[:2] + ["abc"] + r_[3:] for r_ in rows], HEAD), "не число"),
                 (lambda: GS.sections_gencorr(rows[:-1], HEAD), "збалансованим"),
                 (lambda: GS.sections_gencorr([r_[:3] + [""] + r_[4:] if i == 4 else r_ for i, r_ in enumerate(rows)], HEAD), "порожні клітинки")]:
    try: fn(); check(False, f"«{frag}» мало дати помилку")
    except BreedingError as ex: check(frag in str(ex), f"помилка «{frag}»")

print("\n== 4. Експорт (зірочки поруч із числами)")
import openpyxl, docx
tmp = tempfile.mkdtemp(); U.export_xlsx(os.path.join(tmp, "g.xlsx"), S, "Ген. кореляції"); wb = openpyxl.load_workbook(os.path.join(tmp, "g.xlsx"))
ws = next(w for w in wb.worksheets if w.title.startswith("Матриця")); cell = ws["C2"].value
check(isinstance(cell, str) and cell.endswith("*") and "," in cell, f"Excel: rG із зірочкою як текст «{cell}», h² на діагоналі також текст «{ws['B2'].value}»")
U.export_docx(os.path.join(tmp, "g.docx"), "Ген. кореляції", S, True); d = docx.Document(os.path.join(tmp, "g.docx"))
check(len(d.tables) == 3 and d.tables[1].rows[1].cells[2].text.endswith("*"), "Word: 3 таблиці; у матриці rG із зірочкою")
shutil.rmtree(tmp, ignore_errors=True)
print("\n" + "=" * 60)
if FAIL:
    print(f"ПРОВАЛІВ: {len(FAIL)}"); [print("  -", f) for f in FAIL]; sys.exit(1)
print("УСІ ПЕРЕВІРКИ РОЗДІЛІВ ГЕНОТИПНИХ КОРЕЛЯЦІЙ ПРОЙДЕНО")
