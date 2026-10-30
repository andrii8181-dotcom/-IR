# export_probe.py – запуск аналізу в одному модулі й перевірка експорту його звіту (використовується тестом test_sad_modules_export_gui.py)
import os, sys, random, time, json, collections, tempfile
sys.path.insert(0, os.environ.get("SAD_SRC", os.path.dirname(os.path.abspath(__file__))))
import tkinter as tk
from tkinter import messagebox
import numpy as np
import matplotlib; matplotlib.use("Agg")
NAME = sys.argv[1]; OUT = sys.argv[2]
msgs = []
for n in ("showinfo", "showwarning", "showerror"): setattr(messagebox, n, lambda t="", m="", _n=n, **k: msgs.append((_n, str(t), str(m)[:140])))
messagebox.askyesno = lambda *a, **k: True; messagebox.askokcancel = lambda *a, **k: True
root = tk.Tk(); root.geometry("200x100+0+0")
def all_w(x):
    out = [x]
    for c in x.winfo_children(): out += all_w(c)
    return out
PICK = ("Виконати", "Запустити", "Аналіз", "OK", "Ок", "Застосувати", "Продовжити", "Так", "▶", "Обрати", "Далі")
def auto_wait(self, window):
    btns = [b for b in all_w(window) if b.winfo_class() == "Button" and str(b.cget("state")) != "disabled"]
    pick = next((b for p in PICK for b in btns if p in str(b.cget("text"))), None)
    try:
        if pick is not None: pick.invoke()
        else: window.destroy()
    except tk.TclError: pass
tk.Misc.wait_window = auto_wait
import sad_report as R
rg = random.Random(7)
def fill(w, rows):
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            if j < len(w.entries[i]): w.entries[i][j].insert(0, str(v).replace(".", ",") if not isinstance(v, str) else v)
def num(m, s=1.0): return f"{rg.gauss(m, s):.2f}"
before = set()
def new_results_window(match=None):
    for c in root.winfo_children():
        if isinstance(c, tk.Toplevel) and c not in before:
            for d in c.winfo_children():
                pass
    return None
def result_windows():
    out = []
    def walk(w):
        for c in w.winfo_children():
            if isinstance(c, tk.Toplevel): out.append(c); walk(c)
    walk(root); return out
target = None; roots = None
try:
    if NAME == "desc":
        import sad_descriptive_ttest as M; w = M.DescriptiveWindow(root, {}); root.update()
        for j, nm in enumerate(("Урожайність", "Маса плода", "Вміст цукрів")): w.header_vars[j].set(nm)
        fill(w, [[num(25, 2), num(180, 15), num(12, 1)] for _ in range(12)]); w._analyze()
    elif NAME == "ttest":
        import sad_descriptive_ttest as M; w = M.TTestWindow(root); root.update()
        w.e1.insert("1.0", "\n".join(num(20, 1.5).replace(".", ",") for _ in range(10))); w.e2.insert("1.0", "\n".join(num(23, 1.5).replace(".", ",") for _ in range(10))); w._run(); target = w.win; roots = [w.res_txt]
    elif NAME == "reg":
        import sad_regression as M; w = M.RegressionWindow(root, {}); root.update()
        xs = list(range(1, 15)); w.tx.insert("1.0", "\n".join(str(x) for x in xs)); w.ty.insert("1.0", "\n".join(f"{2.5 * x + 3 + rg.gauss(0, 1.2):.2f}".replace(".", ",") for x in xs)); w._run(); target = w.win; roots = [w.res_frame]
    elif NAME == "rm":
        import sad_repeated as M; w = M.RepeatedMeasuresWindow(root, {}); root.update()
        fill(w, [[f"Рослина {i + 1}"] + [num(10 + t * 2 + i * 0.1, 0.8) for t in range(4)] for i in range(10)]); w._run()
    elif NAME == "mixed":
        import sad_repeated as M; w = M.MixedRepeatedWindow(root, {}); root.update()
        fill(w, [[("Контроль" if i < 8 else "Обробка")] + [num(10 + t * 2 + (i >= 8) * 1.5, 0.8) for t in range(4)] for i in range(16)]); w._run()
    elif NAME == "manova":
        import sad_manova as M; w = M.ManovaWindow(root, {}); root.update()
        fill(w, [[["Контроль", "Доза 1", "Доза 2"][i // 8]] + [num(20 + (i // 8) * 2, 1.2), num(50 + (i // 8) * 3, 2), num(8 - (i // 8) * 0.5, 0.6)] for i in range(24)]); w._run()
    elif NAME == "cluster":
        import sad_cluster_pca as M; w = M.ClusterWindow(root, {}); root.update()
        fill(w, [[f"Сорт {i + 1}"] + [num(10 + (i % 3) * 5, 1), num(100 + (i % 3) * 20, 4), num(5 - (i % 3), 0.5), num(30 + (i % 3) * 6, 2)] for i in range(12)]); w._run()
    elif NAME == "pca":
        import sad_cluster_pca as M; w = M.PCAWindow(root, {}); root.update()
        fill(w, [[f"Сорт {i + 1}"] + [num(10 + (i % 3) * 5, 1), num(100 + (i % 3) * 20, 4), num(5 - (i % 3), 0.5), num(30 + (i % 3) * 6, 2)] for i in range(12)]); w._run()
    elif NAME == "stab":
        import sad_stability_ancova as M; w = M.StabilityWindow(root, {}); root.update()
        fill(w, [[f"Генотип {g + 1}"] + [num(40 + g * 2 + e * (g % 3), 2) for e in range(5)] for g in range(8)]); w._run()
    elif NAME == "ancova":
        import sad_stability_ancova as M; w = M.AncovaWindow(root, {}); root.update()
        rows = []
        for i in range(24):
            g = i // 8; x = rg.gauss(10, 2); rows.append([["Контроль", "Доза 1", "Доза 2"][g], f"{x:.2f}", f"{3 + 1.5 * x + g * 2 + rg.gauss(0, 1):.2f}"])
        fill(w, rows); w._run()
    elif NAME == "corr":
        import sad_correlation as M; w = M.CorrelationWindow(root, {}); root.update()
        base = [rg.gauss(0, 1) for _ in range(14)]
        fill(w, [[f"{20 + b * 2 + rg.gauss(0, 0.6):.2f}", f"{100 + b * 8 + rg.gauss(0, 3):.2f}", f"{50 - b * 3 + rg.gauss(0, 2):.2f}", f"{rg.gauss(10, 2):.2f}"] for b in base]); w._run_analysis()
    elif NAME == "anova":
        import sad_app as A
        app = A.SADTk(root); app.open_table(2); root.update(); sorts = ["Голден", "Гала", "Фуджі"]; regs = ["Контроль", "Регалис"]
        for i, (sa, rb) in enumerate((a, b) for a in sorts for b in regs):
            row = app.entries[i]
            for k, v in enumerate((sa, rb)): row[k].insert(0, v)
            for c in range(2, 6): row[c].insert(0, f"{20 + sorts.index(sa) * 2 + regs.index(rb) * 1.5 + rg.gauss(0, 0.8):.2f}".replace(".", ","))
        app.ask_params = lambda: {"ok": True, "indicator": "Урожайність", "units": "т/га", "design": "rcbd", "split_main": "A", "ss_type": "III"}
        app.choose_method = lambda *a, **k: {"ok": True, "method": "lsd"}
        app.analyze(); target = app.report_win
    else: raise SystemExit("невідомий модуль")
    root.update(); time.sleep(0.4); root.update(); root.update_idletasks(); root.update()
    if target is None:
        wins = [t for t in result_windows() if "езультат" in t.title() or "езультат" in t.title().lower()]
        target = wins[-1] if wins else None
    if os.environ.get("PROBE_FIND") and target is not None:
        def _path(w):
            o = []
            while w is not None and w is not root: o.append(w.winfo_class()); w = w.master
            return " < ".join(o[:8])
        for w_ in all_w(target):
            if w_.winfo_class() in ("Label", "TLabel") and os.environ["PROBE_FIND"] in str(w_.cget("text")): print("FIND:", repr(str(w_.cget("text"))[:50]), "| шрифт:", w_.cget("font"), "|", _path(w_), "| видимий:", w_.winfo_ismapped(), flush=True)
    res = {"name": NAME, "found_window": bool(target), "title": target.title() if target else None, "errors": [m for m in msgs if m[0] == "showerror"], "warnings": [m for m in msgs if m[0] == "showwarning"][:3]}
    if target is not None:
        bar = getattr(target, "_sad_export_bar", None); res["bar"] = bar is not None
        if roots is None: res['panels'] = R.visit_panels(target)
        items = R.collect(target, roots=roots); st = R._stats(items); res.update(st); res["kinds"] = dict(collections.Counter(i["kind"] for i in items))
        res["junk"] = [i.get("text", "")[:40] for i in items if i["kind"] in ("heading", "text") and any(k in i.get("text", "") for k in ("Копіювати", "Зберегти PNG", "Налаштування", "Довідка"))]
        res["headings"] = [i["text"][:50] for i in items if i["kind"] in ("heading", "caption")][:14]
        res["tables"] = [(i["headers"][:5], len(i["rows"])) for i in items if i["kind"] == "table"][:12]
        import docx, openpyxl
        dp = os.path.join(OUT, NAME + ".docx"); xp = os.path.join(OUT, NAME + ".xlsx")
        R.export_docx(dp, items, target.title(), "тест"); R.export_xlsx(xp, items, target.title(), "тест")
        d = docx.Document(dp); res["docx"] = {"tables": len(d.tables), "images": len(d.inline_shapes), "grid": all(t.style.name == "Table Grid" for t in d.tables)}
        res["xlsx"] = {"sheets": len(openpyxl.load_workbook(xp).sheetnames)}
    print("RESULT " + json.dumps(res, ensure_ascii=False))
except SystemExit as e: print("RESULT " + json.dumps({"name": NAME, "fatal": str(e)}, ensure_ascii=False))
except Exception as e:
    import traceback; print("RESULT " + json.dumps({"name": NAME, "exception": traceback.format_exc()[-700:], "messages": msgs[:4]}, ensure_ascii=False))
