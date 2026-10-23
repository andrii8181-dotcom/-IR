# test_sad_selection_sections.py — розділи індексів добору
import os, sys, math, tempfile, shutil
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import numpy as np
import sad_selection_core as SC, sad_selection_sections as SS, sad_breeding_ui as U
from sad_gxe_core import BreedingError
FAIL = []
def check(c, msg):
    print(("  [OK]   " if c else "  [FAIL] ") + msg, flush=True)
    if not c: FAIL.append(msg)
def structure_ok(S, label):
    bad = [s["title"] for s in S if s["kind"] == "table" and (len(s["nd"]) != len(s["headers"]) or any(len(r) != len(s["headers"]) for r in s["rows"])
           or (s.get("tags") is not None and len(s["tags"]) != len(s["rows"])))]
    check(not bad, f"{label}: структура таблиць коректна {bad[:2]}")
def figs_ok(S, label):
    from matplotlib.figure import Figure; import matplotlib
    bad = []
    for s in S:
        if s["kind"] == "figure":
            fg = Figure(figsize=s["size"])
            with matplotlib.rc_context(U._RC): m = s["build"](fg)
            if m: bad.append(s["title"])
    check(not bad, f"{label}: графіки будуються {bad[:2]}")

rg = np.random.default_rng(15); g, r = 36, 4
Gt = np.array([[5.0, -2.5, 1.0], [-2.5, 4.0, 0.0], [1.0, 0.0, 3.0]]); Et = np.array([[7.0, 0.5, 0.0], [0.5, 5.0, 0.3], [0.0, 0.3, 6.0]])
gv = rg.multivariate_normal(np.zeros(3), Gt, g); rb = rg.normal(0, 1, (r, 3))
rows = []
for i in range(g):
    for j in range(r):
        y = 50 + gv[i] + rb[j] + rg.multivariate_normal(np.zeros(3), Et)
        rows.append([f"G{i+1}", f"R{j+1}"] + [f"{v:.3f}".replace(".", ",") for v in y])
heads = ["Генотип", "Повторність", "Урожайність", "Білок", "Стійкість"]

print("== 1. Основний сценарій: відносні ваги, три індекси")
S, w = SS.sections_selection(rows, heads, "1; 1; 0.5", "relative", "2; -1; 0", 10)
keys = [s["key"] for s in S]
check(keys == ["par", "rg", "rp", "resp", "coef", "rank", "fig", "help"], f"розділи: {keys}")
structure_ok(S, "індекси добору"); figs_ok(S, "індекси добору")
resp = next(s for s in S if s["key"] == "resp"); names = [r_[0] for r_ in resp["rows"]]
check(names[:3] == ["Сміт–Хейзел", "Базовий", "Песек–Бейкер"] and len(names) == 3 + 3, f"методи: {names}")
ih = resp["headers"].index("ΔH")
check(resp["rows"][0][ih] >= resp["rows"][1][ih] - 1e-9 and all(resp["rows"][0][ih] >= r_[ih] - 1e-9 for r_ in resp["rows"][3:]), "ΔH(Сміт–Хейзел) ≥ ΔH(базовий) і ≥ добору за однією ознакою")
check(abs(resp["rows"][0][-1] - 100.0) < 1e-9, "ΔH Сміта–Хейзела = 100 % від себе")
pb = resp["rows"][2]; dg = np.array(pb[1:4]); check(np.allclose(dg / dg[0], np.array([2, -1, 0]) / 2, atol=1e-6), "Песек–Бейкер: прирости пропорційні бажаним (2; −1; 0)")
co = next(s for s in S if s["key"] == "coef"); mt = SC.multitrait_rcbd([r_[0] for r_ in rows], [r_[1] for r_ in rows],
        np.array([[float(x.replace(",", ".")) for x in r_[2:]] for r_ in rows]), ["a", "b", "c"])
sgv = np.sqrt(np.diag(mt["G"])); a_tab = np.array([r_[1] for r_ in co["rows"]])
check(np.allclose(a_tab, np.array([1, 1, 0.5]) / sgv), "відносні ваги: a = w/σg (ознака вносить пропорційно власній генетичній мінливості)")
rk = next(s for s in S if s["key"] == "rank"); sc_col = [r_[2] for r_ in rk["rows"]]
check(all(sc_col[i] >= sc_col[i + 1] for i in range(len(sc_col) - 1)) and rk["rows"][0][0] == 1, "рейтинг відсортовано за індексом Сміта–Хейзела")
check(sum(1 for t in rk["tags"] if t == "good") == round(0.10 * g), f"відібрано {sum(1 for t in rk['tags'] if t == 'good')} із {g} (10 %)")
par = next(s for s in S if s["key"] == "par")
check(len(par["rows"]) == 3 and all(0 < r_[4] < 1.0001 for r_ in par["rows"]), "H² кожної ознаки в (0; 1]")
rgm = next(s for s in S if s["key"] == "rg")
check(abs(rgm["rows"][0][1] - 1.0) < 1e-9 and rgm["rows"][0][2] < 0, f"генетична кореляція урожайність × білок від'ємна ({rgm['rows'][0][2]:.2f}) — типова селекційна дилема")
# дилема: добір лише за урожайністю знижує другу ознаку, індекс — менше
sy = resp["rows"][3]; check(sy[2] < resp["rows"][0][2] + 1e-9, "відбір лише за урожайністю дає гірший приріст білка, ніж індекс")

print("\n== 2. Економічні ваги й попередження")
S2, w2 = SS.sections_selection(rows, heads, "3; 2; 0", "economic", "", 5)
check([s["key"] for s in S2][:4] == ["par", "rg", "rp", "resp"] and len(next(s for s in S2 if s["key"] == "resp")["rows"]) == 2 + 3, "без бажаних приростів Песека–Бейкера немає")
check(np.allclose([r_[1] for r_ in next(s for s in S2 if s["key"] == "coef")["rows"]], [3, 2, 0]), "економічні ваги використано як задано")
small = [r_ for r_ in rows if int(r_[0][1:]) <= 12]; Sm, wm = SS.sections_selection(small, heads, "1;1;1", "relative", "", 10)
check(any("Лише 12 генотипів" in x for x in wm), "мало генотипів → попередження про нестійкість оцінок G")

print("\n== 3. Помилки")
noise = [r_[:2] + [r_[2]] + [f"{rg.normal(10, 2):.3f}", f"{rg.normal(5, 1):.3f}"] for r_ in rows]
for kw, frag in [(dict(rows=rows, headers=heads, weights_text="1; 1", weights_mode="relative"), "вкажіть 3 значень"),
                 (dict(rows=rows, headers=heads, weights_text="0; 0; 0", weights_mode="economic"), "нульові"),
                 (dict(rows=rows[:-1], headers=heads, weights_text="1;1;1", weights_mode="relative"), "збалансованим"),
                 (dict(rows=[r_[:2] + ["abc"] + r_[3:] for r_ in rows], headers=heads, weights_text="1;1;1", weights_mode="relative"), "не число"),
                 (dict(rows=rows, headers=heads, weights_text="1;1;1", weights_mode="relative", desired_text="1; 2"), "вкажіть 3 значень")]:
    try: SS.sections_selection(**kw); check(False, f"{frag}: мало впасти")
    except BreedingError as ex: check(frag in str(ex), f"помилка «{frag}»")
rows_one = [r_[:3] for r_ in rows]
try: SS.sections_selection(rows_one, heads[:3], "1", "relative"); check(False, "одна ознака")
except BreedingError as ex: check("щонайменше 2" in str(ex), "одна ознака → зрозуміла помилка")

print("\n== 4. Експорт")
import openpyxl, docx
TMP = tempfile.mkdtemp(); px = os.path.join(TMP, "s.xlsx"); U.export_xlsx(px, S, "Індекси"); wb = openpyxl.load_workbook(px)
check(len(wb.sheetnames) >= 6, f"Excel: {len(wb.sheetnames)} аркушів")
pd_ = os.path.join(TMP, "s.docx"); U.export_docx(pd_, "Індекси", S, True); d = docx.Document(pd_)
check(len(d.tables) >= 6 and len(d.inline_shapes) == 1, f"Word: таблиць {len(d.tables)}, рисунків {len(d.inline_shapes)}")
shutil.rmtree(TMP, ignore_errors=True)
print("\n" + "=" * 60)
if FAIL: print(f"ПРОВАЛІВ: {len(FAIL)}"); [print("  -", f) for f in FAIL]; sys.exit(1)
print("УСІ ПЕРЕВІРКИ РОЗДІЛІВ ІНДЕКСІВ ДОБОРУ ПРОЙДЕНО")
