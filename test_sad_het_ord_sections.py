# test_sad_het_ord_sections.py — розділи гетерозису, бальних ознак і AUDPC
import os, sys, math, tempfile, shutil
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import numpy as np
import sad_het_ord_sections as HS, sad_breeding_ui as U
from sad_gxe_core import BreedingError
FAIL = []
def check(c, msg):
    print(("  [OK]   " if c else "  [FAIL] ") + msg, flush=True)
    if not c: FAIL.append(msg)
def structure_ok(S, label):
    bad = [s["title"] for s in S if s["kind"] == "table" and (len(s["nd"]) != len(s["headers"]) or any(len(r) != len(s["headers"]) for r in s["rows"])
           or (s.get("tags") is not None and len(s["tags"]) != len(s["rows"])))]
    check(not bad, f"{label}: структура таблиць коректна {bad[:2]}")
def figs_ok(S, label):
    from matplotlib.figure import Figure; import matplotlib
    bad = []
    for s in S:
        if s["kind"] == "figure":
            fg = Figure(figsize=s["size"])
            with matplotlib.rc_context(U._RC): m = s["build"](fg)
            if m: bad.append(s["title"] + ": " + m)
    check(not bad, f"{label}: графіки будуються {bad[:2]}")
rg = np.random.default_rng(6)

print("== 1. Гетерозис")
spec = {"П1": (40.0, "", ""), "П2": (50.0, "", ""), "Ст": (47.0, "", ""), "Н1": (62.0, "П1", "П2"), "Н2": (45.0, "П1", "П2"), "Н3": (38.0, "П1", "П2")}
rows = []
for nm, (mu, p1, p2) in spec.items():
    for k in range(4): rows.append([nm, p1, p2, f"R{k+1}", f"{mu + rg.normal(0, 1.0):.3f}", "", ""])
S, w = HS.sections_heterosis(rows, 4, "Урожай", "Ст", "higher")
check([s["key"] for s in S] == ["het", "fig", "help"], "розділи: het, fig, help")
structure_ok(S, "гетерозис"); figs_ok(S, "гетерозис")
t = S[0]; hi = {h: i for i, h in enumerate(t["headers"])}; byh = {r[0]: r for r in t["rows"]}
check(all(h in t["headers"] for h in ("MPH, %", "BPH, %", "SH, %")), "є стовпці MPH, BPH і SH")
check(byh["Н1"][hi["MPH, %"]] > 25 and byh["Н1"][hi["BPH, %"]] > 15, f"Н1: MPH = {byh['Н1'][hi['MPH, %']]:.1f} %, BPH = {byh['Н1'][hi['BPH, %']]:.1f} % (закладено 37,8 і 24)")
check("бажаний" in byh["Н1"][hi["Висновок BPH"]] and "небажаний" not in byh["Н1"][hi["Висновок BPH"]], "Н1: достовірний бажаний гетерозис")
check(byh["Н3"][hi["MPH, %"]] < 0 and "небажаний" in byh["Н3"][hi["Висновок BPH"]], "Н3 (нижчий за обох батьків): небажаний")
tg = {r[0]: tag for r, tag in zip(t["rows"], t["tags"])}
check(tg["Н1"] == "good" and tg["Н3"] == "bad", f"кольори: Н1 — зелений, Н3 — червоний ({tg})")
check(any("Холм" in x for x in w), "про поправку Холма повідомлено")
S2, _ = HS.sections_heterosis(rows, 4, "Висота", "Ст", "lower")
t2 = S2[0]; h2 = {h: i for i, h in enumerate(t2["headers"])}; b2 = {r[0]: r for r in t2["rows"]}
check(b2["Н1"][h2["Кращий батько"]] == "П1" and "небажаний" in b2["Н1"][h2["Висновок BPH"]], "«менше — краще»: кращий батько П1; вищий F₁ — небажаний")
S3, _ = HS.sections_heterosis(rows, 4, "Урожай", "", "higher")
check("SH, %" not in S3[0]["headers"], "без стандарту стовпців SH немає")
for kw, frag in [(dict(rows=[r_[:2] + [""] + r_[3:] if r_[0] == "Н2" else r_ for r_ in rows]), "лише одного батька"),
                 (dict(rows=[r_[:1] + ["П2", "П1"] + r_[3:] if (r_[0] == "Н1" and r_[3] == "R2") else r_ for r_ in rows]), "різних батьків"),
                 (dict(rows=[[r_[0], "", ""] + r_[3:] for r_ in rows]), "жодного гібрида"),
                 (dict(rows=[r_[:4] + ["abc", "", ""] for r_ in rows]), "не число")]:
    try: HS.sections_heterosis(trait_col=4, trait_name="Y", **kw); check(False, f"{frag}: мало впасти")
    except BreedingError as ex: check(frag in str(ex), f"помилка «{frag}»")

print("\n== 2. Бальні ознаки")
def scores(n, p):
    return rg.choice(np.arange(1, 10), n, p=np.array(p) / np.sum(p)).astype(float)
lo = [1, 3, 4, 2, 0, 0, 0, 0, 0]; mid = [0, 0, 1, 3, 4, 2, 1, 0, 0]; hi_ = [0, 0, 0, 0, 0, 1, 3, 4, 2]
rows = []
for nm, pr_ in (("Стійкий1", lo), ("Стійкий2", lo), ("Середній", mid), ("Сприйнятливий", hi_), ("Сприйн2", hi_)):
    for v in scores(20, pr_): rows.append([nm, "R1", f"{v:g}", "", ""])
S, w = HS.sections_scores(rows, 2, "Ураження", 1, 9, "holm", 0.05, lower_better=True)
check([s["key"] for s in S] == ["desc", "kw", "dunn", "dist", "help"], "розділи: desc, kw, dunn, dist, help")
structure_ok(S, "бали"); figs_ok(S, "бали")
d = S[0]; names = [r[1] for r in d["rows"]]; lt = {r[1]: r[9] for r in d["rows"]}
check(names[0].startswith("Стійкий") and names[-1].startswith("Сприйн"), f"впорядковано від кращого (менше — краще): {names}")
check(set(lt["Стійкий1"]) & set(lt["Стійкий2"]) and set(lt["Сприйнятливий"]) & set(lt["Сприйн2"]) and not (set(lt["Стійкий1"]) & set(lt["Сприйнятливий"])),
      f"літери: однаково стійкі/сприйнятливі збігаються, стійкі й сприйнятливі різні ({lt})")
kw = S[1]; check(kw["rows"][2][1][0] < 1e-6, "Крускал–Уолліс значущий")
dun = S[2]; pr = {(r[0], r[1]): r[4] for r in dun["rows"]}
check(len(dun["rows"]) == 10 and pr[("Стійкий1", "Сприйнятливий")] < 1e-6 and pr[("Стійкий1", "Стійкий2")] > 0.05, "Данн: 10 пар; стійкий × сприйнятливий — різниця, однакові — ні")
Sh, _ = HS.sections_scores(rows, 2, "Ураження", 1, 9, "holm", 0.05, lower_better=False)
check([r[1] for r in Sh[0]["rows"]][0].startswith("Сприйн"), "«більше — краще»: першими — генотипи з вищими балами")
small = rows[:3] + [r_ for r_ in rows if r_[0] == "Сприйнятливий"]
_, ws = HS.sections_scores(small, 2, "Y", 1, 9)
check(any("Менше 5 спостережень" in x for x in ws), "мало спостережень → попередження про ненадійність літер")
for kw_, frag in [(dict(rows=[r_[:2] + ["12", "", ""] if i == 0 else r_ for i, r_ in enumerate(rows)]), "поза шкалою"), (dict(rows=rows[:20]), "щонайменше 2"),
                  (dict(rows=rows, smin="9", smax="1"), "більшим")]:
    try: HS.sections_scores(trait_col=2, trait_name="Y", **kw_); check(False, f"{frag}: мало впасти")
    except BreedingError as ex: check(frag in str(ex), f"помилка «{frag}»")

print("\n== 3. AUDPC")
heads = ["Генотип", "Повторність", "01.06.2024", "08.06.2024", "15.06.2024", "22.06.2024"]
rows = []
for nm, rate in (("Стійкий", 2.0), ("Стійкий_б", 2.2), ("Середній", 6.0), ("Вразливий", 11.0), ("Вразливий_б", 10.5)):
    for k in range(6):
        rows.append([nm, f"R{k+1}"] + [f"{min(100, rate * t + rg.normal(0, 1) + 1):.1f}" for t in (0, 7, 14, 21)])
S, w = HS.sections_audpc(rows, heads, "holm", 0.05, ymax="100")
check([s["key"] for s in S] == ["audpc", "kw", "dunn", "curves", "help"], "розділи AUDPC")
structure_ok(S, "AUDPC"); figs_ok(S, "AUDPC")
a = S[0]; nm = [r[1] for r in a["rows"]]
check(nm[0].startswith("Стійкий") and nm[-1].startswith("Вразливий"), f"впорядковано від меншої AUDPC: {nm}")
exp_stable = np.mean([(sum(float(x) for x in (r_[2], r_[3])) / 2 * 7 + (float(r_[3]) + float(r_[4])) / 2 * 7 + (float(r_[4]) + float(r_[5])) / 2 * 7) for r_ in rows if r_[0] == "Стійкий"])
got = {r[1]: r[3] for r in a["rows"]}["Стійкий"]
check(abs(got - exp_stable) < 1e-9, f"AUDPC стійкого = {got:.2f} збігається з ручним розрахунком за трапеціями ({exp_stable:.2f})")
raud = {r[1]: r[5] for r in a["rows"]}["Стійкий"]; check(abs(raud - got / (100 * 21)) < 1e-9, "rAUDPC = AUDPC/(ymax·21 діб) при ymax = 100")
_, wa = HS.sections_audpc([r_[:5] + [r_[5]] for r_ in rows[:12]], heads, "holm", 0.05)
for h_, r_, frag in [(heads[:2] + ["15.06.2024", "08.06.2024"], [x[:4] for x in rows], "зростати"), (heads[:3], [x[:3] for x in rows], "щонайменше 2"),
                     (heads[:2] + ["вчора", "завтра"], [x[:4] for x in rows], "розпізнати")]:
    try: HS.sections_audpc(r_, h_); check(False, f"{frag}: мало впасти")
    except BreedingError as ex: check(frag in str(ex), f"помилка «{frag}»")

print("\n== 4. Експорт")
import openpyxl, docx
TMP = tempfile.mkdtemp(); S, _ = HS.sections_scores(rows[:0] or [[f"G{i%3}", "R", f"{(i % 7) + 1}", "", ""] for i in range(30)], 2, "Y", 1, 9)
px = os.path.join(TMP, "o.xlsx"); U.export_xlsx(px, S, "Бали"); check(len(openpyxl.load_workbook(px).sheetnames) >= 3, "Excel")
pd_ = os.path.join(TMP, "o.docx"); U.export_docx(pd_, "Бали", S, True); d_ = docx.Document(pd_); check(len(d_.tables) >= 3 and len(d_.inline_shapes) == 1, f"Word: таблиць {len(d_.tables)}, рисунків {len(d_.inline_shapes)}")
shutil.rmtree(TMP, ignore_errors=True)
print("\n" + "=" * 60)
if FAIL: print(f"ПРОВАЛІВ: {len(FAIL)}"); [print("  -", f) for f in FAIL]; sys.exit(1)
print("УСІ ПЕРЕВІРКИ РОЗДІЛІВ ГЕТЕРОЗИСУ Й БАЛЬНИХ ОЗНАК ПРОЙДЕНО")
