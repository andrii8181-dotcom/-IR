# test_sad_weather_export.py – перевірка експорту: файли записуються й ЧИТАЮТЬСЯ НАЗАД
import os, sys, re, math, tempfile, shutil, datetime as dt
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import numpy as np
import sad_weather_core as W, sad_weather_charts as C, sad_weather_export as X

FAIL = []
def check(cond, msg):
    print(("  [OK]   " if cond else "  [FAIL] ") + msg)
    if not cond: FAIL.append(msg)

def synth(d0, d1, seed=1):
    rng = np.random.default_rng(seed)
    dts = []; cur = d0
    while cur <= d1: dts.append(cur); cur += dt.timedelta(days=1)
    doy = np.array([d.timetuple().tm_yday for d in dts], float)
    tm = 9.5 + 11.5 * np.sin(2*np.pi*(doy-110)/365.0) + rng.normal(0, 2.2, len(dts))
    pr = np.round(np.where(rng.random(len(dts)) < 0.30, rng.gamma(1.1, 5.0, len(dts)), 0.0), 1)
    s = W.DailySeries(dts, {"tmean": tm, "tmax": tm+4.5, "tmin": tm-4.5, "prec": pr,
                            "rh": np.full(len(dts), 70.0), "wind": np.full(len(dts), 3.0),
                            "rad": np.full(len(dts), 15.0), "et0": np.maximum(0, (tm+3)*0.2)})
    s.meta.update({"source": "nasa", "timezone": "Europe/Kyiv",
                   "grid_lat": 48.75, "grid_lon": 30.25, "grid_elevation_m": 215.0,
                   "retrieved_at": "2026-10-09T12:00:00"})
    return s

tmp = tempfile.mkdtemp(prefix="sadw_exp_")
s = synth(dt.date(2021,1,1), dt.date(2024,12,31))
st = W.default_settings()
okw = {y: sub for y, _, _, sub in W.year_windows(s, st["window"])[0]}
nrm, used = W.normal_from_years(okw)
periods = W.intervals_from_dates([dt.date(2024,5,15), dt.date(2024,6,12)], dt.date(2024,5,1)) + [
    ("поза даними", dt.date(2030,1,1), dt.date(2030,1,5)),
    ("частково", dt.date(2024,12,25), dt.date(2025,1,10))]
rep = W.build_report(s, st, norm=(nrm, used, "2021–2024"), periods=periods)
meta = X.make_meta(rep, "Умань, дослідне поле", 48.7484, 30.2218, "nasa")
pub = X.publication_text(rep, "Умань, дослідне поле", 48.7484, 30.2218, "nasa")

print("== 1. Метадані й текст для публікації")
md = dict(meta)
check(md["Джерело даних"].startswith("NASA POWER") and "48.7484" in md["Координати точки"], "джерело й координати")
check("відстань" in md["Центр комірки сітки"] and "км" in md["Центр комірки сітки"], "відстань до центру комірки сітки показано")
dist = W.haversine_km(48.7484, 30.2218, 48.75, 30.25)
check(f"{dist:.1f}" in md["Центр комірки сітки"], f"відстань {dist:.1f} км обчислена коректно")
check(md["Дата отримання даних"] == "09.10.2026 12:00", f"дата отримання даних по-українськи: {md['Дата отримання даних']}")
for need in ("МЕТОДИКА ЗБОРУ МЕТЕОДАНИХ", "ДЖЕРЕЛА ТА ПОДЯКА", "NASA Langley", "УМОВИ ВИКОРИСТАННЯ", "ШКАЛА ВОЛОГОЗАБЕЗПЕЧЕНОСТІ"):
    check(need in pub, f"текст містить «{need}»")
for bad in ("Польов", "Божко", "Вольвач", "Основи агрометеорології", "ТЕС, 2004", "подано за", "Джерело шкали", "Open-Meteo", "ERA5", "некомерційного"):
    check(bad not in pub, f"текст для публікації не містить «{bad}»")

print("\n== 2. Рядкові таблиці (десяткова кома, компактність)")
sp = X.spec_months(rep)
h_full, r_full, c_full = X.to_strings(sp, comma=True, compact=False)
h_cmp, r_cmp, c_cmp = X.to_strings(sp, comma=True, compact=True)
check(len(h_cmp) < len(h_full) and "ГТК" in h_cmp and "Умови зволоження" in h_cmp, "компактна таблиця коротша й містить ГТК")
check(not any(re.fullmatch(r"-?\d+\.\d+", c) for row in r_full for c in row), "жодне число в таблиці не має десяткової крапки (лише кома)")
i_apr = next(i for i, r in enumerate(r_full) if r[0].startswith("кві 2022"))
check(any("," in c for c in r_full[i_apr]), "кома в числах рядка")
jan = next(i for i, r in enumerate(r_full) if r[0].startswith("січ 2022"))
check(r_full[jan][h_full.index("ГТК")] == "–" and c_full[jan] == "na", "січень: ГТК «–» (немає активних температур)")
check(any(c in ("d1","d2","d3","d4","ok","wet") for c in c_full), "для активних місяців є клас ГТК")
check(X.clipboard_text(sp, comma=True).count("\t") >= len(h_full) - 1, "TSV для буфера")
yr_spec = X.spec_years(rep); yh, yr, yc = X.to_strings(yr_spec, True, False)
check("Опади, % норми" in yh and "Δt від норми, °C" in yh, "таблиця за роками містить стовпці норми й аномалій")
ph, pr_, pc = X.to_strings(X.spec_periods(rep), True, False)
check(pr_[-2][4] == "немає даних" and pr_[0][4] != "немає даних", "період повністю поза даними → «немає даних», решта – числа")
check(pr_[-1][3] == "7 з 17" and pr_[-1][4] != "немає даних", f"частково покритий період позначено «{pr_[-1][3]}» (7 діб із 17), а не мовчки")
check(pr_[0][3] == "15", "повністю покритий період – просто число діб")
tot = X.spec_total(rep); th, trs, _ = X.to_strings(tot, True)
check(any("ГТК Селянінова" in r[0] for r in trs) and any("Водний баланс" in r[0] for r in trs), "зведена таблиця: ГТК і водний баланс (є ET₀)")

print("\n== 3. CSV добових даних")
p_csv = os.path.join(tmp, "daily.csv")
X.export_csv_daily(p_csv, s, comma=True, sep=";")
raw = open(p_csv, "rb").read()
check(raw[:3] == b"\xef\xbb\xbf", "UTF-8 із BOM (Excel читає кирилицю)")
lines = raw.decode("utf-8-sig").split("\r\n")
check(lines[0].startswith("Дата;tсер, °C;") and lines[1].startswith("01.01.2021;"), "заголовок і дата ДД.ММ.РРРР")
check(len(lines) - 1 - 1 == len(s), f"рядків даних = {len(s)}")
check("," in lines[1] and "." not in lines[1].split(";", 1)[1], "десяткова кома в CSV")
X.export_csv_daily(p_csv, s, comma=False, sep=",")
check("." in open(p_csv, encoding="utf-8-sig").read().splitlines()[1], "режим «крапка» теж працює")

print("\n== 4. Excel: читаємо назад")
import openpyxl
p_x = os.path.join(tmp, "weather.xlsx")
X.export_xlsx(p_x, rep, meta, pub)
wb = openpyxl.load_workbook(p_x)
check(wb.sheetnames == ["Метадані", "Добові дані", "Підсумок", "Місяці", "Декади", "За роками", "Періоди", "Індекс сухості", "Норма"],
      f"аркуші (є власний ET₀ у даних → з'явився «Індекс сухості»): {wb.sheetnames}")
wd = wb["Добові дані"]
check(wd.max_row == len(s) + 1, "добових рядків стільки, скільки діб")
check(isinstance(wd["B2"].value, float) and isinstance(wd["A2"].value, dt.datetime), "числа – числа, дати – дати (не текст)")
check(wd["A2"].number_format == "DD.MM.YYYY", "формат дати ДД.ММ.РРРР")
wm = wb["Місяці"]
hdr = [c.value for c in wm[1]]
ig = hdr.index("ГТК") + 1
active = [row for row in wm.iter_rows(min_row=2) if isinstance(row[ig-1].value, float)]
check(len(active) > 10, f"ГТК у числовому вигляді для {len(active)} місяців")
filled = [row[ig-1].fill.fgColor.rgb for row in active]
check(all(f and f.endswith(tuple(X.HTC_FILL.values())) for f in filled), "комірки ГТК пофарбовані за класом")
check(wm.freeze_panes == "A2", "закріплено рядок заголовків")
wy = wb["За роками"]; yh_ = [c.value for c in wy[1]]
check("Опади, % норми" in yh_ and wy.max_row == 5, "аркуш «За роками»: 4 роки + норма й аномалії")
check(wb["Норма"].max_row >= 8 and wb["Норма"]["A1"].value == "Показник", "аркуш «Норма»")
check("ШКАЛА ВОЛОГОЗАБЕЗПЕЧЕНОСТІ" in " ".join(str(r[0].value) for r in wb["Метадані"].iter_rows()), "шкала ГТК у метаданих")
wp = wb["Періоди"]; vals = [[c.value for c in r] for r in wp.iter_rows(min_row=2)]
check(vals[-2][4] == "немає даних" and vals[-1][3] == "7 з 17", "аркуш «Періоди»: відсутні дані й неповне покриття позначено")

print("\n== 5. Word: читаємо назад")
import docx
png1 = X.figure_png(rep, "walter", "Умань", C.build_chart, dpi=80)
png2 = X.figure_png(rep, "htc_month", "Умань", C.build_chart, dpi=80)
check(png1 and png1[:8] == b"\x89PNG\r\n\x1a\n", "рисунок побудовано як PNG без вікна")
p_d = os.path.join(tmp, "weather.docx")
X.export_docx(p_d, rep, meta, pub, [("Діаграма Вальтера", png1), ("ГТК за місяцями", png2)], comma=True)
d = docx.Document(p_d)
texts = [p.text for p in d.paragraphs]
check(any(t.startswith("Таблиця 1 –") for t in texts) and any(t.startswith("Таблиця 2 –") for t in texts), "підписи таблиць «Таблиця N – …»")
check(len(d.tables) >= 3, f"таблиць у документі: {len(d.tables)}")
check(len(d.inline_shapes) == 2, "два рисунки вбудовано")
check(any(t.startswith("Рис. 1 –") for t in texts) and any(t.startswith("Рис. 2 –") for t in texts), "підписи рисунків")
tm_ = d.tables[0]
check(tm_.rows[0].cells[0].text == "Місяць" and any("," in c.text for c in tm_.rows[5].cells), "таблиця місяців: заголовок, десяткова кома")
hi = [c.text for c in tm_.rows[0].cells].index("ГТК")
shaded = 0
for row in tm_.rows[1:]:
    tcPr = row.cells[hi]._tc.tcPr
    if tcPr is not None and tcPr.xml.find("w:shd") != -1 and "FFFFFF" not in tcPr.xml and "D9E2F3" not in tcPr.xml: shaded += 1
check(shaded > 5, f"комірки ГТК у Word пофарбовано ({shaded})")
check(any("Джерела даних" == t for t in texts) and any("Шкала вологозабезпеченості за ГТК" == t for t in texts)
      and any("Примітки щодо якості даних" == t for t in texts), "розділи: джерела, примітки, шкала – кожен зі своїм заголовком")
check(not any("УМОВИ ВИКОРИСТАННЯ" in t for t in texts), "службовий блок про ліцензію не потрапив у текст для публікації")
check(any("NASA POWER" in t and "MERRA-2" in t for t in texts), "методика згадує NASA POWER і MERRA-2")
check(not any(("Польов" in t or "Божко" in t or "подано за" in t or "Open-Meteo" in t) for t in texts), "у Word немає згадки джерела градації ГТК та Open-Meteo")
check(len(d.tables) <= 4 and all("Декада" != t.rows[0].cells[0].text for t in d.tables), "декадна таблиця (108 рядків) не роздуває Word – вона в Excel")


print("\n== 6. Міжнародні індекси в експорті")
import re as _re
# повний звіт із широтою: ET₀ розрахується, є вінограрські індекси, індекс сухості та SPI
def synth_full(d0, d1, seed=1):
    rg2 = np.random.default_rng(seed); n = (d1 - d0).days + 1; dts = [d0 + dt.timedelta(days=i) for i in range(n)]
    doy = np.array([d.timetuple().tm_yday for d in dts], float); tm = 10.5 + 11.5 * np.sin(2 * np.pi * (doy - 110) / 365) + rg2.normal(0, 2, n)
    return W.DailySeries(dts, {"tmean": tm, "tmax": tm + 5, "tmin": tm - 5, "prec": np.where(rg2.random(n) < 0.3, rg2.gamma(1.3, 5.0, n), 0.0),
                               "rad": np.full(n, 15.0), "wind": np.full(n, 2.5), "rh": np.full(n, 70.0)}, {"grid_elevation_m": 212.0, "retrieved_at": "2026-10-09T12:00:00"})
import sad_weather_indices as WI
sf = synth_full(dt.date(1988, 1, 1), dt.date(2024, 12, 31)); stf = W.default_settings(); stf["lat"] = 48.75
sp = WI.spi_table(sf, ref=(1991, 2020), scales=(1, 3), study=(dt.date(2023, 1, 1), dt.date(2024, 12, 31)))
subf = sf.slice(dt.date(2021, 1, 1), dt.date(2024, 12, 31)); repf = W.build_report(subf, stf, spi=sp)
metaf = X.make_meta(repf, "Умань", 48.75, 30.22, "nasa"); pubf = X.publication_text(repf, "Умань", 48.75, 30.22, "nasa")
keys = [sp_["key"] for sp_ in X.all_specs(repf)]; check(keys == ["total", "months", "decades", "years", "vine", "aridity", "spi"], f"спеціфікації: {keys}")
check(dict(metaf).get("Еталонне випаровування ET₀", "").startswith("FAO-56"), "метадані: метод ET₀ вказано")
hd, rw, cl, cf = X.to_strings_ex(X.spec_vine(repf), True, False)
check(hd == ["Сезон", "Вінклер (GDD >10 °C), °C·діб", "Регіон за Вінклером", "Індекс Югліна (HI)", "Клас за Югліном", "k (довжина дня)"] and len(rw) == 4, "таблиця вінограрських індексів: заголовки й 4 сезони")
check(all(_re.fullmatch(r"\d{3,4}", r[1]) for r in rw) and all(r[2].split(" ")[0] in ("I", "II", "III", "IV", "V") for r in rw) and all(r[4].startswith("HI") for r in rw) and all(r[5] == "1,06" for r in rw), f"значення в рядках: {rw[0]}")
check(all(set(c) == {1, 2, 3, 4} and all(v.startswith("#") and len(v) == 7 for v in c.values()) for c in cf), "кожен рядок має заливки класів для чотирьох колонок (число й текст класу)")
h2, r2, c2 = X.to_strings(X.spec_vine(repf), True, True); check(h2 == hd[:5] and len(r2) == 4, "стисла форма (для Word) без колонки k; to_strings сумісна зі старим інтерфейсом")
ha, ra, ca, fa = X.to_strings_ex(X.spec_aridity(repf), True, False); check(ha == ["Рік", "Опади, мм", "ET₀ (PET), мм", "Індекс сухості P/ET₀", "Клас ЮНЕП"] and len(ra) == 4 and all(float(r[3].replace(",", ".")) > 0 for r in ra), "таблиця індексу сухості")
hs_, rs_, cs_, fs_ = X.to_strings_ex(X.spec_spi(repf), True, False)
check(hs_ == ["Місяць", "Опади, мм", "SPI-1", "Оцінка SPI-1", "SPI-3", "Оцінка SPI-3"] and len(rs_) == 24, f"таблиця SPI: {hs_}, 24 місяці")
hz, rz, _, fz = X.to_strings_ex(X.spec_spi(repf), True, True); check(hz == ["Місяць", "SPI-1", "SPI-3", "Оцінка SPI-3"], f"стисла форма SPI: {hz}")
row_norm = [r for r in rs_ if r[3] == "норма"]; check(row_norm and all(not f for f, r in zip(fs_, rs_) if r[3] == "норма" and r[5] == "норма"), "клас «норма» не фарбується")
ex_ = [(r, f) for r, f in zip(rs_, fs_) if r[5] in ("помірно волого", "дуже волого", "екстремально волого", "помірно сухо", "дуже сухо", "екстремально сухо")]
check(all(f.get(4) and f.get(5) for r, f in ex_), "для SPI-3 поза нормою фарбуються число й оцінка")
# Excel
px = os.path.join(tmp, "idx.xlsx"); X.export_xlsx(px, repf, metaf, pubf); wbi = openpyxl.load_workbook(px)
check({"Вінклер і Югліна", "Індекс сухості", "SPI"} <= set(wbi.sheetnames), f"Excel: нові аркуші {wbi.sheetnames[-4:]}")
wsv = wbi["Вінклер і Югліна"]; check(isinstance(wsv["B2"].value, float) and wsv["B2"].number_format == "0" and isinstance(wsv["C2"].value, str) and wsv["C2"].fill.fgColor.rgb[-6:] in X.WK_FILL.values(), "Excel: Вінклер – число, регіон – текст із заливкою за класом")
check(isinstance(wsv["D2"].value, float) and wsv["E2"].value.startswith("HI") and wsv["E2"].fill.fgColor.rgb[-6:] in X.HI_FILL.values(), "Excel: Югліна – число, клас – текст із заливкою")
wsa = wbi["Індекс сухості"]; check(wsa["D2"].number_format == "0.00" and wsa["E2"].fill.fgColor.rgb[-6:] in X.AI_FILL.values(), "Excel: індекс сухості з форматом 0,00 і заливкою класу")
wss = wbi["SPI"]; check(wss.max_row == 25 and wss["C1"].value == "SPI-1" and wss["D1"].value == "Оцінка SPI-1" and all(isinstance(wss.cell(r, 3).value, float) for r in range(2, 26)), "Excel: SPI – 24 місяці, числа й оцінки")
wsy = wbi["За роками"]; heads = [c.value for c in wsy[1]]; check("ET₀, мм" in heads and "P/ET₀" in heads, "Excel: у таблиці «За роками» є ET₀ і P/ET₀")
# Word
pd_ = os.path.join(tmp, "idx.docx"); X.export_docx(pd_, repf, metaf, pubf, []); dd = docx.Document(pd_)
th = [[c.text for c in t.rows[0].cells] for t in dd.tables]; check(any(h[0] == "Сезон" for h in th) and any("Індекс сухості" in " ".join(h) for h in th) and any("SPI-1" in h for h in th), "Word: таблиці Вінклера/Югліна, індексу сухості та SPI (24 місяці ≤ 45)")
vt = [t for t in dd.tables if t.rows[0].cells[0].text == "Сезон"][0]
shaded_v = sum(1 for r in vt.rows[1:] for c in r.cells if c._tc.tcPr is not None and "w:shd" in c._tc.tcPr.xml and any(f in c._tc.tcPr.xml for f in list(X.WK_FILL.values()) + list(X.HI_FILL.values())))
check(shaded_v >= 8, f"Word: клітинки класів вінограрських індексів пофарбовано ({shaded_v})")
allp = "\n".join(p.text for p in dd.paragraphs); check("Міжнародні агрокліматичні показники" in allp and "Література до міжнародних показників" in allp and "[1]" in allp, "Word: розділи про методи й літературу")
long_sp = WI.spi_table(sf, ref=(1991, 2020), scales=(1,), study=(dt.date(2015, 1, 1), dt.date(2024, 12, 31))); rep_long = W.build_report(subf, stf, spi=long_sp)
pl = os.path.join(tmp, "long.docx"); X.export_docx(pl, rep_long, metaf, X.publication_text(rep_long, "Умань", 48.75, 30.22, "nasa"), []); dl = docx.Document(pl)
check(not any("SPI-1" in [c.text for c in t.rows[0].cells] for t in dl.tables) and any(c.text == "Місяць" for t in dl.tables for c in t.rows[0].cells[:1]), "SPI на 120 місяців (> 45 рядків) у Word не вставляється (лише в Excel), решта таблиць на місці")
# текст для публікації: нумерація й склад літератури
par, refs = X.intl_text(repf); first = []
for grp in _re.findall(r"\[(\d+(?:, \d+)*)\]", par):
    for x in grp.split(", "):
        if int(x) not in first: first.append(int(x))
check(first == list(range(1, 8)) and len(refs) == 7 and refs[0].startswith("1. Allen") and refs[1].startswith("2. World Atlas") and refs[5].startswith("6. McKee") and refs[6].startswith("7. Standardized"), f"нумерація за порядком першої появи: {first}")
for need in ("Пенмана–Монтейта", "P/ET₀", "гіперарідний", "1389, 1667, 1944 і 2222", "k = 1,06", "Стандартизований індекс опадів", "гамма-розподіл", "1991–2020", "масштабів 1, 3 міс."):
    check(need in par, f"текст методів містить «{need}»")
check(par.count(". ") >= 4 and " . " not in par and ".," not in par and ",," not in par, "розділові знаки в абзаці методів не пошкоджені (крапки між реченнями на місці)")
check(pubf.index("МІЖНАРОДНІ АГРОКЛІМАТИЧНІ ПОКАЗНИКИ") < pubf.index("ДЖЕРЕЛА ТА ПОДЯКА") < pubf.index("ЛІТЕРАТУРА ДО МІЖНАРОДНИХ ПОКАЗНИКІВ"), "порядок блоків: методи індексів → подяка NASA → література")
# залежність від наявних даних
st0 = W.default_settings(); rep0 = W.build_report(subf, st0); par0, refs0 = X.intl_text(rep0)
check(rep0["intl"]["pet_method"] == "none" and not rep0["intl"]["vine"] and par0 == "" and refs0 == [] and "МІЖНАРОДНІ" not in X.publication_text(rep0, "", None, None, "nasa"), "без широти: ET₀ і вінограрські індекси не розраховуються, розділ про методи не з'являється")
nr_ = W.DailySeries(subf.dates, {"tmean": subf.get("tmean"), "tmax": subf.get("tmax"), "tmin": subf.get("tmin"), "prec": subf.get("prec")})
rep_hs = W.build_report(nr_, stf); par_hs, refs_hs = X.intl_text(rep_hs)
check(rep_hs["intl"]["pet_method"] == "hs" and "Харгрівза–Самані" in par_hs and refs_hs[0].startswith("1. Allen") and refs_hs[1].startswith("2. Hargreaves"), "лише tmax/tmin → Харгрівз–Самані, у літературі FAO-56 і Hargreaves & Samani")
check("Пенмана" not in par_hs and "Стандартизований" not in par_hs, "метод, якого не застосовано, у тексті не згадується")
rep_s = W.build_report(subf, stf); par_s, refs_s = X.intl_text(rep_s); check("Стандартизований" not in par_s and all("McKee" not in r_ for r_ in refs_s), "без SPI у тексті немає ані SPI, ані McKee у літературі")
south = synth_full(dt.date(2019, 1, 1), dt.date(2024, 12, 31)); sts = W.default_settings(); sts["lat"] = -34.0; rep_south = W.build_report(south, sts)
par_s2, _ = X.intl_text(rep_south); check("з 1 жовтня по 30 квітня" in par_s2 and "з 1 жовтня по 31 березня" in par_s2 and "k = 1,00" in par_s2, "південна півкуля: сезони з 1 жовтня; k = 1,00 для 34° пд.")
check("1.10–30.04" in X.spec_vine(rep_south)["title"] and X.spec_vine(rep_south)["rows"][0]["label"] == "2019/20", "заголовок і підписи сезонів для півдня: 2019/20, 1.10–30.04")
# погодний набір
wsx = W.weather_set(repf, "Умань"); check("ai" in wsx["years"]["2022"] and wsx["years"]["2022"]["ai"] > 0 and any(i["key"] == "ai" for i in wsx["indicators"]), "погодний набір для аналізів містить P/ET₀ за роками")
# власний ET₀ в даних не перезаписується
gv_ = synth_full(dt.date(2021, 1, 1), dt.date(2022, 12, 31)); data_ = dict(gv_.data); data_["et0"] = np.full(len(gv_), 3.0); gv2 = W.DailySeries(gv_.dates, data_, gv_.meta)
rg_ = W.build_report(gv2, stf); check(rg_["intl"]["pet_method"] == "given" and np.allclose(rg_["series"].get("et0"), 3.0) and abs(rg_["total"]["et0_sum"] - 3.0 * len(gv2)) < 1e-6, "якщо ET₀ є в даних користувача, програма її не замінює власним розрахунком")

shutil.rmtree(tmp, ignore_errors=True)
print("\n" + "=" * 56)
if FAIL:
    print(f"ПРОВАЛІВ: {len(FAIL)}")
    for f in FAIL: print("  -", f)
    sys.exit(1)
print("УСІ ПЕРЕВІРКИ ЕКСПОРТУ ПРОЙДЕНО")
